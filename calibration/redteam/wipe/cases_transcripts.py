"""Red team of wipe mode, dimension transcripts: meeting transcripts and chat logs (Teams, Zoom, Webex and Slack
exports, WebVTT, "[10:03] Name:" lines, Q&A and interview formats, neutral speaker labels, first names, nicknames,
lower-case names, initials, German, English and Russian speakers; the losses are product names, service names,
SQL and labels spoken or written in the transcript).

Every name is from tests/fixtures.py or from INVENTED below. Hosts end in .example. Nothing here is real.
"""
import json
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[3]))
from tests import fixtures as fx  # noqa: E402

# the fixture forms
CUSTOMER, CUSTOMER_SHORT, CUSTOMER_ACRONYM, CUSTOMER_EN = fx.CUSTOMER_FORMS
PERSON, SURNAME = fx.PERSON_FORMS            # a registered person (becomes a PERS code, not a token)
FIRST = fx.PLANTED_PERSON.split()[0]         # Xqarv
LAST = fx.PLANTED_PERSON.split()[1]          # Pomblet
PLANTED_PERSON = fx.PLANTED_PERSON           # Xqarv Pomblet, unregistered
PLANTED_COMPANY = fx.PLANTED_CANDIDATE       # Nrgtz Beratung GmbH, unregistered
BRAND = PLANTED_COMPANY.split()[0]           # Nrgtz
CONTROL = fx.CONTROL_UNREGISTERED            # Musterhausen, unregistered, a place word

# every invented name of this module, so that the owner can see nothing real slipped in
INVENTED = (
    "Tobias Qwertzuio",      # an invented person, first name Tobias, nickname Tobi, initials TQ
    "Anna Zrbl",             # an invented person
    "Peter Pomblet",         # an invented person whose first name is a word of the public corpus
    "Иван Квжрт",            # an invented Russian speaker
    "Ольга",                 # a Russian first name alone
    "Markus",                # a German first name alone, never a speaker
    "tobi",                  # a Slack display name
)

P2 = "Tobias Qwertzuio"
P2_FIRST, P2_LAST = P2.split()
P3 = "Anna Zrbl"
P3_FIRST, P3_LAST = P3.split()
P4 = "Peter " + LAST                          # a known first name before the fixture surname
RU = "Иван Квжрт"
RU_FIRST, RU_LAST = RU.split()

MARK = "stage two of the invented plan"


def txt(name, text):
    """A plain file with `text`; the marker line proves the carrier was read."""
    def build(inbox):
        p = inbox / name
        p.write_text(text + "\n" + MARK + "\n", encoding="utf-8")
        return p
    return build


def raw(name, data: bytes):
    def build(inbox):
        p = inbox / name
        p.write_bytes(data)
        return p
    return build


def docx_paragraphs(name, paragraphs):
    """A Word file of plain paragraphs (python-docx)."""
    def build(inbox):
        import docx
        d = docx.Document()
        for para in paragraphs:
            d.add_paragraph(para)
        d.add_paragraph(MARK)
        p = inbox / name
        d.save(str(p))
        return p
    return build


def docx_table(name, rows, intro=""):
    """A Word file with one table (python-docx); the first row is the header."""
    def build(inbox):
        import docx
        d = docx.Document()
        if intro:
            d.add_paragraph(intro)
        t = d.add_table(rows=len(rows), cols=len(rows[0]))
        for i, row in enumerate(rows):
            for j, cell in enumerate(row):
                t.cell(i, j).text = cell
        d.add_paragraph(MARK)
        p = inbox / name
        d.save(str(p))
        return p
    return build


def case(cid, carrier, values, keep, build, note="", code=None, expect=None):
    c = {"id": cid, "carrier": carrier, "values": list(values), "keep": list(keep), "build": build, "note": note}
    if code is not None:
        c["code"] = code
    if expect:
        c["expect"] = dict(expect)
    return c


PERSONS = {PLANTED_PERSON: "person", P2: "person", P3: "person", P4: "person", RU: "person"}


def persons(*names):
    return {n: "person" for n in names}


CASES = [
    # ------------------------------------------------------------------ Teams, Otter, Slack: name, two spaces, time
    case("teams-fullname", "txt, Teams export: full name, two spaces, m:ss; the surname alone in the speech",
         [PLANTED_PERSON, LAST, P3], ["Elastic Cloud Server"],
         txt("teams-transcript.txt",
             "%s  10:03\nHallo zusammen, wir starten mit dem Elastic Cloud Server.\n\n"
             "%s  10:04\n%s hat recht, die Flavors passen.\n\n"
             "%s  10:05\nDann nehmen wir s3.large.2.\n" % (PLANTED_PERSON, P3, LAST, PLANTED_PERSON)),
         expect=persons(PLANTED_PERSON, P3)),
    case("teams-firstname", "txt, Teams export: a guest shown by first name only, two spaces, time",
         [P2_FIRST, P3_FIRST], ["Cloud Eye"],
         txt("teams-transcript.txt",
             "%s  10:03\nHallo zusammen, der Cloud Eye Alarm ist raus.\n\n"
             "%s  10:04\nDanke %s, ich schaue mir Cloud Eye an.\n\n"
             "%s  10:05\nGut so.\n" % (P2_FIRST, P3_FIRST, P2_FIRST, P2_FIRST)),
         expect=persons(P2_FIRST, P3_FIRST)),
    case("teams-seconds", "txt, Teams transcript after the first hour: name, two spaces, h:mm:ss",
         [PLANTED_PERSON, LAST, P3], ["Direct Connect"],
         txt("teams-transcript.txt",
             "%s  1:03:15\nNach der Pause: der Direct Connect ist bestellt.\n\n"
             "%s  1:04:02\n%s, bitte das Angebot schicken.\n\n"
             "%s  1:05:40\nMache ich.\n" % (PLANTED_PERSON, P3, LAST, PLANTED_PERSON)),
         expect=persons(PLANTED_PERSON, P3)),
    case("teams-tab", "txt, Teams-like export with one tab between the name and the time",
         [PLANTED_PERSON, P2_FIRST, LAST], ["Cloud Eye"],
         txt("teams-transcript.txt",
             "%s\t10:03\nHallo zusammen, der Cloud Eye Alarm ist raus.\n\n"
             "%s\t10:04\n%s, hast du den Link?\n\n"
             "%s\t10:05\nJa.\n" % (PLANTED_PERSON, P2_FIRST, LAST, PLANTED_PERSON)),
         expect=persons(PLANTED_PERSON, P2_FIRST)),
    case("teams-known-firstname", "txt, Teams export: a first name that is a word of the public corpus, alone",
         ["Peter", P3], ["NAT Gateway"],
         txt("teams-transcript.txt",
             "Peter  10:03\nHallo, der NAT Gateway ist fertig.\n\n"
             "%s  10:04\nDanke Peter.\n\n"
             "Peter  10:05\nGern.\n" % P3),
         expect=persons("Peter", P3)),
    case("teams-customer-suffix", "txt, Teams export: the registered customer in brackets after the name",
         [P2, P2_LAST, P3], ["Object Storage Service"],
         txt("teams-transcript.txt",
             "%s (%s)  10:03\nWir brauchen den Object Storage Service.\n\n"
             "%s  10:04\n%s, hast du die Bucket-Policy?\n\n"
             "%s (%s)  10:05\nJa.\n" % (P2, CUSTOMER_SHORT, P3, P2_LAST, P2, CUSTOMER_SHORT)),
         code=True, expect=persons(P2, P3)),
    case("teams-docx", "docx, the Teams transcript download: name, three spaces, m:ss as paragraphs",
         [PLANTED_PERSON, P2_FIRST, LAST], ["Cloud Eye"],
         docx_paragraphs("transcript.docx", [
             "%s   0:03" % PLANTED_PERSON, "Hallo zusammen, wir nehmen Cloud Eye.", "",
             "%s   0:12" % P2_FIRST, "%s, der Alarm steht." % LAST, "",
             "%s   0:40" % PLANTED_PERSON, "Gut so.", ""]),
         expect=persons(PLANTED_PERSON, P2_FIRST)),

    # ------------------------------------------------------------------ WebVTT
    case("vtt-v-fullname", "vtt file, Teams WebVTT: <v Full Name>, the surname alone in a cue",
         [PLANTED_PERSON, LAST, P3], ["Direct Connect", "Terraform"],
         txt("transcript.vtt",
             "WEBVTT\n\n00:00:01.000 --> 00:00:04.000\n<v %s>Hallo zusammen, der Direct Connect ist aktiv.</v>\n\n"
             "00:00:04.500 --> 00:00:08.000\n<v %s>Danke, %s hat recht. Dann Terraform apply.</v>\n\n"
             "00:00:08.500 --> 00:00:10.000\n<v %s>Gut so.</v>\n" % (PLANTED_PERSON, P3, LAST, PLANTED_PERSON)),
         expect=persons(PLANTED_PERSON, P3)),
    # ------------------------------------------------------------------ "[10:03] Name:" and the colon shapes
    case("teams-chat-nocolon", "txt, Teams chat copy: [10:03] Name on its own line, the message indented below",
         [P2_FIRST, P3], ["Cloud Eye"],
         txt("teams-chat.txt",
             "[10:03] %s\n    Der Cloud Eye Alarm ist raus.\n"
             "[10:04] %s\n    Danke %s, ich schaue.\n"
             "[10:05] %s\n    Gern.\n" % (P2_FIRST, P3, P2_FIRST, P2_FIRST)),
         expect=persons(P2_FIRST, P3)),
    case("bracket-time-colon", "md, [10:03] Name: lines, first names and a full name, short surname in the speech",
         [P2_FIRST, P3, P3_LAST], ["Object Storage Service"],
         txt("transcript.md",
             "# Transcript\n\n[10:03] %s: Hallo zusammen, wir starten mit dem Object Storage Service.\n"
             "[10:04] %s: Danke %s. %s hier, ich übernehme das Bucket.\n"
             "[10:05] %s: Gut so, %s.\n" % (P2_FIRST, P3, P2_FIRST, P3_LAST, P2_FIRST, P3_LAST)),
         expect=persons(P2_FIRST, P3)),
    case("colon-newline", "txt, the speaker label alone on its line with the colon, the speech on the next line",
         [P3_FIRST, P2], ["Load Balancer"],
         txt("transcript.txt",
             "%s:\nHallo zusammen, ich habe den Load Balancer umgebaut.\n\n"
             "%s:\nDanke. Wann ist der Test?\n\n"
             "%s:\nMorgen früh.\n" % (P3_FIRST, P2, P3_FIRST)),
         expect=persons(P3_FIRST, P2)),
    case("bold-label-inside", "md, **Name:** with the colon inside the bold",
         [P3_FIRST, P2], ["Terraform"],
         txt("notes.md",
             "**%s:** Hallo zusammen, der Terraform-Plan ist durch.\n"
             "**%s:** Danke %s, dann machen wir das Apply.\n"
             "**%s:** Gut so.\n" % (P3_FIRST, P2, P3_FIRST, P3_FIRST)),
         expect=persons(P3_FIRST, P2)),
    case("label-suffix-paren", "txt, Name (Extern): and Name (Gast): speaker labels",
         [P2_FIRST, P3], ["PostgreSQL"],
         txt("transcript.txt",
             "%s (Extern): Hallo zusammen, PostgreSQL läuft.\n"
             "%s (Gast): Danke %s.\n"
             "%s (Extern): Gern.\n" % (P2_FIRST, P3, P2_FIRST, P2_FIRST)),
         expect=persons(P2_FIRST, P3)),
    case("chat-date-prefix", "txt, WhatsApp export: [date, time] Name: text",
         [P2_FIRST, P3], ["Veeam"],
         txt("chat.txt",
             "[12.03.26, 10:03:12] %s: Hallo, ist der Veeam Job durch?\n"
             "[12.03.26, 10:04:40] %s: Ja, seit 9 Uhr.\n"
             "[12.03.26, 10:05:01] %s: Super.\n" % (P2_FIRST, P3, P2_FIRST)),
         expect=persons(P2_FIRST, P3)),
    case("zoom-chat-from-to", "txt, Zoom saved chat: time, From Name to Everyone : text",
         [P2_FIRST, P3], ["VPC Endpoint", "Everyone"],
         txt("meeting_saved_chat.txt",
             "10:03:15\t From %s  to  Everyone : Hallo, der VPC Endpoint ist da.\n"
             "10:04:02\t From %s  to  Everyone : Danke %s.\n"
             "10:05:30\t From %s  to  Everyone : Gern.\n" % (P2_FIRST, P3, P2_FIRST, P2_FIRST)),
         expect=persons(P2_FIRST, P3)),

    # ------------------------------------------------------------------ chat dumps
    case("slack-json", "json, Slack export: real_name and a lower-case display_name",
         [P2, "tobi", P3], ["Kubernetes"],
         txt("2026-03-12.json", json.dumps([
             {"type": "message", "user": "U0ABC", "ts": "1773302592.000100",
              "user_profile": {"real_name": P2, "display_name": "tobi"},
              "text": "der Kubernetes Cluster ist oben"},
             {"type": "message", "user": "U0DEF", "ts": "1773302650.000200",
              "user_profile": {"real_name": P3, "display_name": "anna"},
              "text": "danke tobi, dann teste ich"},
         ], indent=1, ensure_ascii=False)),
         expect=persons(P2, P3)),
    case("chat-lowercase-nick", "txt, lower-case chat nicks as speaker labels, the name capitalised in the speech",
         [P2_FIRST], ["Cloud Eye"],
         txt("chat.txt",
             "%s: hallo zusammen, der Cloud Eye Alarm ist raus\n"
             "%s: danke %s, ich schaue\n"
             "%s: gern\n"
             "%s: %s, hast du den Link?\n" % (P2_FIRST.lower(), P3_FIRST.lower(), P2_FIRST, P2_FIRST.lower(),
                                              P3_FIRST.lower(), P2_FIRST)),
         expect=persons(P2_FIRST)),
    case("chat-lowercase-mention", "txt, capitalised speaker labels, one message writes the name in lower case",
         [P2_FIRST], ["Cloud Eye"],
         txt("chat.txt",
             "%s: Hallo zusammen, der Cloud Eye Alarm ist raus.\n"
             "%s: danke %s, ich schaue.\n"
             "%s: Gern.\n"
             "Note: %s prüft die Schwellwerte.\n" % (P2_FIRST, P3, P2_FIRST.lower(), P2_FIRST, P2_FIRST)),
         expect=persons(P2_FIRST)),

    # ------------------------------------------------------------------ speaker labels
    case("speaker-known-firstname-full", "txt, Name: lines where the first name is a word of the public corpus",
         [P4, LAST, P3], ["Enterprise Router"],
         txt("transcript.txt",
             "%s: Hallo zusammen, der Enterprise Router ist bestellt.\n"
             "%s: Danke Peter. %s, schickst du das Angebot?\n"
             "%s: Mache ich.\n" % (P4, P3, LAST, P4)),
         expect=persons(P4, P3)),
    case("speaker-known-firstname-alone", "txt, Name: lines, a first name that is a word of the public corpus",
         ["Peter", P3], ["Direct Connect"],
         txt("transcript.txt",
             "Peter: Hallo zusammen, Direct Connect steht.\n"
             "%s: Danke Peter.\n"
             "Peter: Gern.\n" % P3),
         expect=persons("Peter", P3)),
    case("speaker-role-prefix", "txt, a role word before the surname in the speaker label",
         [P2_LAST, P3_LAST], ["Projektleiter", "Architekt", "Direct Connect"],
         txt("transcript.txt",
             "Projektleiter %s: Wir starten mit der Migration.\n"
             "Architekt %s: Der Entwurf ist fertig, Direct Connect kommt dazu.\n"
             "Projektleiter %s: Gut so, dann Freigabe.\n" % (P2_LAST, P3_LAST, P2_LAST)),
         expect=persons(P2_LAST, P3_LAST)),
    case("speaker-company-label", "txt, companies as speaker labels: the customer short form and an unregistered one",
         [BRAND], [],
         txt("transcript.txt",
             "%s: Wir brauchen die Umgebung bis Mai.\n"
             "%s: Wir liefern das Konzept.\n"
             "%s: Gut so.\n"
             "%s: Dann bis nächste Woche.\n" % (CUSTOMER_SHORT, BRAND, CUSTOMER_SHORT, BRAND)),
         code=True, expect={BRAND: "company"}),
    case("speaker-initials", "txt, initials as speaker labels, one pair is also a unit word",
         ["TQ", "TB"], ["Anti-DDoS"],
         txt("transcript.txt",
             "TQ: Hallo zusammen, der Anti-DDoS ist aktiv.\n"
             "TB: Danke, dann ist das erledigt.\n"
             "TQ: Ja.\n"
             "TB: Gut so.\n"),
         expect=persons("TQ", "TB")),
    case("speaker-cyrillic", "txt, Russian speakers: Name: lines, a declined surname and the first name in the speech",
         [RU, RU_LAST + "а", RU_FIRST, "Ольга"], ["Elastic Cloud Server"],
         txt("transcript.txt",
             "%s: Привет всем, Elastic Cloud Server готов.\n"
             "Ольга: Спасибо тебе, %s. У %sа есть доступ?\n"
             "%s: Да.\n"
             "Ольга: Хорошо.\n" % (RU, RU_FIRST, RU_LAST, RU)),
         expect=persons(RU, "Ольга")),
    case("speaker-von-company", "txt, Herr Name von Company: as the speaker label",
         [P2_LAST, BRAND, P3], ["Cloud Eye"],
         txt("protokoll.txt",
             "Herr %s von %s: Das Konzept ist fertig.\n"
             "%s: Danke, dann Cloud Eye dazu.\n"
             "Herr %s von %s: Einverstanden.\n" % (P2_LAST, BRAND, P3, P2_LAST, BRAND)),
         expect=persons(P2_LAST, P3)),
    case("speaker-pipe-company", "txt, Name | Company: speaker labels with the company after a pipe",
         [P2, P2_LAST, PLANTED_COMPANY, P3], [],
         txt("transcript.txt",
             "%s | %s: Hallo zusammen, das Konzept ist fertig.\n"
             "%s | %s: Danke. %s, wann kommt das Angebot?\n"
             "%s | %s: Morgen.\n" % (P2, PLANTED_COMPANY, P3, CUSTOMER, P2_LAST, P2, PLANTED_COMPANY)),
         code=True, expect={P2: "person", P3: "person", PLANTED_COMPANY: "company"}),

    # ------------------------------------------------------------------ Q&A, interviews, labels
    case("qa-paren-name", "md, Q (Name): and A (Full Name): question and answer lines",
         [P3_FIRST, P2], ["FunctionGraph"],
         txt("qa.md",
             "## Q&A\n\nQ (%s): Läuft FunctionGraph auch in eu-de?\n"
             "A (%s): Ja, seit letztem Jahr.\n"
             "Q (%s): Und der Preis?\n"
             "A (%s): Pro Aufruf.\n" % (P3_FIRST, P2, P3_FIRST, P2)),
         expect=persons(P3_FIRST, P2)),
    case("qa-ung-von", "txt, a German noun on -ung, von, first name, colon (Rückmeldung von Anna:)",
         [P3_FIRST, P2_FIRST], ["Relational Database Service"],
         txt("notes.txt",
             "Rückmeldung von %s: Die Relational Database Service Instanz ist zu klein.\n"
             "Einschätzung von %s: Dann s3.xlarge.4.\n"
             "Rückmeldung von %s: Passt.\n" % (P3_FIRST, P2_FIRST, P3_FIRST)),
         expect=persons(P3_FIRST, P2_FIRST)),
    case("interview-interviewee", "txt, interview with the labels Interviewer: and Interviewee:",
         [], ["Interviewee", "Interviewer", "Veeam", "PostgreSQL"],
         txt("interview.txt",
             "Interviewer: Wie läuft der Betrieb heute?\n"
             "Interviewee: Wir haben zwei Rechenzentren und Veeam für die Sicherung.\n"
             "Interviewer: Und die Datenbanken?\n"
             "Interviewee: PostgreSQL auf VMware, sonst nichts.\n")),
    case("interview-ib", "txt, the German interview convention I: and B:",
         [], ["I:", "B:", "Veeam"],
         txt("interview.txt",
             "I: Wie sichern Sie heute?\n"
             "B: Mit Veeam auf ein NAS, einmal am Tag.\n"
             "I: Und die Wiederherstellung?\n"
             "B: Haben wir im Sommer getestet.\n")),
    case("neutral-labels-stay", "txt, Speaker 1, Speaker 2, Note:, Action:, Decision: labels",
         [], ["Speaker 1", "Speaker 2", "Note:", "Action:", "Decision:", "Cloud Eye"],
         txt("transcript.txt",
             "Speaker 1: Hallo zusammen, wir starten.\n"
             "Speaker 2: Der Cloud Eye Alarm ist raus.\n"
             "Note: wir prüfen die Schwellwerte.\n"
             "Action: das Angebot geht morgen raus.\n"
             "Decision: wir nehmen Cloud Eye.\n"
             "Speaker 1: Danke.\n")),
    case("unknown-label-repeats", "txt, Fazit: and Vorschlag: as repeated labels between the speakers",
         [P2, P3], ["Fazit", "Vorschlag", "Cloud Eye"],
         txt("protokoll.txt",
             "%s: Wir nehmen Cloud Eye.\n"
             "Fazit: Cloud Eye reicht für den Anfang.\n"
             "%s: Und die Kosten?\n"
             "Vorschlag: erst die kleine Stufe.\n"
             "Fazit: einverstanden.\n"
             "Vorschlag: nächste Woche entscheiden.\n" % (P2, P3)),
         expect=persons(P2, P3)),
    case("kunde-dialogue", "txt, a sales call labelled by role: Berater: and Kunde:",
         [], ["Landing Zone", "Abnahme", "Migration", "Direct Connect"],
         txt("call.txt",
             "Berater: Guten Morgen, was steht an?\n"
             "Kunde: Wir wollen eine Landing Zone bis Mai, dann die Abnahme.\n"
             "Berater: Gern, mit Direct Connect?\n"
             "Kunde: Ja, und die Migration der Datenbanken im Juni.\n")),
    case("partner-dialogue", "txt, a call labelled by role: Berater: and Partner:",
         [], ["Landing Zone", "Abnahme", "Cloud Eye", "Partner"],
         txt("call.txt",
             "Berater: Was braucht ihr von uns?\n"
             "Partner: Wir brauchen eine Landing Zone bis Mai, dann die Abnahme.\n"
             "Berater: Und der Betrieb?\n"
             "Partner: Den übernehmen wir selbst, mit Cloud Eye.\n")),
    case("speaker-legend", "txt, Speaker N labels with a legend line Speaker N = Name",
         [P2, P3_FIRST], ["Speaker 1", "Speaker 2", "Enterprise Router"],
         txt("transcript.txt",
             "Speaker 1 = %s\nSpeaker 2 = %s\n\n"
             "Speaker 1: Wir starten mit dem Enterprise Router.\n"
             "Speaker 2: Danke, dann bestelle ich.\n"
             "Speaker 1: Gut so.\n" % (P2, P3_FIRST)),
         expect=persons(P2, P3_FIRST)),

    # ------------------------------------------------------------------ names inside the speech
    case("speech-surname-after-run", "txt, a full name only inside the speech, then the surname alone",
         [P2, P2_LAST], ["Enterprise Router"],
         txt("transcript.txt",
             "%s: Ich habe gestern mit %s gesprochen, er baut den Enterprise Router.\n"
             "%s: Gut so. %s meldet sich dann bei mir.\n"
             "%s: Ja.\n" % (P3, P2, PLANTED_PERSON, P2_LAST, P3)),
         expect=persons(P2)),
    case("speech-nickname", "txt, the speaker's nickname inside another speaker's speech",
         ["Tobi"], ["NAT Gateway"],
         txt("transcript.txt",
             "%s: Hallo zusammen, der NAT Gateway ist da.\n"
             "%s: Danke Tobi, dann teste ich.\n"
             "%s: Gern.\n" % (P2, P3, P2))),
    case("speech-first-name-only", "txt, a first name that never speaks, capitalised inside the speech",
         ["Markus"], ["Firewall"],
         txt("transcript.txt",
             "%s: Markus hat die Firewall gebaut, er kennt die Regeln.\n"
             "%s: Dann frag Markus.\n" % (P2, P3))),

    # ------------------------------------------------------------------ terms inside the speech
    case("speech-product-unknown", "txt, product names whose words are not in the public corpus, spoken",
         [P2, P3], ["Paessler PRTG", "Veeam", "Baramundi Management Suite"],
         txt("transcript.txt",
             "%s: Wir monitoren mit Paessler PRTG und sichern mit Veeam.\n"
             "%s: Und die Clients mit Baramundi Management Suite?\n"
             "%s: Ja.\n" % (P2, P3, P2)),
         expect=persons(P2, P3)),
    case("speech-azure-win", "txt, Azure product names of known words, spoken",
         [P2, P3], ["Azure Virtual Desktop", "Site Recovery", "Log Analytics", "Cloud Eye"],
         txt("transcript.txt",
             "%s: Wir kommen von Azure Virtual Desktop und Site Recovery.\n"
             "%s: Und Log Analytics bleibt in Azure?\n"
             "%s: Nein, das geht nach Cloud Eye.\n" % (P2, P3, P2)),
         expect=persons(P2, P3)),
    case("speech-sql-inverted-win", "txt, a SQL statement with Name, Vorname spoken in the transcript",
         [P2, P3], ["SELECT", "FROM", "WHERE", "Vorname", "Kunden"],
         txt("transcript.txt",
             "%s: Die Abfrage ist SELECT Name, Vorname FROM Kunden WHERE Ort = '%s'.\n"
             "%s: Dann fehlt der Index auf Ort.\n" % (P2, CONTROL, P3)),
         expect=persons(P2, P3)),
    case("md-table-speaker", "md, a transcript table with the columns Zeit, Sprecher, Aussage",
         [P2_FIRST, P3], ["Elastic Cloud Server", "Sprecher", "Aussage"],
         txt("transcript.md",
             "| Zeit | Sprecher | Aussage |\n| --- | --- | --- |\n"
             "| 10:03 | %s | Der Elastic Cloud Server ist bestellt. |\n"
             "| 10:04 | %s | Danke %s. |\n"
             "| 10:05 | %s | Gern. |\n" % (P2_FIRST, P3, P2_FIRST, P2_FIRST)),
         expect=persons(P2_FIRST, P3)),
    case("docx-table-speaker", "docx, a transcript table with the columns Zeit, Sprecher, Aussage",
         [P2_FIRST, P3], ["Elastic Cloud Server", "Sprecher", "Aussage"],
         docx_table("protokoll.docx", [
             ["Zeit", "Sprecher", "Aussage"],
             ["10:03", P2_FIRST, "Der Elastic Cloud Server ist bestellt."],
             ["10:04", P3, "Danke %s." % P2_FIRST],
             ["10:05", P2_FIRST, "Gern."]], intro="Gesprächsprotokoll"),
         expect=persons(P2_FIRST, P3)),
    case("teilnehmer-known-first", "txt, the attendee line Teilnehmer: with first names, one a word of the corpus",
         [P2_FIRST, "Peter", P3], ["Cloud Eye"],
         txt("protokoll.txt",
             "Teilnehmer: %s, Peter, %s\n\n"
             "%s: Wir nehmen Cloud Eye.\n"
             "%s: Einverstanden.\n" % (P2_FIRST, P3, P2, P3)),
         expect=persons(P2_FIRST, "Peter", P3)),
    case("owner-single-name", "txt, action items with Owner: last on the line and a first name that never speaks",
         [P3_FIRST, P2], ["Action:", "Owner:", "Due:"],
         txt("protokoll.txt",
             "%s: Dann halten wir fest.\n"
             "Action: Angebot schicken, Due: Freitag, Owner: %s\n"
             "Action: Flavor prüfen, Due: Montag, Owner: %s\n" % (P2, P3_FIRST, P2)),
         expect=persons(P3_FIRST, P2)),
]
