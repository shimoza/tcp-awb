"""Red team of wipe mode, dimension german: German prose and documents.

Declensions and genitives, hyphen compounds and glued compounds, adjectival forms of a surname, title chains
(Prof. Dr. med., Dipl.-Ing., Dipl.-Kfm., Hr., Fr.), salutations and closings, z. Hd. and i. A., Firma and Fa.,
legal footers, tender documents with Ansprechpartner lines, address blocks and the Ort, Datum line, a surname
that is also a German word, a first name alone in prose, ALL CAPS names in headings; losses of German headings,
tender vocabulary, the partner's legal names, laws and standards and German product names.

Every name is from tests/fixtures.py or from INVENTED below. Hosts end in .example. Nothing here is real.
"""
import sys
from email import policy
from email.message import EmailMessage
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[3]))
from tests import fixtures as fx  # noqa: E402

# the fixture forms
CUSTOMER, CUSTOMER_SHORT, CUSTOMER_ACRONYM, CUSTOMER_EN = fx.CUSTOMER_FORMS
PERSON, SURNAME = fx.PERSON_FORMS            # a registered person (becomes a PERS code, not a token)
FIRST = fx.PLANTED_PERSON.split()[0]         # Xqarv
LAST = fx.PLANTED_PERSON.split()[1]          # Pomblet
PLANTED_PERSON = fx.PLANTED_PERSON           # Xqarv Pomblet, unregistered
PLANTED_COMPANY = fx.PLANTED_CANDIDATE       # Nrgtz Beratung GmbH, unregistered
BRAND = PLANTED_COMPANY.split()[0]           # Nrgtz
CONTROL = fx.CONTROL_UNREGISTERED            # Musterhausen, unregistered, a place word

# every invented name of this module, so that the owner can see nothing real slipped in
INVENTED = (
    "Anna Zrbl",                             # an invented person, four-letter surname; appears as "Zrbl" and "Frau Zrbl"
    "Klein",                                 # a German dictionary word used as a bare surname ("Frau Klein"), never with a first name
    "Lang",                                  # a German dictionary word used as a bare surname ("Herr Lang"), never with a first name
    "Thomas",                                # a common first name alone
    "Peter",                                 # a common first name alone, as the brief allows
    "Tobias",                                # a common first name alone
    "Musterstraße 12",                       # an invented street
    "MUSTERHAUSEN",                          # the unregistered control place in capitals
    "nrgtz-beratung.example",                # the mail host of the planted company
    "x.pomblet@nrgtz-beratung.example",      # the planted person's address
    "t." + fx.PERSON_FORMS[1].lower() + "@" + fx.CUSTOMER_DOMAIN,  # the registered person's address at the fixture domain
    "+49 555 1234567",                       # an invented phone number
    "HRB 12345",                             # an invented register number
    "DE123456789",                           # an invented VAT id
)

MARK = "stage two of the invented plan"

ZRBL = "Zrbl"
KLEIN = "Klein"
LANG = "Lang"
STREET = "Musterstraße 12"
TOWN_CAPS = "MUSTERHAUSEN"
HOST1 = "nrgtz-beratung.example"
MAIL1 = "x.pomblet@" + HOST1
MAIL_REG = "t." + fx.PERSON_FORMS[1].lower() + "@" + fx.CUSTOMER_DOMAIN
PHONE = "+49 555 1234567"
HRB = "HRB 12345"
VAT = "DE123456789"


def txt(name, text):
    """A plain file with `text`; the marker line proves the carrier was read."""
    def build(inbox):
        p = inbox / name
        p.write_text(text + "\n" + MARK + "\n", encoding="utf-8")
        return p
    return build


def raw(name, data: bytes):
    def build(inbox):
        p = inbox / name
        p.write_bytes(data)
        return p
    return build


def docx_paragraphs(name, lines):
    def build(inbox):
        import docx
        d = docx.Document()
        for line in lines:
            d.add_paragraph(line)
        d.add_paragraph(MARK)
        p = inbox / name
        d.save(str(p))
        return p
    return build


def xlsx_rows(name, rows, sheet="Kontakte"):
    def build(inbox):
        import openpyxl
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = sheet
        for row in rows:
            ws.append(list(row))
        ws.append([MARK])
        p = inbox / name
        wb.save(str(p))
        return p
    return build


def pptx_slides(name, slides):
    """One slide per entry; every entry is a list of text boxes (the first is the title)."""
    def build(inbox):
        import pptx
        from pptx.util import Inches
        prs = pptx.Presentation()
        for texts in list(slides) + [[MARK]]:
            slide = prs.slides.add_slide(prs.slide_layouts[6])
            for i, t in enumerate(texts):
                box = slide.shapes.add_textbox(Inches(1), Inches(1 + i), Inches(8), Inches(0.8))
                box.text_frame.text = t
        p = inbox / name
        prs.save(str(p))
        return p
    return build


def pdf_letter(name, left_lines, right_lines=(), body_lines=()):
    """A one-page letter: an address block on the left, a reference block on the right, a body below."""
    def build(inbox):
        from reportlab.lib.pagesizes import A4
        from reportlab.pdfgen import canvas
        p = inbox / name
        c = canvas.Canvas(str(p), pagesize=A4)
        c.setFont("Helvetica", 11)
        y = 780
        for line in left_lines:
            c.drawString(60, y, line)
            y -= 16
        y = 780
        for line in right_lines:
            c.drawString(360, y, line)
            y -= 16
        y = 640
        for line in body_lines:
            c.drawString(60, y, line)
            y -= 16
        c.drawString(60, 100, MARK)
        c.save()
        return p
    return build


def eml(name, headers, body):
    def build(inbox):
        msg = EmailMessage(policy=policy.default)
        for k, v in headers:
            msg[k] = v
        msg.set_content(body + "\n" + MARK + "\n")
        p = inbox / name
        p.write_bytes(msg.as_bytes())
        return p
    return build


def case(cid, carrier, values, keep, build, note="", code=None, expect=None):
    c = {"id": cid, "carrier": carrier, "values": list(values), "keep": list(keep), "build": build, "note": note}
    if code is not None:
        c["code"] = code
    if expect:
        c["expect"] = dict(expect)
    return c


ECS = "Elastic Cloud Server"
ELB = "Elastic Load Balancer"
GREETING_REG = "Hallo Herr %s,\n\nder %s ist bestellt.\n\n" % (SURNAME, ECS)

CASES = [
    # ------------------------------------------------------------------ titles and title chains
    case("de-title-chain-prof-dr-med", "txt, Prof. Dr. med. before a full name in prose",
         [PLANTED_PERSON], ["Relational Database Service"],
         txt("gutachten.txt", "Sehr geehrte Damen und Herren,\n\nfür das Gutachten ist Prof. Dr. med. %s zuständig. "
             "Die Daten liegen im Relational Database Service.\n" % PLANTED_PERSON),
         expect={PLANTED_PERSON: "person"},
         note="med. is no title word, so the title rule stops; the run rule takes the name as [name]"),
    case("de-title-dipl-kfm", "txt, Dipl.-Kfm. before a bare surname in prose",
         [LAST], ["Cloud Eye"],
         txt("bewertung.txt", "Die kaufmännische Bewertung übernimmt Dipl.-Kfm. %s, die technische Prüfung läuft "
             "über Cloud Eye.\n" % LAST),
         expect={LAST: "person"}),
    case("de-title-dipl-ing-prose", "txt, Dipl.-Ing. before a bare surname in prose",
         [LAST], [ELB],
         txt("netzplanung.txt", "Die Netzplanung verantwortet Dipl.-Ing. %s, der %s wird von ihm abgenommen.\n"
             % (LAST, ELB)),
         expect={LAST: "person"},
         note="proposed: Ing is a title word and the title cut splits on the hyphen; baseline has no Ing"),
    case("de-title-dipl-ing-salutation", "txt, Sehr geehrter Herr Dipl.-Ing. Surname, as the salutation",
         [LAST], [ELB],
         txt("mail.txt", "Sehr geehrter Herr Dipl.-Ing. %s,\n\nder %s ist abgenommen.\n" % (LAST, ELB)),
         expect={LAST: "person"},
         note="the salutation value ends at the first dot (Herr Dipl); the title rule takes it from Ing"),
    case("de-title-hr-fr-prose", "txt, the abbreviations Hr. and Fr. before bare surnames",
         [LAST, ZRBL], ["Direct Connect"],
         txt("termin.txt", "An dem Termin nahmen Hr. %s und Fr. %s teil. Der Direct Connect wurde freigegeben.\n"
             % (LAST, ZRBL)),
         expect={LAST: "person", ZRBL: "person"}),
    case("de-title-role-projektleiterin", "txt, Frau + role word + surname in prose",
         [ZRBL], ["Projektleiterin", ELB],
         txt("protokoll.txt", "Frau Projektleiterin %s hat den Zeitplan bestätigt, der %s wird in Phase 2 migriert.\n"
             % (ZRBL, ELB)),
         expect={ZRBL: "person"},
         note="the title rule takes the role word into the person"),

    # ------------------------------------------------------------------ address blocks and letters
    case("de-address-caps-town-pdf", "pdf, address block with the town in capitals after the postcode",
         [TOWN_CAPS], ["Cloud Container Engine"],
         pdf_letter("brief.pdf", [PLANTED_COMPANY, "Herrn " + PLANTED_PERSON, STREET, "12345 " + TOWN_CAPS],
                    body_lines=["Sehr geehrter Herr %s," % LAST, "", "der Cloud Container Engine Cluster ist bestellt."]),
         expect={TOWN_CAPS: "place"},
         note="proposed withholds the output: the file name brief.pdf is read as the closing br plus a signer ief.pdf, "
              "the learned pdf then hits the output header; the caps town itself is tested in de-caps-town-txt"),
    case("de-ort-datum", "txt, the Ort, den Datum line that opens a German letter",
         [CONTROL], [ECS],
         txt("brief.txt", "%s, den 5. Oktober 2026\n\nSehr geehrte Damen und Herren,\n\ndas Angebot für den %s liegt bei.\n"
             % (CONTROL, ECS)),
         expect={CONTROL: "place"}),
    case("de-pdf-layout-letter", "pdf, two-column letter head: address block left, reference block right",
         [PLANTED_COMPANY, PLANTED_PERSON, STREET, CONTROL], [ECS, "Angebot Nr. 2026-0811"],
         pdf_letter("angebot.pdf", [PLANTED_COMPANY, "Herrn " + PLANTED_PERSON, STREET, "12345 " + CONTROL],
                    ["Angebot Nr. 2026-0811", "Datum: 05.10.2026", "Ihr Zeichen: XP"],
                    ["Sehr geehrter Herr %s," % LAST, "", "anbei das Angebot für den %s." % ECS]),
         expect={PLANTED_COMPANY: "company", PLANTED_PERSON: "person", STREET: "place", CONTROL: "place"}),

    # ------------------------------------------------------------------ Firma, genitives, compounds, declensions
    case("de-firma-fa", "txt, Firma X and Fa. X in prose",
         [BRAND], ["NAT Gateway"],
         txt("auftrag.txt", "Der Auftrag ging an die Firma %s, die den NAT Gateway betreibt. Fa. %s liefert auch die Hardware.\n"
             % (BRAND, BRAND)),
         expect={BRAND: "company"}),
    case("de-genitive-apostrophe", "txt, the genitive apostrophe after a company learned from Firma X",
         [BRAND], [ELB],
         txt("angebot.txt", "Die Firma %s hat ein Angebot vorgelegt. %s' Angebot umfasst den %s, und %s’ Preise gelten "
             "bis Dezember.\n" % (BRAND, BRAND, ELB, BRAND)),
         expect={BRAND: "company"}),
    case("de-compound-hyphen-lower", "txt, company-seitig and company-intern, no strong shape in the file",
         [BRAND], [ECS],
         txt("notiz.txt", "Die %s-seitige Umsetzung beginnt im November, %s-intern ist der %s bereits bestellt.\n"
             % (BRAND, BRAND, ECS)),
         expect={BRAND: "company"},
         note="a hyphen compound with a lower-case second part: the identifier rule needs two capitalised parts"),
    case("de-compound-hyphen-cap", "txt, Nrgtz-Lösung and Pomblet-Konzept, no strong shape in the file",
         [BRAND, LAST], [ELB],
         txt("notiz.txt", "Die %s-Lösung ersetzt den alten Proxy, das %s-Konzept bleibt die Grundlage für den %s.\n"
             % (BRAND, LAST, ELB)),
         expect={BRAND: "company", LAST: "person"},
         note="the identifier rule skips a two-part hyphen compound as a word (Soll-Konzept), the run rule never sees glued words"),
    case("de-compound-glued-short", "txt, a five-letter company glued into a German compound (Nrgtzlösung)",
         [BRAND], [ECS],
         txt("notiz.txt", "Die Firma %s liefert die %slösung bis Dezember, der %s wird dafür vergrößert.\n"
             % (BRAND, BRAND, ECS)),
         expect={BRAND: "company"},
         note="learned forms inside a plain word are matched from seven letters on"),
    case("de-learned-declensions", "txt, genitive -s and the adjectival -sche forms of a surname learned from Herr X",
         [LAST], [ELB],
         txt("mail.txt", "Sehr geehrter Herr %s,\n\nwir treffen uns in %ss neuem Büro. Die %s'sche Variante und die "
             "%ssche Variante unterscheiden sich nur im %s; die %s’schen Annahmen gelten weiter.\n"
             % (LAST, LAST, LAST, LAST, ELB, LAST)),
         expect={LAST: "person"}),

    # ------------------------------------------------------------------ first names alone
    case("de-first-name-list", "txt, a first name of the list alone at the start of a sentence",
         ["Tobias"], [ELB],
         txt("notiz.txt", "Tobias hat die Firewall-Regeln angepasst, danach lief der %s wieder.\n" % ELB),
         expect={"Tobias": "person"}),
    case("de-first-name-vocab", "txt, a first name that the public vocabulary carries, alone in prose",
         ["Peter"], [ELB],
         txt("notiz.txt", "Peter hat die Firewall-Regeln angepasst, danach lief der %s wieder.\n" % ELB),
         expect={"Peter": "person"},
         note="the lone first name rule uses _known_word, which does not exempt first names"),

    # ------------------------------------------------------------------ headings, tender words, vendors, standards, products
    case("de-headings-known-pptx", "pptx, German slide titles of known words, capitals and mixed case",
         [], ["TECHNISCHE ANFORDERUNGEN", "ALLGEMEINE GESCHÄFTSBEDINGUNGEN", "Netzwerk Architektur", "Kosten Übersicht"],
         pptx_slides("angebot.pptx", [["TECHNISCHE ANFORDERUNGEN", "%s mit 8 vCPU" % ECS],
                                      ["ALLGEMEINE GESCHÄFTSBEDINGUNGEN", "Laufzeit 36 Monate"],
                                      ["Netzwerk Architektur", "Kosten Übersicht"]])),
    case("de-heading-compound-unknown", "md, two-word German headings with compound nouns outside the lists",
         [], ["Leistungsverzeichnis Los 1", "Betriebskonzept Monitoring", "Firewall Regelwerk", "Backup Konzept"],
         txt("lv.md", "# Leistungsverzeichnis Los 1\n\n## Betriebskonzept Monitoring\n\n- Firewall Regelwerk\n- Backup Konzept\n\n"
             "Der %s wird in Los 1 ausgeschrieben.\n" % ECS),
         note="Backup Konzept is the control of known words"),
    case("de-tender-verfahren", "txt, the names of the procurement procedures (adjective + noun)",
         [], ["Offenes Verfahren", "Beschränkte Ausschreibung", "Freihändige Vergabe", "Wettbewerblicher Dialog"],
         txt("vergabe.txt", "Vergabeart: Offenes Verfahren nach VgV. Alternativ kommen eine Beschränkte Ausschreibung, eine "
             "Freihändige Vergabe oder ein Wettbewerblicher Dialog in Betracht. Der %s gehört zu Los 2.\n" % ECS)),
    case("de-vendor-legal-form", "md, the partner's legal names with GmbH and AG in prose",
         [], ["T-Systems International GmbH", "Telekom Deutschland GmbH", "Deutsche Telekom AG"],
         txt("vertrag.md", "Der Vertrag wird mit der T-Systems International GmbH geschlossen, die Anbindung liefert die "
             "Telekom Deutschland GmbH, Konzernmutter ist die Deutsche Telekom AG.\n"),
         note="Deutsche Telekom is a stop phrase, the other two are not"),
    case("de-dsgvo-art", "txt, laws and standards: DSGVO Art. 28, BSI C5, ISO 27001, the BSI's full name",
         [], ["DSGVO Art. 28", "BSI C5", "ISO 27001", "Bundesamt für Sicherheit in der Informationstechnik"],
         txt("compliance.txt", "Die Auftragsverarbeitung richtet sich nach DSGVO Art. 28, die Nachweise nach BSI C5 und ISO 27001 "
             "liegen vor; das Bundesamt für Sicherheit in der Informationstechnik hat das Testat anerkannt. Der %s ist "
             "C5-testiert.\n" % ECS)),
    case("de-product-names", "txt, German product names of two words: DATEV Unternehmen online, ELSTER Online",
         [], ["DATEV Unternehmen online", "ELSTER Online", "Object Storage Service"],
         txt("buchhaltung.txt", "Die Buchhaltung läuft in DATEV Unternehmen online, die Steuermeldungen gehen über ELSTER Online, "
             "die Belege liegen im Object Storage Service.\n"),
         note="DATEV is a vendor word of the lists, ELSTER is not"),
    case("de-role-before-surname", "txt, a role word directly before a bare surname at the start of a sentence",
         [LAST], ["Projektleiter", ELB],
         txt("protokoll.txt", "Projektleiter %s hat den Zeitplan bestätigt, der %s wird in Phase 2 migriert.\n" % (LAST, ELB)),
         expect={LAST: "person"},
         note="the run rule wipes the role word with the name"),
    case("de-heading-swallow", "md, a heading of two known words and one unknown company word",
         [BRAND, LAST], ["Protokoll Abstimmung", ECS],
         txt("protokoll.md", "# Protokoll Abstimmung %s\n\nTeilnehmer: Herr %s\n\nDer %s wird am Montag bestellt.\n"
             % (BRAND, LAST, ECS)),
         expect={BRAND: "company", LAST: "person"},
         note="the run rule wipes the known heading words with the unknown one"),

    # ------------------------------------------------------------------ labels, tables, footers, signatures
    case("de-teilnehmer-paren", "txt, Teilnehmer: Surname (Company), Surname (Company) with parentheses",
         [LAST, ZRBL, BRAND], [ECS],
         txt("protokoll.txt", "Teilnehmer: %s (%s), %s (%s), %s (%s)\n\nDer %s wird am Montag bestellt.\n"
             % (LAST, BRAND, ZRBL, CUSTOMER_SHORT, SURNAME, CUSTOMER_SHORT, ECS)),
         code=True, expect={LAST: "person", ZRBL: "person", BRAND: "company"},
         note="a label item is cut at the parenthesis; the company inside the parentheses is left as it stands"),
    case("de-table-xlsx-name", "xlsx, a contact sheet with the columns Name, Firma, Telefon and bare surnames",
         [LAST, ZRBL, BRAND], ["Telefon"],
         xlsx_rows("kontakte.xlsx", [["Name", "Firma", "Telefon"], [LAST, BRAND, "0555 1234567"],
                                     [ZRBL, CUSTOMER_SHORT, "0555 7654321"], [SURNAME, CUSTOMER_SHORT, "0555 1111111"]]),
         code=True, expect={LAST: "person", ZRBL: "person", BRAND: "company"}),
    case("de-legal-footer-eml", "eml, the legal footer of a German mail signature with Geschäftsführer, Sitz, Registergericht",
         [PLANTED_COMPANY, PLANTED_PERSON, CONTROL], ["Amtsgericht", HRB, "USt-IdNr."],
         eml("mail.eml", [("From", "%s <%s>" % (PLANTED_PERSON, MAIL1)), ("To", MAIL_REG), ("Subject", "Angebot " + ECS)],
             GREETING_REG + "Mit freundlichen Grüßen\n%s\n\n--\n%s | %s | 12345 %s\nGeschäftsführer: %s | Sitz der Gesellschaft: %s | "
             "Registergericht: Amtsgericht %s, %s | USt-IdNr.: %s\n"
             % (PLANTED_PERSON, PLANTED_COMPANY, STREET, CONTROL, PLANTED_PERSON, CONTROL, CONTROL, HRB, VAT)),
         expect={PLANTED_COMPANY: "company", PLANTED_PERSON: "person", CONTROL: "place"}),
    case("de-signature-roles-ia", "txt, a signature block: i. A. before the signer, two German role lines, the company",
         [PLANTED_PERSON, PLANTED_COMPANY], ["Senior Cloud Architect", "Leiter Rechenzentrum"],
         txt("mail.txt", GREETING_REG + "Mit freundlichen Grüßen\n\ni. A. %s\nSenior Cloud Architect\nLeiter Rechenzentrum\n%s\nTel. %s\n"
             % (PLANTED_PERSON, PLANTED_COMPANY, PHONE)),
         expect={PLANTED_PERSON: "person", PLANTED_COMPANY: "company"}),
    case("de-closing-trailing", "txt, a closing with words after it on the same line, the signer on the next line",
         [PLANTED_PERSON], [ECS],
         txt("mail.txt", GREETING_REG + "Mit freundlichen Grüßen und ein schönes Wochenende\n%s\n" % PLANTED_PERSON),
         expect={PLANTED_PERSON: "person"},
         note="the sign-off rule takes the trailing words as the signer and never looks at the next line"),
    case("de-tender-cover-docx", "docx, a tender cover with Ansprechpartner lines, Auskünfte erteilt, Angebotsfrist",
         [ZRBL, LAST], ["Zentrale Vergabestelle", "Angebotsfrist", "Zuschlagskriterien", ECS],
         docx_paragraphs("vergabe.docx", ["Vergabeunterlagen", "Auftraggeber: " + CUSTOMER, "Vergabestelle: Zentrale Vergabestelle",
                                          "Technischer Ansprechpartner: " + ZRBL, "Kaufmännischer Ansprechpartner: " + LAST,
                                          "Auskünfte erteilt: Frau %s, Tel. %s" % (ZRBL, PHONE),
                                          "Angebotsfrist: 05.11.2026, 12:00 Uhr", "Zuschlagskriterien: Preis 70 %, Qualität 30 %",
                                          "Der %s wird in Los 1 ausgeschrieben." % ECS]),
         code=True, expect={ZRBL: "person", LAST: "person"}),

    # ------------------------------------------------------------------ capitals, particles, learned forms, surname words
    case("de-caps-heading-run", "txt, a heading in capitals with the full name, the surname alone in mixed case below",
         [LAST], [ECS],
         txt("angebot.txt", "ANGEBOT %s\n\n%s erhält das Angebot für den %s bis Freitag.\n"
             % (PLANTED_PERSON.upper(), LAST, ECS)),
         expect={LAST: "person"},
         note="a capitalised run is [name] and its words are not learned"),
    case("de-particle-known-first", "txt, Firewall von Nrgtz: a known noun before von and an unknown company after it",
         [BRAND], ["Firewall", ELB],
         txt("notiz.txt", "Die Firewall von %s wird abgelöst, der neue %s übernimmt die Regeln.\n" % (BRAND, ELB)),
         expect={BRAND: "company"},
         note="baseline wipes the noun as part of a person, proposed skips the pair and leaves the company"),
    case("de-learned-standalone", "txt, Herr X once, then the surname alone at a sentence start, mid-sentence and before a colon",
         [LAST], [ECS],
         txt("mail.txt", "Sehr geehrter Herr %s,\n\n%s hatte gefragt, ob der %s reicht. Wie mit %s besprochen, prüfen wir das; "
             "%s: bitte Rückmeldung bis Freitag.\n" % (LAST, LAST, ECS, LAST, LAST)),
         expect={LAST: "person"}),
    case("de-surname-klein", "txt, Frau Klein, then Klein alone at a sentence start, and the adjective klein in the text",
         [KLEIN], ["zu klein", ECS],
         txt("mail.txt", "Sehr geehrte Frau %s,\n\n%s hatte angemerkt, dass der %s zu klein dimensioniert ist.\n"
             % (KLEIN, KLEIN, ECS)),
         expect={KLEIN: "person"},
         note="a surname learned from a title is strong: the adjective in lower case is wiped with it"),
    case("de-surname-lang", "txt, Herr Lang, then Lang alone at a sentence start; lang is a word of the public corpus",
         [LANG], [ECS],
         txt("mail.txt", "Sehr geehrter Herr %s,\n\n%s hatte angemerkt, dass der %s reicht.\n" % (LANG, LANG, ECS)),
         expect={LANG: "person"},
         note="a surname that the corpus carries is never learned, so it leaks where it stands alone"),
    case("de-lone-after-preposition", "txt, a bare company and a bare surname after mit and an, no strong shape in the file",
         [BRAND, LAST], [ECS],
         txt("notiz.txt", "Die Abstimmung mit %s ist erfolgt, Rückfragen bitte an %s. Der %s wird in Los 1 ausgeschrieben.\n"
             % (BRAND, LAST, ECS)),
         expect={BRAND: "company", LAST: "person"},
         note="a lone unknown capitalised word in German prose"),

    # ------------------------------------------------------------------ run 2: variations of the mechanisms found
    case("de-closing-prefix-lines", "txt, German words at a line start that begin with a sign-off abbreviation (Best, Br, Kr, Vg)",
         [], ["Bestellung Nr. 4711", "Kriterien der Wertung", "Vgl. Anlage 3", "Bestätigung des Auftrags",
              "Breite der Anbindung", ECS],
         txt("bestellung.txt", "Bestellung Nr. 4711\nKriterien der Wertung\nVgl. Anlage 3\nBestätigung des Auftrags\n"
             "Breite der Anbindung: 1 Gbit/s\n\nDer %s wird nach Bestellung bereitgestellt.\n" % ECS),
         note="the closing regex has no word boundary: Best + ellung is a sign-off with a signer"),
    case("de-filename-closing-prefix-pdf", "pdf, a file named Bestellung.pdf with a plain German body",
         [], [ECS],
         pdf_letter("Bestellung.pdf", ["Bestellung %s" % ECS], (),
                    ["Sehr geehrte Damen und Herren,", "", "wir bestellen den %s mit 8 vCPU." % ECS]),
         note="the file name is read as the closing best + the signer ellung.pdf; the learned pdf hits the output header"),
    case("de-filename-closing-prefix-txt", "txt, a file named Bestätigung.txt, the same word mid-sentence in the body",
         [], ["Bestätigung", ECS],
         txt("Bestätigung.txt", "Die Bestätigung des Auftrags liegt vor, der %s ist bestellt.\n" % ECS),
         note="the file name teaches the signer ätigung, seven letters, matched inside the body word"),
    case("de-first-names-vocab-more", "txt, two first names the corpus carries in two files or more, alone in prose",
         ["Thomas", "Anna"], [ELB],
         txt("notiz.txt", "Thomas prüft die Anbindung, Anna testet den %s.\n" % ELB),
         expect={"Thomas": "person", "Anna": "person"},
         note="VOCAB_MIN_FILES is 2 now, so Thomas and Anna are known words for the lone first name rule"),
    case("de-heading-learns-los", "md, a heading Leistungsverzeichnis Los 1 teaches Los as a name word",
         [], ["Los 2 die Datenbanken", ECS],
         txt("lv.md", "# Leistungsverzeichnis Los 1\n\nLos 1 umfasst den %s, Los 2 die Datenbanken.\n" % ECS),
         note="the person-like run learns both words; the learned Los is wiped wherever it stands"),
    case("de-caps-town-txt", "txt, address block with the town in capitals after the postcode",
         [TOWN_CAPS], [ECS],
         txt("adresse.txt", "%s\n%s\n12345 %s\n\nDer %s ist bestellt.\n" % (PLANTED_COMPANY, STREET, TOWN_CAPS, ECS)),
         expect={TOWN_CAPS: "place"},
         note="the postcode rule wants a capitalised word with small letters after it"),
    case("de-compound-hyphen-cap-learned", "txt, Nrgtz-Lösung and Pomblet-Konzept after Firma X and Herr Y in the same file",
         [BRAND, LAST], [ELB],
         txt("notiz.txt", "Die Firma %s liefert die %s-Lösung, Herr %s vertritt das %s-Konzept für den %s.\n"
             % (BRAND, BRAND, LAST, LAST, ELB)),
         expect={BRAND: "company", LAST: "person"},
         note="with a strong shape in the file the learned forms reach into the compounds"),
    case("de-stadt-x", "txt, Stadt + unknown town as the public-sector customer",
         ["Stadt " + CONTROL], ["Amt für Digitalisierung", ECS],
         txt("vergabe.txt", "Auftraggeberin ist die Stadt %s, vertreten durch das Amt für Digitalisierung. Der %s wird in "
             "Los 1 ausgeschrieben.\n" % (CONTROL, ECS)),
         expect={"Stadt " + CONTROL: "company"},
         note="Stadt, Landkreis, Gemeinde before an unknown word name an organisation; the run rule gives [name]"),
]
