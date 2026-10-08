"""Red team of wipe mode, dimension tables: tables in docx, xlsx, csv, tsv, Markdown and plain text.

A name split over two cells (first | last), contact lists with headers in German, English and Russian
(Name, Vorname, Nachname, E-Mail, Telefon, Firma, Rolle, Standort, Ansprechpartner, First Name, Last Name,
ФИО, Фамилия, Имя, Отчество), participant and sign-in lists, RACI matrices with initials and with names,
asset and host lists with host names and owners, a company split over cells (brand | legal form), key-value
tables (a label in the first column, the value in the second), tables without a header row, a title row above
the header, a table inside a docx, an xlsx with several sheets, a csv with semicolons, a layout table of
spaces. Losses: product and price tables (service | flavor | price), service mapping tables (Azure, GCP and
VMware products against TCP services), flavor tables, column headers that are capitalised phrases, SQL in a
cell, host and application names under a "Name" column, vendors under "Hersteller", teams under "Owner".

Every name is from tests/fixtures.py or from INVENTED below. Hosts end in .example. Nothing here is real.
"""
import csv
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[3]))
from tests import fixtures as fx  # noqa: E402

# the fixture forms
CUSTOMER, CUSTOMER_SHORT, CUSTOMER_ACRONYM, CUSTOMER_EN = fx.CUSTOMER_FORMS
PERSON, SURNAME = fx.PERSON_FORMS            # a registered person (becomes a PERS code, not a token)
FIRST = fx.PLANTED_PERSON.split()[0]         # Xqarv
LAST = fx.PLANTED_PERSON.split()[1]          # Pomblet
PLANTED_PERSON = fx.PLANTED_PERSON           # Xqarv Pomblet, unregistered
PLANTED_COMPANY = fx.PLANTED_CANDIDATE       # Nrgtz Beratung GmbH, unregistered
BRAND = PLANTED_COMPANY.split()[0]           # Nrgtz
CONTROL = fx.CONTROL_UNREGISTERED            # Musterhausen, unregistered, a place word
ORG, ORG_SHORT = fx.ORG_FORMS
PLACE = fx.PLACE_FORMS[0]

# every invented name of this module, so that the owner can see nothing real slipped in
INVENTED = (
    "Anna Zrbl", "Zrbl",                       # an invented person, four-letter surname, first name of the corpus
    "Timo Zrbko", "Zrbko",                     # an invented person, first name of the list only
    "Max Qwertzung", "Qwertzung",              # an invented surname with a German noun suffix
    "Jan Qwertzuio", "Qwertzuio",              # an invented surname, first name that is a month abbreviation
    "Zrbl Systemhaus", "Zrbl Systemhaus GmbH", # an invented company, brand plus a generic word
    "Qwertzuio Logistik", "Qwertzuio Logistik KG",   # a second invented company
    "tpomblet",                                # a login made of an initial and the planted surname
    "Кхжтов Иван Петрович", "Кхжтов", "Петрович",   # an invented Russian person with a patronymic
    "Зрбл Анна Сергеевна", "Зрбл", "Сергеевна",     # the second, the Cyrillic Zrbl
    "Кхжтов Иван", "Зрбл Анна",
    "Нргтц", "ООО «Нргтц»", "Квёртц", "АО «Квёртц»",  # the brands in Cyrillic
    "Бейшпильштадт", "Мустерхаузен",           # the two places in Cyrillic
    "nrgtz.example", "x.pomblet@nrgtz.example", "anna.zrbl@nrgtz.example", "info@nrgtz.example",
    "srv-pomblet-02", "srv-pomblet-db", "srv-zrbl-01", "srv-zrbl-web", "nrgtz-fw01",
    "203.0.113.5", "203.0.113.6", "203.0.113.7", "203.0.113.9", "+49 555 0113 42", "+49 555 0113 43",
)

P2_FIRST, P2_LAST = "Anna", "Zrbl"
P2 = "%s %s" % (P2_FIRST, P2_LAST)
P3_FIRST, P3_LAST = "Timo", "Zrbko"
P3 = "%s %s" % (P3_FIRST, P3_LAST)
P4_LAST = "Qwertzung"
P5_LAST = "Qwertzuio"
COMPANY2 = "Zrbl Systemhaus"
COMPANY3 = "Qwertzuio Logistik"
LOGIN_T = "t" + LAST.lower()                 # tpomblet
RU_PERSON = "Кхжтов Иван Петрович"
RU_LAST, RU_FIRST, RU_PATRONYM = RU_PERSON.split()
RU_PERSON2 = "Зрбл Анна Сергеевна"
RU_LAST2, RU_FIRST2, RU_PATRONYM2 = RU_PERSON2.split()
RU_BRAND, RU_ORG = "Нргтц", "Квёртц"
RU_PLACE, RU_CONTROL = "Бейшпильштадт", "Мустерхаузен"
DOMAIN = BRAND.lower() + ".example"
MAIL_X = "%s.%s@%s" % (FIRST[0].lower(), LAST.lower(), DOMAIN)        # x.pomblet@nrgtz.example
MAIL_A = "%s.%s@%s" % (P2_FIRST.lower(), P2_LAST.lower(), DOMAIN)     # anna.zrbl@nrgtz.example
MAIL_INFO = "info@" + DOMAIN
HOST_P = "srv-%s-02" % LAST.lower()          # srv-pomblet-02
HOST_P_DB = "srv-%s-db" % LAST.lower()       # srv-pomblet-db
HOST_Z = "srv-%s-01" % P2_LAST.lower()       # srv-zrbl-01
HOST_Z_WEB = "srv-%s-web" % P2_LAST.lower()  # srv-zrbl-web
HOST_FW = BRAND.lower() + "-fw01"            # nrgtz-fw01
IP1, IP2, IP3, IP4 = "203.0.113.5", "203.0.113.6", "203.0.113.7", "203.0.113.9"
TEL1, TEL2 = "+49 555 0113 42", "+49 555 0113 43"

MARK = "stage two of the invented plan"


# --------------------------------------------------------------------------- builders

def txt(name, text, mark=True):
    """A plain, Markdown or tsv file with `text`; the marker line proves the carrier was read."""
    def build(inbox):
        p = inbox / name
        p.write_text(text + ("\n" + MARK + "\n" if mark else ""), encoding="utf-8")
        return p
    return build


def md_table(rows):
    """Markdown table lines from rows of cell strings, the way the extractors render tables."""
    width = max(len(r) for r in rows)
    lines = []
    for i, row in enumerate(rows):
        cells = list(row) + [""] * (width - len(row))
        lines.append("| " + " | ".join(cells) + " |")
        if i == 0:
            lines.append("|" + " --- |" * width)
    return "\n".join(lines) + "\n"


def csvf(name, rows, delim=","):
    def build(inbox):
        p = inbox / name
        with p.open("w", encoding="utf-8", newline="") as fh:
            w = csv.writer(fh, delimiter=delim)
            for r in rows:
                w.writerow(list(r))
        return p
    return build


def docx_table(name, rows, before=(), after=()):
    """A docx with paragraphs `before`, one table of `rows`, paragraphs `after` and the marker."""
    def build(inbox):
        from docx import Document
        doc = Document()
        for text in before:
            doc.add_paragraph(text)
        table = doc.add_table(rows=0, cols=max(len(r) for r in rows))
        for r in rows:
            cells = table.add_row().cells
            for i, value in enumerate(r):
                cells[i].text = value
        for text in after:
            doc.add_paragraph(text)
        doc.add_paragraph(MARK)
        p = inbox / name
        doc.save(str(p))
        return p
    return build


def docx_table_caption(name, caption, rows):
    """A docx table whose first row is one merged caption cell over every column, then the header and rows."""
    def build(inbox):
        from docx import Document
        doc = Document()
        width = max(len(r) for r in rows)
        table = doc.add_table(rows=0, cols=width)
        first = table.add_row().cells
        first[0].text = caption
        first[0].merge(first[width - 1])
        for r in rows:
            cells = table.add_row().cells
            for i, value in enumerate(r):
                cells[i].text = value
        doc.add_paragraph(MARK)
        p = inbox / name
        doc.save(str(p))
        return p
    return build


def xlsx_book(name, sheets):
    """An xlsx with one sheet per (title, rows); an empty row is written as an empty row."""
    def build(inbox):
        from openpyxl import Workbook
        wb = Workbook()
        first = True
        for title, rows in sheets:
            ws = wb.active if first else wb.create_sheet()
            first = False
            ws.title = title
            for r in rows:
                ws.append(list(r))
        p = inbox / name
        wb.save(str(p))
        return p
    return build


def case(cid, carrier, values, keep, build, note="", code=None, expect=None):
    c = {"id": cid, "carrier": carrier, "values": list(values), "keep": list(keep), "build": build, "note": note}
    if code is not None:
        c["code"] = code
    if expect:
        c["expect"] = dict(expect)
    return c


# --------------------------------------------------------------------------- the rows

CONTACTS_DE = [["Vorname", "Nachname", "Firma", "Rolle"],
               [FIRST, LAST, BRAND, "Architekt"],
               [P2_FIRST, P2_LAST, ORG, "PL"]]
CONTACTS_EN = [["First Name", "Last Name", "Company", "Signature"],
               [FIRST, LAST, BRAND, ""],
               [P2_FIRST, P2_LAST, ORG_SHORT, ""],
               [P3_FIRST, P3_LAST, "intern", ""]]
MAPPING_AZURE = [["Azure Service", "TCP Service", "Notes"],
                 ["Virtual Machines", "Elastic Cloud Server", "same sizes"],
                 ["Network Watcher", "Cloud Eye", "flow logs"],
                 ["Azure Virtual Desktop", "n/a", "not offered"],
                 ["Blob Storage", "Object Storage Service", "S3 API"],
                 ["Managed Disks", "Elastic Volume Service", "SSD"],
                 ["Virtual Network", "Virtual Private Cloud", "hub and spoke"]]
MAPPING_OTHER = [["Quelle", "TCP Service", "Hinweis"],
                 ["Cloud Spanner", "GaussDB", "GCP"],
                 ["Horizon View", "Elastic Cloud Server", "VMware"],
                 ["Aria Operations", "Cloud Eye", "VMware"],
                 ["Tanzu Kubernetes Grid", "Cloud Container Engine", "VMware"]]
FLAVORS = [["Flavor", "vCPU", "RAM GiB", "Instance Family"],
           ["s3.large.2", "2", "4", "General Purpose"],
           ["c7n.2xlarge.2", "8", "16", "General Computing-plus"],
           ["m7n.large.8", "2", "16", "Memory-optimized"],
           ["Standard_D4s_v3", "4", "16", "Dsv3-series"],
           ["Large Memory", "32", "512", "Ultra-high Memory"]]
PRICES = [["Service", "Flavor", "Preis pro Monat", "Preis pro Stunde"],
          ["Elastic Cloud Server", "s3.large.2", "48,00 EUR", "0,066 EUR"],
          ["Elastic Cloud Server", "c7n.2xlarge.2", "120,00 EUR", "0,16 EUR"],
          ["Elastic Volume Service", "General Purpose SSD", "0,10 EUR/GB", "n/a"],
          ["Relational Database Service", "rds.pg.c2.large", "95,00 EUR", "0,13 EUR"]]


CASES = [
    # ------------------------------------------------------------------ contact and participant lists, split cells
    case("md-vorname-nachname-firma", "md table, first and last name in two cells under Vorname and Nachname, Firma column",
         [FIRST, LAST, P2_FIRST, P2_LAST, BRAND], ["Vorname", "Nachname", "Firma", "Rolle", "Architekt"],
         txt("teilnehmer.md", md_table(CONTACTS_DE)),
         expect={FIRST: "person", LAST: "person", P2_LAST: "person", BRAND: "company"},
         note="the first header cell is a person label: the key-value pass of the table rule may take the second header cell for a value"),
    case("docx-signin-first-last-en", "docx table of a sign-in sheet, First Name | Last Name | Company, no mail column",
         [FIRST, LAST, P2_FIRST, P2_LAST, P3_LAST, BRAND], ["First Name", "Last Name", "Company", "Signature"],
         docx_table("anwesenheit.docx", CONTACTS_EN, before=["Training TCP Basics, 12.10.2026"]),
         expect={LAST: "person"}),
    case("csv-outlook-contacts-en", "csv, an Outlook contacts export: First Name, Last Name, Company, E-mail Address with first.last addresses",
         [FIRST, LAST, P2_FIRST, P2_LAST, BRAND, MAIL_X, MAIL_A],
         ["First Name", "Last Name", "Company", "E-mail Address", "Business Phone"],
         csvf("contacts.csv", [["First Name", "Middle Name", "Last Name", "Company", "E-mail Address", "Business Phone"],
                               [FIRST, "", LAST, BRAND, MAIL_X, TEL1],
                               [P2_FIRST, "", P2_LAST, BRAND, MAIL_A, TEL2]]),
         expect={LAST: "person", P2_LAST: "person"},
         note="the surnames are learned from the local parts of the addresses; the first name of an initial-only local part has no source"),
    case("md-vor-und-nachname", "md table, the full name in one cell under the header Vor- und Nachname",
         [PLANTED_PERSON, P2, BRAND], ["Vor- und Nachname", "Firma", "Telefon"],
         txt("liste.md", md_table([["Vor- und Nachname", "Firma", "Telefon"], [PLANTED_PERSON, BRAND, TEL1], [P2, ORG_SHORT, TEL2]])),
         expect={PLANTED_PERSON: "person", P2: "person", BRAND: "company"}),
    case("tsv-teams-attendance", "tsv, a Teams attendance export: Full Name, Join Time, Email (a UPN without a dot), Role",
         [PLANTED_PERSON, P2, "x" + LAST.lower() + "@" + DOMAIN], ["Full Name", "Join Time", "Leave Time", "Duration", "Presenter", "Attendee"],
         txt("attendance.tsv", "\n".join("\t".join(r) for r in [
             ["Full Name", "Join Time", "Leave Time", "Duration", "Email", "Role"],
             [PLANTED_PERSON, "10/7/2026, 10:03:15 AM", "10/7/2026, 11:00:02 AM", "56m 47s", "x" + LAST.lower() + "@" + DOMAIN, "Presenter"],
             [P2, "10/7/2026, 10:04:01 AM", "10/7/2026, 11:00:02 AM", "56m 1s", P2_FIRST[0].lower() + P2_LAST.lower() + "@" + DOMAIN, "Attendee"]]),
             mark=False),
         expect={PLANTED_PERSON: "person", P2: "person"}),
    case("xlsx-title-row-above-header", "xlsx, a title in A1, an empty row, then the header row Vorname | Nachname | Firma",
         [FIRST, LAST, P2_FIRST, P2_LAST, BRAND], ["Vorname", "Nachname", "Firma", "Teilnehmerliste Workshop"],
         xlsx_book("teilnehmer.xlsx", [("Teilnehmer", [["Teilnehmerliste Workshop 12.10.2026"], [], ["Vorname", "Nachname", "Firma"],
                                                       [FIRST, LAST, BRAND], [P2_FIRST, P2_LAST, ORG_SHORT]])]),
         expect={LAST: "person"},
         note="the extractor puts the table separator after the first non-empty row, so the real header is a data row"),
    case("xlsx-two-sheets-contacts-hosts", "xlsx with two sheets: Kontakte (Nachname, Vorname, Firma, E-Mail) and Hosts (Hostname, Owner, IP) with the surname alone as the owner and in a host name",
         [FIRST, LAST, P2_FIRST, P2_LAST, BRAND, MAIL_X, MAIL_A, HOST_P_DB, HOST_Z_WEB],
         ["Kontakte", "Hosts", "Hostname", "Owner", "E-Mail"],
         xlsx_book("inventar.xlsx", [("Kontakte", [["Nachname", "Vorname", "Firma", "E-Mail"], [LAST, FIRST, BRAND, MAIL_X], [P2_LAST, P2_FIRST, BRAND, MAIL_A]]),
                                     ("Hosts", [["Hostname", "Owner", "IP"], [HOST_P_DB, LAST, IP1], [HOST_Z_WEB, P2_LAST, IP2]])]),
         expect={LAST: "person", P2_LAST: "person", BRAND: "company"}),
    case("xlsx-sheet-named-after-person", "xlsx, a timesheet with one sheet per person, the sheet name is the surname and a second sheet is the company brand; no name column",
         [LAST, BRAND], ["Datum", "Stunden", "Netzkonzept", "web01"],
         xlsx_book("stunden.xlsx", [(LAST, [["Datum", "Stunden", "Tätigkeit"], ["2026-10-01", "8", "Netzkonzept"], ["2026-10-02", "6", "Migration"]]),
                                    (BRAND, [["Host", "Owner", "IP"], ["web01", "netops", IP1]])]),
         note="the sheet name becomes a Markdown heading of the output"),
    case("md-keyvalue-steckbrief", "md two-column table, labels in the first column (Ansprechpartner, Firma, Ort), values in the second",
         [PLANTED_PERSON, BRAND, CONTROL, MAIL_INFO], ["Cloud-Migration", "Ansprechpartner", "Firma", "Ort", "Telefon", "E-Mail"],
         txt("steckbrief.md", md_table([["Projekt", "Cloud-Migration"], ["Ansprechpartner", PLANTED_PERSON], ["Firma", BRAND],
                                        ["Ort", CONTROL], ["Telefon", TEL1], ["E-Mail", MAIL_INFO]])),
         expect={PLANTED_PERSON: "person", BRAND: "company", CONTROL: "place"}),
    case("md-table-no-outer-pipes", "md table written without the outer pipes (GitHub style), Vorname | Nachname | Firma",
         [FIRST, LAST, P2_FIRST, P2_LAST, BRAND], ["Vorname", "Nachname", "Firma"],
         txt("liste.md", "Vorname | Nachname | Firma\n--- | --- | ---\n%s | %s | %s\n%s | %s | %s\n" % (FIRST, LAST, BRAND, P2_FIRST, P2_LAST, ORG_SHORT)),
         expect={LAST: "person"}),
    case("txt-layout-table-wide-gaps", "txt, a fixed-width table (pandoc simple table): columns aligned with runs of spaces, a dashed line under the header",
         [FIRST, LAST, P2_FIRST, P2_LAST, BRAND], ["Vorname", "Nachname", "Firma"],
         txt("liste.txt", "Vorname      Nachname     Firma\n-----------  -----------  -------------\n%-12s %-12s %s\n%-12s %-12s %s\n"
             % (FIRST, LAST, BRAND, P2_FIRST, P2_LAST, ORG_SHORT))),

    # ------------------------------------------------------------------ Russian headers
    case("md-ru-fio-company-city", "md table with Russian headers ФИО | Компания | Город | Телефон, a full name with patronymic, a company in guillemets",
         [RU_PERSON, RU_LAST, RU_PERSON2, RU_LAST2, RU_BRAND, RU_ORG, RU_PLACE, RU_CONTROL],
         ["Компания", "Город", "Телефон"],
         txt("kontakty.md", md_table([["ФИО", "Компания", "Город", "Телефон"],
                                      [RU_PERSON, "ООО «%s»" % RU_BRAND, RU_PLACE, "+7 555 011 34 20"],
                                      [RU_PERSON2, "АО «%s»" % RU_ORG, RU_CONTROL, "+7 555 011 34 30"]])),
         expect={RU_PERSON: "person", RU_BRAND: "company", RU_PLACE: "place"}),
    case("csv-ru-split-patronymic", "csv with semicolons and Russian headers Фамилия;Имя;Отчество;Отдел, one word per cell",
         [RU_LAST, RU_PATRONYM, RU_LAST2, RU_PATRONYM2], ["Фамилия", "Имя", "Отчество", "Отдел", "Закупки"],
         csvf("sotrudniki.csv", [["Фамилия", "Имя", "Отчество", "Отдел"], [RU_LAST, RU_FIRST, RU_PATRONYM, "ИТ"],
                                 [RU_LAST2, RU_FIRST2, RU_PATRONYM2, "Закупки"]], delim=";"),
         expect={RU_LAST: "person"},
         note="the patronymic column has no label of its own"),

    # ------------------------------------------------------------------ companies split over cells
    case("md-brand-legal-form-no-header", "md table without a header row: the brand in one cell, the legal form in the next, the first row is read as the header",
         [BRAND, COMPANY2, COMPANY3], ["GmbH", "KG"],
         txt("firmen.md", md_table([[BRAND, "GmbH"], [COMPANY2, "GmbH"], [COMPANY3, "KG"]])),
         expect={COMPANY2: "company", COMPANY3: "company"},
         note="the row of the brand alone is read as the header row"),
    case("csv-semicolon-firma-rechtsform", "csv with semicolons: Firma;Rechtsform;Sitz, the brand alone under Firma",
         [BRAND, COMPANY2, CONTROL], ["Rechtsform", "GmbH", "Sitz"],
         csvf("firmen.csv", [["Firma", "Rechtsform", "Sitz"], [BRAND, "GmbH", CONTROL], [COMPANY2, "GmbH", PLACE]], delim=";"),
         expect={BRAND: "company", COMPANY2: "company", CONTROL: "place"}),
    case("md-name-rechtsform", "md table Name | Rechtsform with company brands under Name",
         [BRAND, COMPANY2], ["Rechtsform", "GmbH"],
         txt("firmen.md", md_table([["Name", "Rechtsform"], [BRAND, "GmbH"], [COMPANY2, "GmbH"]])),
         expect={BRAND: "company", COMPANY2: "company"}),

    # ------------------------------------------------------------------ headers the label list does not carry
    case("docx-mitarbeiter-column", "docx table Mitarbeiter | Abteilung | Standort, the surname alone under Mitarbeiter",
         [LAST, P2_LAST, CONTROL], ["Abteilung", "Einkauf", "Standort"],
         docx_table("orga.docx", [["Mitarbeiter", "Abteilung", "Standort"], [LAST, "IT", CONTROL], [P2_LAST, "Einkauf", PLACE]]),
         expect={CONTROL: "place"}),
    case("docx-ansprechpartner-kunde-column", "docx table Projekt | Ansprechpartner Kunde | Ansprechpartner T-Systems, surnames alone",
         [LAST, P2_LAST, P5_LAST, P3_LAST], ["Cloud-Migration", "Backup", "Ansprechpartner Kunde"],
         docx_table("projekte.docx", [["Projekt", "Ansprechpartner Kunde", "Ansprechpartner T-Systems"],
                                      ["Cloud-Migration", LAST, P2_LAST], ["Backup", P5_LAST, P3_LAST]])),
    case("xlsx-raci-en-accountable", "xlsx RACI with the English role headers Responsible | Accountable | Consulted | Informed and surnames in the cells",
         [LAST, P2_LAST, P5_LAST, P4_LAST], ["Responsible", "Accountable", "Consulted", "Informed", "Network design"],
         xlsx_book("raci.xlsx", [("RACI", [["Task", "Responsible", "Accountable", "Consulted", "Informed"],
                                           ["Network design", LAST, P2_LAST, P4_LAST, P5_LAST],
                                           ["Migration", P2_LAST, LAST, P5_LAST, P4_LAST]])]),
         expect={LAST: "person"}),
    case("md-site-column-en", "md host table Hostname | Site | Rack with an unregistered town under Site",
         [CONTROL], ["Rack", "R12", "Hostname"],
         txt("racks.md", md_table([["Hostname", "Site", "Rack"], ["srv-app", CONTROL, "R12"], ["srv-db", PLACE, "R13"]]))),

    # ------------------------------------------------------------------ RACI matrices
    case("md-raci-names-in-header", "md RACI with the full names as column headers and R/A/C/I in the cells, then the surname alone in a sentence",
         [PLANTED_PERSON, LAST, P2, P2_LAST], ["Netzkonzept", "Migration", "Freitag", "Aufgabe"],
         txt("raci.md", md_table([["Aufgabe", PLANTED_PERSON, P2, PERSON], ["Netzkonzept", "R", "A", "C"], ["Migration", "A", "R", "I"]])
             + "\n%s klärt das mit %s bis Freitag.\n" % (LAST, P2_LAST)),
         expect={PLANTED_PERSON: "person", P2: "person"}),
    case("xlsx-raci-initials-legend", "xlsx RACI with initials as headers and a second sheet Kürzel | Name as the legend",
         [PLANTED_PERSON, FIRST, LAST, P2], ["Kürzel", "Netzkonzept", "Aufgabe", "XP", "AZ"],
         xlsx_book("raci.xlsx", [("RACI", [["Aufgabe", "XP", "AZ", "TB"], ["Netzkonzept", "R", "A", "C"], ["Migration", "A", "R", "I"]]),
                                 ("Legende", [["Kürzel", "Name"], ["XP", PLANTED_PERSON], ["AZ", P2], ["TB", PERSON]])]),
         expect={PLANTED_PERSON: "person", P2: "person"}),

    # ------------------------------------------------------------------ asset and host lists
    case("csv-hosts-owner-standort", "csv host list hostname,owner,standort,ip with the surname in a host name and the owner's full name",
         [HOST_P, PLANTED_PERSON, LAST, P2, CONTROL], ["hostname", "owner", "standort", "ip", "netops"],
         csvf("hosts.csv", [["hostname", "owner", "standort", "ip"], [HOST_P, PLANTED_PERSON, CONTROL, IP1],
                            [HOST_Z, P2, PLACE, IP2], [HOST_FW, "netops", CONTROL, IP3]]),
         expect={PLANTED_PERSON: "person", CONTROL: "place"}),
    case("md-host-number-before-ip-cell", "md host table, a host name ending in a number in the cell before the IP cell, no name planted",
         [], ["srv-web-02", "srv-db-01", "Web", "Firewall"],
         txt("hosts.md", md_table([["Hostname", "IP", "Rolle"], ["srv-web-02", IP4, "Web"], ["srv-db-01", "10.0.0.5", "DB"], ["fw-01", IP3, "Firewall"]])),
         note="the structured patterns read a number and the IP across the cell border as one phone number"),
    case("xlsx-hersteller-column", "xlsx asset list Hersteller | Modell | Anzahl with vendors the lists do not carry",
         [], ["Supermicro", "Kyocera", "Dell", "Fujitsu", "Hersteller"],
         xlsx_book("assets.xlsx", [("Assets", [["Hersteller", "Modell", "Anzahl"], ["Supermicro", "SYS-1029", "4"], ["Kyocera", "ECOSYS M3145", "2"],
                                               ["Dell", "R740", "6"], ["Fujitsu", "RX2540", "1"]])])),
    case("csv-name-column-hostnames", "csv inventory Name,Typ,IP with host names under the Name column",
         [], ["SRV-APP-01", "Fileserver-Nord", "Exchange-Server", "Jumphost", "web-prod-01"],
         csvf("vms.csv", [["Name", "Typ", "IP"], ["SRV-APP-01", "ECS", IP1], ["Fileserver-Nord", "ECS", IP2], ["Exchange-Server", "ECS", IP3],
                          ["Jumphost", "ECS", IP4], ["web-prod-01", "ECS", "203.0.113.10"]])),
    case("md-owner-column-teams", "md host table Hostname | Owner with teams and departments as owners",
         [], ["Netzwerkteam", "IT-Betrieb", "Cloud Operations", "Team Plattform", "netops"],
         txt("owners.md", md_table([["Hostname", "Owner"], ["srv-app-01", "Netzwerkteam"], ["srv-db-01", "IT-Betrieb"], ["srv-web-01", "Cloud Operations"],
                                    ["srv-mon-01", "Team Plattform"], ["srv-fw-01", "netops"]]))),

    # ------------------------------------------------------------------ product, price, flavor and mapping tables
    case("xlsx-cost-headers-de", "xlsx price sheet with German cost headers of two capitalised words (Monatliche Kosten, Einmalige Kosten, Stückpreis Netto)",
         [], ["Monatliche Kosten", "Einmalige Kosten", "Stückpreis Netto", "Preis pro Monat", "Summe Netto"],
         xlsx_book("kosten.xlsx", [("Kosten", [["Position", "Monatliche Kosten", "Einmalige Kosten", "Stückpreis Netto", "Preis pro Monat", "Summe Netto"],
                                               ["ECS s3.large.2", "48,00", "0,00", "0,066", "48,00", "48,00"],
                                               ["EVS 100 GB", "10,00", "0,00", "0,10", "10,00", "10,00"]])])),
    case("md-service-price-table", "md product and price table service | flavor | price",
         [], ["Elastic Cloud Server", "Elastic Volume Service", "Relational Database Service", "s3.large.2", "General Purpose SSD", "rds.pg.c2.large", "Preis pro Monat"],
         txt("preise.md", md_table(PRICES))),
    case("md-flavor-table", "md flavor table with instance families (General Purpose, Memory-optimized, Large Memory)",
         [], ["s3.large.2", "Standard_D4s_v3", "General Purpose", "General Computing-plus", "Memory-optimized", "Instance Family", "Large Memory", "Ultra-high Memory"],
         txt("flavors.md", md_table(FLAVORS))),
    case("md-service-mapping-azure", "md service mapping table Azure Service -> TCP Service",
         [], ["Virtual Machines", "Network Watcher", "Azure Virtual Desktop", "Blob Storage", "Managed Disks", "Virtual Network",
              "Elastic Cloud Server", "Cloud Eye", "Object Storage Service", "Virtual Private Cloud"],
         txt("mapping.md", md_table(MAPPING_AZURE))),
    case("md-service-mapping-gcp-vmware", "md service mapping table with GCP and VMware products whose words the lists do not carry",
         [], ["Cloud Spanner", "Horizon View", "Aria Operations", "Tanzu Kubernetes Grid", "GaussDB", "Cloud Eye"],
         txt("mapping2.md", md_table(MAPPING_OTHER))),

    # ------------------------------------------------------------------ SQL in a cell
    case("md-sql-qualified-columns", "md table of queries, qualified column names k.Name and a.id in a cell",
         [], ["k.Name", "a.id", "FROM Kundenstamm", "JOIN Artikelstamm"],
         txt("queries.md", md_table([["Query", "Laufzeit"], ["SELECT k.Name, k.Ort FROM Kundenstamm k JOIN Artikelstamm a ON a.id = k.id", "40 s"],
                                     ["SELECT * FROM Kundenstamm", "2 s"]]))),
    case("docx-sql-keyword-outside-stoplist", "docx table of statements, a SQL keyword outside the stop list followed by a German table name",
         [], ["TRUNCATE Auftragskopf", "EXEC Auftragsabschluss", "INSERT INTO Auftragskopf", "MERGE INTO Kundenstamm"],
         docx_table("statements.docx", [["Statement", "Dauer"], ["TRUNCATE Auftragskopf", "1 s"], ["EXEC Auftragsabschluss", "3 s"],
                                        ["INSERT INTO Auftragskopf VALUES (1)", "1 s"], ["MERGE INTO Kundenstamm USING Import", "9 s"]])),
    case("md-sql-name-vorname-outer-join", "md table of queries with the columns Name, Vorname and LEFT OUTER JOIN",
         [], ["Name, Vorname", "FROM Mitarbeiter", "LEFT OUTER JOIN", "WHERE Abteilung"],
         txt("queries.md", md_table([["Query", "Laufzeit"], ["SELECT Name, Vorname FROM Mitarbeiter WHERE Abteilung = 'IT'", "2 s"],
                                     ["SELECT * FROM Orders o LEFT OUTER JOIN Shipments s ON s.order_no = o.order_no", "12 s"]]))),

    # ------------------------------------------------------------------ roles, departments and teams as values
    case("docx-role-department-cells", "docx table Name | Rolle | Bereich with roles and German department names as cells",
         [PLANTED_PERSON, P2], ["Projektleiter Kunde", "Lead Architect", "Teamleiter Netzwerk", "Fachbereich Logistik", "Abteilung Rechnungswesen", "Stabsstelle Datenschutz"],
         docx_table("rollen.docx", [["Name", "Rolle", "Bereich"], [PLANTED_PERSON, "Projektleiter Kunde", "Fachbereich Logistik"],
                                    [P2, "Lead Architect", "Abteilung Rechnungswesen"], [PERSON, "Teamleiter Netzwerk", "Stabsstelle Datenschutz"]]),
         expect={PLANTED_PERSON: "person", P2: "person"}),
    case("docx-keyvalue-team-values", "docx two-column table with labels and colons (Owner:, Kontakt:, Verantwortlich:, Standort:, Name:) and teams or systems as values",
         [], ["Bestellportal", "Plattform-Team", "Servicedesk", "Fachbereich Logistik", "Rechenzentrum Nord", "Jumphost Verwaltung"],
         docx_table("service.docx", [["Service:", "Bestellportal"], ["Owner:", "Plattform-Team"], ["Kontakt:", "Servicedesk"],
                                     ["Verantwortlich:", "Fachbereich Logistik"], ["Standort:", "Rechenzentrum Nord"], ["Name:", "Jumphost Verwaltung"]])),

    # ------------------------------------------------------------------ run 2: the mechanisms of run 1, varied
    case("md-person-column", "md table Person | Aufgabe, the surname alone under Person",
         [LAST, P2_LAST], ["Aufgabe", "Netzkonzept", "Migration"],
         txt("aufgaben.md", md_table([["Person", "Aufgabe"], [LAST, "Netzkonzept"], [P2_LAST, "Migration"]])),
         note="varies the headers the label list does not carry: Person is the most generic one"),
    case("docx-caption-row-above-header", "docx table whose first row is one merged caption cell over all columns, the header row below it",
         [FIRST, LAST, P2_FIRST, P2_LAST, BRAND], ["Vorname", "Nachname", "Firma", "Teilnehmerliste Workshop"],
         docx_table_caption("teilnehmer.docx", "Teilnehmerliste Workshop 12.10.2026",
                            [["Vorname", "Nachname", "Firma"], [FIRST, LAST, BRAND], [P2_FIRST, P2_LAST, ORG_SHORT]]),
         note="varies xlsx-title-row-above-header: the caption row inside the table of a Word document"),
    case("md-table-without-separator-line", "md rows with pipes but no separator line under the header (typed in a chat or a mail)",
         [FIRST, LAST, P2_FIRST, P2_LAST, BRAND], ["Vorname", "Nachname", "Firma"],
         txt("liste.md", "| Vorname | Nachname | Firma |\n| %s | %s | %s |\n| %s | %s | %s |\n" % (FIRST, LAST, BRAND, P2_FIRST, P2_LAST, ORG_SHORT)),
         note="varies md-table-no-outer-pipes: the column rule needs the separator line"),
    case("md-kunde-projekt-header", "md ticket table Kunde | Projekt | Status, the brand alone under Kunde, the header Projekt next to the label",
         [BRAND], ["Projekt", "Status", "Cloud-Migration", "offen"],
         txt("tickets.md", md_table([["Kunde", "Projekt", "Status"], [BRAND, "Cloud-Migration", "offen"], [CUSTOMER_SHORT, "Backup", "erledigt"]])),
         expect={BRAND: "company"}, code=True,
         note="varies the header loss of md-vorname-nachname-firma: a generic company word as the second header cell"),
    case("md-name-unternehmen-generic-values", "md table Name | Unternehmen with the words Partner and Kunde as the company of a row",
         [PLANTED_PERSON, P2, BRAND], ["Partner", "Kunde", "Unternehmen"],
         txt("liste.md", md_table([["Name", "Unternehmen"], [PLANTED_PERSON, BRAND], [P2, "Partner"], [PERSON, "Kunde"]])),
         expect={PLANTED_PERSON: "person", BRAND: "company"},
         note="a generic company word standing alone as the value of a company column"),
    case("csv-name-column-applications", "csv application list Name,Typ,Owner with German application names under Name and a login under Owner",
         [LOGIN_T], ["Kundenportal", "Ticketsystem", "Fileserver", "Zeiterfassung", "Intranet", "Typ"],
         csvf("apps.csv", [["Name", "Typ", "Owner"], ["Kundenportal", "Web", LOGIN_T], ["Ticketsystem", "Web", "netops"],
                           ["Fileserver", "SMB", LOGIN_T], ["Zeiterfassung", "Web", "netops"], ["Intranet", "Web", LOGIN_T]]),
         expect={LOGIN_T: "person"},
         note="varies csv-name-column-hostnames: application names instead of host names"),
    case("docx-price-table-de-headers", "docx price table with the headers Einzelpreis Netto, Gesamtpreis Brutto, Jährliche Kosten",
         [], ["Einzelpreis Netto", "Gesamtpreis Brutto", "Jährliche Kosten", "Menge", "Position"],
         docx_table("angebot.docx", [["Position", "Menge", "Einzelpreis Netto", "Gesamtpreis Brutto", "Jährliche Kosten"],
                                     ["ECS s3.large.2", "4", "48,00", "228,48", "2.741,76"], ["EVS 500 GB", "1", "50,00", "59,50", "714,00"]]),
         note="varies xlsx-cost-headers-de: more German two-word cost headers"),
    case("xlsx-host-number-before-ip", "xlsx host sheet Hostname | IP | Rolle with host names ending in -01 and -02 before the IP cell",
         [], ["srv-db-01", "srv-app-02", "Firewall"],
         xlsx_book("hosts.xlsx", [("Hosts", [["Hostname", "IP", "Rolle"], ["srv-db-01", IP4, "DB"], ["srv-app-02", IP3, "App"], ["fw-a", IP1, "Firewall"]])]),
         note="varies md-host-number-before-ip-cell in the xlsx carrier"),
]
