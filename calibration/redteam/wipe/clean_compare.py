"""The wipe red team through `awb clean` of the laptop next to today's intake (task lc-laptop-clean).

    .venv/bin/python calibration/redteam/wipe/clean_compare.py --work DIR [--json OUT] [--only DIMENSION ...]
    .venv/bin/python calibration/redteam/wipe/clean_compare.py --work DIR --bench 50

Every case of the pack (cases_<dimension>.py) is built twice: once through the intake as `harness.py` runs it (a
throw-away Workbench, `intake.run`), once through `clean.run(..., kind="docs")` in a throw-away laptop folder with the
same invented register, and the cases whose file is a log, config or ticket file also through `kind="logs"`. Leaks
and losses are read with the readers' views of calibration/redteam/harness.py over every output, its file name and
the report. `--bench N` writes N generated customer documents (docx with a logo, a header, a footer, properties,
tables and invented names) and times `awb clean` and the intake over them.

Invented values only. Nothing here prints a planted value: lines carry ids, counts and seconds.
"""
from __future__ import annotations

import argparse
import json
import os
import shutil
import sys
import tempfile
import time
from pathlib import Path

REPO = Path(os.environ.get("AWB_REPO") or Path(__file__).resolve().parents[3])
sys.path.insert(0, str(REPO))

from calibration.redteam import harness as rt  # noqa: E402
from calibration.redteam.wipe import harness as wh  # noqa: E402

PROJECT = "tcp-q7m4"


def fresh_laptop(root: Path) -> Path:
    from awb import laptop
    from tests import fixtures

    base = Path(tempfile.mkdtemp(prefix="lap-", dir=str(root)))
    os.environ[laptop.HOME_ENV] = str(base / "laptop")
    laptop.init()
    laptop.register_path().write_text("\n".join(fixtures.register_lines()) + "\n", encoding="utf-8")
    os.chmod(laptop.register_path(), 0o600)
    return base


def clean_once(case: dict, root: Path, kind: str) -> dict:
    from awb import clean

    base = fresh_laptop(root)
    rec = {"error": None, "outputs": 0, "withheld": False, "skipped": False, "leaks": [], "losses": [], "kept": 0,
           "code": None}
    texts: list[tuple[str, str]] = []
    outputs: list[str] = []
    try:
        inbox = base / "inbox"
        inbox.mkdir()
        src = Path(case["build"](inbox))
        res = clean.run([src], PROJECT, kind, base / "out", say=lambda s: None)
        rec["outputs"] = len(res.outputs)
        rec["withheld"] = any(f.state == "withheld" for f in res.files)
        rec["skipped"] = any(f.state == "skipped" for f in res.files)
        for o in res.outputs:
            body = o.read_text(encoding="utf-8", errors="replace")
            outputs.append(body)
            texts += [("name", o.name), ("copy", body)]
        if res.report:
            texts.append(("report", res.report.read_text(encoding="utf-8")))
    except Exception as err:  # noqa: BLE001
        rec["error"] = type(err).__name__
        texts.append(("exception text", str(err)))
    for i, value in enumerate(case.get("values") or []):
        for label, text in texts:
            how = rt.readable(value, text)
            if how:
                rec["leaks"].append({"value": i, "file": label, "how": how})
    joined = "\n".join(outputs)
    for i, term in enumerate(case.get("keep") or []):
        if rt.readable(term, joined):
            rec["kept"] += 1
        else:
            rec["losses"].append({"term": i})
    if case.get("code"):
        rec["code"] = "CUST-Q7M4" in joined
    shutil.rmtree(base, ignore_errors=True)
    return rec


def _tally(runs: list[dict]) -> dict:
    return {"cases": len(runs), "leak_cases": sum(1 for r in runs if r["leaks"]),
            "leaked_values": sum(len({l["value"] for l in r["leaks"]}) for r in runs),
            "loss_cases": sum(1 for r in runs if r["losses"]), "lost_terms": sum(len(r["losses"]) for r in runs),
            "kept_terms": sum(r["kept"] for r in runs), "withheld": sum(1 for r in runs if r["withheld"]),
            "no_output": sum(1 for r in runs if not r.get("outputs") and not r["error"]),
            "no_code": sum(1 for r in runs if r.get("code") is False),
            "errors": sum(1 for r in runs if r["error"])}


def compare(work: Path, dimensions) -> dict:
    from awb import clean

    out: dict = {"dimensions": {}, "cases": []}
    for dim in dimensions:
        cases = wh.load_cases(wh.module_path(dim))
        intake_runs, docs_runs, log_runs = [], [], []
        t_intake = t_docs = t_logs = 0.0
        for case in cases:
            started = time.time()
            rec = wh.run_once(case, work)
            shutil.rmtree(rec.pop("workbench"), ignore_errors=True)
            t_intake += time.time() - started
            intake_runs.append(rec)
            started = time.time()
            d = clean_once(case, work, "docs")
            t_docs += time.time() - started
            docs_runs.append(d)
            probe = Path(tempfile.mkdtemp(dir=str(work)))
            suffix = Path(case["build"](probe)).suffix.lower()
            shutil.rmtree(probe, ignore_errors=True)
            entry = {"id": case["id"], "dimension": dim, "intake": wh.outcome({"run": rec}),
                     "docs": "leak" if d["leaks"] else ("loss" if d["losses"] else "clean")}
            if suffix in clean.LOG_SUFFIXES:
                started = time.time()
                lg = clean_once(case, work, "logs")
                t_logs += time.time() - started
                log_runs.append(lg)
                entry["logs"] = "leak" if lg["leaks"] else ("loss" if lg["losses"] else "clean")
            out["cases"].append(entry)
        out["dimensions"][dim] = {"intake": _tally(intake_runs), "docs": _tally(docs_runs),
                                  "logs": _tally(log_runs) if log_runs else None,
                                  "seconds": {"intake": round(t_intake, 1), "docs": round(t_docs, 1),
                                              "logs": round(t_logs, 1)}}
        d = out["dimensions"][dim]
        print("%-12s intake leak %3d loss %3d | clean docs leak %3d loss %3d | logs %s"
              % (dim, d["intake"]["leak_cases"], d["intake"]["loss_cases"], d["docs"]["leak_cases"],
                 d["docs"]["loss_cases"], "-" if not d["logs"] else "leak %d loss %d of %d" % (
                     d["logs"]["leak_cases"], d["logs"]["loss_cases"], d["logs"]["cases"])), flush=True)
    return out


# --------------------------------------------------------------------------- the 50 documents

FIRSTS = ("Ulvrad", "Xqarv", "Jonas", "Anna", "Tobias", "Peter", "Thomas")
LASTS = ("Qesmotz", "Pomblet", "Vrelkam", "Zrbl", "Brinkhove", "Taldrup", "Wenzmar")
COMPANIES = ("Grvetz Systemhaus AG", "Nrgtz Beratung GmbH", "Taldrup Data KG", "Wenzmar Netz GmbH")


def make_docs(folder: Path, n: int) -> list[Path]:
    import base64

    import docx
    from docx.shared import Inches

    from tests import fixtures as fx

    png = base64.b64decode("iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mP8z8BQDwAEhQGAhKmMIQAAAABJRU5ErkJggg==")
    logo = folder / "logo.png"
    logo.write_bytes(png)
    files = []
    for i in range(n):
        first, last = FIRSTS[i % len(FIRSTS)], LASTS[(i * 3) % len(LASTS)]
        company = COMPANIES[i % len(COMPANIES)]
        d = docx.Document()
        d.core_properties.author = "%s %s" % (first, last)
        d.core_properties.title = "Konzept %d für %s" % (i, fx.CUSTOMER_FORMS[0])
        d.sections[0].header.paragraphs[0].text = "%s, vertraulich" % fx.CUSTOMER_FORMS[0]
        d.sections[0].footer.paragraphs[0].text = "%s, Musterstraße 12, 12345 Musterhausen" % company
        d.add_picture(str(logo), width=Inches(0.5))
        d.add_heading("Migrationskonzept %d: %s nach T Cloud Public" % (i, fx.CUSTOMER_FORMS[1]), 1)
        for k in range(12):
            d.add_paragraph(
                "Abschnitt %d.%d: Die Anwendung läuft heute auf VMware vSphere 7 mit 24 VMs und PostgreSQL 14. "
                "Ansprechpartner ist Herr %s %s (%s.%s@%s, +49 555 12345%02d). Wir planen ECS, CCE 1.29, RDS und "
                "OBS in eu-de; das Netz geht über Direct Connect mit 1 Gbit/s. %s liefert die Firewall."
                % (i, k, first, last, first.lower(), last.lower(), fx.CUSTOMER_DOMAIN, k, company))
        table = d.add_table(rows=6, cols=4)
        for r, row in enumerate(table.rows):
            cells = ("Host", "Rolle", "IP", "Verantwortlich") if r == 0 else (
                "srv-app-%02d" % r, "Applikation", "10.1.%d.%d" % (i % 250, r), "%s %s" % (first, last))
            for c, value in enumerate(cells):
                row.cells[c].text = value
        d.add_paragraph("Mit freundlichen Grüßen\n%s %s\n%s" % (first, last, company))
        path = folder / ("konzept-%02d.docx" % i)
        d.save(path)
        files.append(path)
    logo.unlink()
    return files


def bench(work: Path, n: int) -> dict:
    from awb import clean, intake

    folder = Path(tempfile.mkdtemp(prefix="bench-", dir=str(work)))
    files = make_docs(folder, n)
    base = fresh_laptop(work)
    started = time.time()
    res = clean.run(files, PROJECT, "docs", base / "out", say=lambda s: None)
    t_clean = time.time() - started
    copies = [o.read_text(encoding="utf-8") for o in res.outputs]
    leaks = sum(1 for t in copies for v in FIRSTS[:2] + LASTS + ("Grvetz", "Nrgtz", "Musterhausen")
                if rt.readable(v, t))
    # today's intake over the same documents, then the host's trusted import of the clean copies
    wb, p = rt.fresh_workbench(work)
    inbox_files = []
    for f in files:
        dst = p.inbox / f.name
        shutil.copy(f, dst)
        inbox_files.append(dst)
    started = time.time()
    intake.run(inbox_files, "CUST-Q7M4", p)
    t_intake = time.time() - started
    wb2, p2 = rt.fresh_workbench(work)
    host_files = []
    for o in res.outputs:
        dst = p2.inbox / o.name
        shutil.copy(o, dst)
        host_files.append(dst)
    started = time.time()
    intake.run(host_files, "CUST-Q7M4", p2, trusted=True)
    t_trusted = time.time() - started
    for d in (folder, base, wb, wb2):
        shutil.rmtree(d, ignore_errors=True)
    return {"documents": n, "clean_seconds": round(t_clean, 1), "copies": len(res.outputs),
            "readable_planted_values_in_copies": leaks, "intake_seconds": round(t_intake, 1),
            "host_trusted_import_seconds": round(t_trusted, 1)}


def main(argv=None) -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--work", required=True)
    ap.add_argument("--json")
    ap.add_argument("--only", nargs="*")
    ap.add_argument("--bench", type=int, default=0)
    args = ap.parse_args(argv)
    work = Path(args.work)
    work.mkdir(parents=True, exist_ok=True)
    rt.code_under_test()
    result = {}
    if args.bench:
        result["bench"] = bench(work, args.bench)
        print(json.dumps(result["bench"]))
    else:
        result = compare(work, args.only or wh.DIMENSIONS)
    if args.json:
        Path(args.json).write_text(json.dumps(result, indent=1), encoding="utf-8")
    return 0


if __name__ == "__main__":
    sys.exit(main())
