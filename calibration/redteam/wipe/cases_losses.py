"""Red team of wipe mode, dimension losses: the false-positive side of typical T Cloud Public (TCP) architecture
texts. Service names singular and plural, vendor and product names, source-platform mapping tables, SQL, flavors,
HTTP and log texts, headings in Markdown and in capitals, standards, roles, legal phrases, table headers, and the
mixed shapes (a term next to a planted person, a service as a speaker label, a term after Hallo). The second half
plants the shapes the prototype's rules take for names: all-caps abbreviations that are also title words (DR, MS,
MX, HR, NAT, LLM, Fr), lettered sections (A. Einleitung), a noun before an initial (Szenario B.), first names of
the list as technical words (Grace Period, Max Pods, Bill of Materials), products whose words the public corpus
does not carry (an HSM vendor, a ticket system, a backup tool), product labels of a key-value list, "Am Standort 2",
five-digit quantities, one-word headings and subjects, teams as signers and addressees, and the learned forms: a
vendor word wiped once at a label, a documentation URL or a support address, and then wiped everywhere. The
variations of run 1: a German sentence that opens with a preposition outside the lists, a lettered table of
contents, a learned word that the intake report itself carries, participle headings of slides, and the slide
heading the pptx extractor writes.

Every name is from tests/fixtures.py or from INVENTED below. Hosts end in .example. The technical terms are real
and public. Nothing here is real as a customer, partner or person.
"""
import csv
import sys
from email import policy
from email.message import EmailMessage
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[3]))
from tests import fixtures as fx  # noqa: E402

# the fixture forms
CUSTOMER, CUSTOMER_SHORT, CUSTOMER_ACRONYM, CUSTOMER_EN = fx.CUSTOMER_FORMS
PERSON, SURNAME = fx.PERSON_FORMS            # a registered person (becomes a PERS code, not a token)
FIRST = fx.PLANTED_PERSON.split()[0]         # Xqarv
LAST = fx.PLANTED_PERSON.split()[1]          # Pomblet
PLANTED_PERSON = fx.PLANTED_PERSON           # Xqarv Pomblet, unregistered
PLANTED_COMPANY = fx.PLANTED_CANDIDATE       # Nrgtz Beratung GmbH, unregistered
BRAND = PLANTED_COMPANY.split()[0]           # Nrgtz
CONTROL = fx.CONTROL_UNREGISTERED            # Musterhausen, unregistered, a place word

# every invented name and host of this module, so that the owner can see nothing real slipped in; the product
# words inside the hosts (zammad, bareos, kemp) are public product names, the hosts themselves are invented
INVENTED = (
    "docs.zammad.example", "support@bareos.example",      # a documentation host and a support address under .example
    "zammad-web-01", "zammad-db-01", "bareos-sd-01",       # host names built from a product word
    "kemp-lb-01", "srv-bkp-01", "srv-bkp-02",
    "203.0.113.20",
)

MARK = "stage two of the invented plan"


def txt(name, text):
    """A plain or Markdown file with `text`; the marker line proves the carrier was read."""
    def build(inbox):
        p = inbox / name
        p.write_text(text + "\n" + MARK + "\n", encoding="utf-8")
        return p
    return build


def csvf(name, rows, delim=";"):
    def build(inbox):
        p = inbox / name
        with p.open("w", encoding="utf-8", newline="") as fh:
            w = csv.writer(fh, delimiter=delim)
            for r in rows:
                w.writerow(r)
            w.writerow([MARK])
        return p
    return build


def docx_paragraphs(name, lines):
    def build(inbox):
        import docx
        d = docx.Document()
        for line in lines:
            d.add_paragraph(line)
        d.add_paragraph(MARK)
        p = inbox / name
        d.save(str(p))
        return p
    return build


def xlsx_rows(name, rows, sheet="Tabelle1"):
    def build(inbox):
        import openpyxl
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = sheet
        for row in rows:
            ws.append(list(row))
        ws.append([MARK])
        p = inbox / name
        wb.save(str(p))
        return p
    return build


def pptx_slides(name, slides):
    """A deck of python-pptx: one blank slide per item, each item a list of text boxes, each box a list of lines."""
    def build(inbox):
        from pptx import Presentation
        from pptx.util import Inches
        prs = Presentation()
        for boxes in slides:
            s = prs.slides.add_slide(prs.slide_layouts[6])
            for i, lines in enumerate(boxes):
                tf = s.shapes.add_textbox(Inches(1), Inches(1 + 2 * i), Inches(7), Inches(1.5)).text_frame
                tf.text = lines[0]
                for line in lines[1:]:
                    tf.add_paragraph().text = line
        s = prs.slides.add_slide(prs.slide_layouts[6])
        s.shapes.add_textbox(Inches(1), Inches(1), Inches(7), Inches(1)).text_frame.text = MARK
        p = inbox / name
        prs.save(str(p))
        return p
    return build


def eml(name, headers, body):
    def build(inbox):
        msg = EmailMessage(policy=policy.default)
        for k, v in headers:
            msg[k] = v
        msg.set_content(body + "\n" + MARK + "\n")
        p = inbox / name
        p.write_bytes(msg.as_bytes())
        return p
    return build


def case(cid, carrier, values, keep, build, note="", code=None, expect=None):
    c = {"id": cid, "carrier": carrier, "values": list(values), "keep": list(keep), "build": build, "note": note}
    if code is not None:
        c["code"] = code
    if expect:
        c["expect"] = dict(expect)
    return c


ECS = "Elastic Cloud Server"
OBS = "Object Storage Service"
CBR = "Cloud Backup and Recovery"

CASES = [
    # ------------------------------------------------------------------ typical TCP texts: the terms must stay
    case("fp-service-plurals", "md, service names in the plural and a service with 'and' in prose", [],
         ["Elastic Cloud Servers", "Virtual Private Clouds", "Cloud Backup and Recovery", "Elastic Volumes"],
         txt("zielbild.md", "## Zielbild\n\nWir bauen drei Elastic Cloud Servers in zwei Virtual Private Clouds. "
             "Cloud Backup and Recovery sichert die Elastic Volumes jede Nacht nach eu-nl.\n")),
    case("fp-vendor-products-prose", "txt, vendor and product names of the source estate in prose", [],
         ["Cisco Meraki vMX", "Palo Alto Networks", "Veeam Backup & Replication", "SUSE Linux Enterprise Server",
          "Red Hat Enterprise Linux", "Windows Server 2022", "Microsoft SQL Server", "Oracle Database 19c", "SAP HANA",
          "VMware vSphere", "Citrix Virtual Apps"],
         txt("ist-aufnahme.txt", "Der Kunde setzt Cisco Meraki vMX und Palo Alto Networks ein, Backup mit Veeam Backup & "
             "Replication.\nBetriebssysteme: SUSE Linux Enterprise Server 15, Red Hat Enterprise Linux 9 und Windows Server "
             "2022.\nDatenbanken: Microsoft SQL Server 2019, Oracle Database 19c und SAP HANA auf VMware vSphere, Citrix "
             "Virtual Apps für die Clients.\n")),
    case("fp-source-mapping-table", "md, a service mapping table Azure, AWS and GCP -> TCP", [],
         ["Azure Virtual Desktop", "Blob Storage", "Site Recovery", "Log Analytics", "Key Vault", "Virtual Network",
          "Elastic Beanstalk", "Transit Gateway", "Compute Engine", "Cloud Spanner"],
         txt("mapping.md", "| Quelle | Service | TCP |\n| --- | --- | --- |\n| Azure | Azure Virtual Desktop | Workspace |\n"
             "| Azure | Blob Storage | %s |\n| Azure | Site Recovery | Storage Disaster Recovery Service |\n"
             "| Azure | Log Analytics | Log Tank Service |\n| Azure | Key Vault | Key Management Service |\n"
             "| Azure | Virtual Network | Virtual Private Cloud |\n| AWS | Elastic Beanstalk | Cloud Container Engine |\n"
             "| AWS | Transit Gateway | Enterprise Router |\n| GCP | Compute Engine | %s |\n"
             "| GCP | Cloud Spanner | Relational Database Service |\n" % (OBS, ECS))),
    case("fp-sql-statements", "md, SQL statements in a fenced block", [],
         ["PRIMARY KEY", "DEFAULT NULL", "LEFT OUTER JOIN", "ORDER BY"],
         txt("schema.md", "## Schema\n\n```sql\nCREATE TABLE auftrag (id SERIAL PRIMARY KEY, kunde VARCHAR(80) DEFAULT NULL);\n"
             "SELECT a.id FROM auftrag a LEFT OUTER JOIN position p ON p.auftrag = a.id ORDER BY a.id;\n```\n")),
    case("fp-price-table-csv", "csv, a price table with flavor names, an Azure reference and German cost headers", [],
         ["s3.large.2", "c7n.2xlarge.2", "General Purpose", "Standard_D4s_v3", "Monatliche Kosten", "Availability Zone"],
         csvf("preise.csv", [("Service", "Flavor", "Azure Referenz", "Monatliche Kosten", "Availability Zone"),
                             (ECS, "s3.large.2", "Standard_D4s_v3", "52,00 EUR", "eu-de-01"),
                             (ECS, "c7n.2xlarge.2", "Standard_D8s_v5", "210,00 EUR", "eu-de-02"),
                             ("Elastic Volume Service", "General Purpose", "Standard SSD", "9,80 EUR", "eu-de-01")])),
    case("fp-http-log-lines", "txt, HTTP status texts and log lines with levels", [],
         ["404 Not Found", "Internal Server Error", "Bad Gateway", "WARNING Connection refused", "ERROR Access denied"],
         txt("app.log", "2026-10-01 10:03:11 WARNING Connection refused by upstream 203.0.113.20\n"
             "2026-10-01 10:03:12 ERROR Access denied for the service account\n"
             "GET /api/v1/orders 404 Not Found\nGET /api/v1/orders 500 Internal Server Error\nGET /health 502 Bad Gateway\n")),
    case("fp-headings-md", "md, capitalised English headings with # and landing zone terms in prose", [],
         ["Executive Summary", "Technical Requirements", "Network Architecture Overview", "Landing Zone",
          "Availability Zone 1", "Resource Group"],
         txt("konzept.md", "# Executive Summary\n\nKurzfassung.\n\n## Technical Requirements\n\nSiehe unten.\n\n"
             "## Network Architecture Overview\n\nDie Landing Zone hat eine Resource Group je Mandant in Availability Zone 1.\n")),
    case("fp-allcaps-headings-docx", "docx, headings in capitals, German and English", [],
         ["TECHNISCHE ANFORDERUNGEN", "NETWORK ARCHITECTURE OVERVIEW", "EXECUTIVE SUMMARY", "ZIELARCHITEKTUR UND BETRIEB"],
         docx_paragraphs("konzept.docx", ["EXECUTIVE SUMMARY", "Kurzfassung des Vorhabens.", "TECHNISCHE ANFORDERUNGEN",
                                          "Die Anforderungen stehen im Lastenheft.", "NETWORK ARCHITECTURE OVERVIEW",
                                          "Hub und Spoke.", "ZIELARCHITEKTUR UND BETRIEB", "Betrieb durch den Kunden."])),
    case("fp-offer-boilerplate", "txt, standards, laws, role titles and legal phrases of an offer", [],
         ["ISO 27001", "BSI C5", "GDPR Article 28", "DSGVO Art. 28", "TISAX", "KRITIS", "Senior Cloud Architect",
          "Key Account Manager", "Terms and Conditions", "Service Level Agreement"],
         txt("angebot.txt", "Nachweise: ISO 27001, BSI C5, GDPR Article 28 und DSGVO Art. 28, dazu TISAX und KRITIS.\n"
             "Ihr Senior Cloud Architect und Ihr Key Account Manager begleiten das Projekt.\n"
             "Es gelten die Terms and Conditions und das Service Level Agreement der Leistungsbeschreibung.\n")),
    case("fp-person-term-sentence", "txt, a planted person and two terms in one sentence", [FIRST, LAST],
         [ECS, "Veeam Backup & Replication"],
         txt("notiz.txt", "%s hat den %s für Veeam Backup & Replication bestellt.\n" % (PLANTED_PERSON, ECS))),
    case("fp-term-after-person-role", "txt, a role title right after a planted person, then a planted company",
         [FIRST, LAST, BRAND], ["Senior Cloud Architect"],
         txt("notiz.txt", "Verantwortlich ist %s Senior Cloud Architect bei %s.\n" % (PLANTED_PERSON, PLANTED_COMPANY))),
    case("fp-service-speaker-labels", "txt, service names as labels of a sizing list next to a repeated person label",
         [LAST], [ECS, OBS, CBR],
         txt("sizing.txt", "%s: 4 vCPU, 16 GB RAM\n%s: 2 TB\n%s: so passt es.\n%s: täglich, 30 Tage\n%s: Flavor s3.large.2\n%s: bestellen.\n"
             % (ECS, OBS, LAST, CBR, ECS, LAST))),
    case("fp-hallo-generic-and-term", "txt, Hallo zusammen, Hallo Team and Hallo plus a planted first name, terms in the body",
         [FIRST], [ECS, "Landing Zone"],
         txt("mails.txt", "Hallo zusammen,\n\nder %s steht.\n\nHallo Team,\n\ndie Landing Zone ist abgenommen.\n\n"
             "Hallo %s,\n\nbitte die Rechnung freigeben.\n" % (ECS, FIRST)), expect={FIRST: "person"}),

    # ------------------------------------------------------------------ abbreviations that are also title words
    case("fp-title-dr", "txt, DR (disaster recovery) before a capitalised noun", [], ["DR Konzept", "DR Site", "DR Test"],
         txt("dr.txt", "Das DR Konzept sieht eine DR Site in eu-nl vor. Der DR Test findet im Oktober statt.\n")),
    case("fp-title-ms", "txt, MS (Microsoft) before a product word", [], ["MS Teams", "MS Exchange Online", "MS Defender"],
         txt("clients.txt", "Die Clients nutzen MS Teams und MS Exchange Online, Schutz durch MS Defender.\n")),
    case("fp-title-extra-acronyms", "txt, MX, HR, NAT and LLM before a capitalised noun (the extra titles of proposed)", [],
         ["MX Record", "MX Eintrag", "HR Portal", "HR System", "NAT Regel", "NAT Rules", "LLM Gateway", "LLM Proxy"],
         txt("netz.txt", "Der MX Record zeigt auf das Secure Mail Gateway, der MX Eintrag wird am Cutover umgestellt.\n"
             "Das HR Portal und das HR System bleiben on premises.\n"
             "Pro Subnetz gilt eine NAT Regel, die NAT Rules liegen am NAT Gateway.\n"
             "Ein LLM Gateway vor dem LLM Proxy begrenzt die Tokens.\n")),
    case("fp-title-weekday-fr", "txt, the weekday abbreviation Fr before a capitalised noun (service hours)", [],
         ["Supportzeiten", "Wartungsfenster"],
         txt("sla.txt", "Erreichbarkeit: Mo bis Fr Supportzeiten 8 bis 18 Uhr, Mo. bis Fr. Wartungsfenster ab 22 Uhr.\n")),

    # ------------------------------------------------------------------ a single letter with a dot
    case("fp-lettered-sections", "docx, sections numbered A. B. C. and I. V.", [],
         ["A. Einleitung", "B. Zielarchitektur", "C. Kostenschätzung", "V. Betrieb"],
         docx_paragraphs("konzept.docx", ["A. Einleitung", "Das Vorhaben.", "B. Zielarchitektur", "Hub und Spoke.",
                                          "C. Kostenschätzung", "Siehe Preisblatt.", "Teil V. Betrieb", "Durch den Kunden."])),
    case("fp-noun-before-initial", "txt, an unknown noun before a letter with a dot (Szenario B., Teilprojekt B., Exhibit B.)", [],
         ["Szenario B.", "Teilprojekt B.", "Exhibit B."],
         txt("varianten.txt", "Szenario B. Lift and Shift der VMs.\nTeilprojekt B. Netzwerk startet danach.\n"
             "Die Preise stehen in Exhibit B. des Vertrags.\n")),

    # ------------------------------------------------------------------ first names of the list as technical words
    case("fp-firstname-first-term", "txt, a term that starts with a first name of the list (Grace Period, Max Pods)", [],
         ["Grace Period", "Max Pods"],
         txt("limits.txt", "Nach Ablauf der Grace Period wird die Ressource freigegeben.\nMax Pods pro Node: 110.\n")),
    case("fp-bill-of-materials", "xlsx, a sheet titled Bill of Materials", [], ["Bill of Materials"],
         xlsx_rows("bom.xlsx", [("Bill of Materials",), ("Position", "Menge", "Beschreibung"), ("1", "3", ECS),
                                ("2", "1", "Elastic Load Balancer")], sheet="BOM")),

    # ------------------------------------------------------------------ products whose words the corpus does not carry
    case("fp-hsm-vendor-run", "txt, HSM vendor and product names in prose", [],
         ["Thales Luna Network HSM", "Utimaco CryptoServer"],
         txt("keys.txt", "Die Schlüssel liegen im Thales Luna Network HSM, alternativ im Utimaco CryptoServer.\n")),
    case("fp-product-unknown-first-word", "md bullets, a product whose first word is unknown and the rest known", [],
         ["Zammad Ticketsystem", "Gitea Runner", "Kemp Technologies", "Fabasoft Business Process Cloud"],
         txt("tools.md", "## Werkzeuge\n\n- Zammad Ticketsystem\n- Gitea Runner\n- Kemp Technologies\n- Fabasoft Business Process Cloud\n")),
    case("fp-nda-addendum", "docx, contract names with a hyphenated or rare word", [],
         ["Non-Disclosure Agreement", "Data Processing Addendum"],
         docx_paragraphs("vertrag.docx", ["Das Non-Disclosure Agreement und das Data Processing Addendum sind unterschrieben."])),
    case("fp-migration-jargon", "pptx, migration and agile jargon as bullets (Rehosting, Replatforming, Tribe, Guild)", [],
         ["Rehosting Wave 1", "Replatforming Candidates", "Tribe Lead", "Guild Master"],
         pptx_slides("plan.pptx", [[["Roadmap"], ["Rehosting Wave 1", "Replatforming Candidates", "Tribe Lead informiert",
                                                  "Guild Master informiert"]]])),

    # ------------------------------------------------------------------ product labels of a key-value list
    case("fp-speaker-product-role-word", "txt, a product label with a role word (Bareos Director:) in a key-value list", [],
         ["Bareos Director", "Bareos Storage Daemon", OBS],
         txt("backup.txt", "Bareos Director: srv-bkp-01\nBareos Storage Daemon: srv-bkp-02\n%s: 20 TB\n" % OBS)),
    case("fp-speaker-product-repeated", "txt, a one-word product label that repeats (Zammad:)", [], ["Zammad", "Keycloak"],
         txt("tools.txt", "Zammad: Ticketsystem der IT\nZammad: Version 6.3\nKeycloak: SSO\n")),

    # ------------------------------------------------------------------ address and quantity shapes
    case("fp-address-am-standort", "txt, Am Standort 2 and Im Rechenzentrum 1 in prose", [],
         ["Am Standort 2", "Im Rechenzentrum 1"],
         txt("standorte.txt", "Am Standort 2 laufen 40 VMs. Im Rechenzentrum 1 steht die Primärseite.\n")),
    case("fp-quantity-five-digits", "txt, a five-digit quantity before a capitalised German noun", [],
         ["30000 Postfächer", "80000 Datensätze", "12000 Mandanten"],
         txt("mengen.txt", "Es gibt rund 30000 Postfächer und etwa 80000 Datensätze, dazu 12000 Mandanten.\n")),

    # ------------------------------------------------------------------ one-word headings and subjects
    case("fp-heading-one-word", "md, one-word headings: a product, a German compound, a short German noun", [],
         ["Zammad", "Netzplan", "Härtung"],
         txt("konzept.md", "# Zammad\n\nTicketsystem der IT.\n\n# Netzplan\n\nSiehe Anlage.\n\n# Härtung\n\nCIS Benchmarks.\n")),
    case("fp-subject-one-word", "eml, a one-word product subject", [], ["Tailscale"],
         eml("mail.eml", [("From", "it@" + fx.CUSTOMER_DOMAIN), ("To", "cloud@" + fx.CUSTOMER_DOMAIN), ("Subject", "Tailscale")],
             "Hallo zusammen,\n\nwir testen das Mesh-VPN im Lab, Ergebnis nächste Woche.\n\nViele Grüße\nIT\n")),

    # ------------------------------------------------------------------ teams as signers and addressees
    case("fp-signoff-team", "eml, a team as the signer under the closing (Das Projektteam, TCP Architekturteam)", [],
         ["Das Projektteam", "TCP Architekturteam"],
         eml("mail.eml", [("From", "projekt@" + fx.CUSTOMER_DOMAIN), ("To", "cloud@" + fx.CUSTOMER_DOMAIN), ("Subject", "Abnahme")],
             "Hallo zusammen,\n\ndie Abnahme ist erfolgt.\n\nViele Grüße\nDas Projektteam\n\n"
             "> Hallo zusammen,\n>\n> der Termin steht.\n>\n> Mit freundlichen Grüßen\n> TCP Architekturteam\n")),
    case("fp-hallo-department", "txt, a team or a department after Hallo (Hallo Netzwerkteam, Hallo Einkauf)", [],
         ["Netzwerkteam", "Einkauf"],
         txt("mails.txt", "Hallo Netzwerkteam,\n\nbitte die Firewall prüfen.\n\nHallo Einkauf,\n\nanbei die Bestellung.\n")),

    # ------------------------------------------------------------------ learned forms: one wipe spreads
    case("fp-url-label-learned", "md, a documentation URL whose host label is the product, the product in prose and in a host name", [],
         ["Zammad", "zammad-web-01"],
         txt("betrieb.md", "Doku: https://docs.zammad.example/de/\n\nZammad läuft auf zammad-web-01 und zammad-db-01.\n")),
    case("fp-mail-label-learned", "txt, a vendor support address whose host label is the product, the product in prose and in a host name", [],
         ["Bareos", "bareos-sd-01"],
         txt("support.txt", "Support: support@bareos.example\n\nBareos sichert auf bareos-sd-01.\n")),
    case("fp-label-learned-vendor", "txt, a vendor under Hersteller:, then the vendor in prose and in a host name", [],
         ["Kemp", "Kemp LoadMaster", "kemp-lb-01"],
         txt("lb.txt", "Hersteller: Kemp\nProdukt: LoadMaster\n\nDer Kemp LoadMaster steht als kemp-lb-01 in der DMZ.\n")),

    # ------------------------------------------------------------------ a product on a line of its own
    case("fp-twoline-tool-list", "pptx, a two-word title box, then a list of tools one per line without bullets; the extractor's own slide heading", [],
         ["Zammad", "Gitea", "Slide 1"],
         pptx_slides("toolchain.pptx", [[["Werkzeuge Übersicht"], ["Zammad", "Gitea", "Monitoring"]]]),
         note="Slide 1 is the heading the pptx extractor writes over every slide, not planted text"),
    case("fp-lone-line-tool-purpose", "pptx, a product over a role line (Bareos over Backup Administrator)", [],
         ["Bareos", "Netbox"],
         pptx_slides("rollen.pptx", [[["Bareos", "Backup Administrator"], ["Netbox", "Netzwerk Team"]]])),

    # ------------------------------------------------------------------ two products joined by a comma or a particle
    case("fp-inverted-two-tools", "txt, exactly two unknown product words around a comma", [], ["Gitea, Zammad"],
         txt("tools.txt", "Tools: Gitea, Zammad. Beide laufen im Cluster.\n")),
    case("fp-particle-umstieg", "txt, an unknown noun before von and a product after it (Umstieg von Gitea)", [],
         ["Umstieg von Gitea", "Wechsel von Bareos"],
         txt("plan.txt", "Umstieg von Gitea auf Forgejo im Q1. Wechsel von Bareos zu Veeam im Q2.\n")),

    # ------------------------------------------------------------------ by-lines with a board or a department
    case("fp-byline-board", "docx, Reviewed by Architecture Board and Erstellt von Abteilung Netze", [],
         ["Architecture Board", "Abteilung Netze"],
         docx_paragraphs("konzept.docx", ["Zielarchitektur", "Reviewed by Architecture Board", "Erstellt von Abteilung Netze"])),

    # ------------------------------------------------------------------ variations of run 1
    case("fp-sentence-opener-noun", "txt, a German sentence that opens with a capitalised preposition outside the lists, then a noun", [],
         ["Pro Subnetz", "Je Mandant", "Innerhalb Europas", "Zwischen Mandanten", "Statt Backup", "Sofern Kapazität"],
         txt("regeln.txt", "Pro Subnetz gilt eine Regel. Je Mandant gibt es ein Subnetz. Innerhalb Europas bleiben die Daten.\n"
             "Zwischen Mandanten gibt es keine Verbindung. Statt Backup nutzen wir Snapshots. Sofern Kapazität frei ist, starten wir.\n"),
         note="varies fp-title-extra-acronyms, where Pro Subnetz lost its noun as a by-catch"),
    case("fp-lettered-toc", "md, a table of contents of consecutive lettered lines (A. B. C. D.)", [],
         ["A. Einleitung", "B. Zielarchitektur", "C. Kostenschätzung", "D. Betrieb"],
         txt("inhalt.md", "Inhalt\n\nA. Einleitung\nB. Zielarchitektur\nC. Kostenschätzung\nD. Betrieb\n"),
         note="varies fp-lettered-sections: no sentence ends before the letter, so the z. B. window cannot save it"),
    case("fp-learned-report-word", "txt, a two-word run of unknown words whose second word is a word of the intake report (Candidates)", [],
         ["Rehosting Candidates"],
         txt("welle1.txt", "Rehosting Candidates\n\nDie Liste der VMs der ersten Welle.\n"),
         note="varies fp-migration-jargon, whose proposed run ended in an error of the public report"),
    case("fp-heading-participle-noun", "pptx, slide titles of a German participle and a known noun (Eingesetzte Werkzeuge)", [],
         ["Eingesetzte Werkzeuge", "Betroffene Systeme", "Verwendete Komponenten", "Empfohlene Maßnahmen", "Migrierte Workloads"],
         pptx_slides("titel.pptx", [[["Eingesetzte Werkzeuge"], ["Terraform und Ansible."]], [["Betroffene Systeme"], ["Zwölf VMs."]],
                                    [["Verwendete Komponenten"], ["Hub und Spoke."]], [["Empfohlene Maßnahmen"], ["Backup zuerst."]],
                                    [["Migrierte Workloads"], ["Welle 1."]]]),
         note="varies fp-twoline-tool-list, where the title word Eingesetzte was wiped as a by-catch"),
]
