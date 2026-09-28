"""Red team, dimension office: hidden and half-hidden parts of office containers.

Every case builds one office file with python-docx, openpyxl, python-pptx or by hand with zipfile, plants a
registered form or a structured value in one carrier and an innocent marker next to it. Same interface as
smoke.py. Nothing here is printed; the values live in this file only.
"""
import io
import re
import sys
import zipfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[2]))   # the repository that holds this pack
from tests import fixtures as fx  # noqa: E402

FULL, SHORT, ACRONYM, ENGLISH = fx.CUSTOMER_FORMS
PERSON, SURNAME = fx.PERSON_FORMS
ORG, ORG_SHORT = fx.ORG_FORMS
IBAN = "DE89 3704 0044 0532 0130 00"
IBAN_TIGHT = "DE89370400440532013000"
PHONE = "+49 30 1234567"
MAIL = "info@" + fx.CUSTOMER_DOMAIN

# --------------------------------------------------------------------------- helpers

W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
R_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
MC_NS = "http://schemas.openxmlformats.org/markup-compatibility/2006"
V_NS = "urn:schemas-microsoft-com:vml"
A_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"
C_NS = "http://schemas.openxmlformats.org/drawingml/2006/chart"
DGM_NS = "http://schemas.openxmlformats.org/drawingml/2006/diagram"
WP_NS = "http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing"
P_NS = "http://schemas.openxmlformats.org/presentationml/2006/main"
X_NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
REL_T = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/"


def esc(s):
    return s.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;").replace('"', "&quot;")


def zip_edit(path, changes, first=None):
    """Rewrite the zip at `path`: changes maps member name to bytes, to None (remove) or to a callable that
    takes the old bytes and returns the new bytes. `first` names a member written first and stored (odf mimetype)."""
    path = Path(path)
    with zipfile.ZipFile(path) as zf:
        members = {info.filename: zf.read(info.filename) for info in zf.infolist() if not info.is_dir()}
    for name, change in changes.items():
        if change is None:
            members.pop(name, None)
        elif callable(change):
            members[name] = change(members.get(name, b""))
        else:
            members[name] = change if isinstance(change, bytes) else change.encode("utf-8")
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as out:
        if first and first in members:
            out.writestr(zipfile.ZipInfo(first), members[first], compress_type=zipfile.ZIP_STORED)
        for name, data in members.items():
            if name == first:
                continue
            out.writestr(name, data)
    path.write_bytes(buf.getvalue())
    return path


def add_rel(rels_bytes, rid, rtype, target, external=False):
    text = rels_bytes.decode("utf-8") if rels_bytes else (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
        '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"></Relationships>')
    rel = '<Relationship Id="%s" Type="%s" Target="%s"%s/>' % (
        rid, rtype, esc(target), ' TargetMode="External"' if external else "")
    return text.replace("</Relationships>", rel + "</Relationships>").encode("utf-8")


def add_override(ct_bytes, part, ctype):
    text = ct_bytes.decode("utf-8")
    return text.replace("</Types>", '<Override PartName="%s" ContentType="%s"/></Types>' % (part, ctype)).encode("utf-8")


def add_default(ct_bytes, ext, ctype):
    text = ct_bytes.decode("utf-8")
    if 'Extension="%s"' % ext in text:
        return ct_bytes
    return text.replace("<Default ", '<Default Extension="%s" ContentType="%s"/><Default ' % (ext, ctype), 1).encode("utf-8")


def before_sectpr(doc_bytes, xml):
    text = doc_bytes.decode("utf-8")
    i = text.rfind("<w:sectPr")
    if i < 0:
        i = text.rfind("</w:body>")
    return (text[:i] + xml + text[i:]).encode("utf-8")


def w_p(text, rpr=""):
    return '<w:p><w:r>%s<w:t xml:space="preserve">%s</w:t></w:r></w:p>' % (rpr, esc(text))


def new_docx(inbox, name, paragraphs):
    import docx
    d = docx.Document()
    for p in paragraphs:
        d.add_paragraph(p)
    path = inbox / name
    d.save(str(path))
    return path


def new_xlsx(inbox, name, rows, sheet="Sheet1"):
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet
    for row in rows:
        ws.append(row)
    path = inbox / name
    wb.save(str(path))
    return path


def new_pptx(inbox, name, texts):
    import pptx
    from pptx.util import Inches
    prs = pptx.Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    for i, t in enumerate(texts):
        box = slide.shapes.add_textbox(Inches(1), Inches(1 + i), Inches(6), Inches(0.8))
        box.text_frame.text = t
    path = inbox / name
    prs.save(str(path))
    return path


ODF_NS = (
    'xmlns:office="urn:oasis:names:tc:opendocument:xmlns:office:1.0" '
    'xmlns:text="urn:oasis:names:tc:opendocument:xmlns:text:1.0" '
    'xmlns:table="urn:oasis:names:tc:opendocument:xmlns:table:1.0" '
    'xmlns:style="urn:oasis:names:tc:opendocument:xmlns:style:1.0" '
    'xmlns:number="urn:oasis:names:tc:opendocument:xmlns:datastyle:1.0" '
    'xmlns:draw="urn:oasis:names:tc:opendocument:xmlns:drawing:1.0" '
    'xmlns:presentation="urn:oasis:names:tc:opendocument:xmlns:presentation:1.0" '
    'xmlns:meta="urn:oasis:names:tc:opendocument:xmlns:meta:1.0" '
    'xmlns:config="urn:oasis:names:tc:opendocument:xmlns:config:1.0" '
    'xmlns:ooow="http://openoffice.org/2004/writer" '
    'xmlns:of="urn:oasis:names:tc:opendocument:xmlns:of:1.2" '
    'xmlns:dc="http://purl.org/dc/elements/1.1/" '
    'xmlns:xlink="http://www.w3.org/1999/xlink" office:version="1.2"'
)


def new_odf(inbox, name, kind, body_xml, extra=None, decls="", styles_xml=None, meta_xml=None, settings_xml=None):
    """An odt, ods or odp by hand. `body_xml` is the content of office:text, office:spreadsheet or
    office:presentation; `decls` goes before it inside office:body/office:<kind> (declarations)."""
    mime = {"odt": "application/vnd.oasis.opendocument.text",
            "ods": "application/vnd.oasis.opendocument.spreadsheet",
            "odp": "application/vnd.oasis.opendocument.presentation"}[kind]
    inner = {"odt": "text", "ods": "spreadsheet", "odp": "presentation"}[kind]
    content = ('<?xml version="1.0" encoding="UTF-8"?><office:document-content %s>%s<office:body><office:%s>%s%s'
               '</office:%s></office:body></office:document-content>'
               % (ODF_NS, extra or "", inner, decls, body_xml, inner))
    manifest = ('<?xml version="1.0" encoding="UTF-8"?><manifest:manifest '
                'xmlns:manifest="urn:oasis:names:tc:opendocument:xmlns:manifest:1.0" manifest:version="1.2">'
                '<manifest:file-entry manifest:full-path="/" manifest:media-type="%s"/>'
                '<manifest:file-entry manifest:full-path="content.xml" manifest:media-type="text/xml"/>'
                '%s%s%s</manifest:manifest>' % (
                    mime,
                    '<manifest:file-entry manifest:full-path="styles.xml" manifest:media-type="text/xml"/>' if styles_xml else "",
                    '<manifest:file-entry manifest:full-path="meta.xml" manifest:media-type="text/xml"/>' if meta_xml else "",
                    '<manifest:file-entry manifest:full-path="settings.xml" manifest:media-type="text/xml"/>' if settings_xml else ""))
    path = inbox / name
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as zf:
        zf.writestr(zipfile.ZipInfo("mimetype"), mime, compress_type=zipfile.ZIP_STORED)
        zf.writestr("content.xml", content)
        zf.writestr("META-INF/manifest.xml", manifest)
        if styles_xml:
            zf.writestr("styles.xml", '<?xml version="1.0" encoding="UTF-8"?><office:document-styles %s>%s</office:document-styles>' % (ODF_NS, styles_xml))
        if meta_xml:
            zf.writestr("meta.xml", '<?xml version="1.0" encoding="UTF-8"?><office:document-meta %s><office:meta>%s</office:meta></office:document-meta>' % (ODF_NS, meta_xml))
        if settings_xml:
            zf.writestr("settings.xml", '<?xml version="1.0" encoding="UTF-8"?><office:document-settings %s><office:settings>%s</office:settings></office:document-settings>' % (ODF_NS, settings_xml))
    return path


# --------------------------------------------------------------------------- docx builders


def docx_sdt_listitem(inbox):
    p = new_docx(inbox, "form.docx", ["the invented form, part one"])
    sdt = ('<w:sdt><w:sdtPr><w:alias w:val="%s"/><w:tag w:val="kunde"/><w:dropDownList>'
           '<w:listItem w:displayText="%s" w:value="1"/><w:listItem w:displayText="IBAN %s" w:value="2"/>'
           '<w:listItem w:displayText="%s" w:value="3"/></w:dropDownList></w:sdtPr>'
           '<w:sdtContent>%s</w:sdtContent></w:sdt>' % (esc("marker alpha list alias"), esc(FULL), esc(IBAN),
                                                        esc("marker beta list item"), w_p("choose an item")))
    return zip_edit(p, {"word/document.xml": lambda b: before_sectpr(b, sdt)})


def docx_fallback_only(inbox):
    p = new_docx(inbox, "alt.docx", ["the invented plan, alternate content"])
    alt = ('<w:p><mc:AlternateContent xmlns:mc="%s" xmlns:zz="http://example.invalid/zz/2026"><mc:Choice Requires="zz">'
           '<w:r><w:t>nothing to see in the choice branch</w:t></w:r></mc:Choice><mc:Fallback>'
           '<w:r><w:t xml:space="preserve">marker gamma fallback branch %s IBAN %s</w:t></w:r></mc:Fallback>'
           '</mc:AlternateContent></w:p>' % (MC_NS, esc(FULL), esc(IBAN)))
    return zip_edit(p, {"word/document.xml": lambda b: before_sectpr(b, alt)})


def docx_altchunk_html(inbox):
    p = new_docx(inbox, "chunk.docx", ["the invented plan with an imported chunk"])
    html = ("<html><body><p>marker delta imported chunk</p><p>%s</p><p>IBAN %s</p><p>%s</p></body></html>"
            % (esc(FULL), IBAN, esc(MAIL)))
    return zip_edit(p, {
        "word/afchunk.html": html.encode("utf-8"),
        "word/document.xml": lambda b: before_sectpr(b, '<w:altChunk r:id="rIdChunk1"/>'),
        "word/_rels/document.xml.rels": lambda b: add_rel(b, "rIdChunk1", REL_T + "aFChunk", "afchunk.html"),
        "[Content_Types].xml": lambda b: add_default(b, "html", "text/html"),
    })


def docx_vml_textpath(inbox):
    p = new_docx(inbox, "wordart.docx", ["the invented plan with word art"])
    art = ('<w:p><w:r><w:pict><v:shape xmlns:v="%s" id="wa1" type="#_x0000_t136" style="width:400pt;height:40pt" '
           'fillcolor="black"><v:shadow on="t"/><v:textpath style="font-family:&quot;Arial&quot;" fitpath="t" '
           'string="%s"/></v:shape></w:pict></w:r></w:p>' % (V_NS, esc("marker epsilon word art %s IBAN %s" % (FULL, IBAN))))
    return zip_edit(p, {"word/document.xml": lambda b: before_sectpr(b, art)})


def chart_xml(title, series_name, cats, iban):
    return ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
            '<c:chartSpace xmlns:c="%s" xmlns:a="%s" xmlns:r="%s"><c:chart><c:title><c:tx><c:rich><a:bodyPr/><a:p><a:r>'
            '<a:t>%s</a:t></a:r></a:p></c:rich></c:tx></c:title><c:plotArea><c:layout/><c:barChart><c:barDir val="col"/>'
            '<c:ser><c:idx val="0"/><c:order val="0"/><c:tx><c:strRef><c:f>Sheet1!$B$1</c:f><c:strCache><c:ptCount val="1"/>'
            '<c:pt idx="0"><c:v>%s</c:v></c:pt></c:strCache></c:strRef></c:tx><c:cat><c:strRef><c:f>Sheet1!$A$2:$A$3</c:f>'
            '<c:strCache><c:ptCount val="2"/><c:pt idx="0"><c:v>%s</c:v></c:pt><c:pt idx="1"><c:v>%s</c:v></c:pt></c:strCache>'
            '</c:strRef></c:cat><c:val><c:numRef><c:f>Sheet1!$B$2:$B$3</c:f><c:numCache><c:formatCode>General</c:formatCode>'
            '<c:ptCount val="2"/><c:pt idx="0"><c:v>1</c:v></c:pt><c:pt idx="1"><c:v>2</c:v></c:pt></c:numCache></c:numRef>'
            '</c:val></c:ser><c:axId val="1"/><c:axId val="2"/></c:barChart><c:catAx><c:axId val="1"/><c:scaling/>'
            '<c:crossAx val="2"/></c:catAx><c:valAx><c:axId val="2"/><c:scaling/><c:crossAx val="1"/></c:valAx>'
            '</c:plotArea></c:chart></c:chartSpace>'
            % (C_NS, A_NS, R_NS, esc(title), esc(series_name), esc(cats[0]), esc(cats[1])))


CHART_CT = "application/vnd.openxmlformats-officedocument.drawingml.chart+xml"


def docx_chart(inbox):
    p = new_docx(inbox, "chart.docx", ["the invented plan with a chart"])
    drawing = ('<w:p><w:r><w:drawing><wp:inline xmlns:wp="%s"><wp:extent cx="5000000" cy="3000000"/>'
               '<wp:docPr id="7" name="Chart 7"/><a:graphic xmlns:a="%s"><a:graphicData uri="%s">'
               '<c:chart xmlns:c="%s" xmlns:r="%s" r:id="rIdChart1"/></a:graphicData></a:graphic></wp:inline>'
               '</w:drawing></w:r></w:p>' % (WP_NS, A_NS, C_NS, C_NS, R_NS))
    return zip_edit(p, {
        "word/charts/chart1.xml": chart_xml("marker zeta chart title " + FULL, SHORT, ["marker eta category", "IBAN " + IBAN], IBAN).encode("utf-8"),
        "word/document.xml": lambda b: before_sectpr(b, drawing),
        "word/_rels/document.xml.rels": lambda b: add_rel(b, "rIdChart1", REL_T + "chart", "charts/chart1.xml"),
        "[Content_Types].xml": lambda b: add_override(b, "/word/charts/chart1.xml", CHART_CT),
    })


def diagram_data_xml(texts):
    pts = "".join('<dgm:pt modelId="{%08d-0000-0000-0000-000000000000}" type="node"><dgm:prSet/><dgm:spPr/>'
                  '<dgm:t><a:bodyPr/><a:p><a:r><a:t>%s</a:t></a:r></a:p></dgm:t></dgm:pt>' % (i + 1, esc(t))
                  for i, t in enumerate(texts))
    return ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?><dgm:dataModel xmlns:dgm="%s" xmlns:a="%s">'
            '<dgm:ptLst>%s</dgm:ptLst><dgm:cxnLst/></dgm:dataModel>' % (DGM_NS, A_NS, pts))


DGM_CT = "application/vnd.openxmlformats-officedocument.drawingml.diagramData+xml"


def docx_smartart(inbox):
    p = new_docx(inbox, "smartart.docx", ["the invented plan with a diagram"])
    drawing = ('<w:p><w:r><w:drawing><wp:inline xmlns:wp="%s"><wp:extent cx="5000000" cy="3000000"/>'
               '<wp:docPr id="8" name="Diagram 8"/><a:graphic xmlns:a="%s"><a:graphicData uri="%s">'
               '<dgm:relIds xmlns:dgm="%s" xmlns:r="%s" r:dm="rIdDm1"/></a:graphicData></a:graphic></wp:inline>'
               '</w:drawing></w:r></w:p>' % (WP_NS, A_NS, DGM_NS, DGM_NS, R_NS))
    return zip_edit(p, {
        "word/diagrams/data1.xml": diagram_data_xml(["marker theta diagram node", FULL, "IBAN " + IBAN]).encode("utf-8"),
        "word/document.xml": lambda b: before_sectpr(b, drawing),
        "word/_rels/document.xml.rels": lambda b: add_rel(b, "rIdDm1", REL_T + "diagramData", "diagrams/data1.xml"),
        "[Content_Types].xml": lambda b: add_override(b, "/word/diagrams/data1.xml", DGM_CT),
    })


def docx_table_caption(inbox):
    import docx
    d = docx.Document()
    d.add_paragraph("the invented plan with a table")
    t = d.add_table(rows=1, cols=2)
    t.rows[0].cells[0].text = "stage"
    t.rows[0].cells[1].text = "two"
    path = inbox / "tblalt.docx"
    d.save(str(path))
    cap = '<w:tblCaption w:val="%s"/><w:tblDescription w:val="%s"/>' % (esc("marker iota table caption " + FULL), esc("IBAN " + IBAN))
    return zip_edit(path, {"word/document.xml": lambda b: b.replace(b"</w:tblPr>", cap.encode("utf-8") + b"</w:tblPr>", 1)})


def docx_field_instr(inbox):
    p = new_docx(inbox, "field.docx", ["the invented plan with a field"])
    fld = ('<w:p><w:r><w:fldChar w:fldCharType="begin"/></w:r><w:r><w:instrText xml:space="preserve"> QUOTE "%s" </w:instrText>'
           '</w:r><w:r><w:fldChar w:fldCharType="separate"/></w:r><w:r><w:fldChar w:fldCharType="end"/></w:r></w:p>'
           % esc("marker kappa quoted field %s IBAN %s" % (FULL, IBAN)))
    return zip_edit(p, {
        "word/document.xml": lambda b: before_sectpr(b, fld),
        "word/settings.xml": lambda b: b.replace(b"<w:zoom", b'<w:updateFields w:val="true"/><w:zoom', 1),
    })


def docx_first_page_header(inbox):
    import docx
    d = docx.Document()
    d.add_paragraph("the invented plan with a first page header")
    sec = d.sections[0]
    sec.different_first_page_header_footer = True
    sec.first_page_header.paragraphs[0].text = "marker lambda first page header %s" % FULL
    sec.even_page_header.paragraphs[0].text = "marker mu even page header IBAN %s" % IBAN
    path = inbox / "hdr.docx"
    d.save(str(path))
    return path


def docx_hidden_runs(inbox):
    p = new_docx(inbox, "hidden.docx", ["the invented plan with hidden runs"])
    vanish = w_p("marker nu hidden run " + FULL, '<w:rPr><w:vanish/></w:rPr>')
    tiny = w_p("marker xi white tiny run IBAN " + IBAN, '<w:rPr><w:color w:val="FFFFFF"/><w:sz w:val="2"/></w:rPr>')
    return zip_edit(p, {"word/document.xml": lambda b: before_sectpr(b, vanish + tiny)})


GLOSSARY_CT = "application/vnd.openxmlformats-officedocument.wordprocessingml.document.glossary+xml"


def docx_glossary(inbox):
    p = new_docx(inbox, "glossary.docx", ["the invented plan with building blocks"])
    gl = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?><w:glossaryDocument xmlns:w="%s"><w:docParts><w:docPart>'
          '<w:docPartPr><w:name w:val="%s"/><w:category><w:name w:val="General"/><w:gallery w:val="placeholder"/></w:category>'
          '</w:docPartPr><w:docPartBody>%s%s</w:docPartBody></w:docPart></w:docParts></w:glossaryDocument>'
          % (W_NS, esc("marker omicron quick part"), w_p("signature block of " + FULL), w_p("IBAN " + IBAN)))
    return zip_edit(p, {
        "word/glossary/document.xml": gl.encode("utf-8"),
        "word/_rels/document.xml.rels": lambda b: add_rel(b, "rIdGl1", REL_T + "glossaryDocument", "glossary/document.xml"),
        "[Content_Types].xml": lambda b: add_override(b, "/word/glossary/document.xml", GLOSSARY_CT),
    })


def docx_customxml(inbox):
    p = new_docx(inbox, "cx.docx", ["the invented plan with custom xml"])
    item = ('<?xml version="1.0" encoding="UTF-8"?><contract xmlns="http://example.invalid/contract"><note>marker pi custom xml</note>'
            '<customer>%s</customer><iban>%s</iban></contract>' % (esc(FULL), IBAN))
    return zip_edit(p, {
        "customXml/item2.xml": item.encode("utf-8"),
        "customXml/itemProps2.xml": ('<?xml version="1.0" encoding="UTF-8"?><ds:datastoreItem ds:itemID="{6C3C8BC8-F283-45AE' '-878A-BAB7291924A2}" '
                                     'xmlns:ds="http://schemas.openxmlformats.org/officeDocument/2006/customXml"/>').encode("utf-8"),
        "customXml/_rels/item2.xml.rels": add_rel(b"", "rId1", REL_T + "customXmlProps", "itemProps2.xml"),
        "word/_rels/document.xml.rels": lambda b: add_rel(b, "rIdCx2", REL_T + "customXml", "../customXml/item2.xml"),
        "[Content_Types].xml": lambda b: add_override(b, "/customXml/itemProps2.xml", "application/vnd.openxmlformats-officedocument.customXmlProperties+xml"),
    })


def docx_docvars(inbox):
    p = new_docx(inbox, "vars.docx", ["the invented plan with document variables"])
    dv = '<w:docVars><w:docVar w:name="kunde" w:val="%s"/><w:docVar w:name="konto" w:val="%s"/></w:docVars>' % (
        esc("marker rho document variable " + FULL), esc("IBAN " + IBAN))
    return zip_edit(p, {"word/settings.xml": lambda b: b.replace(b"<w:zoom", dv.encode("utf-8") + b"<w:zoom", 1)})


def docx_people(inbox):
    p = new_docx(inbox, "people.docx", ["the invented plan with people"])
    ppl = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?><w15:people xmlns:w15="http://schemas.microsoft.com/office/word/2012/wordml">'
           '<w15:person w15:author="%s"><w15:presenceInfo w15:providerId="AD" w15:userId="marker sigma people part %s"/></w15:person>'
           '</w15:people>' % (esc(PERSON), esc(MAIL)))
    return zip_edit(p, {
        "word/people.xml": ppl.encode("utf-8"),
        "word/_rels/document.xml.rels": lambda b: add_rel(b, "rIdPpl", "http://schemas.microsoft.com/office/2011/relationships/people", "people.xml"),
        "[Content_Types].xml": lambda b: add_override(b, "/word/people.xml", "application/vnd.openxmlformats-officedocument.wordprocessingml.people+xml"),
    })


def docx_custom_prop_blob(inbox):
    import base64
    p = new_docx(inbox, "props.docx", ["the invented plan with custom properties"])
    blob = base64.b64encode(("marker tau blob property %s IBAN %s" % (FULL, IBAN)).encode("utf-16-le")).decode()
    custom = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?><Properties xmlns="http://schemas.openxmlformats.org/officeDocument/2006/custom-properties" '
              'xmlns:vt="http://schemas.openxmlformats.org/officeDocument/2006/docPropsVTypes">'
              '<property fmtid="{D5CDD505-2E9C-101B' '-9397-08002B2CF9AE}" pid="2" name="Kontakt"><vt:blob>%s</vt:blob></property>'
              '<property fmtid="{D5CDD505-2E9C-101B' '-9397-08002B2CF9AE}" pid="3" name="Stufe"><vt:i4>2</vt:i4></property>'
              '<property fmtid="{D5CDD505-2E9C-101B' '-9397-08002B2CF9AE}" pid="4" name="%s"><vt:bool>true</vt:bool></property>'
              '</Properties>' % (blob, esc("marker upsilon property name " + SHORT)))
    return zip_edit(p, {
        "docProps/custom.xml": custom.encode("utf-8"),
        "_rels/.rels": lambda b: add_rel(b, "rIdCust", REL_T + "custom-properties", "docProps/custom.xml"),
        "[Content_Types].xml": lambda b: add_override(b, "/docProps/custom.xml", "application/vnd.openxmlformats-officedocument.custom-properties+xml"),
    })


def docx_embedded_font(inbox):
    p = new_docx(inbox, "font.docx", ["the invented plan with an embedded font"])
    # a fake obfuscated font part: the bytes carry a name table in plain text after the 32 obfuscated bytes
    font = b"\x00" * 32 + b"name\x00\x00" + ("marker phi font family %s IBAN %s" % (FULL, IBAN)).encode("utf-16-be") + b"\x00" * 64
    ft = ('<w:font w:name="%s"><w:embedRegular r:id="rIdFont1" w:fontKey="{00000000-0000-0000' '-0000-000000000001}"/></w:font>'
          % esc(SHORT + " Sans"))
    return zip_edit(p, {
        "word/fonts/font1.odttf": font,
        "word/fontTable.xml": lambda b: b.replace(b"</w:fonts>", ft.encode("utf-8") + b"</w:fonts>", 1),
        "word/_rels/fontTable.xml.rels": add_rel(b"", "rIdFont1", REL_T + "font", "fonts/font1.odttf"),
        "[Content_Types].xml": lambda b: add_default(b, "odttf", "application/vnd.openxmlformats-officedocument.obfuscatedFont"),
    })


def docx_split_cells_tail(inbox):
    import docx
    d = docx.Document()
    d.add_paragraph("marker chi address table of the invented plan")
    t = d.add_table(rows=2, cols=2)
    t.rows[0].cells[0].text = "contact"
    t.rows[0].cells[1].text = "street"
    t.rows[1].cells[0].text = SURNAME[:8]
    t.rows[1].cells[1].text = SURNAME[8:] + "str. 5"
    path = inbox / "addr.docx"
    d.save(str(path))
    return path


def docx_split_hostname(inbox):
    low = SHORT.lower()
    return new_docx(inbox, "hosts.docx", ["marker psi host list of the invented plan",
                                          "host %s-%s01 is down" % (low[:3], low[3:]),
                                          "host %s.%s02 is up" % (low[:3], low[3:])])


def docx_rlo(inbox):
    return new_docx(inbox, "rtl.docx", ["marker omega right to left of the invented plan",
                                        "\u202e" + FULL[::-1] + "\u202c"])


# --------------------------------------------------------------------------- xlsx builders


def xlsx_rph(inbox):
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = "stage two of the invented sheet"
    path = inbox / "rph.xlsx"
    wb.save(str(path))
    shared = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?><sst xmlns="%s" count="1" uniqueCount="1"><si><t>stage two of the invented sheet</t>'
              '<rPh sb="0" eb="5"><t>%s</t></rPh><rPh sb="6" eb="9"><t>%s</t></rPh><phoneticPr fontId="1" type="noConversion"/></si></sst>'
              % (X_NS, esc("marker one phonetic " + FULL), esc("IBAN " + IBAN)))

    def sheet(b):
        return b.replace(b'<c r="A1" t="inlineStr"><is><t>stage two of the invented sheet</t></is></c>', b'<c r="A1" t="s"><v>0</v></c>')
    return zip_edit(path, {
        "xl/sharedStrings.xml": shared.encode("utf-8"),
        "xl/worksheets/sheet1.xml": sheet,
        "xl/_rels/workbook.xml.rels": lambda b: add_rel(b, "rIdSst", REL_T + "sharedStrings", "sharedStrings.xml"),
        "[Content_Types].xml": lambda b: add_override(b, "/xl/sharedStrings.xml", "application/vnd.openxmlformats-officedocument.spreadsheetml.sharedStrings+xml"),
    })


def xlsx_numfmt(inbox):
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = "stage two of the invented sheet"
    ws["A2"] = 5
    ws["A2"].number_format = '"marker two number format %s "0' % FULL
    ws["A3"] = 7
    ws["A3"].number_format = '"IBAN %s "0' % IBAN
    path = inbox / "numfmt.xlsx"
    wb.save(str(path))
    return path


def xlsx_dv_list(inbox):
    import openpyxl
    from openpyxl.worksheet.datavalidation import DataValidation
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = "stage two of the invented sheet"
    dv = DataValidation(type="list", formula1='"marker three validation,%s,IBAN %s"' % (FULL, IBAN), allow_blank=True)
    dv.prompt = "marker four prompt " + SHORT
    dv.promptTitle = "customer"
    dv.error = "IBAN " + IBAN
    ws.add_data_validation(dv)
    dv.add("B1")
    path = inbox / "dv.xlsx"
    wb.save(str(path))
    return path


def xlsx_cf_formula(inbox):
    import openpyxl
    from openpyxl.formatting.rule import FormulaRule
    from openpyxl.styles import PatternFill
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = "stage two of the invented sheet"
    ws.conditional_formatting.add("A1:A9", FormulaRule(formula=['A1="marker five conditional %s"' % FULL], fill=PatternFill("solid", fgColor="FFFF00")))
    ws.conditional_formatting.add("B1:B9", FormulaRule(formula=['B1="IBAN %s"' % IBAN], fill=PatternFill("solid", fgColor="FF0000")))
    path = inbox / "cf.xlsx"
    wb.save(str(path))
    return path


def xlsx_hyperlink_tooltip(inbox):
    import openpyxl
    from openpyxl.worksheet.hyperlink import Hyperlink
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = "stage two of the invented sheet"
    ws["A2"] = "portal"
    ws["A2"].hyperlink = Hyperlink(ref="A2", target="https://portal.example/offer", tooltip="marker six tooltip %s IBAN %s" % (FULL, IBAN),
                                   display="marker seven display " + SHORT)
    path = inbox / "link.xlsx"
    wb.save(str(path))
    return path


def xlsx_extlink_cache(inbox):
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = "stage two of the invented sheet"
    path = inbox / "ext.xlsx"
    wb.save(str(path))
    ext = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?><externalLink xmlns="%s" xmlns:r="%s"><externalBook r:id="rId1">'
           '<sheetNames><sheetName val="%s"/></sheetNames><sheetDataSet><sheetData sheetId="0"><row r="1">'
           '<cell r="A1" t="str"><v>%s</v></cell><cell r="B1" t="str"><v>IBAN %s</v></cell></row></sheetData></sheetDataSet>'
           '</externalBook></externalLink>' % (X_NS, R_NS, esc("marker eight external sheet"), esc(FULL), IBAN))

    def sheet(b):
        cells = b'<c r="B1"><f>[1]%s!$A$1</f></c><c r="C1"><f>[1]%s!$B$1</f></c>' % (b"marker eight external sheet", b"marker eight external sheet")
        return b.replace(b"</row>", cells + b"</row>", 1)

    def workbook(b):
        return b.replace(b"<definedNames/>", b'<definedNames/><externalReferences><externalReference xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" r:id="rIdExt1"/></externalReferences>', 1)
    return zip_edit(path, {
        "xl/externalLinks/externalLink1.xml": ext.encode("utf-8"),
        "xl/externalLinks/_rels/externalLink1.xml.rels": add_rel(b"", "rId1", REL_T + "externalLinkPath", "file:///Q:/plan/source.xlsx", external=True),
        "xl/worksheets/sheet1.xml": sheet,
        "xl/workbook.xml": workbook,
        "xl/_rels/workbook.xml.rels": lambda b: add_rel(b, "rIdExt1", REL_T + "externalLink", "externalLinks/externalLink1.xml"),
        "[Content_Types].xml": lambda b: add_override(b, "/xl/externalLinks/externalLink1.xml", "application/vnd.openxmlformats-officedocument.spreadsheetml.externalLink+xml"),
    })


def xlsx_chart_cache(inbox):
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = "stage two of the invented sheet"
    path = inbox / "chart.xlsx"
    wb.save(str(path))
    drawing = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?><xdr:wsDr xmlns:xdr="http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing" '
               'xmlns:a="%s" xmlns:r="%s" xmlns:c="%s"><xdr:twoCellAnchor><xdr:from><xdr:col>2</xdr:col><xdr:colOff>0</xdr:colOff><xdr:row>2</xdr:row><xdr:rowOff>0</xdr:rowOff></xdr:from>'
               '<xdr:to><xdr:col>10</xdr:col><xdr:colOff>0</xdr:colOff><xdr:row>20</xdr:row><xdr:rowOff>0</xdr:rowOff></xdr:to><xdr:graphicFrame macro="">'
               '<xdr:nvGraphicFramePr><xdr:cNvPr id="2" name="Chart 1"/><xdr:cNvGraphicFramePr/></xdr:nvGraphicFramePr><xdr:xfrm><a:off x="0" y="0"/><a:ext cx="0" cy="0"/></xdr:xfrm>'
               '<a:graphic><a:graphicData uri="%s"><c:chart r:id="rId1"/></a:graphicData></a:graphic></xdr:graphicFrame><xdr:clientData/></xdr:twoCellAnchor></xdr:wsDr>'
               % (A_NS, R_NS, C_NS, C_NS))
    return zip_edit(path, {
        "xl/charts/chart1.xml": chart_xml("marker nine chart title " + SHORT, "marker ten series " + FULL, ["IBAN " + IBAN, "second"], IBAN).encode("utf-8"),
        "xl/drawings/drawing1.xml": drawing.encode("utf-8"),
        "xl/drawings/_rels/drawing1.xml.rels": add_rel(b"", "rId1", REL_T + "chart", "../charts/chart1.xml"),
        "xl/worksheets/_rels/sheet1.xml.rels": add_rel(b"", "rIdDr1", REL_T + "drawing", "../drawings/drawing1.xml"),
        "xl/worksheets/sheet1.xml": lambda b: b.replace(b"</worksheet>", b'<drawing xmlns:r="%s" r:id="rIdDr1"/></worksheet>' % R_NS.encode()),
        "[Content_Types].xml": lambda b: add_override(add_override(b, "/xl/charts/chart1.xml", CHART_CT), "/xl/drawings/drawing1.xml", "application/vnd.openxmlformats-officedocument.drawing+xml"),
    })


def xlsx_vml_textbox(inbox):
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = "stage two of the invented sheet"
    path = inbox / "vml.xlsx"
    wb.save(str(path))
    vml = ('<xml xmlns:v="%s" xmlns:o="urn:schemas-microsoft-com:office:office" xmlns:x="urn:schemas-microsoft-com:office:excel">'
           '<o:shapelayout v:ext="edit"><o:idmap v:ext="edit" data="1"/></o:shapelayout>'
           '<v:shapetype id="_x0000_t202" coordsize="21600,21600" o:spt="202" path="m,l,21600r21600,l21600,xe"><v:stroke joinstyle="miter"/>'
           '<v:path gradientshapeok="t" o:connecttype="rect"/></v:shapetype>'
           '<v:shape id="_x0000_s1025" type="#_x0000_t202" style="position:absolute;margin-left:100pt;margin-top:20pt;width:300pt;height:60pt;z-index:1" fillcolor="#ffffe1" o:insetmode="auto">'
           '<v:fill color2="#ffffe1"/><v:textbox style="mso-direction-alt:auto"><div style="text-align:left">marker eleven legacy text box %s<br>IBAN %s</div></v:textbox>'
           '<x:ClientData ObjectType="Rect"><x:MoveWithCells/><x:SizeWithCells/><x:Anchor>2, 15, 1, 10, 6, 15, 5, 4</x:Anchor><x:AutoFill>False</x:AutoFill></x:ClientData>'
           '</v:shape></xml>' % (V_NS, esc(FULL), IBAN))
    return zip_edit(path, {
        "xl/drawings/vmlDrawing1.vml": vml.encode("utf-8"),
        "xl/worksheets/_rels/sheet1.xml.rels": add_rel(b"", "rIdVml1", REL_T + "vmlDrawing", "../drawings/vmlDrawing1.vml"),
        "xl/worksheets/sheet1.xml": lambda b: b.replace(b"</worksheet>", b'<legacyDrawing xmlns:r="%s" r:id="rIdVml1"/></worksheet>' % R_NS.encode()),
        "[Content_Types].xml": lambda b: add_default(b, "vml", "application/vnd.openxmlformats-officedocument.vmlDrawing"),
    })


def xlsx_connection(inbox):
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = "stage two of the invented sheet"
    path = inbox / "conn.xlsx"
    wb.save(str(path))
    conn = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?><connections xmlns="%s"><connection id="1" name="%s" description="%s" type="1" refreshedVersion="6" savePassword="0">'
            '<dbPr connection="Provider=SQLOLEDB.1;Data Source=%s-db01.%s;User ID=svc_report;Initial Catalog=Plan" command="select * from dbo.plan where iban = &apos;%s&apos;"/>'
            '</connection></connections>' % (X_NS, esc("marker twelve connection name"), esc(SHORT + " reporting"), SHORT.lower(), fx.CUSTOMER_DOMAIN, IBAN_TIGHT))
    return zip_edit(path, {
        "xl/connections.xml": conn.encode("utf-8"),
        "xl/_rels/workbook.xml.rels": lambda b: add_rel(b, "rIdConn", REL_T + "connections", "connections.xml"),
        "[Content_Types].xml": lambda b: add_override(b, "/xl/connections.xml", "application/vnd.openxmlformats-officedocument.spreadsheetml.connections+xml"),
    })


def xlsx_pivot_records(inbox):
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = "stage two of the invented sheet"
    path = inbox / "pivot.xlsx"
    wb.save(str(path))
    defn = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?><pivotCacheDefinition xmlns="%s" xmlns:r="%s" r:id="rId1" refreshedBy="%s" recordCount="2">'
            '<cacheSource type="worksheet"><worksheetSource ref="A1:B3" sheet="%s"/></cacheSource><cacheFields count="1">'
            '<cacheField name="%s" numFmtId="0"><sharedItems count="2"><s v="%s"/><s v="IBAN %s"/></sharedItems></cacheField></cacheFields>'
            '</pivotCacheDefinition>' % (X_NS, R_NS, esc(PERSON), esc("marker thirteen source sheet"), esc("marker fourteen cache field"), esc(FULL), IBAN))
    recs = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?><pivotCacheRecords xmlns="%s" count="2"><r><x v="0"/></r><r><s v="%s"/></r></pivotCacheRecords>'
            % (X_NS, esc("marker fifteen record " + ORG)))
    return zip_edit(path, {
        "xl/pivotCache/pivotCacheDefinition1.xml": defn.encode("utf-8"),
        "xl/pivotCache/pivotCacheRecords1.xml": recs.encode("utf-8"),
        "xl/pivotCache/_rels/pivotCacheDefinition1.xml.rels": add_rel(b"", "rId1", REL_T + "pivotCacheRecords", "pivotCacheRecords1.xml"),
        "xl/workbook.xml": lambda b: b.replace(b"</workbook>", b'<pivotCaches><pivotCache xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" cacheId="1" r:id="rIdPc1"/></pivotCaches></workbook>'),
        "xl/_rels/workbook.xml.rels": lambda b: add_rel(b, "rIdPc1", REL_T + "pivotCacheDefinition", "pivotCache/pivotCacheDefinition1.xml"),
        "[Content_Types].xml": lambda b: add_override(add_override(b, "/xl/pivotCache/pivotCacheDefinition1.xml", "application/vnd.openxmlformats-officedocument.spreadsheetml.pivotCacheDefinition+xml"),
                                                      "/xl/pivotCache/pivotCacheRecords1.xml", "application/vnd.openxmlformats-officedocument.spreadsheetml.pivotCacheRecords+xml"),
    })


def xlsx_veryhidden(inbox):
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = "stage two of the invented sheet"
    ws2 = wb.create_sheet("marker sixteen hidden tab")
    ws2["A1"] = "marker seventeen very hidden cell " + FULL
    ws2["A2"] = "IBAN " + IBAN
    ws2.sheet_state = "veryHidden"
    path = inbox / "vh.xlsx"
    wb.save(str(path))
    return path


def xlsx_cached_str(inbox):
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = "stage two of the invented sheet"
    path = inbox / "cached.xlsx"
    wb.save(str(path))
    cells = ('<c r="B1" t="str"><f>C1&amp;D1</f><v>%s</v></c><c r="B2" t="str"><f>C2&amp;D2</f><v>marker eighteen cached IBAN %s</v></c>' % (esc(FULL), IBAN))
    return zip_edit(path, {"xl/worksheets/sheet1.xml": lambda b: b.replace(b"</row>", cells.encode("utf-8") + b"</row>", 1)})


def xlsx_split_rows_tail(inbox):
    return new_xlsx(inbox, "addr.xlsx", [[SURNAME[:8]], [SURNAME[8:] + "str. 5"], ["marker nineteen address rows"]])


def xlsx_split_cells_tail(inbox):
    low = SHORT.lower()
    return new_xlsx(inbox, "hosts.xlsx", [["host", "state"], [low[:3], low[3:] + "01"], ["marker twenty host cells", "up"]])


# --------------------------------------------------------------------------- pptx builders


def pptx_fallback_only(inbox):
    p = new_pptx(inbox, "alt.pptx", ["slide one of the invented deck"])
    sp = ('<mc:AlternateContent xmlns:mc="%s" xmlns:zz="http://example.invalid/zz/2026"><mc:Choice Requires="zz">'
          '<p:sp><p:nvSpPr><p:cNvPr id="20" name="box"/><p:cNvSpPr txBox="1"/><p:nvPr/></p:nvSpPr><p:spPr/><p:txBody><a:bodyPr/><a:p><a:r><a:t>nothing in the choice branch</a:t></a:r></a:p></p:txBody></p:sp>'
          '</mc:Choice><mc:Fallback><p:sp><p:nvSpPr><p:cNvPr id="21" name="box"/><p:cNvSpPr txBox="1"/><p:nvPr/></p:nvSpPr><p:spPr/><p:txBody><a:bodyPr/>'
          '<a:p><a:r><a:t>%s</a:t></a:r></a:p><a:p><a:r><a:t>IBAN %s</a:t></a:r></a:p></p:txBody></p:sp></mc:Fallback></mc:AlternateContent>'
          % (MC_NS, esc("marker aa fallback shape " + FULL), IBAN))
    return zip_edit(p, {"ppt/slides/slide1.xml": lambda b: b.replace(b"</p:spTree>", sp.encode("utf-8") + b"</p:spTree>", 1)})


def pptx_tags(inbox):
    p = new_pptx(inbox, "tags.pptx", ["slide one of the invented deck"])
    tag = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?><p:tagLst xmlns:p="%s"><p:tag name="KUNDE" val="%s"/><p:tag name="KONTO" val="IBAN %s"/></p:tagLst>'
           % (P_NS, esc("marker bb slide tag " + FULL), IBAN))
    return zip_edit(p, {
        "ppt/tags/tag1.xml": tag.encode("utf-8"),
        "ppt/slides/slide1.xml": lambda b: b.replace(b"</p:spTree>", b'</p:spTree><p:custDataLst><p:tags r:id="rIdTag1"/></p:custDataLst>', 1),
        "ppt/slides/_rels/slide1.xml.rels": lambda b: add_rel(b, "rIdTag1", REL_T + "tags", "../tags/tag1.xml"),
        "[Content_Types].xml": lambda b: add_override(b, "/ppt/tags/tag1.xml", "application/vnd.openxmlformats-officedocument.presentationml.tags+xml"),
    })


def pptx_section_names(inbox):
    p = new_pptx(inbox, "sections.pptx", ["slide one of the invented deck"])
    ext = ('<p:extLst><p:ext uri="{521415D9-36F7-43E2' '-AB2F-B90AF26B5E84}"><p14:sectionLst xmlns:p14="http://schemas.microsoft.com/office/powerpoint/2010/main">'
           '<p14:section name="%s" id="{1F7A0F2B-5A3E-4C1D' '-9B2E-000000000001}"><p14:sldIdLst><p14:sldId id="256"/></p14:sldIdLst></p14:section>'
           '<p14:section name="IBAN %s" id="{1F7A0F2B-5A3E-4C1D' '-9B2E-000000000002}"><p14:sldIdLst/></p14:section></p14:sectionLst></p:ext></p:extLst>'
           % (esc("marker cc section " + FULL), IBAN))
    return zip_edit(p, {"ppt/presentation.xml": lambda b: b.replace(b"</p:presentation>", ext.encode("utf-8") + b"</p:presentation>")})


def pptx_smartart(inbox):
    p = new_pptx(inbox, "smartart.pptx", ["slide one of the invented deck"])
    frame = ('<p:graphicFrame><p:nvGraphicFramePr><p:cNvPr id="30" name="Diagram 30"/><p:cNvGraphicFramePr/><p:nvPr/></p:nvGraphicFramePr>'
             '<p:xfrm><a:off x="1000000" y="1000000"/><a:ext cx="6000000" cy="4000000"/></p:xfrm><a:graphic><a:graphicData uri="%s">'
             '<dgm:relIds xmlns:dgm="%s" r:dm="rIdDm1"/></a:graphicData></a:graphic></p:graphicFrame>' % (DGM_NS, DGM_NS))
    return zip_edit(p, {
        "ppt/diagrams/data1.xml": diagram_data_xml(["marker dd diagram node", FULL, "IBAN " + IBAN]).encode("utf-8"),
        "ppt/slides/slide1.xml": lambda b: b.replace(b"</p:spTree>", frame.encode("utf-8") + b"</p:spTree>", 1),
        "ppt/slides/_rels/slide1.xml.rels": lambda b: add_rel(b, "rIdDm1", REL_T + "diagramData", "../diagrams/data1.xml"),
        "[Content_Types].xml": lambda b: add_override(b, "/ppt/diagrams/data1.xml", DGM_CT),
    })


def pptx_chart(inbox):
    p = new_pptx(inbox, "chart.pptx", ["slide one of the invented deck"])
    frame = ('<p:graphicFrame><p:nvGraphicFramePr><p:cNvPr id="31" name="Chart 31"/><p:cNvGraphicFramePr/><p:nvPr/></p:nvGraphicFramePr>'
             '<p:xfrm><a:off x="1000000" y="1000000"/><a:ext cx="6000000" cy="4000000"/></p:xfrm><a:graphic><a:graphicData uri="%s">'
             '<c:chart xmlns:c="%s" r:id="rIdChart1"/></a:graphicData></a:graphic></p:graphicFrame>' % (C_NS, C_NS))
    return zip_edit(p, {
        "ppt/charts/chart1.xml": chart_xml("marker ee chart title " + FULL, "marker ff series " + SHORT, ["IBAN " + IBAN, "second"], IBAN).encode("utf-8"),
        "ppt/slides/slide1.xml": lambda b: b.replace(b"</p:spTree>", frame.encode("utf-8") + b"</p:spTree>", 1),
        "ppt/slides/_rels/slide1.xml.rels": lambda b: add_rel(b, "rIdChart1", REL_T + "chart", "../charts/chart1.xml"),
        "[Content_Types].xml": lambda b: add_override(b, "/ppt/charts/chart1.xml", CHART_CT),
    })


def pptx_slide_oddname(inbox):
    p = new_pptx(inbox, "deck.pptx", ["marker gg renamed slide part " + FULL, "IBAN " + IBAN])
    with zipfile.ZipFile(p) as zf:
        slide = zf.read("ppt/slides/slide1.xml")
        rels = zf.read("ppt/slides/_rels/slide1.xml.rels")
    return zip_edit(p, {
        "ppt/slides/slide1.xml": None,
        "ppt/slides/_rels/slide1.xml.rels": None,
        "ppt/slides/deck-a.xml": slide,
        "ppt/slides/_rels/deck-a.xml.rels": rels,
        "ppt/_rels/presentation.xml.rels": lambda b: b.replace(b"slides/slide1.xml", b"slides/deck-a.xml"),
        "[Content_Types].xml": lambda b: b.replace(b"/ppt/slides/slide1.xml", b"/ppt/slides/deck-a.xml"),
    })


def pptx_notes_master(inbox):
    import pptx
    prs = pptx.Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.shapes.add_textbox(0, 0, 100, 100).text_frame.text = "slide one of the invented deck"
    slide.notes_slide.notes_text_frame.text = "marker hh speaker notes"
    path = inbox / "notes.pptx"
    prs.save(str(path))
    sp = ('<p:sp><p:nvSpPr><p:cNvPr id="40" name="box"/><p:cNvSpPr txBox="1"/><p:nvPr/></p:nvSpPr><p:spPr/><p:txBody><a:bodyPr/>'
          '<a:p><a:r><a:t>%s</a:t></a:r></a:p><a:p><a:r><a:t>IBAN %s</a:t></a:r></a:p></p:txBody></p:sp>' % (esc("marker ii notes master " + FULL), IBAN))
    names = [n for n in zipfile.ZipFile(path).namelist() if n.startswith("ppt/notesMasters/notesMaster")]
    return zip_edit(path, {names[0]: lambda b: b.replace(b"</p:spTree>", sp.encode("utf-8") + b"</p:spTree>", 1)})


def pptx_split_shapes_tail(inbox):
    return new_pptx(inbox, "addr.pptx", ["marker jj address on the invented slide", SURNAME[:8], SURNAME[8:] + "str. 5"])


# --------------------------------------------------------------------------- odf builders


def odt_userfield_decl(inbox):
    decls = ('<text:user-field-decls><text:user-field-decl office:value-type="string" office:string-value="%s" text:name="Kunde"/>'
             '<text:user-field-decl office:value-type="string" office:string-value="IBAN %s" text:name="Konto"/></text:user-field-decls>'
             % (esc("marker odt one user field " + FULL), IBAN))
    body = '<text:p>stage two of the invented text</text:p><text:p>customer: <text:user-field-get text:name="Kunde"/> account: <text:user-field-get text:name="Konto"/></text:p>'
    return new_odf(inbox, "field.odt", "odt", body, decls=decls)


def odt_hidden_text(inbox):
    body = ('<text:p>stage two of the invented text</text:p><text:p>price for <text:hidden-text text:condition="ooow:FALSE" '
            'text:string-value="%s" text:is-hidden="false"/> paid to <text:hidden-text text:condition="ooow:FALSE" text:string-value="IBAN %s" text:is-hidden="false"/></text:p>'
            % (esc("marker odt two hidden text " + FULL), IBAN))
    return new_odf(inbox, "hidden.odt", "odt", body)


def odt_conditional_text(inbox):
    body = ('<text:p>stage two of the invented text</text:p><text:p>partner: <text:conditional-text text:condition="ooow:1 == 1" '
            'text:string-value-if-true="%s" text:string-value-if-false="nobody" text:current-value="true"/> account <text:conditional-text '
            'text:condition="ooow:1 == 1" text:string-value-if-true="IBAN %s" text:string-value-if-false="none" text:current-value="true"/></text:p>'
            % (esc("marker odt three conditional " + FULL), IBAN))
    return new_odf(inbox, "cond.odt", "odt", body)


def odt_meta_userdefined(inbox):
    meta = ('<meta:user-defined meta:name="Kunde">%s</meta:user-defined><meta:user-defined meta:name="%s">IBAN %s</meta:user-defined>'
            % (esc("marker odt four user meta " + FULL), esc("marker odt five meta name " + SHORT), IBAN))
    return new_odf(inbox, "meta.odt", "odt", "<text:p>stage two of the invented text</text:p>", meta_xml=meta)


def odt_settings(inbox):
    settings = ('<config:config-item-set config:name="ooo:configuration-settings"><config:config-item config:name="PrinterName" config:type="string">%s</config:config-item>'
                '<config:config-item config:name="PrinterSetup" config:type="base64Binary">%s</config:config-item></config:config-item-set>'
                % (esc("marker odt six printer " + FULL), "SUJBTiBERTg5IDM3MDQgMDA0NCAwNTMyIDAxMzAgMDA="))
    return new_odf(inbox, "settings.odt", "odt", "<text:p>stage two of the invented text</text:p>", settings_xml=settings)


def odt_tracked_deletion(inbox):
    body = ('<text:tracked-changes><text:changed-region text:id="ct1"><text:deletion><office:change-info><dc:creator>%s</dc:creator>'
            '<dc:date>2026-09-20T10:00:00</dc:date></office:change-info><text:p>%s</text:p><text:p>IBAN %s</text:p></text:deletion></text:changed-region></text:tracked-changes>'
            '<text:p>stage two of the invented text<text:change text:change-id="ct1"/> continues</text:p>'
            % (esc(PERSON), esc("marker odt seven deleted " + FULL), IBAN))
    return new_odf(inbox, "tracked.odt", "odt", body)


def ods_validation(inbox):
    decls = ('<table:content-validations><table:content-validation table:name="val1" table:condition="of:cell-content-is-in-list(&quot;%s&quot;;&quot;IBAN %s&quot;)" '
             'table:allow-empty-cell="true" table:display-list="unsorted"><table:help-message table:title="hint" table:display="true"><text:p>%s</text:p></table:help-message>'
             '</table:content-validation></table:content-validations>' % (esc(FULL), IBAN, esc("marker ods one help message " + SHORT)))
    body = ('<table:table table:name="Plan"><table:table-column/><table:table-row><table:table-cell office:value-type="string"><text:p>stage two of the invented sheet</text:p></table:table-cell></table:table-row>'
            '<table:table-row><table:table-cell table:content-validation-name="val1" office:value-type="string"><text:p>pick</text:p></table:table-cell></table:table-row></table:table>')
    return new_odf(inbox, "val.ods", "ods", body, decls=decls)


def ods_numfmt(inbox):
    extra = ('<office:automatic-styles><number:number-style style:name="N77"><number:text>%s </number:text><number:number number:decimal-places="0" number:min-integer-digits="1"/></number:number-style>'
             '<number:number-style style:name="N78"><number:text>IBAN %s </number:text><number:number number:decimal-places="0" number:min-integer-digits="1"/></number:number-style>'
             '<style:style style:name="ce1" style:family="table-cell" style:data-style-name="N77"/><style:style style:name="ce2" style:family="table-cell" style:data-style-name="N78"/></office:automatic-styles>'
             % (esc("marker ods two number text " + FULL), IBAN))
    body = ('<table:table table:name="Plan"><table:table-column/><table:table-row><table:table-cell office:value-type="string"><text:p>stage two of the invented sheet</text:p></table:table-cell></table:table-row>'
            '<table:table-row><table:table-cell table:style-name="ce1" office:value-type="float" office:value="5"/></table:table-row>'
            '<table:table-row><table:table-cell table:style-name="ce2" office:value-type="float" office:value="7"/></table:table-row></table:table>')
    return new_odf(inbox, "numfmt.ods", "ods", body, extra=extra)


def odp_notes(inbox):
    body = ('<draw:page draw:name="page1"><draw:frame presentation:class="title"><draw:text-box><text:p>slide one of the invented deck</text:p></draw:text-box></draw:frame>'
            '<presentation:notes><draw:page-thumbnail/><draw:frame presentation:class="notes"><draw:text-box><text:p>%s</text:p><text:p>IBAN %s</text:p></draw:text-box></draw:frame></presentation:notes></draw:page>'
            % (esc("marker odp one speaker notes " + FULL), IBAN))
    return new_odf(inbox, "notes.odp", "odp", body)


# --------------------------------------------------------------------------- cases

CASES = [
    # docx carriers
    {"id": "docx-sdt-listitem", "cls": "name", "carrier": "docx content control drop-down list items (w:listItem displayText) and alias",
     "values": [FULL, IBAN], "visible": ["marker alpha list alias", "marker beta list item"], "build": docx_sdt_listitem},
    {"id": "docx-fallback-only", "cls": "name", "carrier": "docx mc:AlternateContent whose Choice requires an unknown namespace, text in Fallback only",
     "values": [FULL, IBAN], "visible": ["marker gamma fallback branch"], "build": docx_fallback_only},
    {"id": "docx-altchunk-html", "cls": "name", "carrier": "docx w:altChunk to an html part", "values": [FULL, IBAN, MAIL],
     "visible": ["marker delta imported chunk"], "build": docx_altchunk_html},
    {"id": "docx-vml-textpath", "cls": "name", "carrier": "docx legacy WordArt v:textpath string attribute", "values": [FULL, IBAN],
     "visible": ["marker epsilon word art"], "build": docx_vml_textpath},
    {"id": "docx-chart", "cls": "name", "carrier": "docx chart part: title rich text, series name and category string cache",
     "values": [FULL, SHORT, IBAN], "visible": ["marker zeta chart title", "marker eta category"], "build": docx_chart},
    {"id": "docx-smartart", "cls": "name", "carrier": "docx SmartArt diagrams/data1.xml node text", "values": [FULL, IBAN],
     "visible": ["marker theta diagram node"], "build": docx_smartart},
    {"id": "docx-table-caption", "cls": "name", "carrier": "docx table alt text (w:tblCaption, w:tblDescription)", "values": [FULL, IBAN],
     "visible": ["marker iota table caption"], "build": docx_table_caption},
    {"id": "docx-field-instr", "cls": "name", "carrier": "docx field instruction (QUOTE) with an empty result and updateFields on open",
     "values": [FULL, IBAN], "visible": ["marker kappa quoted field"], "build": docx_field_instr},
    {"id": "docx-first-even-header", "cls": "name", "carrier": "docx first-page and even-page headers", "values": [FULL, IBAN],
     "visible": ["marker lambda first page header", "marker mu even page header"], "build": docx_first_page_header},
    {"id": "docx-hidden-runs", "cls": "name", "carrier": "docx w:vanish run and a white 1pt run in the body", "values": [FULL, IBAN],
     "visible": ["marker nu hidden run", "marker xi white tiny run"], "build": docx_hidden_runs},
    {"id": "docx-glossary", "cls": "name", "carrier": "docx glossary document (quick parts)", "values": [FULL, IBAN],
     "visible": ["marker omicron quick part"], "build": docx_glossary},
    {"id": "docx-customxml", "cls": "name", "carrier": "docx customXml/item part", "values": [FULL, IBAN],
     "visible": ["marker pi custom xml"], "build": docx_customxml},
    {"id": "docx-docvars", "cls": "name", "carrier": "docx settings.xml document variables", "values": [FULL, IBAN],
     "visible": ["marker rho document variable"], "build": docx_docvars},
    {"id": "docx-people", "cls": "name", "carrier": "docx people.xml author and presence user id", "values": [PERSON, MAIL],
     "visible": ["marker sigma people part"], "build": docx_people},
    {"id": "docx-custom-prop-blob", "cls": "name", "carrier": "docx custom property of type vt:blob (base64 of utf-16) and a property name",
     "values": [FULL, IBAN, SHORT], "visible": ["marker tau blob property", "marker upsilon property name"], "build": docx_custom_prop_blob},
    {"id": "docx-embedded-font", "cls": "name", "carrier": "docx embedded font part (word/fonts/font1.odttf) with a name table",
     "values": [FULL, IBAN], "visible": ["marker phi font family"], "build": docx_embedded_font},
    # docx shapes in the body
    {"id": "docx-split-cells-tail", "cls": "name", "carrier": "docx table: surname split over two cells, the second half glued to a suffix (street)",
     "values": [SURNAME], "visible": ["marker chi address table"], "build": docx_split_cells_tail},
    {"id": "docx-split-hostname", "cls": "name", "carrier": "docx paragraph: short form split by a hyphen or dot inside a host name with a digit tail",
     "values": [SHORT], "visible": ["marker psi host list"], "build": docx_split_hostname},
    {"id": "docx-rlo", "cls": "name", "carrier": "docx paragraph: the form written reversed behind a right-to-left override",
     "values": [FULL], "visible": ["marker omega right to left"], "build": docx_rlo},
    # xlsx carriers
    {"id": "xlsx-rph", "cls": "name", "carrier": "xlsx shared string phonetic runs (rPh)", "values": [FULL, IBAN],
     "visible": ["marker one phonetic"], "build": xlsx_rph},
    {"id": "xlsx-numfmt", "cls": "name", "carrier": "xlsx number format with literal text, cell shows the text", "values": [FULL, IBAN],
     "visible": ["marker two number format"], "build": xlsx_numfmt},
    {"id": "xlsx-dv-list", "cls": "name", "carrier": "xlsx data validation list, prompt and error text", "values": [FULL, SHORT, IBAN],
     "visible": ["marker three validation", "marker four prompt"], "build": xlsx_dv_list},
    {"id": "xlsx-cf-formula", "cls": "name", "carrier": "xlsx conditional formatting formula", "values": [FULL, IBAN],
     "visible": ["marker five conditional"], "build": xlsx_cf_formula},
    {"id": "xlsx-hyperlink-tooltip", "cls": "name", "carrier": "xlsx hyperlink tooltip and display", "values": [FULL, SHORT, IBAN],
     "visible": ["marker six tooltip", "marker seven display"], "build": xlsx_hyperlink_tooltip},
    {"id": "xlsx-extlink-cache", "cls": "name", "carrier": "xlsx external link cached sheet data (shown after recalculation on load)",
     "values": [FULL, IBAN], "visible": ["marker eight external sheet"], "build": xlsx_extlink_cache},
    {"id": "xlsx-chart-cache", "cls": "name", "carrier": "xlsx chart title (rich), series name and categories in string caches",
     "values": [SHORT, FULL, IBAN], "visible": ["marker nine chart title", "marker ten series"], "build": xlsx_chart_cache},
    {"id": "xlsx-vml-textbox", "cls": "name", "carrier": "xlsx legacy vml drawing text box", "values": [FULL, IBAN],
     "visible": ["marker eleven legacy text box"], "build": xlsx_vml_textbox},
    {"id": "xlsx-connection", "cls": "name", "carrier": "xlsx data connection: name, description, connection string, command",
     "values": [SHORT, fx.CUSTOMER_DOMAIN, IBAN_TIGHT], "visible": ["marker twelve connection name"], "build": xlsx_connection},
    {"id": "xlsx-pivot-records", "cls": "name", "carrier": "xlsx pivot cache definition (refreshedBy, shared items) and records",
     "values": [PERSON, FULL, ORG, IBAN], "visible": ["marker thirteen source sheet", "marker fourteen cache field", "marker fifteen record"],
     "build": xlsx_pivot_records},
    {"id": "xlsx-veryhidden", "cls": "name", "carrier": "xlsx very hidden sheet cells and its name", "values": [FULL, IBAN],
     "visible": ["marker sixteen hidden tab", "marker seventeen very hidden cell"], "build": xlsx_veryhidden},
    {"id": "xlsx-cached-str", "cls": "name", "carrier": "xlsx formula cell with a cached string value", "values": [FULL, IBAN],
     "visible": ["marker eighteen cached"], "build": xlsx_cached_str},
    {"id": "xlsx-split-rows-tail", "cls": "name", "carrier": "xlsx surname split over two rows (Markdown header separator between), second half glued to a suffix",
     "values": [SURNAME], "visible": ["marker nineteen address rows"], "build": xlsx_split_rows_tail},
    {"id": "xlsx-split-cells-tail", "cls": "name", "carrier": "xlsx short form split over two cells with a digit tail", "values": [SHORT],
     "visible": ["marker twenty host cells"], "build": xlsx_split_cells_tail},
    # pptx carriers
    {"id": "pptx-fallback-only", "cls": "name", "carrier": "pptx mc:AlternateContent Fallback shape (Choice requires an unknown namespace)",
     "values": [FULL, IBAN], "visible": ["marker aa fallback shape"], "build": pptx_fallback_only},
    {"id": "pptx-tags", "cls": "name", "carrier": "pptx slide tags part (p:tag val)", "values": [FULL, IBAN],
     "visible": ["marker bb slide tag"], "build": pptx_tags},
    {"id": "pptx-section-names", "cls": "name", "carrier": "pptx section names (p14:sectionLst)", "values": [FULL, IBAN],
     "visible": ["marker cc section"], "build": pptx_section_names},
    {"id": "pptx-smartart", "cls": "name", "carrier": "pptx SmartArt diagrams/data1.xml", "values": [FULL, IBAN],
     "visible": ["marker dd diagram node"], "build": pptx_smartart},
    {"id": "pptx-chart", "cls": "name", "carrier": "pptx chart part: title, series name, categories", "values": [FULL, SHORT, IBAN],
     "visible": ["marker ee chart title", "marker ff series"], "build": pptx_chart},
    {"id": "pptx-slide-oddname", "cls": "name", "carrier": "pptx slide part not named slideN.xml (referenced by rels only)",
     "values": [FULL, IBAN], "visible": ["marker gg renamed slide part"], "build": pptx_slide_oddname},
    {"id": "pptx-notes-master", "cls": "name", "carrier": "pptx notes master text box", "values": [FULL, IBAN],
     "visible": ["marker hh speaker notes", "marker ii notes master"], "build": pptx_notes_master},
    {"id": "pptx-split-shapes-tail", "cls": "name", "carrier": "pptx surname split over two text boxes, second half glued to a suffix",
     "values": [SURNAME], "visible": ["marker jj address on the invented slide"], "build": pptx_split_shapes_tail},
    # odf carriers
    {"id": "odt-userfield-decl", "cls": "name", "carrier": "odt user field declaration (office:string-value), empty field in the body",
     "values": [FULL, IBAN], "visible": ["marker odt one user field"], "build": odt_userfield_decl},
    {"id": "odt-hidden-text", "cls": "name", "carrier": "odt text:hidden-text string-value with a false condition (shown)", "values": [FULL, IBAN],
     "visible": ["marker odt two hidden text"], "build": odt_hidden_text},
    {"id": "odt-conditional-text", "cls": "name", "carrier": "odt text:conditional-text string-value-if-true", "values": [FULL, IBAN],
     "visible": ["marker odt three conditional"], "build": odt_conditional_text},
    {"id": "odt-meta-userdefined", "cls": "name", "carrier": "odt meta.xml user-defined fields (value and name)", "values": [FULL, SHORT, IBAN],
     "visible": ["marker odt four user meta", "marker odt five meta name"], "build": odt_meta_userdefined},
    {"id": "odt-settings", "cls": "name", "carrier": "odt settings.xml config items (printer name, base64 printer setup)", "values": [FULL, IBAN],
     "visible": ["marker odt six printer"], "build": odt_settings},
    {"id": "odt-tracked-deletion", "cls": "name", "carrier": "odt tracked deletion text and its creator", "values": [PERSON, FULL, IBAN],
     "visible": ["marker odt seven deleted"], "build": odt_tracked_deletion},
    {"id": "ods-validation", "cls": "name", "carrier": "ods content validation list condition and help message", "values": [FULL, SHORT, IBAN],
     "visible": ["marker ods one help message"], "build": ods_validation},
    {"id": "ods-numfmt", "cls": "name", "carrier": "ods number style text, cell without cached display text", "values": [FULL, IBAN],
     "visible": ["marker ods two number text"], "build": ods_numfmt},
    {"id": "odp-notes", "cls": "name", "carrier": "odp presentation notes", "values": [FULL, IBAN],
     "visible": ["marker odp one speaker notes"], "build": odp_notes},
]


# --------------------------------------------------------------------------- round 2

HRB = "HRB 123456"
VAT = "DE123456789"
PHONE_NAT = "030 1234567"


def docx_split_softbreak_tail(inbox):
    p = new_docx(inbox, "softbreak.docx", ["marker r2a soft break address"])
    para = ('<w:p><w:r><w:t>%s</w:t></w:r><w:r><w:br/></w:r><w:r><w:t xml:space="preserve">%sstr. 5</w:t></w:r></w:p>'
            % (esc(SURNAME[:8]), esc(SURNAME[8:])))
    return zip_edit(p, {"word/document.xml": lambda b: before_sectpr(b, para)})


def docx_split_cells_tail_sep(inbox):
    import docx
    d = docx.Document()
    d.add_paragraph("marker r2b control, tail separated")
    t = d.add_table(rows=1, cols=2)
    t.rows[0].cells[0].text = SURNAME[:8]
    t.rows[0].cells[1].text = SURNAME[8:] + " str. 5"
    path = inbox / "control.docx"
    d.save(str(path))
    return path


def xlsx_split_one_digit(inbox):
    low = SHORT.lower()
    return new_xlsx(inbox, "one.xlsx", [["host", "state"], [low[:3], low[3:] + "1"], ["marker r2c one digit tail", "up"]])


def docx_split_head_glued(inbox):
    low = SHORT.lower()
    return new_docx(inbox, "head.docx", ["marker r2d head glued", "asset 01%s-%s retired" % (low[:3], low[3:])])


def xlsx_split_hyphen_cell_end(inbox):
    return new_xlsx(inbox, "hyph.xlsx", [[SURNAME[:8] + "-"], [SURNAME[8:] + "str. 5"], ["marker r2e hyphen at the cell end"]])


def docx_phone_columns(inbox):
    import docx
    d = docx.Document()
    d.add_paragraph("marker r2f contact table")
    t = d.add_table(rows=2, cols=3)
    for c, v in zip(t.rows[0].cells, ["Name", "Vorwahl", "Rufnummer"]):
        c.text = v
    for c, v in zip(t.rows[1].cells, ["Zentrale", "030", "1234567"]):
        c.text = v
    path = inbox / "contacts.docx"
    d.save(str(path))
    return path


def xlsx_phone_columns_intl(inbox):
    return new_xlsx(inbox, "contacts.xlsx", [["Land", "Vorwahl", "Rufnummer"], ["+49", "30", "1234567"], ["marker r2g contact sheet", "", ""]])


def xlsx_iban_rows(inbox):
    return new_xlsx(inbox, "bank.xlsx", [["Konto"], ["DE89 3704 0044 0532"], ["0130 00"], ["marker r2h bank rows"]])


def docx_hrb_cells(inbox):
    import docx
    d = docx.Document()
    d.add_paragraph("marker r2i register table")
    t = d.add_table(rows=2, cols=3)
    for c, v in zip(t.rows[0].cells, ["Registergericht", fx.PLACE_FORMS[0], ""]):
        c.text = v
    for c, v in zip(t.rows[1].cells, ["Registernummer", "HRB", "123456"]):
        c.text = v
    path = inbox / "register.docx"
    d.save(str(path))
    return path


def docx_vat_cells(inbox):
    import docx
    d = docx.Document()
    d.add_paragraph("marker r2j vat table")
    t = d.add_table(rows=1, cols=3)
    for c, v in zip(t.rows[0].cells, ["USt-IdNr.", "DE", "123456789"]):
        c.text = v
    path = inbox / "vat.docx"
    d.save(str(path))
    return path


def docx_footnotes_renamed(inbox):
    p = new_docx(inbox, "fn.docx", ["marker r2k body with a note"])
    fn = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?><w:footnotes xmlns:w="%s"><w:footnote w:type="separator" w:id="-1"><w:p><w:r><w:separator/></w:r></w:p></w:footnote>'
          '<w:footnote w:type="continuationSeparator" w:id="0"><w:p><w:r><w:continuationSeparator/></w:r></w:p></w:footnote>'
          '<w:footnote w:id="1">%s%s</w:footnote></w:footnotes>' % (W_NS, w_p("marker r2l footnote text " + FULL), w_p("IBAN " + IBAN)))
    ref = '<w:p><w:r><w:t>see the note</w:t></w:r><w:r><w:rPr><w:vertAlign w:val="superscript"/></w:rPr><w:footnoteReference w:id="1"/></w:r></w:p>'
    return zip_edit(p, {
        "word/fn.xml": fn.encode("utf-8"),
        "word/document.xml": lambda b: before_sectpr(b, ref),
        "word/_rels/document.xml.rels": lambda b: add_rel(b, "rIdFn1", REL_T + "footnotes", "fn.xml"),
        "[Content_Types].xml": lambda b: add_override(b, "/word/fn.xml", "application/vnd.openxmlformats-officedocument.wordprocessingml.footnotes+xml"),
    })


def xlsx_sst_renamed(inbox):
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = "placeholder"
    path = inbox / "sst.xlsx"
    wb.save(str(path))
    sst = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?><sst xmlns="%s" count="3" uniqueCount="3"><si><t>marker r2m shared string part</t></si>'
           '<si><t>%s</t></si><si><t>IBAN %s</t></si></sst>' % (X_NS, esc(FULL), IBAN))

    def sheet(b):
        return b.replace(b'<c r="A1" t="inlineStr"><is><t>placeholder</t></is></c>',
                         b'<c r="A1" t="s"><v>0</v></c><c r="B1" t="s"><v>1</v></c><c r="C1" t="s"><v>2</v></c>')
    return zip_edit(path, {
        "xl/strings.xml": sst.encode("utf-8"),
        "xl/worksheets/sheet1.xml": sheet,
        "xl/_rels/workbook.xml.rels": lambda b: add_rel(b, "rIdSst", REL_T + "sharedStrings", "strings.xml"),
        "[Content_Types].xml": lambda b: add_override(b, "/xl/strings.xml", "application/vnd.openxmlformats-officedocument.spreadsheetml.sharedStrings+xml"),
    })


def ods_shapes(inbox):
    body = ('<table:table table:name="Plan"><table:shapes><draw:frame draw:z-index="0"><draw:text-box><text:p>%s</text:p><text:p>IBAN %s</text:p></draw:text-box></draw:frame></table:shapes>'
            '<table:table-column/><table:table-row><table:table-cell office:value-type="string"><text:p>stage two of the invented sheet</text:p></table:table-cell></table:table-row></table:table>'
            % (esc("marker r2n sheet text box " + FULL), IBAN))
    return new_odf(inbox, "shapes.ods", "ods", body)


def odt_forms(inbox):
    extra = ""
    decls = ('<office:forms form:automatic-focus="false" form:apply-design-mode="false" xmlns:form="urn:oasis:names:tc:opendocument:xmlns:form:1.0">'
             '<form:form form:name="f1"><form:text form:name="kunde" form:control-implementation="ooo:com.sun.star.form.component.TextField" xml:id="c1" '
             'form:current-value="%s"/><form:text form:name="konto" form:control-implementation="ooo:com.sun.star.form.component.TextField" xml:id="c2" '
             'form:current-value="IBAN %s"/></form:form></office:forms>' % (esc("marker r2o form control " + FULL), IBAN))
    body = ('<text:p>stage two of the invented text</text:p><text:p><draw:control draw:control="c1"/></text:p><text:p><draw:control draw:control="c2"/></text:p>')
    return new_odf(inbox, "forms.odt", "odt", body, extra=extra, decls=decls)


def ods_embedded_chart(inbox):
    body = ('<table:table table:name="Plan"><table:shapes><draw:frame draw:z-index="0"><draw:object xlink:href="./Object 1" xlink:type="simple" xlink:show="embed" xlink:actuate="onLoad"/></draw:frame></table:shapes>'
            '<table:table-column/><table:table-row><table:table-cell office:value-type="string"><text:p>stage two of the invented sheet</text:p></table:table-cell></table:table-row></table:table>')
    path = new_odf(inbox, "chartobj.ods", "ods", body)
    chart = ('<?xml version="1.0" encoding="UTF-8"?><office:document-content %s xmlns:chart="urn:oasis:names:tc:opendocument:xmlns:chart:1.0"><office:body><office:chart>'
             '<chart:chart chart:class="chart:bar"><chart:title><text:p>%s</text:p></chart:title><chart:plot-area><chart:series chart:label-cell-address="Plan.B1">'
             '<chart:data-point/></chart:series></chart:plot-area></chart:chart><table:table table:name="local-table"><table:table-header-rows><table:table-row>'
             '<table:table-cell office:value-type="string"><text:p>IBAN %s</text:p></table:table-cell></table:table-row></table:table-header-rows></table:table>'
             '</office:chart></office:body></office:document-content>' % (ODF_NS, esc("marker r2p chart title " + FULL), IBAN))

    def manifest(b):
        return b.replace(b"</manifest:manifest>", b'<manifest:file-entry manifest:full-path="Object 1/" manifest:media-type="application/vnd.oasis.opendocument.chart"/>'
                         b'<manifest:file-entry manifest:full-path="Object 1/content.xml" manifest:media-type="text/xml"/></manifest:manifest>')
    return zip_edit(path, {"Object 1/content.xml": chart.encode("utf-8"), "META-INF/manifest.xml": manifest}, first="mimetype")


def docx_numbering_lvltext(inbox):
    p = new_docx(inbox, "num.docx", ["marker r2q numbered list"])
    with zipfile.ZipFile(p) as zf:
        numbering = zf.read("word/numbering.xml").decode("utf-8")
    num_id = re.search(r'<w:num w:numId="(\d+)"', numbering).group(1)
    para = ('<w:p><w:pPr><w:numPr><w:ilvl w:val="0"/><w:numId w:val="%s"/></w:numPr></w:pPr><w:r><w:t>first item</w:t></w:r></w:p>' % num_id)

    def edit_num(b):
        t = b.decode("utf-8")
        return re.sub(r'<w:lvlText w:val="[^"]*"/>', '<w:lvlText w:val="%s"/>' % esc("marker r2r level text %s IBAN %s %%1." % (FULL, IBAN)), t, count=1).encode("utf-8")
    return zip_edit(p, {"word/document.xml": lambda b: before_sectpr(b, para), "word/numbering.xml": edit_num})


CASES += [
    {"id": "docx-split-softbreak-tail", "cls": "name", "carrier": "docx paragraph: surname split by a soft line break (w:br), second half glued to a suffix",
     "values": [SURNAME], "visible": ["marker r2a soft break address"], "build": docx_split_softbreak_tail},
    {"id": "docx-split-cells-tail-sep", "cls": "name", "carrier": "control: docx table, surname split over two cells, suffix separated by a space",
     "values": [SURNAME], "visible": ["marker r2b control, tail separated"], "build": docx_split_cells_tail_sep},
    {"id": "xlsx-split-one-digit", "cls": "name", "carrier": "xlsx short form split over two cells, one digit glued to the second half",
     "values": [SHORT], "visible": ["marker r2c one digit tail"], "build": xlsx_split_one_digit},
    {"id": "docx-split-head-glued", "cls": "name", "carrier": "docx paragraph: digits glued in front of the first half, hyphen between the halves",
     "values": [SHORT], "visible": ["marker r2d head glued"], "build": docx_split_head_glued},
    {"id": "xlsx-split-hyphen-cell-end", "cls": "name", "carrier": "xlsx surname hyphenated at the end of one cell, continued with a suffix in the next row",
     "values": [SURNAME], "visible": ["marker r2e hyphen at the cell end"], "build": xlsx_split_hyphen_cell_end},
    {"id": "docx-phone-columns", "cls": "phone", "carrier": "docx contact table: area code and number in two columns",
     "values": [PHONE_NAT], "visible": ["marker r2f contact table"], "build": docx_phone_columns},
    {"id": "xlsx-phone-columns-intl", "cls": "phone", "carrier": "xlsx contact sheet: country code, area code and number in three columns",
     "values": [PHONE], "visible": ["marker r2g contact sheet"], "build": xlsx_phone_columns_intl},
    {"id": "xlsx-iban-rows", "cls": "iban", "carrier": "xlsx IBAN typed over two rows of one column", "values": [IBAN],
     "visible": ["marker r2h bank rows"], "build": xlsx_iban_rows},
    {"id": "docx-hrb-cells", "cls": "hrb", "carrier": "docx register table: HRB and the number in two columns", "values": [HRB, fx.PLACE_FORMS[0]],
     "visible": ["marker r2i register table"], "build": docx_hrb_cells},
    {"id": "docx-vat-cells", "cls": "vat", "carrier": "docx table: DE and the nine digits in two columns", "values": [VAT],
     "visible": ["marker r2j vat table"], "build": docx_vat_cells},
    {"id": "docx-footnotes-renamed", "cls": "name", "carrier": "docx footnotes part named word/fn.xml, reached through the rels only",
     "values": [FULL, IBAN], "visible": ["marker r2k body with a note", "marker r2l footnote text"], "build": docx_footnotes_renamed},
    {"id": "xlsx-sst-renamed", "cls": "name", "carrier": "xlsx shared strings part named xl/strings.xml, reached through the rels only",
     "values": [FULL, IBAN], "visible": ["marker r2m shared string part"], "build": xlsx_sst_renamed},
    {"id": "ods-shapes", "cls": "name", "carrier": "ods text box in table:shapes", "values": [FULL, IBAN],
     "visible": ["marker r2n sheet text box"], "build": ods_shapes},
    {"id": "odt-forms", "cls": "name", "carrier": "odt form controls (office:forms current-value)", "values": [FULL, IBAN],
     "visible": ["marker r2o form control"], "build": odt_forms},
    {"id": "ods-embedded-chart", "cls": "name", "carrier": "ods embedded chart object (Object 1/content.xml) title and local table",
     "values": [FULL, IBAN], "visible": ["marker r2p chart title"], "build": ods_embedded_chart},
    {"id": "docx-numbering-lvltext", "cls": "name", "carrier": "docx numbering level text (shown before every list item)", "values": [FULL, IBAN],
     "visible": ["marker r2q numbered list", "marker r2r level text"], "build": docx_numbering_lvltext},
]


# --------------------------------------------------------------------------- round 3


def _ovba_compress_literal(data):
    """MS-OVBA compressed container made of literal tokens only: a flag byte 0x00 before every eight bytes."""
    out = bytearray(b"\x01")
    for start in range(0, len(data), 4096):
        chunk = data[start:start + 4096]
        body = bytearray()
        for i in range(0, len(chunk), 8):
            body += b"\x00" + chunk[i:i + 8]
        header = ((len(body) + 2 - 3) & 0x0FFF) | 0xB000
        out += header.to_bytes(2, "little") + body
    return bytes(out)


def docx_vba(inbox):
    p = new_docx(inbox, "macro.docm", ["marker r3a macro document"])
    src = "Attribute VB_Name = \"Module1\"\r\nSub Fill()\r\n"
    # align the short form so that a flag byte falls after its third letter, the surname so that two letters
    # stand before a flag byte and two after the next one
    while len(src) % 8 != 5:
        src += " "
    src += SHORT + " GmbH: "
    while len(src) % 8 != 6:
        src += " "
    src += SURNAME + "\r\n  ' marker r3b macro comment\r\nEnd Sub\r\n"
    stream = _ovba_compress_literal(src.encode("cp1252"))
    ole = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\x00" * 504 + b"PROJECT\x00" + b"ID=\"{00000000-0000-0000-0000-000000000000}\"\r\nModule=Module1\r\n" + b"\x00" * 64 + stream + b"\x00" * 128
    return zip_edit(p, {
        "word/vbaProject.bin": ole,
        "word/_rels/document.xml.rels": lambda b: add_rel(b, "rIdVba", "http://schemas.microsoft.com/office/2006/relationships/vbaProject", "vbaProject.bin"),
        "[Content_Types].xml": lambda b: add_default(b, "bin", "application/vnd.ms-office.vbaProject"),
    })


def docx_header_renamed(inbox):
    p = new_docx(inbox, "hdr2.docx", ["marker r3c body under a header"])
    hdr = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?><w:hdr xmlns:w="%s">%s%s</w:hdr>'
           % (W_NS, w_p("marker r3d header text " + FULL), w_p("IBAN " + IBAN)))

    def doc(b):
        t = b.decode("utf-8")
        return re.sub(r"<w:sectPr([^>]*)>", r'<w:sectPr\1><w:headerReference w:type="default" r:id="rIdHdrX"/>', t, count=1).encode("utf-8")
    return zip_edit(p, {
        "word/hdr-first.xml": hdr.encode("utf-8"),
        "word/document.xml": doc,
        "word/_rels/document.xml.rels": lambda b: add_rel(b, "rIdHdrX", REL_T + "header", "hdr-first.xml"),
        "[Content_Types].xml": lambda b: add_override(b, "/word/hdr-first.xml", "application/vnd.openxmlformats-officedocument.wordprocessingml.header+xml"),
    })


def docx_wps_textbox(inbox):
    p = new_docx(inbox, "txbx.docx", ["marker r3e body with a text box"])
    box = ('<w:p><w:r><mc:AlternateContent xmlns:mc="%s"><mc:Choice Requires="wps"><w:drawing><wp:anchor xmlns:wp="%s" distT="0" distB="0" distL="0" distR="0" simplePos="0" '
           'relativeHeight="1" behindDoc="0" locked="0" layoutInCell="1" allowOverlap="1"><wp:simplePos x="0" y="0"/><wp:positionH relativeFrom="column"><wp:posOffset>0</wp:posOffset></wp:positionH>'
           '<wp:positionV relativeFrom="paragraph"><wp:posOffset>0</wp:posOffset></wp:positionV><wp:extent cx="3000000" cy="800000"/><wp:wrapNone/><wp:docPr id="50" name="Text Box 50"/>'
           '<a:graphic xmlns:a="%s"><a:graphicData uri="http://schemas.microsoft.com/office/word/2010/wordprocessingShape"><wps:wsp xmlns:wps="http://schemas.microsoft.com/office/word/2010/wordprocessingShape">'
           '<wps:cNvSpPr txBox="1"/><wps:spPr><a:xfrm><a:off x="0" y="0"/><a:ext cx="3000000" cy="800000"/></a:xfrm><a:prstGeom prst="rect"><a:avLst/></a:prstGeom></wps:spPr>'
           '<wps:txbx><w:txbxContent>%s%s</w:txbxContent></wps:txbx><wps:bodyPr/></wps:wsp></a:graphicData></a:graphic></wp:anchor></w:drawing></mc:Choice>'
           '<mc:Fallback><w:pict><v:shape xmlns:v="%s" id="tb50" type="#_x0000_t202" style="position:absolute;width:200pt;height:60pt"><v:textbox><w:txbxContent>%s%s</w:txbxContent></v:textbox></v:shape></w:pict></mc:Fallback>'
           '</mc:AlternateContent></w:r></w:p>' % (MC_NS, WP_NS, A_NS, w_p("marker r3f modern text box " + FULL), w_p("IBAN " + IBAN), V_NS,
                                                     w_p("marker r3f modern text box " + FULL), w_p("IBAN " + IBAN)))
    return zip_edit(p, {"word/document.xml": lambda b: before_sectpr(b, box)})


def docx_legacy_pict_textbox(inbox):
    p = new_docx(inbox, "pict.docx", ["marker r3g body with a legacy text box"])
    box = ('<w:p><w:r><w:pict><v:shape xmlns:v="%s" id="tb60" type="#_x0000_t202" style="position:absolute;width:200pt;height:60pt"><v:textbox>'
           '<w:txbxContent>%s%s</w:txbxContent></v:textbox></v:shape></w:pict></w:r></w:p>'
           % (V_NS, w_p("marker r3h legacy text box " + FULL), w_p("IBAN " + IBAN)))
    return zip_edit(p, {"word/document.xml": lambda b: before_sectpr(b, box)})


def xlsx_hyperlink_mailto(inbox):
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = "marker r3i sheet with a mail link"
    ws["A2"] = "write to us"
    ws["A2"].hyperlink = "mailto:info@plan-office.example?subject=marker%20r3j%20mail%20link"
    path = inbox / "mailto.xlsx"
    wb.save(str(path))
    return path


def xlsx_defined_name_const(inbox):
    import openpyxl
    from openpyxl.workbook.defined_name import DefinedName
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = "marker r3k sheet with names"
    wb.defined_names["Kunde"] = DefinedName("Kunde", attr_text='"%s"' % FULL)
    wb.defined_names["Konto"] = DefinedName("Konto", attr_text='"IBAN %s"' % IBAN)
    wb.defined_names["Notiz"] = DefinedName("Notiz", attr_text='"marker r3l defined name"')
    path = inbox / "names.xlsx"
    wb.save(str(path))
    return path


def xlsx_table_part(inbox):
    import openpyxl
    from openpyxl.worksheet.table import Table
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Kunde", "Betrag"])
    ws.append(["a", 1])
    ws.append(["b", 2])
    ws["D1"] = "marker r3m sheet with a table"
    ws.add_table(Table(displayName="Kunden_" + SHORT, ref="A1:B3"))
    path = inbox / "table.xlsx"
    wb.save(str(path))
    names = [n for n in zipfile.ZipFile(path).namelist() if n.startswith("xl/tables/")]

    def edit(b):
        t = b.decode("utf-8")
        t = t.replace('name="Kunde"', 'name="%s" totalsRowLabel="IBAN %s"' % (esc(SURNAME), IBAN), 1)
        t = t.replace('name="Betrag"', 'name="marker r3n table column"', 1)
        return t.encode("utf-8")
    return zip_edit(path, {names[0]: edit})


def xlsx_querytable(inbox):
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = "marker r3o sheet with a query"
    path = inbox / "qt.xlsx"
    wb.save(str(path))
    qt = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?><queryTable xmlns="%s" name="%s" connectionId="1" autoFormatId="16" applyNumberFormats="0" '
          'applyBorderFormats="0" applyFontFormats="0" applyPatternFormats="0" applyAlignmentFormats="0" applyWidthHeightFormats="0"><queryTableRefresh nextId="3">'
          '<queryTableFields count="2"><queryTableField id="1" name="%s"/><queryTableField id="2" name="IBAN %s"/></queryTableFields></queryTableRefresh></queryTable>'
          % (X_NS, esc("marker r3p query " + SHORT), esc(FULL), IBAN))
    return zip_edit(path, {
        "xl/queryTables/queryTable1.xml": qt.encode("utf-8"),
        "xl/worksheets/_rels/sheet1.xml.rels": add_rel(b"", "rIdQt1", REL_T + "queryTable", "../queryTables/queryTable1.xml"),
        "[Content_Types].xml": lambda b: add_override(b, "/xl/queryTables/queryTable1.xml", "application/vnd.openxmlformats-officedocument.spreadsheetml.queryTable+xml"),
    })


def xlsx_slicer(inbox):
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = "marker r3q sheet with a slicer"
    path = inbox / "slicer.xlsx"
    wb.save(str(path))
    sc = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?><slicerCacheDefinition xmlns="http://schemas.microsoft.com/office/spreadsheetml/2009/9/main" '
          'name="Slicer_Kunde" sourceName="%s"><tabular pivotCacheId="1"><items count="2"><i x="0" s="1"/><i x="1"/></items></tabular>'
          '<extLst><ext uri="{2F2917AC-EB37-4324' '-AD4E-5DD8C200BD13}"><x15:tableSlicerCache xmlns:x15="http://schemas.microsoft.com/office/spreadsheetml/2010/11/main" tableId="1" column="1"/></ext></extLst>'
          '</slicerCacheDefinition>' % esc("marker r3r slicer source " + FULL))
    sl = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?><slicers xmlns="http://schemas.microsoft.com/office/spreadsheetml/2009/9/main"><slicer name="Kunde" cache="Slicer_Kunde" '
          'caption="IBAN %s" rowHeight="241300"/></slicers>' % IBAN)
    return zip_edit(path, {
        "xl/slicerCaches/slicerCache1.xml": sc.encode("utf-8"),
        "xl/slicers/slicer1.xml": sl.encode("utf-8"),
        "xl/worksheets/_rels/sheet1.xml.rels": add_rel(b"", "rIdSl1", "http://schemas.microsoft.com/office/2007/relationships/slicer", "../slicers/slicer1.xml"),
        "xl/_rels/workbook.xml.rels": lambda b: add_rel(b, "rIdSc1", "http://schemas.microsoft.com/office/2007/relationships/slicerCache", "slicerCaches/slicerCache1.xml"),
        "[Content_Types].xml": lambda b: add_override(add_override(b, "/xl/slicerCaches/slicerCache1.xml", "application/vnd.ms-excel.slicerCache+xml"), "/xl/slicers/slicer1.xml", "application/vnd.ms-excel.slicer+xml"),
    })


def odt_phone_columns(inbox):
    def cell(t):
        return '<table:table-cell office:value-type="string"><text:p>%s</text:p></table:table-cell>' % esc(t)
    body = ('<text:p>marker r3s contact table in a text</text:p><table:table table:name="Kontakte"><table:table-column table:number-columns-repeated="3"/>'
            '<table:table-row>%s%s%s</table:table-row><table:table-row>%s%s%s</table:table-row></table:table>'
            % (cell("Name"), cell("Vorwahl"), cell("Rufnummer"), cell("Zentrale"), cell("040"), cell("5550123")))
    return new_odf(inbox, "contacts.odt", "odt", body)


def xlsx_phone_leading_zero(inbox):
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = "marker r3t phone stored as a number"
    ws["A2"] = 301234567
    ws["A2"].number_format = "0000000000"
    path = inbox / "leadzero.xlsx"
    wb.save(str(path))
    return path


def docx_iban_softbreak(inbox):
    p = new_docx(inbox, "bank.docx", ["marker r3u bank block"])
    para = ('<w:p><w:r><w:t xml:space="preserve">IBAN DE89 3704 0044 0532</w:t></w:r><w:r><w:br/></w:r><w:r><w:t>0130 00</w:t></w:r></w:p>')
    return zip_edit(p, {"word/document.xml": lambda b: before_sectpr(b, para)})


PHONE_040 = "040 5550123"

CASES += [
    {"id": "docx-vba", "cls": "name", "carrier": "docm vbaProject.bin: names inside MS-OVBA compressed module source (flag byte every eight bytes)",
     "values": [SHORT, SURNAME], "visible": ["marker r3a macro document", "marker r3b macro comment"], "build": docx_vba},
    {"id": "docx-header-renamed", "cls": "name", "carrier": "docx header part named word/hdr-first.xml, reached through the rels only",
     "values": [FULL, IBAN], "visible": ["marker r3c body under a header", "marker r3d header text"], "build": docx_header_renamed},
    {"id": "docx-wps-textbox", "cls": "name", "carrier": "docx modern text box (wps:txbx) with a vml fallback", "values": [FULL, IBAN],
     "visible": ["marker r3e body with a text box", "marker r3f modern text box"], "build": docx_wps_textbox},
    {"id": "docx-legacy-pict-textbox", "cls": "name", "carrier": "docx legacy text box (w:pict/v:textbox) outside AlternateContent", "values": [FULL, IBAN],
     "visible": ["marker r3g body with a legacy text box", "marker r3h legacy text box"], "build": docx_legacy_pict_textbox},
    {"id": "xlsx-hyperlink-mailto", "cls": "mail", "carrier": "xlsx hyperlink target mailto: with a subject", "values": ["info@plan-office.example"],
     "visible": ["marker r3i sheet with a mail link", "marker r3j mail link"], "build": xlsx_hyperlink_mailto},
    {"id": "xlsx-defined-name-const", "cls": "name", "carrier": "xlsx defined names holding string constants", "values": [FULL, IBAN],
     "visible": ["marker r3k sheet with names", "marker r3l defined name"], "build": xlsx_defined_name_const},
    {"id": "xlsx-table-part", "cls": "name", "carrier": "xlsx table part: displayName, column names, totalsRowLabel", "values": [SHORT, SURNAME, IBAN],
     "visible": ["marker r3m sheet with a table", "marker r3n table column"], "build": xlsx_table_part},
    {"id": "xlsx-querytable", "cls": "name", "carrier": "xlsx query table name and field names", "values": [SHORT, FULL, IBAN],
     "visible": ["marker r3o sheet with a query", "marker r3p query"], "build": xlsx_querytable},
    {"id": "xlsx-slicer", "cls": "name", "carrier": "xlsx slicer cache source name and slicer caption", "values": [FULL, IBAN],
     "visible": ["marker r3q sheet with a slicer", "marker r3r slicer source"], "build": xlsx_slicer},
    {"id": "odt-phone-columns", "cls": "phone", "carrier": "odt contact table: area code and number in two columns", "values": [PHONE_040],
     "visible": ["marker r3s contact table in a text"], "build": odt_phone_columns},
    {"id": "xlsx-phone-leading-zero", "cls": "phone", "carrier": "xlsx phone stored as a number with a 0000000000 format (Excel shows the leading zero, the value has none)",
     "values": [PHONE_NAT], "visible": ["marker r3t phone stored as a number"], "build": xlsx_phone_leading_zero},
    {"id": "docx-iban-softbreak", "cls": "iban", "carrier": "docx IBAN split by a soft line break in a paragraph", "values": [IBAN],
     "visible": ["marker r3u bank block"], "build": docx_iban_softbreak},
]


# --------------------------------------------------------------------------- round 4: the boundary of both families


def _docx_table(inbox, name, marker, rows):
    import docx
    d = docx.Document()
    d.add_paragraph(marker)
    t = d.add_table(rows=len(rows), cols=max(len(r) for r in rows))
    for r, row in enumerate(rows):
        for c, v in enumerate(row):
            t.rows[r].cells[c].text = v
    path = inbox / name
    d.save(str(path))
    return path


def odt_split_cells_tail(inbox):
    def cell(t):
        return '<table:table-cell office:value-type="string"><text:p>%s</text:p></table:table-cell>' % esc(t)
    body = ('<text:p>marker r4a address table in a text</text:p><table:table table:name="Adressen"><table:table-column table:number-columns-repeated="2"/>'
            '<table:table-row>%s%s</table:table-row></table:table>' % (cell(SURNAME[:8]), cell(SURNAME[8:] + "str. 5")))
    return new_odf(inbox, "addr.odt", "odt", body)


CASES += [
    {"id": "odt-split-cells-tail", "cls": "name", "carrier": "odt table: surname split over two cells, second half glued to a suffix",
     "values": [SURNAME], "visible": ["marker r4a address table in a text"], "build": odt_split_cells_tail},
    {"id": "docx-iban-boxes", "cls": "iban", "carrier": "docx bank form: IBAN in six boxes of a table row", "values": [IBAN],
     "visible": ["marker r4b bank form"], "build": lambda inbox: _docx_table(inbox, "boxes.docx", "marker r4b bank form", [["IBAN", "DE89", "3704", "0044", "0532", "0130", "00"]])},
    {"id": "docx-vat-group-spaces", "cls": "vat", "carrier": "docx table: VAT id written with group spaces in one cell", "values": [VAT],
     "visible": ["marker r4c vat with spaces"], "build": lambda inbox: _docx_table(inbox, "vat2.docx", "marker r4c vat with spaces", [["USt-IdNr.", "DE 123 456 789"]])},
    {"id": "docx-split-cells-plain", "cls": "name", "carrier": "control: surname split over two cells, nothing glued", "values": [SURNAME],
     "visible": ["marker r4d plain split"], "build": lambda inbox: _docx_table(inbox, "plain.docx", "marker r4d plain split", [[SURNAME[:8], SURNAME[8:]]])},
    {"id": "xlsx-split-rows-plain", "cls": "name", "carrier": "control: surname split over two xlsx rows, nothing glued (header separator between)",
     "values": [SURNAME], "visible": ["marker r4e plain rows"], "build": lambda inbox: new_xlsx(inbox, "plainrows.xlsx", [[SURNAME[:8]], [SURNAME[8:]], ["marker r4e plain rows"]])},
    {"id": "docx-full-form-tail", "cls": "name", "carrier": "control: full form split over two cells, digits glued to the legal form", "values": [FULL],
     "visible": ["marker r4f full form"], "build": lambda inbox: _docx_table(inbox, "fulltail.docx", "marker r4f full form", [[FULL.rsplit(" ", 1)[0], FULL.rsplit(" ", 1)[1] + "01"]])},
    {"id": "docx-glued-token", "cls": "name", "carrier": "control: short form glued in one host name token", "values": [SHORT],
     "visible": ["marker r4g glued token"], "build": lambda inbox: new_docx(inbox, "glued.docx", ["marker r4g glued token", "host %s01 is down" % SHORT.lower()])},
    {"id": "docx-acronym-tail", "cls": "name", "carrier": "control: acronym with a hyphen and digits", "values": [ACRONYM],
     "visible": ["marker r4h acronym"], "build": lambda inbox: new_docx(inbox, "acr.docx", ["marker r4h acronym", "asset %s-01 retired" % ACRONYM])},
    {"id": "docx-phone-one-cell", "cls": "phone", "carrier": "control: phone number in one cell", "values": [PHONE_NAT],
     "visible": ["marker r4i one cell phone"], "build": lambda inbox: _docx_table(inbox, "phone1.docx", "marker r4i one cell phone", [["Telefon", PHONE_NAT]])},
    {"id": "xlsx-iban-one-cell", "cls": "iban", "carrier": "control: IBAN in one cell", "values": [IBAN],
     "visible": ["marker r4j one cell iban"], "build": lambda inbox: new_xlsx(inbox, "iban1.xlsx", [["Konto", IBAN], ["marker r4j one cell iban", ""]])},
]


# --------------------------------------------------------------------------- round 5: controls and covered carriers


def docx_footnote_standard(inbox):
    p = new_docx(inbox, "fn2.docx", ["marker r5a body with a standard note"])
    fn = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?><w:footnotes xmlns:w="%s"><w:footnote w:type="separator" w:id="-1"><w:p><w:r><w:separator/></w:r></w:p></w:footnote>'
          '<w:footnote w:type="continuationSeparator" w:id="0"><w:p><w:r><w:continuationSeparator/></w:r></w:p></w:footnote>'
          '<w:footnote w:id="1">%s</w:footnote></w:footnotes>' % (W_NS, w_p("marker r5b footnote " + FULL + " IBAN " + IBAN)))
    en = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?><w:endnotes xmlns:w="%s"><w:endnote w:id="1">%s</w:endnote></w:endnotes>'
          % (W_NS, w_p("marker r5c endnote " + SURNAME)))
    ref = '<w:p><w:r><w:t>see the notes</w:t></w:r><w:r><w:footnoteReference w:id="1"/></w:r><w:r><w:endnoteReference w:id="1"/></w:r></w:p>'
    return zip_edit(p, {
        "word/footnotes.xml": fn.encode("utf-8"), "word/endnotes.xml": en.encode("utf-8"),
        "word/document.xml": lambda b: before_sectpr(b, ref),
        "word/_rels/document.xml.rels": lambda b: add_rel(add_rel(b, "rIdFn", REL_T + "footnotes", "footnotes.xml"), "rIdEn", REL_T + "endnotes", "endnotes.xml"),
        "[Content_Types].xml": lambda b: add_override(add_override(b, "/word/footnotes.xml", "application/vnd.openxmlformats-officedocument.wordprocessingml.footnotes+xml"),
                                                      "/word/endnotes.xml", "application/vnd.openxmlformats-officedocument.wordprocessingml.endnotes+xml"),
    })


def xlsx_sst_standard(inbox):
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = "placeholder"
    path = inbox / "sst2.xlsx"
    wb.save(str(path))
    sst = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?><sst xmlns="%s" count="2" uniqueCount="2"><si><t>marker r5d standard shared strings</t></si><si><t>%s IBAN %s</t></si></sst>'
           % (X_NS, esc(FULL), IBAN))
    return zip_edit(path, {
        "xl/sharedStrings.xml": sst.encode("utf-8"),
        "xl/worksheets/sheet1.xml": lambda b: b.replace(b'<c r="A1" t="inlineStr"><is><t>placeholder</t></is></c>', b'<c r="A1" t="s"><v>0</v></c><c r="B1" t="s"><v>1</v></c>'),
        "xl/_rels/workbook.xml.rels": lambda b: add_rel(b, "rIdSst", REL_T + "sharedStrings", "sharedStrings.xml"),
        "[Content_Types].xml": lambda b: add_override(b, "/xl/sharedStrings.xml", "application/vnd.openxmlformats-officedocument.spreadsheetml.sharedStrings+xml"),
    })


def pptx_slide_orphan(inbox):
    p = new_pptx(inbox, "orphan.pptx", ["marker r5e the shown slide"])
    with zipfile.ZipFile(p) as zf:
        slide = zf.read("ppt/slides/slide1.xml").decode("utf-8")
        rels = zf.read("ppt/slides/_rels/slide1.xml.rels")
    orphan = slide.replace("marker r5e the shown slide", esc("marker r5f orphan slide " + FULL + " IBAN " + IBAN))
    return zip_edit(p, {
        "ppt/slides/slide9.xml": orphan.encode("utf-8"),
        "ppt/slides/_rels/slide9.xml.rels": rels,
        "[Content_Types].xml": lambda b: add_override(b, "/ppt/slides/slide9.xml", "application/vnd.openxmlformats-officedocument.presentationml.slide+xml"),
    })


def docx_duplicate_entry(inbox):
    p = new_docx(inbox, "dup.docx", ["nothing to see in the second copy"])
    with zipfile.ZipFile(p) as zf:
        members = [(info.filename, zf.read(info.filename)) for info in zf.infolist()]
    real = None
    for name, data in members:
        if name == "word/document.xml":
            real = data.replace(b"nothing to see in the second copy", esc("marker r5g first copy " + FULL + " IBAN " + IBAN).encode("utf-8"))
    buf = io.BytesIO()
    import warnings
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as out:
            for name, data in members:
                if name == "word/document.xml":
                    out.writestr(name, real)
                out.writestr(name, data)
    p.write_bytes(buf.getvalue())
    return p


def xlsx_hidden_row_col(inbox):
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = "marker r5h visible cell"
    ws["A2"] = "marker r5i hidden row " + FULL
    ws["C1"] = "marker r5j hidden column IBAN " + IBAN
    ws.row_dimensions[2].hidden = True
    ws.column_dimensions["C"].hidden = True
    path = inbox / "hidden.xlsx"
    wb.save(str(path))
    return path


def odt_styles_header(inbox):
    styles = ('<office:master-styles><style:master-page style:name="Standard" style:page-layout-name="pm1"><style:header><text:p>%s</text:p></style:header>'
              '<style:footer><text:p>IBAN %s</text:p></style:footer></style:master-page></office:master-styles>' % (esc("marker r5k page header " + FULL), IBAN))
    return new_odf(inbox, "hdr.odt", "odt", "<text:p>marker r5l body under a page header</text:p>", styles_xml=styles)


def pptx_layout_placeholder(inbox):
    p = new_pptx(inbox, "layout.pptx", ["marker r5m the shown slide"])
    sp = ('<p:sp><p:nvSpPr><p:cNvPr id="70" name="box"/><p:cNvSpPr txBox="1"/><p:nvPr/></p:nvSpPr><p:spPr/><p:txBody><a:bodyPr/><a:p><a:r><a:t>%s</a:t></a:r></a:p></p:txBody></p:sp>'
          % esc("marker r5n layout text " + FULL + " IBAN " + IBAN))
    return zip_edit(p, {"ppt/slideLayouts/slideLayout7.xml": lambda b: b.replace(b"</p:spTree>", sp.encode("utf-8") + b"</p:spTree>", 1)})


def docx_picture_alt(inbox):
    import docx
    from PIL import Image
    img = inbox / "pic.png"
    Image.new("RGB", (40, 30), (200, 200, 200)).save(str(img))
    d = docx.Document()
    d.add_paragraph("marker r5o body with a picture")
    d.add_picture(str(img))
    path = inbox / "pic.docx"
    d.save(str(path))
    img.unlink()

    def doc(b):
        return re.sub(rb'<wp:docPr ([^>]*)/>', lambda m: b'<wp:docPr ' + m.group(1) + (' descr="%s" title="IBAN %s"' % (esc("marker r5p picture alt text " + FULL), IBAN)).encode("utf-8") + b'/>', b, count=1)
    return zip_edit(path, {"word/document.xml": doc})


def docx_tracked_deletion(inbox):
    p = new_docx(inbox, "del.docx", ["marker r5q body with a tracked change"])
    para = ('<w:p><w:r><w:t xml:space="preserve">the offer goes to </w:t></w:r><w:del w:id="1" w:author="%s" w:date="2026-09-20T10:00:00Z"><w:r><w:delText>%s</w:delText></w:r>'
            '<w:r><w:delText xml:space="preserve"> IBAN %s</w:delText></w:r></w:del><w:ins w:id="2" w:author="%s" w:date="2026-09-20T10:01:00Z"><w:r><w:t>the customer</w:t></w:r></w:ins></w:p>'
            % (esc(PERSON), esc("marker r5r deleted " + FULL), IBAN, esc(PERSON)))
    return zip_edit(p, {"word/document.xml": lambda b: before_sectpr(b, para)})


def xlsx_comment_legacy(inbox):
    import openpyxl
    from openpyxl.comments import Comment
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = "marker r5s cell with a comment"
    ws["A1"].comment = Comment("marker r5t comment text %s IBAN %s" % (FULL, IBAN), PERSON)
    path = inbox / "comment.xlsx"
    wb.save(str(path))
    return path


def docx_math(inbox):
    p = new_docx(inbox, "math.docx", ["marker r5u body with an equation"])
    para = ('<w:p><m:oMathPara xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math"><m:oMath><m:r><m:t>%s</m:t></m:r></m:oMath></m:oMathPara></w:p>'
            % esc("marker r5v equation " + FULL + " IBAN " + IBAN))
    return zip_edit(p, {"word/document.xml": lambda b: before_sectpr(b, para)})


CASES += [
    {"id": "docx-footnote-standard", "cls": "name", "carrier": "control: footnotes.xml and endnotes.xml under their standard names", "values": [FULL, IBAN, SURNAME],
     "visible": ["marker r5a body with a standard note", "marker r5b footnote", "marker r5c endnote"], "build": docx_footnote_standard},
    {"id": "xlsx-sst-standard", "cls": "name", "carrier": "control: shared strings under the standard name", "values": [FULL, IBAN],
     "visible": ["marker r5d standard shared strings"], "build": xlsx_sst_standard},
    {"id": "pptx-slide-orphan", "cls": "name", "carrier": "pptx slide part present in the package but not in the presentation's slide list",
     "values": [FULL, IBAN], "visible": ["marker r5e the shown slide", "marker r5f orphan slide"], "build": pptx_slide_orphan},
    {"id": "docx-duplicate-entry", "cls": "name", "carrier": "docx container with two zip entries named word/document.xml (first carries the value)",
     "values": [FULL, IBAN], "visible": ["marker r5g first copy"], "build": docx_duplicate_entry},
    {"id": "xlsx-hidden-row-col", "cls": "name", "carrier": "xlsx hidden row and hidden column", "values": [FULL, IBAN],
     "visible": ["marker r5h visible cell", "marker r5i hidden row", "marker r5j hidden column"], "build": xlsx_hidden_row_col},
    {"id": "odt-styles-header", "cls": "name", "carrier": "odt page header and footer in styles.xml", "values": [FULL, IBAN],
     "visible": ["marker r5k page header", "marker r5l body under a page header"], "build": odt_styles_header},
    {"id": "pptx-layout-placeholder", "cls": "name", "carrier": "pptx slide layout text", "values": [FULL, IBAN],
     "visible": ["marker r5m the shown slide", "marker r5n layout text"], "build": pptx_layout_placeholder},
    {"id": "docx-picture-alt", "cls": "name", "carrier": "docx picture alt text (descr and title)", "values": [FULL, IBAN],
     "visible": ["marker r5o body with a picture", "marker r5p picture alt text"], "build": docx_picture_alt},
    {"id": "docx-tracked-deletion", "cls": "name", "carrier": "docx tracked deletion text and its author", "values": [PERSON, FULL, IBAN],
     "visible": ["marker r5q body with a tracked change", "marker r5r deleted"], "build": docx_tracked_deletion},
    {"id": "xlsx-comment-legacy", "cls": "name", "carrier": "xlsx cell comment text and author", "values": [PERSON, FULL, IBAN],
     "visible": ["marker r5s cell with a comment", "marker r5t comment text"], "build": xlsx_comment_legacy},
    {"id": "docx-math", "cls": "name", "carrier": "docx equation text (OMML m:t)", "values": [FULL, IBAN],
     "visible": ["marker r5u body with an equation", "marker r5v equation"], "build": docx_math},
]


# What tests/test_redteam_pack.py accepts besides the fixture forms. DERIVED: strings built from a fixture form
# (a typo, an encoding, a transliteration). INVENTED: structured values this module builds, invented but valid
# in shape (documentation ranges, fake digit families). A value may carry one of them whole or in part.
DERIVED = ()
INVENTED = (IBAN, IBAN_TIGHT, PHONE, PHONE_NAT, PHONE_040, HRB, VAT, "info@plan-office.example")
