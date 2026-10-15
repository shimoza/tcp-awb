"""check: positions and classes only, never a code, a form or a matched value.

Every name comes from tests/fixtures.py; every number is a documentation example.
"""
from __future__ import annotations

import base64
import io
import json
import os
import shutil
import zipfile
from pathlib import Path

import docx
import pytest

from awb import check, codes, config, normalize, register, vault
from tests import fixtures
from tests.test_vault import encrypt_vault, fast_gpg, serving, sockets  # noqa: F401 (fixtures)

CUSTOMER = fixtures.CUSTOMER_FORMS[0]
SHORT = fixtures.CUSTOMER_FORMS[1]
PERSON = fixtures.PERSON_FORMS[0]
SURNAME = fixtures.PERSON_FORMS[1]
ORG = fixtures.ORG_FORMS[0]
IBAN = "DE89 3704 0044 0532 0130 00"
PUBLIC_KEYS = {"start", "length", "cls"}
CLEAN = "release 1.4.2 of 2026-09-22 for %s in tcp-q7m4, version v0.1.0\n" % fixtures.CUSTOMER_CODE
RETIRED = [line.split("\t")[2] for line in fixtures.register_lines() if line.endswith("\tretired")][0]


def hit_values(text: str, hits: list[dict]) -> list[tuple[str, str]]:
    """(class, value) per hit, read from the normalised text as the positions promise."""
    n = normalize.normalize(text).text
    return [(h["cls"], n[h["start"]:h["start"] + h["length"]]) for h in hits]


# --------------------------------------------------------------------------- check_text


def test_registered_forms_are_found_in_order(register_path):
    text = "offer to %s,\nattention %s, copy to %s.\n" % (CUSTOMER, PERSON, ORG)
    hits = check.check_text(text, register_path)
    assert hit_values(text, hits) == [("name", CUSTOMER), ("name", PERSON), ("name", ORG)]


def test_hits_carry_start_length_and_class_only(register_path):
    text = "%s pays from %s, ping 203.0.113.7" % (SHORT, IBAN)
    hits = check.check_text(text, register_path)
    assert hits
    assert all(set(h) == PUBLIC_KEYS for h in hits)
    assert all(isinstance(h["start"], int) and isinstance(h["length"], int) for h in hits)
    assert [h["cls"] for h in hits] == ["name", "iban", "ip"]


def test_positions_are_in_the_normalised_text(register_path):
    broken = SHORT[:3] + "\u00ad" + SHORT[3:]          # soft hyphen inside the name
    entity = ORG.replace("ö", "&ouml;")                 # html entity inside the name
    text = "a " + broken + " and " + entity + " end"
    hits = check.check_text(text, register_path)
    assert hit_values(text, hits) == [("name", SHORT), ("name", ORG)]
    assert hits[0] == {"start": 2, "length": len(SHORT), "cls": "name"}


def test_structured_data_is_found(register_path):
    text = "account %s, host 198.51.100.20, mac 00:1A:2B:3C:4D:5E, call +49 40 123456-78" % IBAN
    assert [h["cls"] for h in check.check_text(text, register_path)] == ["iban", "ip", "mac", "phone"]


def test_a_registered_domain_is_one_name_hit_not_also_a_url(register_path):
    text = "host %s answers" % fixtures.CUSTOMER_DOMAIN
    assert hit_values(text, check.check_text(text, register_path)) == [("name", fixtures.CUSTOMER_DOMAIN)]


def test_clean_text_unregistered_and_retired_forms_give_nothing(register_path):
    text = CLEAN + "near %s, formerly %s\n" % (fixtures.CONTROL_UNREGISTERED, RETIRED)
    assert check.check_text(text, register_path) == []


def test_without_a_register_only_structured_data_is_found(tmp_path):
    text = "%s pays from %s" % (SHORT, IBAN)
    hits = check.check_text(text, tmp_path / "missing.tsv")
    assert [h["cls"] for h in hits] == ["iban"]


def test_a_broken_register_raises_without_its_content(tmp_path):
    bad = tmp_path / "register.tsv"
    bad.write_text(
        "code\tkind\tform\tadded\tstatus\n%s\tCUST\t%s (note)\t2026-09-22\tactive\n" % (fixtures.CUSTOMER_CODE, SHORT),
        encoding="utf-8",
    )
    with pytest.raises(Exception) as err:
        check.check_text("anything", bad)
    fixtures.assert_no_fixture_name(str(err.value), "the register error")
    assert "line 2" in str(err.value)


# --------------------------------------------------------------------------- check_file


def test_text_file_is_checked_directly(tmp_path, register_path):
    text = "minutes\n%s joined, account %s\n" % (PERSON, IBAN)
    f = tmp_path / "minutes.md"
    f.write_text(text, encoding="utf-8")
    assert check.check_file(f, register_path) == check.check_text(text, register_path)
    assert [h["cls"] for h in check.check_file(f, register_path)] == ["name", "iban"]


def test_base64_block_in_a_text_file_is_decoded(tmp_path, register_path):
    hidden = base64.b64encode(("contact %s for the details of the second round" % PERSON).encode()).decode()
    f = tmp_path / "blob.txt"
    f.write_text("data follows\n%s\n" % hidden, encoding="utf-8")
    hits = check.check_file(f, register_path)
    assert [h["cls"] for h in hits] == ["name"]
    assert hits[0]["start"] > len("data follows\n%s\n" % hidden)


def test_latin1_file_goes_through_the_reader(tmp_path, register_path):
    f = tmp_path / "legacy.txt"
    f.write_bytes(("supplier: %s\n" % ORG).encode("latin-1"))
    assert [h["cls"] for h in check.check_file(f, register_path)] == ["name"]


def test_docx_metadata_is_checked(tmp_path, register_path):
    d = docx.Document()
    d.add_paragraph("architecture draft, second round")
    d.core_properties.author = PERSON
    f = tmp_path / "draft.docx"
    d.save(f)
    hits = check.check_file(f, register_path)
    assert "name" in [h["cls"] for h in hits]
    assert "opaque" not in [h["cls"] for h in hits]


def test_quoted_printable_mail_is_decoded_before_the_check(tmp_path, register_path):
    split = SURNAME[:6] + "=\r\n" + SURNAME[6:]         # soft line break inside the name
    raw = (
        "From: sender@example.org\r\nTo: receiver@example.org\r\nSubject: minutes\r\n"
        "Date: Tue, 22 Sep 2026 09:12:00 +0200\r\nMIME-Version: 1.0\r\n"
        "Content-Type: text/plain; charset=utf-8\r\nContent-Transfer-Encoding: quoted-printable\r\n\r\n"
        "the call with %s went well\r\n" % split
    )
    f = tmp_path / "minutes.eml"
    f.write_bytes(raw.encode("ascii"))
    assert "name" in [h["cls"] for h in check.check_file(f, register_path)]


def test_archive_members_are_checked(tmp_path, register_path):
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as zf:
        zf.writestr("notes/readme.txt", "the plan for %s\n" % SHORT)
    f = tmp_path / "bundle.zip"
    f.write_bytes(buf.getvalue())
    assert "name" in [h["cls"] for h in check.check_file(f, register_path)]


def test_image_is_opaque_never_clean(tmp_path, register_path):
    f = tmp_path / "scan.png"
    f.write_bytes(b"\x89PNG\r\n\x1a\n" + b"\x00" * 64)
    assert check.check_file(f, register_path) == [{"start": 0, "length": 0, "cls": "opaque"}]


def test_missing_file_raises(tmp_path, register_path):
    with pytest.raises(OSError):
        check.check_file(tmp_path / "nothing-here.txt", register_path)


# --------------------------------------------------------------------------- command line


def test_cli_prints_four_keys_and_never_a_form_or_a_code(tmp_path, register_path, capsys):
    f = tmp_path / "offer.md"
    f.write_text(
        "offer for %s (%s)\nattention %s, mail %s\naccount %s, ref %s\n"
        % (CUSTOMER, SHORT, PERSON, "info@" + fixtures.CUSTOMER_DOMAIN, IBAN, fixtures.TENDER_ID),
        encoding="utf-8",
    )
    code = check.main(["--register", str(register_path), str(f)])
    out, err = capsys.readouterr()
    assert code == 1
    lines = out.splitlines()
    assert len(lines) >= 6
    for line in lines:
        record = json.loads(line)
        assert list(record) == ["file", "start", "length", "cls"]
        assert record["file"] == str(f)
    classes = {json.loads(line)["cls"] for line in lines}
    assert {"name", "iban"} <= classes
    for stream in (out, err):
        fixtures.assert_no_fixture_name(stream, "check output")
        assert not codes.PLACEHOLDER_RE.search(stream)
        assert fixtures.CUSTOMER_CODE not in stream
        assert IBAN not in stream and IBAN.replace(" ", "") not in stream


def test_cli_clean_file_exits_0_and_prints_nothing(tmp_path, register_path, capsys):
    f = tmp_path / "clean.md"
    f.write_text(CLEAN, encoding="utf-8")
    assert check.main(["--register", str(register_path), str(f)]) == 0
    assert capsys.readouterr().out == ""


def test_cli_uses_the_vault_register_by_default(tmp_path, register_path, capsys):
    f = tmp_path / "note.md"
    f.write_text("meeting with %s\n" % SURNAME, encoding="utf-8")
    assert check.main([str(f)]) == 1
    assert json.loads(capsys.readouterr().out.splitlines()[0])["cls"] == "name"


def test_cli_missing_file_is_an_error(tmp_path, register_path, capsys):
    assert check.main(["--register", str(register_path), str(tmp_path / "nothing-here.txt")]) == 2


def test_cli_broken_register_is_an_error_without_content(tmp_path, capsys):
    bad = tmp_path / "register.tsv"
    bad.write_text("code\tkind\tform\n%s\tCUST\t%s\n" % (fixtures.CUSTOMER_CODE, SHORT), encoding="utf-8")
    f = tmp_path / "note.md"
    f.write_text("hello\n", encoding="utf-8")
    assert check.main(["--register", str(bad), str(f)]) == 2
    out, err = capsys.readouterr()
    assert out == ""
    fixtures.assert_no_fixture_name(err, "check error")
    assert fixtures.CUSTOMER_CODE not in err


# --------------------------------------------------------------------------- review findings (2026-09-22)


def test_cli_prints_file_n_for_a_path_that_carries_a_form(tmp_path, register_path, capsys):
    folder = tmp_path / SHORT
    folder.mkdir()
    f = folder / ("%s Angebot.txt" % SURNAME)
    f.write_text("kontakt 192.0.2.44\n", encoding="utf-8")
    clean = tmp_path / "clean.txt"
    clean.write_text("kontakt 192.0.2.45\n", encoding="utf-8")
    code = check.main(["--register", str(register_path), str(clean), str(f), str(folder / "fehlt.txt"),
                       str(folder)])
    out, err = capsys.readouterr()
    assert code == 2
    files = [json.loads(line)["file"] for line in out.splitlines()]
    assert str(clean) in files and "file 2" in files
    assert "file 3 is not a readable file" in err and "file 4 is not a readable file" in err
    for stream in (out, err):
        fixtures.assert_no_fixture_name(stream, "check output")


def test_cli_never_echoes_an_unknown_option(tmp_path, register_path, capsys):
    assert check.main(["--register", str(register_path), "--" + SHORT, str(tmp_path)]) == 2
    out, err = capsys.readouterr()
    fixtures.assert_no_fixture_name(out + err, "usage error")


def test_cli_without_a_register_is_an_error_unless_asked_for(tmp_path, capsys):
    f = tmp_path / "note.md"
    f.write_text("meeting with %s, account %s\n" % (SURNAME, IBAN), encoding="utf-8")
    missing = tmp_path / "no-register.tsv"
    assert check.main(["--register", str(missing), str(f)]) == 2
    assert "no register" in capsys.readouterr().err
    assert check.main(["--register", str(missing), "--no-register", str(f)]) == 1
    assert [json.loads(line)["cls"] for line in capsys.readouterr().out.splitlines()] == ["iban"]


def test_a_form_in_the_file_name_gives_a_path_record(tmp_path, register_path):
    f = tmp_path / ("angebot-%s.md" % "-".join(CUSTOMER.lower().split()))
    f.write_text(CLEAN, encoding="utf-8")
    assert check.check_file(f, register_path) == [{"start": 0, "length": 0, "cls": "path"}]


def test_html_with_a_form_split_by_tags_is_found(tmp_path, register_path):
    f = tmp_path / "page.html"
    f.write_text("<html><body><p>%s<b>%s</b> ok</p></body></html>" % (SHORT[:3], SHORT[3:]), encoding="utf-8")
    # found in the raw page (split by the tag) and in the stripped text
    assert {h["cls"] for h in check.check_file(f, register_path)} == {"name"}


def test_an_archive_that_was_not_read_in_full_is_opaque(tmp_path, register_path):
    f = tmp_path / "e.zip"
    with zipfile.ZipFile(f, "w") as zf:
        zf.writestr("a.txt", "harmlos")
    raw = bytearray(f.read_bytes())
    for sig, off in ((b"PK\x03\x04", 6), (b"PK\x01\x02", 8)):
        raw[raw.find(sig) + off] |= 0x01
    f.write_bytes(bytes(raw))
    assert check.check_file(f, register_path)[0] == {"start": 0, "length": 0, "cls": "opaque"}


def test_raw_parts_of_an_office_file_are_checked_for_names(tmp_path, register_path):
    d = docx.Document()
    d.add_paragraph("angebot ohne namen")
    buf = io.BytesIO()
    d.save(buf)
    src = zipfile.ZipFile(io.BytesIO(buf.getvalue()))
    f = tmp_path / "a.docx"
    with zipfile.ZipFile(f, "w") as zf:
        for info in src.infolist():
            data = src.read(info.filename)
            if info.filename == "word/document.xml":
                data = data.replace(b"</w:body>", ('<w:p><w:hyperlink w:tooltip="%s"><w:r><w:t>x</w:t></w:r>'
                                                   '</w:hyperlink></w:p></w:body>' % CUSTOMER).encode())
            zf.writestr(info.filename, data)
    assert [h["cls"] for h in check.check_file(f, register_path)] == ["name"]


# --------------------------------------------------------------------------- the vault daemon (release 2)

needs_permissions = pytest.mark.skipif(os.geteuid() == 0, reason="root reads a file of mode 000")
NAMED = "offer to %s, attention %s, account %s\n" % (CUSTOMER, PERSON, IBAN)


def unreadable_copy(register_path, tmp_path):
    copy = tmp_path / "locked-away" / "register.tsv"
    copy.parent.mkdir()
    shutil.copy(register_path, copy)
    os.chmod(copy, 0)
    return copy


@needs_permissions
def test_check_text_goes_to_the_daemon_when_the_register_is_unreadable(tmp_path, register_path, sockets):
    local = check.check_text(NAMED, register_path)
    copy = unreadable_copy(register_path, tmp_path)
    assert check.register_source(copy)[0] == check.REMOTE
    with serving(config.paths()):
        assert check.check_text(NAMED, copy) == local
        assert check.check_text(NAMED, None) == local
        # a folder that cannot be seen and an encrypted register go the same way
        assert check.check_text(NAMED, tmp_path / "no-such-folder" / "register.tsv") == local
        (tmp_path / "enc").mkdir()
        (tmp_path / "enc" / "register.tsv.gpg").write_bytes(b"x")
        assert check.check_text(NAMED, tmp_path / "enc" / "register.tsv") == local
    assert [h["cls"] for h in local] == ["name", "name", "iban"]


@needs_permissions
def test_check_text_without_a_daemon_raises_check_unavailable(tmp_path, register_path, sockets):
    copy = unreadable_copy(register_path, tmp_path)
    for where in (copy, None, tmp_path / "no-such-folder" / "register.tsv"):
        with pytest.raises(check.CheckUnavailable) as err:
            check.check_text(NAMED, where)
        assert str(err.value) == "name check unavailable: no vault daemon"
    assert issubclass(check.CheckUnavailable, register.RegisterError)
    f = tmp_path / "note.md"
    f.write_text(NAMED, encoding="utf-8")
    with pytest.raises(check.CheckUnavailable):
        check.check_file(f, copy)


def test_a_folder_without_any_register_stays_an_empty_register(tmp_path, sockets):
    assert check.register_source(tmp_path / "register.tsv") == (check.MISSING, None)
    assert [h["cls"] for h in check.check_text(NAMED, tmp_path / "register.tsv")] == ["iban"]


@needs_permissions
def test_check_file_through_the_daemon_matches_the_local_check(tmp_path, register_path, sockets):
    f = tmp_path / ("notes-%s.md" % SHORT.lower())
    f.write_text(NAMED, encoding="utf-8")
    d = docx.Document()
    d.add_paragraph("architecture draft")
    d.core_properties.author = PERSON
    office = tmp_path / "draft.docx"
    d.save(office)
    copy = unreadable_copy(register_path, tmp_path)
    with serving(config.paths()):
        for x in (f, office):
            assert check.check_file(x, copy) == check.check_file(x, register_path)
    assert check.check_file(f, register_path)[0] == {"start": 0, "length": 0, "cls": "path"}


@needs_permissions
def test_cli_exits_2_and_prints_no_hit_list_when_the_check_is_unavailable(tmp_path, register_path, sockets, capsys):
    copy = unreadable_copy(register_path, tmp_path)
    f = tmp_path / "offer.md"
    f.write_text(NAMED, encoding="utf-8")
    assert check.main(["--register", str(copy), str(f), str(f)]) == 2
    out, err = capsys.readouterr()
    assert out == ""
    assert err.strip() == "awb check: name check unavailable: no vault daemon"


@needs_permissions
def test_cli_checks_through_the_daemon(tmp_path, register_path, sockets, capsys):
    f = tmp_path / "offer.md"
    f.write_text(NAMED, encoding="utf-8")
    assert check.main(["--register", str(register_path), str(f)]) == 1
    local = capsys.readouterr().out
    copy = unreadable_copy(register_path, tmp_path)
    with serving(config.paths()):
        assert check.main(["--register", str(copy), str(f)]) == 1
    out, err = capsys.readouterr()
    assert out == local and err == ""
    fixtures.assert_no_fixture_name(out, "check output")


def test_cli_with_a_locked_vault_exits_2(home, sockets, fast_gpg, monkeypatch, tmp_path, capsys):
    f = tmp_path / "offer.md"
    f.write_text(NAMED, encoding="utf-8")
    with serving(config.paths()) as d:
        assert encrypt_vault(monkeypatch) == 0
        assert d.state() == "unlocked"
        capsys.readouterr()
        assert check.main([str(f)]) == 1
        assert [json.loads(line)["cls"] for line in capsys.readouterr().out.splitlines()] == ["name", "name", "iban"]
        assert vault.main(["lock"]) == 0
        capsys.readouterr()
        assert check.main([str(f)]) == 2
        out, err = capsys.readouterr()
    assert out == ""
    assert err.strip() == "awb check: name check unavailable: vault locked"


@needs_permissions
def test_cli_prints_no_hit_list_when_the_check_fails_after_the_first_file(tmp_path, register_path, sockets,
                                                                          monkeypatch, capsys):
    copy = unreadable_copy(register_path, tmp_path)
    files = [tmp_path / "a.md", tmp_path / "b.md"]
    for f in files:
        f.write_text(NAMED, encoding="utf-8")
    calls = []

    def flaky(text, sock):
        calls.append(len(text))
        if len(calls) > 3:
            raise vault.VaultLocked("vault locked")
        return check.check_text(text, register_path)

    monkeypatch.setattr(vault, "ping", lambda sock: "unlocked")
    monkeypatch.setattr(vault, "check_remote", flaky)
    assert check.main(["--register", str(copy)] + [str(f) for f in files]) == 2
    out, err = capsys.readouterr()
    assert out == "" and err.strip() == "awb check: name check unavailable: vault locked"
    assert len(calls) == 4


# --------------------------------------------------------------------------- the five doors (B-021)
# The same invented inputs go through every door that screens a text before it leaves or is shown: the Ask
# question (awb/tcp/ask.py), the project chat (chat_service.py), the working copy of an imported material
# (materials_api.py), the project read API (projects_api.py, which masks) and the owner status (owner.py). The
# tests named "as_before" passed on the code before check.screen and must keep passing: no door accepts what it
# refused before and no door masks less. The table "One text screen" in INTERFACES.md says the same.

def _secret_assignment() -> str:
    return "db_" + "pass" + "word" + ' = "' + "t7Kq2Zp9" + "Lw4Rx8Vn3Bm6Yc" + '"'


DOOR_INPUTS = {
    "name": "Notes for %s on the cluster." % CUSTOMER,
    "hidden-name": "Blob " + base64.b64encode(
        ("The workshop with %s starts on Monday morning." % CUSTOMER).encode()).decode(),
    "ip": "Is the server 203.0.113.40 still running?",
    "mail": "Write to ops@example.org today.",
    "phone": "Call +49 30 1234567 now.",
    "iban": "Pay to %s today." % IBAN,
    "secret": _secret_assignment(),
    "private-key": "-----BEGIN RSA " + "PRI" + "VATE KEY-----",
    "token": "push with " + "gh" + "p_" + "Zx8Kq3Lw7Rt2Vn5Bm9Yc4Hd6Jf1Gs0Pa3Ue7Wi",
    "homepath": "written to /ho" + "me/builder/notes.txt",
}
DOOR_CLEAN = "Which node flavors does the cluster service offer?"
STRUCTURED = ("ip", "mail", "phone", "iban")
SECRETS = ("secret", "private-key", "token", "homepath")


@pytest.fixture
def daemon(home, monkeypatch):
    """The check socket of the vault daemon, stood in by the scan the daemon runs over the fixture register."""
    from awb.matcher import Matcher

    matcher = Matcher(register.forms_for_matching(register.parse(home.register.read_bytes())))
    calls = []

    def check_remote(text, sock, *args, **kwargs):
        calls.append(text)
        return check._scan(text, matcher)

    monkeypatch.setattr(vault, "ping", lambda sock, *a, **k: "unlocked")
    monkeypatch.setattr(vault, "check_remote", check_remote)
    return calls


def _ask_refuses(home, text) -> bool:
    from awb.tcp import ask

    try:
        ask.screen(text, home)
    except ask.AskError:
        return True
    return False


def _chat_refuses(home, text) -> bool:
    from awb.tcp.web import chat_service

    try:
        out = chat_service.screen(text, home)
    except chat_service.ChatError as err:
        assert err.status == 422
        return True
    assert out == normalize.normalize(text).text
    return False


def _materials_refuses(home, tmp_path, text) -> bool:
    from awb.tcp.web import materials_api

    store = materials_api.Store(home, tmp_path / "materials", sources=object(), publisher=lambda row, text: None)
    try:
        assert store.checked(text) == text
    except materials_api.Problem as err:
        assert err.status == 422
        return True
    return False


def _owner_refuses(text) -> bool:
    from awb import status_shape
    from awb.tcp import owner

    try:
        assert owner.name_check(text) == "passed"
    except status_shape.StatusError:
        return True
    return False


def test_the_ask_door_refuses_as_before(home):
    for cls in ("name", "hidden-name") + STRUCTURED:
        assert _ask_refuses(home, DOOR_INPUTS[cls]), cls
    assert not _ask_refuses(home, DOOR_CLEAN)


def test_the_chat_door_refuses_as_before(home, daemon):
    from awb.tcp.web import chat_service

    for cls in ("name",) + STRUCTURED + SECRETS:
        assert _chat_refuses(home, DOOR_INPUTS[cls]), cls
    assert not _chat_refuses(home, DOOR_CLEAN)
    # the answer of the public price tool may carry a number that reads as a phone number, nothing else
    assert chat_service.screen(DOOR_INPUTS["phone"], home, numeric_price=True) == DOOR_INPUTS["phone"]
    for cls in ("name", "ip", "mail", "iban") + SECRETS:
        with pytest.raises(chat_service.ChatError):
            chat_service.screen(DOOR_INPUTS[cls], home, numeric_price=True)


def test_the_materials_door_refuses_as_before(home, tmp_path):
    for cls in ("name", "hidden-name") + STRUCTURED + SECRETS:
        assert _materials_refuses(home, tmp_path, DOOR_INPUTS[cls]), cls
    assert not _materials_refuses(home, tmp_path, DOOR_CLEAN)


def test_the_projects_door_masks_as_before(home, daemon):
    from awb.tcp.web import projects_api

    assert projects_api.checked(DOOR_INPUTS["name"], home) == ("Notes for [withheld] on the cluster.", True)
    assert projects_api.checked(DOOR_INPUTS["ip"], home) == ("Is the server [withheld] still running?", True)
    assert projects_api.checked(DOOR_INPUTS["mail"], home) == ("Write to [withheld] today.", True)
    assert projects_api.checked(DOOR_INPUTS["phone"], home) == ("Call [withheld] now.", True)
    assert projects_api.checked(DOOR_INPUTS["iban"], home) == ("Pay to [withheld] today.", True)
    assert projects_api.checked(DOOR_CLEAN, home) == (DOOR_CLEAN, False)


def test_the_owner_door_refuses_as_before(home):
    from awb.tcp import owner

    for cls in ("name", "hidden-name", "secret", "private-key", "token"):
        assert _owner_refuses(DOOR_INPUTS[cls]), cls
    for cls in STRUCTURED + ("homepath",):
        assert not _owner_refuses(DOOR_INPUTS[cls]), cls
    assert not _owner_refuses(DOOR_CLEAN)
    screen = owner.summary_screen()
    for cls in ("name", "hidden-name", "secret", "private-key", "token"):
        assert screen(DOOR_INPUTS[cls]) == (None, "name" if cls == "hidden-name" else cls), cls
    for cls in STRUCTURED + ("homepath",):
        assert screen(DOOR_INPUTS[cls]) == (DOOR_INPUTS[cls], None), cls
    assert screen(DOOR_CLEAN) == (DOOR_CLEAN, None)


def _local(home) -> check.LocalCheck:
    return check.checker(home.register)


def test_screen_finds_the_same_classes_for_every_door(home, daemon):
    for names in (_local(home), check.RemoteCheck(home.check_socket)):
        for cls, text in DOOR_INPUTS.items():
            found = check.screen(text, names).found
            assert found == ("name",) if cls == "hidden-name" else found == (cls,), cls
        assert check.screen(DOOR_CLEAN, names) == check.Screened(DOOR_CLEAN, ())
    assert check.screen(DOOR_INPUTS["phone"], _local(home), allow=("phone",)).found == ()
    assert check.screen(DOOR_INPUTS["token"], _local(home), allow=("phone",)).found == ("token",)
    # a Matcher is taken as the local check over it
    assert check.screen(DOOR_INPUTS["name"], _local(home).matcher).found == ("name",)


def test_screen_with_mask_masks_what_it_can_and_withholds_the_rest(home):
    names = _local(home)
    assert check.screen(DOOR_INPUTS["name"], names, mask=True) == \
        check.Screened("Notes for [withheld] on the cluster.", ("name",))
    both = "Mail ops@example.org about %s." % CUSTOMER
    assert check.screen(both, names, mask=True) == check.Screened("Mail [withheld] about [withheld].",
                                                                  ("mail", "name"))
    for cls in ("hidden-name",) + SECRETS:
        assert check.screen(DOOR_INPUTS[cls], names, mask=True).text is None, cls
    assert check.screen(DOOR_INPUTS["phone"], names, mask=True, allow=("phone",)) == \
        check.Screened(DOOR_INPUTS["phone"], ())
    assert check.screen(DOOR_CLEAN, names, mask=True) == check.Screened(DOOR_CLEAN, ())


def test_screen_fails_closed_when_the_daemon_does_not_answer(home, monkeypatch):
    from awb.tcp.web import chat_service, projects_api

    def down(text, sock, *args, **kwargs):
        raise vault.VaultUnavailable("no vault daemon")

    monkeypatch.setattr(vault, "ping", lambda sock, *a, **k: "unlocked")
    monkeypatch.setattr(vault, "check_remote", down)
    with pytest.raises(check.CheckUnavailable):
        check.screen(DOOR_CLEAN, check.RemoteCheck(home.check_socket))
    with pytest.raises(chat_service.ChatError) as err:
        chat_service.screen(DOOR_CLEAN, home)
    assert err.value.status == 503
    with pytest.raises(projects_api.Unavailable):
        projects_api.checked(DOOR_CLEAN, home)


def test_the_doors_now_refuse_what_another_door_already_refused(home, daemon):
    """The gaps of F1 1b and 1c: the Ask question ran no secret detector, the chat and the project read API did not
    look behind base64 or hex, the project read API ran no secret detector."""
    from awb.tcp.web import projects_api

    for cls in SECRETS:
        assert _ask_refuses(home, DOOR_INPUTS[cls]), cls
    assert _chat_refuses(home, DOOR_INPUTS["hidden-name"])
    for cls in ("hidden-name",) + SECRETS:
        assert projects_api.checked(DOOR_INPUTS[cls], home) == ("[Content withheld by the data check.]", True), cls


def test_the_table_of_the_doors_names_every_door():
    text = (Path(__file__).resolve().parent.parent / "INTERFACES.md").read_text(encoding="utf-8")
    table = text.split("### One text screen", 1)[1].split("\n\n|", 1)[1].split("\n\n", 1)[0]
    rows = table.splitlines()[2:]
    assert len(rows) == 5
    for where in ("awb/tcp/ask.py", "chat_service.py", "materials_api.py", "projects_api.py", "awb/tcp/owner.py"):
        assert sum(where in row for row in rows) == 1, where
