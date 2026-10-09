"""The pictures of the intake (awb/images.py): held on the vault side, listed, released by number after a look.

Every picture is generated here: a few pixels, no text. The documents carry only invented text.
"""
from __future__ import annotations

import hashlib
import io
import os
import stat
import struct
import zlib
from pathlib import Path

import zipfile

import docx
import pytest
from reportlab.lib.utils import ImageReader
from reportlab.pdfgen import canvas

from awb import cli, config, images, intake
from tests import fixtures as fx


def png(color: tuple[int, int, int], size: int = 4) -> bytes:
    """A small PNG of one color, built by hand."""
    raw = b"".join(b"\x00" + bytes(color) * size for _ in range(size))

    def chunk(kind: bytes, data: bytes) -> bytes:
        return struct.pack(">I", len(data)) + kind + data + struct.pack(">I", zlib.crc32(kind + data) & 0xffffffff)

    return (b"\x89PNG\r\n\x1a\n" + chunk(b"IHDR", struct.pack(">IIBBBBB", size, size, 8, 2, 0, 0, 0))
            + chunk(b"IDAT", zlib.compress(raw)) + chunk(b"IEND", b""))


def deck(path: Path, pictures: list[bytes]) -> Path:
    d = docx.Document()
    d.add_paragraph("The landing zone plan of the invented test, two stages.")
    for pic in pictures:
        d.add_picture(io.BytesIO(pic))
    d.save(str(path))
    return path


def members(path: Path) -> list[bytes]:
    """Every picture member of an office file, the page thumbnail of the template included: a thumbnail shows the
    page and is held like any other picture."""
    with zipfile.ZipFile(path) as z:
        return [z.read(n) for n in z.namelist() if n.rsplit(".", 1)[-1].lower() in images.IMAGE_EXTS]


def test_pictures_of_an_office_file_an_image_file_and_a_pdf(tmp_path):
    red, blue = png((200, 0, 0)), png((0, 0, 200))
    src = deck(tmp_path / "plan.docx", [red, blue])
    found, more = images.pictures(src, "docx")
    assert sorted(data for _, data in found) == sorted(members(src)) and more == 0
    assert {red, blue} <= {data for _, data in found}
    (tmp_path / "logo.png").write_bytes(red)
    assert images.pictures(tmp_path / "logo.png", "image") == ([("png", red)], 0)
    pdf = tmp_path / "plan.pdf"
    c = canvas.Canvas(str(pdf))
    c.drawString(72, 720, "The network plan of the invented test.")
    c.drawImage(ImageReader(io.BytesIO(png((0, 150, 0), 16))), 72, 500, width=64, height=64)
    c.save()
    found, _ = images.pictures(pdf, "pdf")
    assert len(found) == 1
    assert images.pictures(tmp_path / "missing.docx", "docx") == ([], 0)


def test_hold_list_and_release_by_number_and_a_project_code_is_a_box(home, tmp_path, capsys):
    """Replaces test_hold_list_and_release_by_number (2026-10-09): the project code is the one identifier, so the
    pictures of a project without a customer are held under its code; another kind of code is still refused."""
    red, blue, green = png((200, 0, 0)), png((0, 0, 200)), png((0, 200, 0))
    src = deck(tmp_path / "plan.docx", [red, blue, green])
    total = len(members(src))
    assert total >= 3
    assert images.hold(home, fx.CUSTOMER_CODE, "f01", src, "docx", today="2026-09-25") == total
    where = images.folder(home, fx.CUSTOMER_CODE, "f01")
    assert stat.S_IMODE(where.stat().st_mode) == 0o700
    assert all(stat.S_IMODE(f.stat().st_mode) == 0o600 for f in where.iterdir())
    manifest = images.load_manifest(where)
    assert [i["state"] for i in manifest["images"]] == ["held"] * total
    assert "plan" not in (where / images.MANIFEST).read_text(encoding="utf-8")      # no file or member name
    assert not (home.outbox / fx.CUSTOMER_CODE / images.IMAGES).exists()

    assert cli.main(["images", "list"]) == 0
    assert "held %d  released 0" % total in capsys.readouterr().out
    out = images.release(home, fx.CUSTOMER_CODE, "f01", [2], today="2026-09-25")
    assert [f.name for f in out] == ["f01-img-002.png"]
    assert hashlib.sha256(out[0].read_bytes()).hexdigest() == manifest["images"][1]["sha256"]
    assert [i["state"] for i in images.load_manifest(where)["images"]][:3] == ["held", "released", "held"]
    with pytest.raises(images.ImageError, match="not held pictures"):
        images.release(home, fx.CUSTOMER_CODE, "f01", [99])
    assert cli.main(["images", "release", fx.CUSTOMER_CODE, "f01", "--all"]) == 0
    assert "%d picture(s) of f01 released" % total in capsys.readouterr().out
    assert cli.main(["images", "release", fx.CUSTOMER_CODE, "f01"]) == 2
    assert images.folder(home, "tcp-abcd") == home.vault / images.QUARANTINE / "tcp-abcd"
    for other in ("PERS-ABCD", "CUST-Q7M4-X1", "../tcp-abcd"):
        with pytest.raises(images.ImageError, match="not a customer code"):
            images.folder(home, other)


def test_a_picture_that_changed_on_disk_is_not_released(home, tmp_path):
    src = deck(tmp_path / "plan.docx", [png((10, 20, 30))])
    images.hold(home, fx.CUSTOMER_CODE, "f02", src, "docx")
    where = images.folder(home, fx.CUSTOMER_CODE, "f02")
    (where / "img-001.png").write_bytes(png((30, 20, 10)))
    with pytest.raises(images.ImageError, match="does not match its hash"):
        images.release(home, fx.CUSTOMER_CODE, "f02", [1])


def test_the_intake_holds_the_pictures_and_names_them_in_the_public_report(home, register_path, tmp_path):
    src = deck(home.inbox / "plan.docx", [png((200, 0, 0)), png((0, 0, 200))])
    total = len(members(src))
    res = intake.run([src], fx.CUSTOMER_CODE, home)
    assert not res.blocked and len(res.outputs) == 1
    rows = images.held(home, fx.CUSTOMER_CODE)
    assert len(rows) == 1 and rows[0]["held"] == total
    assert "%d picture(s) held on the vault side" % total in res.public_report.read_text(encoding="utf-8")
    assert not (home.outbox / fx.CUSTOMER_CODE / images.IMAGES).exists()
    text = src.with_suffix(".txt")
    text.write_text("text only, no picture\n", encoding="utf-8")
    intake.run([text], fx.CUSTOMER_CODE, home)
    assert len(images.held(home, fx.CUSTOMER_CODE)) == 1


def test_in_an_encrypted_vault_pictures_are_sealed_and_released_through_the_daemon(home, tmp_path, monkeypatch):
    import base64
    from awb import vault
    sealed_store: dict[str, bytes] = {}

    def seal(p, path):
        path = Path(path)
        sealed_store[str(path) + ".gpg"] = path.read_bytes()
        path.with_name(path.name + ".gpg").write_bytes(b"sealed")
        path.unlink()
        return None

    def admin_call(op, sock=None, **fields):
        assert op == "open_file"
        return {"ok": True, "data": base64.b64encode(sealed_store[fields["path"]]).decode("ascii")}

    monkeypatch.setattr(intake, "seal", seal)
    monkeypatch.setattr(vault, "admin_call", admin_call)
    red = png((200, 0, 0))
    (tmp_path / "logo.png").write_bytes(red)
    assert images.hold(home, fx.CUSTOMER_CODE, "f03", tmp_path / "logo.png", "image", encrypted=True) == 1
    where = images.folder(home, fx.CUSTOMER_CODE, "f03")
    assert sorted(f.name for f in where.iterdir()) == ["img-001.png.gpg", "manifest.json"]
    out = images.release(home, fx.CUSTOMER_CODE, "f03", [1])
    assert out[0].read_bytes() == red

    def refuse(p, path):
        return "vault locked"

    monkeypatch.setattr(intake, "seal", refuse)
    assert images.hold(home, fx.CUSTOMER_CODE, "f04", tmp_path / "logo.png", "image", encrypted=True) == 0
    where = images.folder(home, fx.CUSTOMER_CODE, "f04")
    assert sorted(f.name for f in where.iterdir()) == ["manifest.json"]          # nothing left in plaintext
    assert images.load_manifest(where)["not_held"] == 1


# --------------------------------------------------------------------------- the review of 2026-09-27


def chunk(kind: bytes, data: bytes) -> bytes:
    return struct.pack(">I", len(data)) + kind + data + struct.pack(">I", zlib.crc32(kind + data) & 0xffffffff)


def segment(marker: int, data: bytes) -> bytes:
    return bytes([0xFF, marker]) + struct.pack(">H", len(data) + 2) + data


def test_a_released_picture_is_scrubbed_of_its_metadata(home, tmp_path):
    """sec-private-1: EXIF, XMP, comments and the text chunks of a PNG go before a picture leaves the vault
    side; the hash in the manifest is of the picture as it came."""
    name = fx.CUSTOMER_FORMS[0].encode("utf-8")
    plain = png((1, 2, 3))
    noisy = (plain[:-12] + chunk(b"tEXt", b"Comment\x00" + name) + chunk(b"zTXt", b"k\x00\x00" + zlib.compress(name))
             + chunk(b"iTXt", b"k\x00\x00\x00\x00\x00" + name) + chunk(b"eXIf", b"MM\x00*" + name) + plain[-12:])
    assert images.scrubbed("png", noisy) == plain
    jpeg = (b"\xff\xd8" + segment(0xE0, b"JFIF\x00\x01\x01\x00\x00\x01\x00\x01\x00\x00")
            + segment(0xE1, b"Exif\x00\x00MM\x00*" + name) + segment(0xE1, b"http://ns.adobe.com/xap/1.0/\x00" + name)
            + segment(0xED, b"Photoshop 3.0\x00" + name) + segment(0xFE, b"a comment about " + name)
            + segment(0xDB, b"\x00" + bytes(64)) + segment(0xEE, b"Adobe\x00" + bytes(7))
            + b"\xff\xda\x00\x08\x01\x01\x00\x00\x3f\x00" + b"scan data " + name + b"\xff\xd9")
    out = images.scrubbed("jpg", jpeg)
    assert out.startswith(b"\xff\xd8\xff\xe0\x00\x10JFIF") and out.endswith(b"scan data " + name + b"\xff\xd9")
    assert out.count(name) == 1 and segment(0xDB, b"\x00" + bytes(64)) in out and segment(0xEE, b"Adobe\x00" + bytes(7)) in out
    assert images.scrubbed("gif", b"GIF89a" + name) == b"GIF89a" + name        # a format without a scrub is left as it is
    (tmp_path / "shot.png").write_bytes(noisy)
    images.hold(home, fx.CUSTOMER_CODE, "f03", tmp_path / "shot.png", "image")
    where = images.folder(home, fx.CUSTOMER_CODE, "f03")
    assert images.load_manifest(where)["images"][0]["sha256"] == hashlib.sha256(noisy).hexdigest()
    out = images.release(home, fx.CUSTOMER_CODE, "f03", [1])
    assert out[0].read_bytes() == plain


def test_a_picture_whose_text_or_metadata_carries_a_name_is_not_released(home, tmp_path):
    """sec-private-1: what a picture carries besides its pixels goes through the name check; a registered
    name in the text of an svg or in the strings of a format that is not scrubbed keeps it on the vault side."""
    name = fx.CUSTOMER_FORMS[0]
    (tmp_path / "diagram.svg").write_text(
        '<svg xmlns="http://www.w3.org/2000/svg"><text x="1" y="1">%s</text></svg>' % name, encoding="utf-8")
    images.hold(home, fx.CUSTOMER_CODE, "f04", tmp_path / "diagram.svg", "svg")
    with pytest.raises(images.ImageError, match="registered name") as err:
        images.release(home, fx.CUSTOMER_CODE, "f04", [1])
    assert name not in str(err.value)
    assert not (home.outbox / fx.CUSTOMER_CODE / images.IMAGES).exists()
    assert [i["state"] for i in images.load_manifest(images.folder(home, fx.CUSTOMER_CODE, "f04"))["images"]] == ["held"]
    comment = name.encode("utf-8")
    gif = (b"GIF89a\x01\x00\x01\x00\x00\x00\x00" + b"\x21\xfe" + bytes([len(comment)]) + comment + b"\x00"
           + b"\x2c\x00\x00\x00\x00\x01\x00\x01\x00\x00\x02\x02\x44\x01\x00\x3b")
    (tmp_path / "logo.gif").write_bytes(gif)
    images.hold(home, fx.CUSTOMER_CODE, "f05", tmp_path / "logo.gif", "image")
    with pytest.raises(images.ImageError, match="registered name"):
        images.release(home, fx.CUSTOMER_CODE, "f05", [1])
    clean = (tmp_path / "clean.svg")
    clean.write_text('<svg xmlns="http://www.w3.org/2000/svg"><text x="1" y="1">stage two</text></svg>', encoding="utf-8")
    images.hold(home, fx.CUSTOMER_CODE, "f06", clean, "svg")
    assert [f.name for f in images.release(home, fx.CUSTOMER_CODE, "f06", [1])] == ["f06-img-001.svg"]


def test_the_images_command_refuses_the_work_user(monkeypatch, capsys):
    """sec-private-2: the pictures are held on the vault side; the work user is told so, without a path."""
    monkeypatch.setattr(config, "is_work_user", lambda: True)
    assert cli.main(["images", "list"]) == 1
    err = capsys.readouterr().err
    assert "vault side" in err and "/" not in err
