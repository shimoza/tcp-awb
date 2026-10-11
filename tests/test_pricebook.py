"""awb/tcp/pricebook.py and awb/xlsxlite.py: the price workbook, against a stand-in price API with invented rates.

openpyxl reads what the standard library writer made. No test reaches the real price API: the fake answers on
127.0.0.1 and a guard refuses any other address.
"""
from __future__ import annotations

import ast
import datetime
import json
import re
import socket
import zipfile
from decimal import Decimal
from pathlib import Path

import openpyxl
import pytest

import awb
from awb import cli, jobs, offered, xlsxlite
from awb.tcp import price, pricebook, throttle
from tests import fixtures as fx
from tests.tcp_fake import FakePriceAPI, default_records, record

ROOT = Path(awb.__file__).resolve().parent.parent


def extra_records() -> list[dict]:
    return default_records() + [
        record("TEST_TRAP_GB", service="evs", flavor="trap.gb", os="Standard", payg="1.005000", unit="GB/month",
               r12="0.000000", r24="0.000000"),
        record("TEST_DIME_GB", service="evs", flavor="dime.gb", os="Standard", payg="0.100000", unit="GB/month",
               r12="0.000000", r24="0.000000"),
        record("TEST_TWO_DIMES_GB", service="evs", flavor="twodimes.gb", os="Standard", payg="0.200000",
               unit="GB/month", r12="0.000000", r24="0.000000"),
        record("TEST_FLEX_LNX", service="ecsflex", flavor="s3.flex.1", payg="0.040000", r12="20.000000"),
        record("TEST_TIERS_LNX", flavor="tiers.1", os="Open Linux", payg="0.050000"),
        record("TEST_TIERS_SUSE", flavor="tiers.1", os="SUSE Linux", payg="0.060000"),
    ]


@pytest.fixture
def api(monkeypatch):
    with FakePriceAPI(extra_records(), cached_at="2026-10-09 14:12:40") as f:
        monkeypatch.setenv("AWB_PRICE_API", f.url)
        yield f


@pytest.fixture(autouse=True)
def quiet(monkeypatch):
    monkeypatch.setattr(jobs.time, "sleep", lambda s: None)
    monkeypatch.setattr(throttle, "wait_turn", lambda host, *a, **k: 0.0)
    real = socket.socket.connect

    def only_here(self, address):
        host = address[0] if isinstance(address, tuple) else address
        assert host in ("127.0.0.1", "::1", "localhost"), "a test reached an address outside this host"
        return real(self, address)

    monkeypatch.setattr(socket.socket, "connect", only_here)


@pytest.fixture
def folder(tmp_path):
    (tmp_path / "SCOPE.md").write_text("# scope\n", encoding="utf-8")
    return tmp_path


HEAD = "title: Sizing for %s\nregion: eu-de\n" % fx.CUSTOMER_CODE

WEB = """
[web servers]
service: ecs
flavor: s3.large.2
os: Linux
quantity: %s
mode: payg
assumption: Two zones and all day.
assumption.de: Rund um die Uhr in zwei Zonen.
"""

DISKS = """
[data disks]
service: evs
flavor: evs.ssd
quantity: 2000
assumption: 500 GB per server.
"""


def spec_file(folder: Path, body: str, name: str = "est.txt") -> Path:
    path = folder / name
    path.write_text(body, encoding="utf-8")
    return path


def run(argv) -> int:
    return cli.main(["pricebook"] + [str(a) for a in argv])


def sheet_cells(path: Path, name: str | None = None, values: bool = False) -> dict:
    wb = openpyxl.load_workbook(path, data_only=values)
    ws = wb[name] if name else wb.worksheets[0]
    return {c.coordinate: c.value for row in ws.iter_rows() for c in row if c.value is not None}


def find_row(cells: dict, text: str) -> int:
    return next(int(k[1:]) for k, v in cells.items() if v == text and k.startswith("B"))


# --------------------------------------------------------------------------- the specification


def test_text_and_json_forms_read_the_same():
    text = HEAD + WEB % "4" + "\n# a comment\n" + DISKS
    spec = pricebook.parse_spec(text.encode(), "est.txt")
    data = {"title": "Sizing for %s" % fx.CUSTOMER_CODE, "region": "eu-de", "lines": [
        {"label": "web servers", "service": "ecs", "flavor": "s3.large.2", "os": "Linux", "quantity": 4,
         "mode": "payg", "assumption": "Two zones and all day.", "assumption_de": "Rund um die Uhr in zwei Zonen."},
        {"label": "data disks", "service": "evs", "flavor": "evs.ssd", "quantity": 2000,
         "assumption": "500 GB per server."}]}
    other = pricebook.parse_spec(json.dumps(data).encode(), "est.json")
    assert spec.lines == other.lines
    assert [l.label for l in spec.lines] == ["web servers", "data disks"]
    assert spec.lines[0].quantity == Decimal("4") and spec.lines[0].term == "PAYG"
    assert spec.lines[0].assumption_for("de") == "Rund um die Uhr in zwei Zonen."
    assert spec.lines[1].assumption_for("de") == "500 GB per server."


def test_a_value_continues_on_an_indented_line():
    spec = pricebook.parse_spec((HEAD + "[a]\nservice: ecs\nflavor: s3.large.2\nquantity: 1\n"
                                 "assumption: first part\n    second part\n").encode())
    assert spec.lines[0].assumption == "first part second part"


@pytest.mark.parametrize("text, term", [
    ("payg", "PAYG"), ("Pay as you go", "PAYG"), ("pay-as-you-go", "PAYG"), ("r12", "R12"), ("reserved 24", "R24"),
    ("Reserved 36 months", "R36"), ("ru12", "RU12"), ("up front 24", "RU24"), ("upfront-36", "RU36"),
    ("reserved 12 up front", "RU12"), ("reserved up front 36 months", "RU36"),
])
def test_modes_are_read(text, term):
    assert pricebook.mode_term(text) == term


@pytest.mark.parametrize("text", ["", "monthly", "r13", "reserved", "spot", "r12 r24"])
def test_unknown_modes_are_not_guessed(text):
    assert pricebook.mode_term(text) is None


def test_a_bad_specification_lists_every_problem():
    body = HEAD + """
[a]
service: ecs
flavor: s3.large.2
quantity: 1,5
mode: spot
colour: red

[a]
service: ecs
flavor: s3.large.2
quantity: 1

[b]
service: Elastic Cloud Server
flavor: s3.large.2
quantity: 0

[c]
service: ecs
quantity: 1
region: us-east

[d]
service: ecs
id: TEST_S3_LARGE_2_LNX
flavor: s3.large.2
quantity: 1
hours: 800
"""
    with pytest.raises(pricebook.SpecError) as err:
        pricebook.parse_spec(body.encode())
    text = "\n".join(err.value.problems)
    for want in ("colour", "[a] quantity", "[a] mode", "used twice", "[b] give api", "[b] quantity",
                 "[c] a line needs a flavor or an id", "[c] the region", "[d] id replaces", "[d] hours"):
        assert want in text, want


def test_the_text_form_refuses_what_is_no_setting_and_no_head():
    with pytest.raises(pricebook.SpecError) as err:
        pricebook.parse_spec((HEAD + "this is a sentence\n[a]\nservice: ecs\nflavor: x\nquantity: 1\n").encode())
    assert "line 3" in err.value.problems[0]


def test_a_specification_needs_a_line_and_stays_small():
    with pytest.raises(pricebook.SpecError):
        pricebook.parse_spec(HEAD.encode())
    many = HEAD + "".join("[l%d]\nservice: ecs\nflavor: s3.large.2\nquantity: 1\n" % i
                          for i in range(pricebook.MAX_LINES + 1))
    with pytest.raises(pricebook.SpecError) as err:
        pricebook.parse_spec(many.encode())
    assert "more than" in err.value.problems[0]
    with pytest.raises(pricebook.SpecError):
        pricebook.parse_spec(b"\xff\xfe\x00")


# --------------------------------------------------------------------------- the workbook


def test_build_writes_lines_formulas_and_sources(folder, api, capsys):
    spec = spec_file(folder, HEAD + WEB % "4" + DISKS)
    assert run(["build", spec]) == 0
    out = capsys.readouterr().out
    assert "2 line(s)" in out and "monthly total 484.80 EUR" in out
    xlsx = folder / "est.xlsx"
    assert xlsx.is_file() and (folder / "est.md").is_file()

    cells = sheet_cells(xlsx)
    wb = openpyxl.load_workbook(xlsx)
    assert wb.sheetnames == ["Lines", "Sources"]
    heads = [cells["%s4" % c] for c in "ABCDEFGHIJKLMNO"]
    assert heads == ["No.", "Item", "Service", "Flavor or product", "Operating system", "Region", "Pricing mode",
                     "Unit", "Unit price (EUR)", "Quantity", "Hours per month", "Monthly (EUR)", "Yearly (EUR)",
                     "Assumption", "Price record"]
    # the first line: every price is the record's, none is typed
    assert (cells["B5"], cells["C5"], cells["D5"], cells["E5"], cells["F5"], cells["G5"], cells["H5"]) == (
        "web servers", "ecs", "s3.large.2", "Open Linux", "eu-de", "Pay as you go", "h")
    assert cells["I5"] == 0.085 and cells["J5"] == 4 and cells["O5"] == "TEST_S3_LARGE_2_LNX"
    assert cells["N5"] == "Two zones and all day."
    assert cells["I6"] == 0.12 and cells["J6"] == 2000 and cells["N6"] == "500 GB per server."
    assert "K6" not in cells                                    # a monthly unit has no hours
    # the figures are formulas, the hours of line 1 come from the cell of the totals block
    assert cells["L5"] == '=ROUND(I5*J5*IF(K5="",1,K5),2)' and cells["M5"] == "=L5*$C$11"
    assert cells["K5"] == "=$C$12"
    assert cells["C9"] == "=SUM(L5:L6)" and cells["C10"] == "=SUM(M5:M6)"
    assert cells["B9"] == "Monthly total (EUR)" and cells["B10"] == "Yearly total (EUR)"
    assert (cells["B11"], cells["C11"], cells["C12"]) == ("Months per year", 12, 720)
    # the stored results are the exact ones
    values = sheet_cells(xlsx, values=True)
    assert (values["K5"], values["L5"], values["M5"]) == (720, 244.8, 2937.6)
    assert (values["L6"], values["M6"], values["C9"], values["C10"]) == (240, 2880, 484.8, 5817.6)

    src = sheet_cells(xlsx, "Sources")
    rows = {src["A%d" % r]: src.get("B%d" % r) for r in range(1, 40) if "A%d" % r in src}
    snapshot = next(v for k, v in rows.items() if k.startswith("Price API snapshot"))
    assert re.fullmatch(r"\d{4}-\d\d-\d\d \d\d:\d\d UTC", snapshot)
    assert rows["Region"] == "eu-de"
    catalog = offered.load()
    assert rows["Service description revision"].startswith(offered.revision_label(catalog))
    assert rows["Price API"] == api.url
    assert "Elastic Cloud Server" in " ".join(str(v) for v in src.values())
    assert any(k.startswith("Served from the cache") and v == "2026-10-09 14:12:40" for k, v in rows.items())


def test_the_price_comes_from_the_api_and_follows_it(folder, api):
    spec = spec_file(folder, HEAD + WEB % "4")
    assert run(["build", spec]) == 0
    assert sheet_cells(folder / "est.xlsx")["I5"] == 0.085
    for r in api.records:
        if r["id"] == "TEST_S3_LARGE_2_LNX":
            r["priceAmount"] = "0.100000 EUR"
    assert run(["build", spec]) == 0
    cells = sheet_cells(folder / "est.xlsx", values=True)
    assert cells["I5"] == 0.1 and cells["L5"] == 288
    assert len(api.calls) == 2                                              # asked at each build, not remembered


def test_a_service_the_description_does_not_offer_is_refused_before_the_api_is_asked(folder, api, capsys):
    body = HEAD + WEB % "1" + "\n[bare metal]\nservice: bms\nflavor: x\nquantity: 1\n"
    spec = spec_file(folder, body)
    assert run(["build", spec]) == 1
    err = capsys.readouterr().err
    assert "[bare metal]" in err and "not offered" in err and "Do not propose it" in err
    assert "web servers" not in err
    assert api.calls == []
    assert not (folder / "est.xlsx").exists() and not (folder / "est.md").exists()


def test_a_withdrawn_service_is_refused(folder, api, capsys):
    spec = spec_file(folder, HEAD + "\n[backup]\nservice: csbs\nflavor: x\nquantity: 1\n")
    assert run(["build", spec]) == 1
    assert "no longer" in capsys.readouterr().err


def test_the_api_name_cannot_hide_a_service_that_is_not_offered(folder, api, capsys):
    spec = spec_file(folder, HEAD + "\n[hidden]\nservice: ecs\napi: bms\nflavor: x\nquantity: 1\n")
    assert run(["build", spec]) == 1
    assert "does not offer" in capsys.readouterr().err
    assert api.calls == []


def test_a_name_the_description_does_not_list_gets_a_hint(folder, api, capsys):
    spec = spec_file(folder, HEAD + "\n[flex]\nservice: ecsflex\nflavor: s3.flex.1\nquantity: 1\n")
    assert run(["build", spec]) == 1
    assert "in `service`" in capsys.readouterr().err


def test_a_family_of_the_api_is_priced_under_the_service(folder, api):
    spec = spec_file(folder, HEAD + "\n[flex]\nservice: ecs\napi: ecsflex\nflavor: s3.flex.1\nquantity: 2\n")
    assert run(["build", spec]) == 0
    cells = sheet_cells(folder / "est.xlsx", values=True)
    assert cells["C5"] == "ecs" and cells["O5"] == "TEST_FLEX_LNX" and cells["L5"] == 57.6


def test_lines_the_api_cannot_price_are_refused_with_the_reason(folder, api, capsys):
    body = HEAD + """
[ambiguous]
service: ecs
flavor: s3.large.2
quantity: 1

[other family]
service: ecs
flavor: s3.flex.1
quantity: 1

[missing]
service: ecs
flavor: s9.nothing.1
quantity: 1

[no reserved rate]
service: ecs
flavor: s3.large.2
os: Linux
quantity: 1
mode: r36

[hours on reserved]
service: ecs
flavor: s3.large.2
os: Linux
quantity: 1
mode: r12
hours: 100

[hours on a monthly unit]
service: evs
flavor: evs.ssd
quantity: 1
hours: 100

[wrong id family]
service: ecs
id: TEST_EVS_SSD
quantity: 1

[no such os]
service: ecs
flavor: s3.large.2
os: Plan
quantity: 1
"""
    spec = spec_file(folder, body)
    assert run(["build", spec]) == 1
    err = capsys.readouterr().err
    for want in ("[ambiguous] 2 records match", "[other family] no record under sn=ecs", "sn=ecsflex",
                 "[missing] no record for service ecs and flavor s9.nothing.1", "[no reserved rate] the API carries "
                 "no R36 rate", "[hours on reserved] reserved rates are flat monthly values",
                 "[hours on a monthly unit] hours apply to an hourly rate", "[wrong id family] the record "
                 "TEST_EVS_SSD belongs to the price API service evs", "[no such os] no record of the flavor"):
        assert want in err, want
    assert "8 line(s) refused" in err
    assert not (folder / "est.xlsx").exists()


def test_a_part_of_the_os_unit_that_fits_several_records_is_not_guessed(folder, api, capsys):
    head = "\n[tiers]\nservice: ecs\nflavor: tiers.1\nquantity: 1\n"
    assert run(["build", spec_file(folder, HEAD + head + "os: Linux\n")]) == 1
    assert "2 records match (OS units: Open Linux, SUSE Linux" in capsys.readouterr().err
    assert run(["build", spec_file(folder, HEAD + head + "os: open linux\n")]) == 0
    assert sheet_cells(folder / "est.xlsx")["O5"] == "TEST_TIERS_LNX"
    assert run(["build", spec_file(folder, HEAD + head + "os: SUSE\n")]) == 0
    assert sheet_cells(folder / "est.xlsx")["O5"] == "TEST_TIERS_SUSE"


def test_a_reserved_line_is_a_flat_monthly_value(folder, api):
    spec = spec_file(folder, HEAD + "\n[node]\nservice: ecs\nflavor: s3.large.2\nos: Linux\nquantity: 3\nmode: r12\n"
                     "assumption: Booked for a year.\n")
    assert run(["build", spec]) == 0
    values = sheet_cells(folder / "est.xlsx", values=True)
    assert values["G5"] == "Reserved 12 months" and values["I5"] == 45 and "K5" not in values
    assert values["L5"] == 135 and values["M5"] == 1620
    md = (folder / "est.md").read_text(encoding="utf-8")
    assert "Reserved rates are monthly values" in md


def test_a_record_is_chosen_by_id_and_a_tiered_rate_is_noted(folder, api):
    spec = spec_file(folder, HEAD + "\n[uplink]\nservice: vpc\nid: TEST_BW_1\nquantity: 10\nhours: 720\n"
                     "\n[disk]\nservice: evs\nid: TEST_EVS_SSD\nquantity: 10\n")
    assert run(["build", spec]) == 0
    values = sheet_cells(folder / "est.xlsx", values=True)
    assert values["L5"] == 216 and values["K5"] == 720 and values["L6"] == 1.2
    src = sheet_cells(folder / "est.xlsx", "Sources")
    assert any("tiered rate" in str(v) for v in src.values())


def test_the_price_api_not_answering_writes_nothing(folder, api, capsys):
    api.mode = "html"
    spec = spec_file(folder, HEAD + WEB % "1")
    assert run(["build", spec]) == 2
    assert "could not be read" in capsys.readouterr().err
    assert not (folder / "est.xlsx").exists()


def test_an_empty_region_is_refused_for_every_line(folder, api, capsys):
    spec = spec_file(folder, "title: x\nregion: eu-ch2\n" + WEB % "1")
    assert run(["build", spec]) == 1
    assert "holds no records for eu-ch2" in capsys.readouterr().err


def test_a_line_in_another_region_uses_that_region(folder, api):
    body = HEAD + WEB % "1" + "\n[nl servers]\nservice: ecs\nflavor: s3.large.2\nregion: eu-nl\nquantity: 1\n"
    spec = spec_file(folder, body)
    assert run(["build", spec]) == 0
    values = sheet_cells(folder / "est.xlsx", values=True)
    assert values["F6"] == "eu-nl" and values["O6"] == "TEST_NL_S3_LARGE_2_LNX" and values["L6"] == 64.8
    src = sheet_cells(folder / "est.xlsx", "Sources")
    assert "eu-de, eu-nl" in src.values()
    assert sum(1 for v in src.values() if str(v).startswith("20") and v.endswith(" UTC")) == 2


# --------------------------------------------------------------------------- German


def test_german_headers_and_fixed_texts(folder, api):
    spec = spec_file(folder, HEAD + WEB % "4" + DISKS)
    assert run(["build", spec, "--lang", "de"]) == 0
    xlsx = folder / "est.xlsx"
    wb = openpyxl.load_workbook(xlsx)
    assert wb.sheetnames == ["Positionen", "Quellen"]
    cells = sheet_cells(xlsx)
    heads = [cells["%s4" % c] for c in "ABCDEFGHIJKLMNO"]
    assert heads == ["Pos.", "Position", "Dienst", "Flavor oder Produkt", "Betriebssystem", "Region", "Preismodell",
                     "Einheit", "Einzelpreis (EUR)", "Menge", "Stunden pro Monat", "Monatlich (EUR)",
                     "Jährlich (EUR)", "Annahme", "Preis-ID"]
    assert cells["N5"] == "Rund um die Uhr in zwei Zonen."          # the German variant of the assumption
    assert cells["N6"] == "500 GB per server."                       # no variant: the plain text
    assert (cells["B9"], cells["B10"], cells["B11"]) == ("Summe monatlich (EUR)", "Summe jährlich (EUR)",
                                                         "Monate pro Jahr")
    assert "Preis-API" in cells["A2"] and re.search(r"\d\d\.\d\d\.\d{4} \d\d:\d\d UTC", cells["A2"])
    src = sheet_cells(xlsx, "Quellen")
    assert src["A3"] == "Angabe" and "Stand der Leistungsbeschreibung" in src.values()
    assert any(str(v).startswith("Stand der Preis-API") for v in src.values())
    md = (folder / "est.md").read_text(encoding="utf-8")
    assert "Summen" in md and "Monatlich: 484,80 EUR" in md and "Jährlich: 5.817,60 EUR" in md
    assert "Annahmen" in md and "Rund um die Uhr in zwei Zonen." in md


def test_the_language_of_the_specification_is_the_default(folder, api):
    spec = spec_file(folder, HEAD + "language: de\n" + WEB % "1")
    assert run(["build", spec]) == 0
    assert openpyxl.load_workbook(folder / "est.xlsx").sheetnames == ["Positionen", "Quellen"]
    assert run(["build", spec, "--lang", "en"]) == 0
    assert openpyxl.load_workbook(folder / "est.xlsx").sheetnames == ["Lines", "Sources"]


# --------------------------------------------------------------------------- exact arithmetic


def test_a_sum_that_floats_get_wrong_is_exact(folder, api):
    body = HEAD + """
[a]
service: evs
flavor: dime.gb
quantity: 1

[b]
service: evs
flavor: twodimes.gb
quantity: 1

[c]
service: evs
flavor: trap.gb
quantity: 1
"""
    spec = spec_file(folder, body)
    assert run(["build", spec]) == 0
    assert 0.1 + 0.2 != 0.3 and round(1.005, 2) == 1.0                # what floats do
    with zipfile.ZipFile(folder / "est.xlsx") as zf:
        sheet = zf.read("xl/worksheets/sheet1.xml").decode()
    assert "<v>0.1</v>" in sheet and "<v>0.2</v>" in sheet
    assert "<v>1.01</v>" in sheet                                      # 1.005 rounds half up to the cent
    assert "<f>SUM(L5:L7)</f><v>1.31</v>" in sheet                     # 0.10 + 0.20 + 1.01, not 1.3100000000000001
    assert "<f>SUM(M5:M7)</f><v>15.72</v>" in sheet
    assert "0000000" not in sheet                                      # no float noise anywhere
    values = sheet_cells(folder / "est.xlsx", values=True)
    assert values["L7"] == 1.01 and values["C10"] == 1.31
    md = (folder / "est.md").read_text(encoding="utf-8")
    assert "Monthly: 1.31 EUR" in md and "Yearly: 15.72 EUR" in md


def test_the_hours_and_the_rounding_follow_the_calculation_rules(folder, api):
    body = HEAD + "\n[a]\nservice: ecs\nflavor: s3.xlarge.2\nquantity: 3\nhours: 730\n"
    spec = spec_file(folder, body)
    assert run(["build", spec]) == 0
    values = sheet_cells(folder / "est.xlsx", values=True)
    assert values["K5"] == 730 and values["L5"] == 372.3 and values["M5"] == 4467.6      # 0.17 * 3 * 730


# --------------------------------------------------------------------------- the summary for a mail


def test_the_summary_names_totals_three_biggest_lines_and_assumptions(folder, api):
    body = (HEAD + WEB % "4" + DISKS + "\n[small]\nservice: evs\nflavor: dime.gb\nquantity: 10\n"
            "assumption: Scratch space.\n\n[tiny]\nservice: evs\nflavor: dime.gb\nquantity: 1\n")
    spec = spec_file(folder, body)
    assert run(["build", spec]) == 0
    md = (folder / "est.md").read_text(encoding="utf-8")
    assert md.startswith("# Sizing for %s\n" % fx.CUSTOMER_CODE)
    assert "- Monthly: 485.90 EUR" in md and "- Yearly: 5,830.80 EUR" in md
    biggest = md.split("## Biggest lines")[1].split("##")[0]
    lines = [l for l in biggest.splitlines() if l.strip()]
    assert [l.split(":")[0] for l in lines] == ["1. web servers", "2. data disks", "3. small"]
    assert "244.80 EUR per month (50.4 % of the monthly total)" in lines[0]
    assumptions = md.split("## Assumptions")[1]
    assert "- web servers: Two zones and all day." in assumptions and "- small: Scratch space." in assumptions
    assert "tiny:" not in assumptions
    assert "720 hours a month" in assumptions
    assert chr(0x2014) not in md
    fx.assert_no_fixture_name(md, "the summary")


@pytest.mark.parametrize("lang", ["en", "de"])
def test_the_summary_passes_the_writing_check_for_a_mail(folder, api, lang, capsys):
    spec = spec_file(folder, HEAD + WEB % "4" + DISKS)
    assert run(["build", spec, "--lang", lang, "--no-record"]) == 0
    capsys.readouterr()
    assert cli.main(["write", "check", str(folder / "est.md"), "--mode", "mail"]) == 0
    assert "0 blocking" in capsys.readouterr().out


def test_the_texts_the_command_writes_have_no_em_dash_and_no_comma_before_and_or():
    for lang, table in pricebook.TEXT.items():
        for key, text in table.items():
            assert chr(0x2014) not in text and chr(0x2013) not in text, (lang, key)
            assert not re.search(r",\s+(?:and|or|und|oder)\b", text), (lang, key)


# --------------------------------------------------------------------------- the totals in the project


def test_the_totals_are_recorded_for_the_review(folder, api, capsys):
    spec = spec_file(folder, HEAD + WEB % "4" + DISKS)
    assert run(["build", spec]) == 0
    assert "recorded  K-1, K-2" in capsys.readouterr().out
    from awb import calc
    first, second = calc.find(folder, "K-1"), calc.find(folder, "K-2")
    assert first.result == "484.80" and first.expression == "sum(244.80, 240.00)"
    assert second.result == "5817.60" and second.expression == "sum(244.80, 240.00) * 12"
    assert calc.appears_in(first.value(), "The monthly total is 484.80 EUR.")
    src = sheet_cells(folder / "est.xlsx", "Sources")
    assert "K-1, K-2" in src.values()
    assert "Recorded calculations: K-1, K-2." in (folder / "est.md").read_text(encoding="utf-8")
    assert run(["build", spec, "--no-record"]) == 0
    assert [r.id for r in calc.load(folder)] == ["K-1", "K-2"]


def test_outside_a_project_nothing_is_recorded(tmp_path, api, capsys):
    spec = spec_file(tmp_path, HEAD + WEB % "1")
    assert run(["build", spec]) == 0
    out = capsys.readouterr().out
    assert "not recorded" in out and not (tmp_path / "calc").exists()


# --------------------------------------------------------------------------- the files


def test_out_names_the_workbook_and_the_summary_goes_next_to_it(folder, api):
    (folder / "deliverables").mkdir()
    spec = spec_file(folder, HEAD + WEB % "1")
    assert run(["build", spec, "--out", folder / "deliverables" / "offer.xlsx"]) == 0
    assert (folder / "deliverables" / "offer.xlsx").is_file() and (folder / "deliverables" / "offer.md").is_file()
    assert not (folder / "est.xlsx").exists()
    assert run(["build", spec, "--out", folder / "offer.csv"]) == 2
    assert run(["build", spec, "--out", folder / "missing" / "offer.xlsx"]) == 2
    assert run(["build", spec, "--out", spec]) == 2


def test_the_same_specification_and_snapshot_give_the_same_bytes(folder, api):
    spec = spec_file(folder, HEAD + WEB % "4" + DISKS)
    when = datetime.datetime(2026, 10, 11, 12, 0, tzinfo=datetime.timezone.utc)

    def fetcher(service, region):
        return price.fetch(service, region, now=when)

    for name in ("one", "two"):
        pricebook.build(spec, out=folder / (name + ".xlsx"), fetcher=fetcher, record=False)
    assert (folder / "one.xlsx").read_bytes() == (folder / "two.xlsx").read_bytes()
    assert sheet_cells(folder / "one.xlsx", "Sources")["B5"] == "2026-10-11 12:00 UTC"


def test_a_bad_specification_writes_nothing_and_exits_two(folder, api, capsys):
    spec = spec_file(folder, HEAD + "[a]\nservice: ecs\n")
    assert run(["build", spec]) == 2
    assert "problem(s)" in capsys.readouterr().err
    assert run(["build", folder / "absent.txt"]) == 2
    assert api.calls == []


# --------------------------------------------------------------------------- the diff


def test_diff_of_a_changed_quantity(folder, api, capsys):
    old = spec_file(folder, HEAD + WEB % "2" + DISKS, "v1.txt")
    new = spec_file(folder, HEAD + WEB % "4" + DISKS, "v2.txt")
    assert run(["diff", old, new]) == 1
    out = capsys.readouterr().out
    assert "changed  web servers  quantity  2 -> 4" in out
    assert "changed  web servers  monthly  122.40 -> 244.80  (+100.0 %)" in out
    assert "changed  web servers  yearly  1468.80 -> 2937.60" in out
    assert "totals   monthly  362.40 -> 484.80  (+122.40 EUR, +33.8 %)" in out
    assert "totals   yearly  4348.80 -> 5817.60  (+1468.80 EUR, +33.8 %)" in out
    assert "data disks" not in out
    assert len(api.calls) == 1                                          # one snapshot serves both versions


def test_diff_without_a_change_exits_zero(folder, api, capsys):
    one = spec_file(folder, HEAD + WEB % "2", "v1.txt")
    two = spec_file(folder, HEAD + WEB % "2", "v2.txt")
    assert run(["diff", one, two]) == 0
    assert "no difference" in capsys.readouterr().out


def test_diff_lists_added_and_removed_lines_and_other_fields(folder, api, capsys):
    old = spec_file(folder, HEAD + WEB % "2" + DISKS, "v1.txt")
    new = spec_file(folder, HEAD + (WEB % "2").replace("mode: payg", "mode: r12").replace(
        "assumption: Two zones and all day.", "assumption: Three zones.") + "\n[backup disks]\nservice: evs\n"
        "flavor: dime.gb\nquantity: 500\n", "v2.txt")
    assert run(["diff", old, new]) == 1
    out = capsys.readouterr().out
    assert "removed  data disks  (evs evs.ssd, quantity 2000, 240.00 EUR per month)" in out
    assert "added    backup disks  (evs dime.gb, quantity 500, 50.00 EUR per month)" in out
    assert "web servers  pricing mode  PAYG -> R12" in out
    assert "web servers  assumption  Two zones and all day. -> Three zones." in out
    assert "web servers  unit price  0.085 -> 45" in out
    assert "web servers  hours  720 -> -" in out


def test_diff_without_prices_compares_the_settings_and_asks_no_api(folder, api, capsys):
    old = spec_file(folder, HEAD + WEB % "2", "v1.txt")
    new = spec_file(folder, HEAD + WEB % "3", "v2.txt")
    assert run(["diff", old, new, "--no-prices"]) == 1
    out = capsys.readouterr().out
    assert "web servers  quantity  2 -> 3" in out and "monthly" not in out.split("totals")[0]
    assert "not compared" in out
    assert api.calls == []


def test_diff_of_two_workbooks_and_of_a_workbook_against_a_specification(folder, api, capsys):
    v1 = spec_file(folder, HEAD + WEB % "2" + DISKS, "v1.txt")
    v2 = spec_file(folder, HEAD + WEB % "4" + DISKS, "v2.txt")
    assert run(["build", v1, "--no-record"]) == 0 and run(["build", v2, "--no-record"]) == 0
    capsys.readouterr()
    assert run(["diff", folder / "v1.xlsx", folder / "v2.xlsx"]) == 1
    out = capsys.readouterr().out
    assert "changed  web servers  quantity  2 -> 4" in out and "totals   monthly  362.40 -> 484.80" in out
    assert run(["diff", folder / "v2.xlsx", v2]) == 0                   # the same estimate, built and specified
    assert run(["diff", folder / "v1.xlsx", v2]) == 1
    assert "quantity  2 -> 4" in capsys.readouterr().out


def test_diff_reads_a_german_workbook_and_an_english_one_alike(folder, api, capsys):
    v = spec_file(folder, HEAD + WEB % "2" + DISKS, "v1.txt")
    assert run(["build", v, "--lang", "de", "--out", folder / "de.xlsx", "--no-record"]) == 0
    assert run(["build", v, "--lang", "en", "--out", folder / "en.xlsx", "--no-record"]) == 0
    capsys.readouterr()
    assert run(["diff", folder / "de.xlsx", folder / "en.xlsx", "--lang", "en"]) == 1
    out = capsys.readouterr().out
    assert "assumption" in out and "quantity" not in out and "totals   monthly  362.40 -> 362.40" in out


def test_diff_reports_a_price_that_moved_between_two_workbooks(folder, api, capsys):
    spec = spec_file(folder, HEAD + WEB % "2", "v1.txt")
    assert run(["build", spec, "--out", folder / "old.xlsx", "--no-record"]) == 0
    for r in api.records:
        if r["id"] == "TEST_S3_LARGE_2_LNX":
            r["priceAmount"] = "0.100000 EUR"
    assert run(["build", spec, "--out", folder / "new.xlsx", "--no-record"]) == 0
    capsys.readouterr()
    assert run(["diff", folder / "old.xlsx", folder / "new.xlsx"]) == 1
    out = capsys.readouterr().out
    assert "unit price  0.085 -> 0.1" in out and "monthly  122.40 -> 144.00" in out


def test_diff_with_the_price_api_unknown_exits_two(folder, api, capsys):
    one = spec_file(folder, HEAD + WEB % "2", "v1.txt")
    two = spec_file(folder, HEAD + WEB % "3", "v2.txt")
    api.mode = "html"
    assert run(["diff", one, two]) == 2
    assert "could not be read" in capsys.readouterr().err


def test_diff_marks_a_line_that_cannot_be_priced_and_compares_no_totals(folder, api, capsys):
    one = spec_file(folder, HEAD + WEB % "2" + "\n[backup]\nservice: csbs\nflavor: x\nquantity: 1\n", "v1.txt")
    two = spec_file(folder, HEAD + WEB % "3", "v2.txt")
    assert run(["diff", one, two]) == 1
    out = capsys.readouterr().out
    assert "not priced" in out and "totals   not compared" in out


# --------------------------------------------------------------------------- build against an older version


def test_build_against_a_workbook_adds_the_sheet_of_changes(folder, api, capsys):
    v1 = spec_file(folder, HEAD + WEB % "2" + DISKS, "v1.txt")
    v2 = spec_file(folder, HEAD + WEB % "4" + DISKS, "v2.txt")
    assert run(["build", v1, "--no-record"]) == 0
    assert run(["build", v2, "--against", folder / "v1.xlsx", "--no-record"]) == 0
    assert "against   3 change(s) in the lines" in capsys.readouterr().out
    xlsx = folder / "v2.xlsx"
    assert openpyxl.load_workbook(xlsx).sheetnames == ["Lines", "Sources", "Changes"]
    cells = sheet_cells(xlsx, "Changes")
    assert [cells["%s3" % c] for c in "ABCDE"] == ["Change", "Item", "Field", "Old", "New"]
    rows = {(cells["B%d" % r], cells["C%d" % r]): (cells["D%d" % r], cells["E%d" % r]) for r in (4, 5, 6)}
    assert rows[("web servers", "Quantity")] == (2, 4)
    assert rows[("web servers", "Monthly (EUR)")] == (122.4, 244.8)
    assert cells["A4"] == "changed"
    top = find_row(cells, "Monthly total (EUR)")
    assert cells["C%d" % top] == 362.4
    assert cells["D%d" % top] == "='Lines'!$C$9" and cells["E%d" % top] == "=D%d-C%d" % (top, top)
    values = sheet_cells(xlsx, "Changes", values=True)
    assert values["D%d" % top] == 484.8 and values["E%d" % top] == 122.4
    yearly = find_row(cells, "Yearly total (EUR)")
    assert cells["D%d" % yearly] == "='Lines'!$C$10" and values["E%d" % yearly] == 1468.8
    md = (folder / "v2.md").read_text(encoding="utf-8")
    assert "## Changes against the old version" in md and "web servers, Quantity: 2 -> 4" in md
    assert "Monthly: from 362.40 to 484.80 EUR (difference +122.40)" in md


def test_build_against_a_specification_and_in_german(folder, api):
    v1 = spec_file(folder, HEAD + WEB % "2" + DISKS, "v1.txt")
    v2 = spec_file(folder, HEAD + (WEB % "2") + "\n[extra]\nservice: evs\nflavor: dime.gb\nquantity: 100\n", "v2.txt")
    assert run(["build", v2, "--against", v1, "--lang", "de", "--no-record"]) == 0
    xlsx = folder / "v2.xlsx"
    assert openpyxl.load_workbook(xlsx).sheetnames == ["Positionen", "Quellen", "Änderungen"]
    cells = sheet_cells(xlsx, "Änderungen")
    assert [cells["%s3" % c] for c in "ABCDE"] == ["Änderung", "Position", "Feld", "Alt", "Neu"]
    kinds = {cells["B%d" % r]: cells["A%d" % r] for r in range(4, 8) if "B%d" % r in cells}
    assert kinds == {"data disks": "entfallen", "extra": "neu"}
    assert "Summen gegenüber der alten Fassung" in cells.values()


def test_build_against_the_same_version_says_so(folder, api, capsys):
    v1 = spec_file(folder, HEAD + WEB % "2", "v1.txt")
    assert run(["build", v1, "--no-record"]) == 0
    assert run(["build", v1, "--against", folder / "v1.xlsx", "--no-record"]) == 0
    assert "against   no difference" in capsys.readouterr().out
    assert "Nothing changed." in (folder / "v1.md").read_text(encoding="utf-8")


def test_build_against_a_file_that_is_no_workbook_is_an_error(folder, api, capsys):
    v1 = spec_file(folder, HEAD + WEB % "2", "v1.txt")
    (folder / "old.xlsx").write_bytes(b"not a workbook")
    assert run(["build", v1, "--against", folder / "old.xlsx"]) == 2
    assert "nothing written" in capsys.readouterr().err
    assert not (folder / "v1.xlsx").exists()


# --------------------------------------------------------------------------- the writer and the reader


def test_the_writer_and_the_reader_agree_and_a_text_stays_a_text(tmp_path):
    a = xlsxlite.Sheet("One")
    a.put(1, 1, "=1+1", style="title")
    a.put(2, 1, Decimal("0.1"))
    a.put(2, 2, Decimal("0.2"))
    a.put(2, 3, Decimal("0.3"), formula="A2+B2", style="money")
    a.put(3, 1, "tab\there\x01 & <b>")
    a.put(4, 1, "")
    b = xlsxlite.Sheet("Two words")
    b.put(1, 1, "x")
    path = tmp_path / "t.xlsx"
    xlsxlite.write(path, [a, b], title="t", created="2026-10-11T12:00:00Z")
    wb = openpyxl.load_workbook(path)
    assert wb.sheetnames == ["One", "Two words"]
    ws = wb["One"]
    assert ws["A1"].data_type == "s" and ws["A1"].value == "=1+1"
    assert ws["C2"].value == "=A2+B2" and openpyxl.load_workbook(path, data_only=True)["One"]["C2"].value == 0.3
    read = xlsxlite.read(path)
    assert [s.name for s in read] == ["One", "Two words"]
    assert read[0].rows[2] == {1: Decimal("0.1"), 2: Decimal("0.2"), 3: Decimal("0.3")}
    assert read[0].rows[3][1] == "tab\there  & <b>" and read[0].rows[4][1] == ""
    with zipfile.ZipFile(path) as zf:
        assert 'fullCalcOnLoad="1"' in zf.read("xl/workbook.xml").decode()


def test_the_writer_refuses_what_a_workbook_cannot_hold():
    for bad in ("", "a/b", "x" * 32, "[x]", "'x"):
        with pytest.raises(xlsxlite.XlsxError):
            xlsxlite.build([xlsxlite.Sheet(bad)])
    with pytest.raises(xlsxlite.XlsxError):
        xlsxlite.build([xlsxlite.Sheet("A"), xlsxlite.Sheet("a")])
    with pytest.raises(xlsxlite.XlsxError):
        xlsxlite.build([])
    with pytest.raises(xlsxlite.XlsxError):
        xlsxlite.Sheet("A").put(0, 1, "x")
    with pytest.raises(xlsxlite.XlsxError):
        xlsxlite.Sheet("A").put(1, 1, "x", style="loud")
    with pytest.raises(xlsxlite.XlsxError):
        xlsxlite.build([xlsxlite.Sheet("A")], created="yesterday")


def test_the_reader_refuses_a_file_that_is_no_workbook_or_carries_a_declaration(tmp_path):
    (tmp_path / "a.xlsx").write_bytes(b"plain")
    with pytest.raises(xlsxlite.XlsxError):
        xlsxlite.read(tmp_path / "a.xlsx")
    with pytest.raises(xlsxlite.XlsxError):
        xlsxlite.read(tmp_path / "missing.xlsx")
    path = tmp_path / "b.xlsx"
    sheet = xlsxlite.Sheet("A")
    sheet.put(1, 1, "x")
    xlsxlite.write(path, [sheet])
    evil = tmp_path / "evil.xlsx"
    with zipfile.ZipFile(path) as src, zipfile.ZipFile(evil, "w") as out:
        for info in src.infolist():
            data = src.read(info.filename)
            if info.filename == "xl/sharedStrings.xml":
                data = b'<?xml version="1.0"?><!DOCTYPE x [<!ENTITY a "aaaa">]>' + data.split(b"?>", 1)[1]
            out.writestr(info, data)
    with pytest.raises(xlsxlite.XlsxError, match="declaration"):
        xlsxlite.read(evil)


def test_a_workbook_that_is_no_lines_sheet_is_refused(tmp_path):
    sheet = xlsxlite.Sheet("A")
    sheet.put(1, 1, "x")
    xlsxlite.write(tmp_path / "w.xlsx", [sheet])
    with pytest.raises(pricebook.PricebookError, match="not a lines sheet"):
        pricebook.book_of_workbook(tmp_path / "w.xlsx")


# --------------------------------------------------------------------------- the rules of the repository


def test_runtime_code_uses_the_standard_library_and_the_package_only():
    import sys
    stdlib = set(sys.stdlib_module_names)
    for rel in ("awb/tcp/pricebook.py", "awb/xlsxlite.py"):
        tree = ast.parse((ROOT / rel).read_text(encoding="utf-8"))
        for node in ast.walk(tree):
            names = []
            if isinstance(node, ast.Import):
                names = [a.name.split(".")[0] for a in node.names]
            elif isinstance(node, ast.ImportFrom) and node.level == 0:
                names = [(node.module or "").split(".")[0]]
            for name in names:
                assert name in stdlib or name == "awb", (rel, name)


def test_the_command_is_wired_and_documented():
    assert "pricebook" in cli.DELEGATED
    assert cli.main(["pricebook"]) == 2
    for doc in ("COMMANDS.md", "README.md"):
        text = (ROOT / doc).read_text(encoding="utf-8")
        assert "awb pricebook build" in text and "awb pricebook diff" in text, doc
    changelog = (ROOT / "CHANGELOG.md").read_text(encoding="utf-8")
    assert "awb pricebook build" in changelog.split("\n## Week 40")[0]


def test_the_rows_of_this_file_are_in_the_test_list():
    rows = json.loads((ROOT / "build" / "tests.json").read_text(encoding="utf-8"))
    mine = [t for t in rows["tests"] if t["id"].startswith("PB-")]
    assert len(mine) >= 8 and all("tests/test_pricebook.py" in t["run"] for t in mine)
    assert rows["total_tests"] == len(rows["tests"])
