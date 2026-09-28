"""Office readers: docx, xlsx, pptx and odt, read straight from the zip container.

Every document is built here with invented fixture names. The docx, xlsx and pptx files are made with
python-docx, openpyxl and python-pptx; parts those libraries cannot write (tracked deletions, cached
formula values, the company property) are written into the XML by hand afterwards.
"""
from __future__ import annotations

import re
import zipfile
from pathlib import Path

import docx
import openpyxl
import pptx
import pytest
from openpyxl.comments import Comment
from openpyxl.workbook.defined_name import DefinedName

from awb.extract import Extraction, extract, sniff
from awb.extract import office
from tests import fixtures as fx

CUSTOMER = fx.CUSTOMER_FORMS[0]        # full legal name
CUSTOMER_SHORT = fx.CUSTOMER_FORMS[1]
CUSTOMER_EN = fx.CUSTOMER_FORMS[3]
PERSON = fx.PERSON_FORMS[0]
PERSON_SHORT = fx.PERSON_FORMS[1]
ORG = fx.ORG_FORMS[0]
ORG_SHORT = fx.ORG_FORMS[1]
LAWFIRM_SHORT = fx.LAWFIRM_FORMS[1]
PLACE = fx.PLACE_FORMS[0]


# --------------------------------------------------------------------------- helpers


def rewrite_zip(path: Path, edit) -> None:
    """Rewrite every member of the zip through `edit(name, data) -> data`, keeping the order."""
    with zipfile.ZipFile(path) as zf:
        members = [(i, zf.read(i.filename)) for i in zf.infolist()]
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as zf:
        for info, data in members:
            zf.writestr(info.filename, edit(info.filename, data))


def assert_clean_notes(ex: Extraction) -> None:
    fx.assert_no_fixture_name("\n".join(ex.notes), "notes")


def lines_of(text: str) -> list[str]:
    return [line.strip() for line in text.splitlines()]


# --------------------------------------------------------------------------- docx


def make_docx(path: Path) -> Path:
    d = docx.Document()
    d.add_heading("angebot", level=1)
    p = d.add_paragraph()
    p.add_run("vertrag mit ")
    # the legal name split over two runs, as Word does after a spelling check or a format change
    p.add_run(CUSTOMER[:3])
    run = p.add_run(CUSTOMER[3:])
    p.add_run(" vom 1. oktober")
    d.add_comment(run, text="bitte mit %s abstimmen" % ORG_SHORT, author=PERSON, initials="tb")

    keep = d.add_paragraph()
    keep.add_run("es bleibt MARKDEL")

    table = d.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "rolle"
    table.cell(0, 1).text = "firma"
    table.cell(1, 0).text = "kunde"
    table.cell(1, 1).text = CUSTOMER

    section = d.sections[0]
    section.header.paragraphs[0].text = "vertraulich für %s" % LAWFIRM_SHORT
    section.footer.paragraphs[0].text = "%s, seite 1" % PLACE

    d.core_properties.author = PERSON
    d.core_properties.last_modified_by = PERSON_SHORT
    d.core_properties.title = "angebot %s" % fx.TENDER_ID
    d.save(path)

    deletion = (
        '</w:t></w:r><w:del w:id="91" w:author="%s" w:date="2026-09-22T08:00:00Z">'
        "<w:r><w:delText>%s</w:delText></w:r><w:r><w:delText>%s</w:delText></w:r></w:del>"
        % (PERSON_SHORT, ORG[:4], ORG[4:])
    )

    def edit(name: str, data: bytes) -> bytes:
        if name == "word/document.xml":
            xml = data.decode("utf-8")
            assert xml.count("MARKDEL</w:t></w:r>") == 1
            return xml.replace("MARKDEL</w:t></w:r>", deletion).encode("utf-8")
        if name == "docProps/app.xml":
            xml = re.sub(r"<Company>.*?</Company>|<Company/>", "", data.decode("utf-8"))
            return xml.replace("</Properties>", "<Company>%s</Company></Properties>" % CUSTOMER_EN).encode("utf-8")
        return data

    rewrite_zip(path, edit)
    return path


@pytest.fixture
def docx_ex(tmp_path) -> Extraction:
    return extract(make_docx(tmp_path / "angebot.docx"))


def test_docx_is_sniffed_by_content(tmp_path):
    path = make_docx(tmp_path / "angebot.docx")
    renamed = path.rename(tmp_path / "anhang.dat")
    assert sniff(renamed) == "docx"
    assert extract(renamed).kind == "docx"


def test_docx_split_runs_are_joined(docx_ex):
    assert docx_ex.state == "ok"
    assert docx_ex.kind == "docx"
    assert "vertrag mit %s vom 1. oktober" % CUSTOMER in docx_ex.text
    assert "# angebot" in lines_of(docx_ex.text)
    assert_clean_notes(docx_ex)


def test_docx_header_and_footer_in_detect_text_only(docx_ex):
    header = "vertraulich für %s" % LAWFIRM_SHORT
    footer = "%s, seite 1" % PLACE
    assert header in docx_ex.detect_text
    assert footer in docx_ex.detect_text
    assert LAWFIRM_SHORT not in docx_ex.text
    assert PLACE not in docx_ex.text
    assert docx_ex.meta["header_footer_parts"] >= 2


def test_docx_comment_with_author(docx_ex):
    assert "comment by %s" % PERSON in docx_ex.detect_text
    assert "bitte mit %s abstimmen" % ORG_SHORT in docx_ex.detect_text
    assert ORG_SHORT not in docx_ex.text
    assert docx_ex.meta["comment_count"] == 1


def test_docx_core_and_app_properties(docx_ex):
    meta = docx_ex.meta
    assert meta["author"] == PERSON
    assert meta["creator"] == PERSON
    assert meta["last_modified_by"] == PERSON_SHORT
    assert meta["title"] == "angebot %s" % fx.TENDER_ID
    assert meta["company"] == CUSTOMER_EN
    for value in (PERSON, PERSON_SHORT, fx.TENDER_ID, CUSTOMER_EN):
        assert value in docx_ex.detect_text
    assert CUSTOMER_EN not in docx_ex.text


def test_docx_tracked_deletion_in_detect_text(docx_ex):
    # the deletion was split over two runs: detection must see it as one string
    assert "deleted: %s" % ORG in docx_ex.detect_text
    assert "revision by %s" % PERSON_SHORT in docx_ex.detect_text
    assert "es bleibt" in docx_ex.text
    assert ORG not in docx_ex.text
    assert docx_ex.meta["deleted_count"] == 1


def test_docx_table_row_as_markdown(docx_ex):
    rows = lines_of(docx_ex.text)
    assert "| rolle | firma |" in rows
    assert "| kunde | %s |" % CUSTOMER in rows
    header_index = rows.index("| rolle | firma |")
    assert rows[header_index + 1] == "| --- | --- |"


def test_docx_text_is_part_of_detect_text(docx_ex):
    assert docx_ex.detect_text.startswith(docx_ex.text)


def test_docx_separate_deletions_are_not_glued(tmp_path):
    d = docx.Document()
    d.add_paragraph("eins MARKA zwei MARKB drei")
    path = tmp_path / "zwei.docx"
    d.save(path)

    def edit(name, data):
        if name != "word/document.xml":
            return data
        xml = data.decode("utf-8")
        xml = xml.replace("MARKA", '</w:t></w:r><w:del w:id="1"><w:r><w:delText>%s</w:delText></w:r></w:del><w:r><w:t xml:space="preserve">' % PERSON_SHORT)
        xml = xml.replace("MARKB", '</w:t></w:r><w:del w:id="2"><w:r><w:delText>und</w:delText></w:r></w:del><w:r><w:t xml:space="preserve">')
        return xml.encode("utf-8")

    rewrite_zip(path, edit)
    ex = extract(path)
    assert "deleted: %s" % PERSON_SHORT in lines_of(ex.detect_text)
    assert "deleted: und" in lines_of(ex.detect_text)
    assert PERSON_SHORT + "und" not in ex.detect_text
    assert ex.text == "eins  zwei  drei"


def test_docx_nested_table_rows_are_not_repeated(tmp_path):
    d = docx.Document()
    outer = d.add_table(rows=1, cols=2)
    outer.cell(0, 0).text = "aussen"
    inner = outer.cell(0, 1).add_table(rows=1, cols=1)
    inner.cell(0, 0).text = "innen"
    path = tmp_path / "nested.docx"
    d.save(path)
    ex = extract(path)
    assert ex.text.count("innen") == 1
    assert ex.text.count("aussen") == 1


def test_docx_without_main_part_fails(tmp_path):
    path = tmp_path / "broken.docx"
    with zipfile.ZipFile(path, "w") as zf:
        zf.writestr("word/styles.xml", "<styles/>")
    assert sniff(path) == "docx"
    ex = extract(path)
    assert ex.state == "failed"
    assert ex.text == ""


# --------------------------------------------------------------------------- xlsx

FORMULA = 'CONCATENATE("%s","-",Z99)' % CUSTOMER_SHORT
CACHED = "%s-%s" % (CUSTOMER_SHORT, fx.FILE_NUMBER)   # stale cached value: Z99 is empty now


def make_xlsx(path: Path) -> Path:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "%s plan" % ORG_SHORT
    ws["A1"] = "kunde"
    ws["B1"] = "=" + FORMULA
    ws["C1"] = LAWFIRM_SHORT
    ws["A2"] = "summe"
    ws["B2"] = 42
    ws.column_dimensions["C"].hidden = True
    ws["A1"].comment = Comment("rückfrage zu %s" % fx.TENDER_ID, PERSON)

    hidden = wb.create_sheet("geheim")
    hidden["A1"] = PLACE
    hidden.sheet_state = "hidden"

    name = "%s_kunden" % CUSTOMER_SHORT
    wb.defined_names[name] = DefinedName(name, attr_text="geheim!$A$1")
    wb.save(path)

    def edit(member: str, data: bytes) -> bytes:
        if member != "xl/worksheets/sheet1.xml":
            return data
        xml = data.decode("utf-8")
        new, n = re.subn(
            r'<c r="B1"([^>]*)><f>(.*?)</f>(?:<v\s*/>|<v>\s*</v>)?</c>',
            lambda m: '<c r="B1"%s t="str"><f>%s</f><v>%s</v></c>' % (m.group(1), m.group(2), CACHED),
            xml,
        )
        assert n == 1
        return new.encode("utf-8")

    rewrite_zip(path, edit)
    return path


@pytest.fixture
def xlsx_ex(tmp_path) -> Extraction:
    return extract(make_xlsx(tmp_path / "plan.xlsx"))


def test_xlsx_visible_cells_as_markdown(xlsx_ex):
    assert xlsx_ex.state == "ok"
    assert xlsx_ex.kind == "xlsx"
    rows = lines_of(xlsx_ex.text)
    assert "| kunde | %s |" % CACHED in rows
    assert "| summe | 42 |" in rows
    assert_clean_notes(xlsx_ex)


def test_xlsx_hidden_sheet_in_detect_text_only(xlsx_ex):
    assert "hidden sheet: geheim" in xlsx_ex.detect_text
    assert PLACE in xlsx_ex.detect_text
    assert PLACE not in xlsx_ex.text
    assert xlsx_ex.meta["hidden_sheet_count"] == 1
    assert xlsx_ex.meta["sheet_count"] == 2


def test_xlsx_hidden_column_in_detect_text_only(xlsx_ex):
    assert LAWFIRM_SHORT in xlsx_ex.detect_text
    assert LAWFIRM_SHORT not in xlsx_ex.text


def test_xlsx_comment_with_author(xlsx_ex):
    assert "comment on A1 by %s: rückfrage zu %s" % (PERSON, fx.TENDER_ID) in xlsx_ex.detect_text
    assert fx.TENDER_ID not in xlsx_ex.text
    assert xlsx_ex.meta["comment_count"] == 1


def test_xlsx_formula_text_and_cached_value(xlsx_ex):
    assert "formula B1: =%s -> %s" % (FORMULA, CACHED) in xlsx_ex.detect_text
    # the string literal of the formula and the stale cached value are both seen
    assert '"%s"' % CUSTOMER_SHORT in xlsx_ex.detect_text
    assert fx.FILE_NUMBER in xlsx_ex.detect_text


def test_xlsx_defined_name_and_sheet_name(xlsx_ex):
    assert "defined name %s_kunden: geheim!$A$1" % CUSTOMER_SHORT in xlsx_ex.detect_text
    assert "sheet: %s plan" % ORG_SHORT in xlsx_ex.detect_text
    assert "sheet: geheim" in xlsx_ex.detect_text


def test_xlsx_wide_row_keeps_far_columns(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = "links"
    ws.cell(row=1, column=3000, value=PERSON_SHORT)
    path = tmp_path / "breit.xlsx"
    wb.save(path)
    ex = extract(path)
    assert PERSON_SHORT in ex.text
    assert "links" in ex.text


def test_xlsx_hidden_row_in_detect_text_only(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = "sichtbar"
    ws["A2"] = PERSON
    ws.row_dimensions[2].hidden = True
    path = tmp_path / "zeilen.xlsx"
    wb.save(path)
    ex = extract(path)
    assert "sichtbar" in ex.text
    assert PERSON not in ex.text
    assert PERSON in ex.detect_text


# --------------------------------------------------------------------------- pptx


def make_pptx(path: Path) -> Path:
    prs = pptx.Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "angebot für %s" % CUSTOMER
    slide.placeholders[1].text = "migration nach %s" % PLACE
    slide.notes_slide.notes_text_frame.text = "%s vorher anrufen" % PERSON
    second = prs.slides.add_slide(prs.slide_layouts[1])
    second.shapes.title.text = "zeitplan"
    prs.save(path)
    return path


def test_pptx_slide_text_and_notes(tmp_path):
    ex = extract(make_pptx(tmp_path / "vortrag.pptx"))
    assert ex.state == "ok"
    assert ex.kind == "pptx"
    assert "angebot für %s" % CUSTOMER in ex.text
    assert "migration nach %s" % PLACE in ex.text
    assert "%s vorher anrufen" % PERSON in ex.text
    assert "zeitplan" in ex.text
    assert ex.text.index(CUSTOMER) < ex.text.index("zeitplan")
    assert ex.meta["slide_count"] == 2
    assert ex.meta["notes_count"] == 1
    assert ex.detect_text.startswith(ex.text)
    assert_clean_notes(ex)


# --------------------------------------------------------------------------- odt

ODF_NS = (
    'xmlns:office="urn:oasis:names:tc:opendocument:xmlns:office:1.0" '
    'xmlns:text="urn:oasis:names:tc:opendocument:xmlns:text:1.0" '
    'xmlns:style="urn:oasis:names:tc:opendocument:xmlns:style:1.0" '
    'xmlns:table="urn:oasis:names:tc:opendocument:xmlns:table:1.0" '
    'xmlns:meta="urn:oasis:names:tc:opendocument:xmlns:meta:1.0" '
    'xmlns:dc="http://purl.org/dc/elements/1.1/"'
)


def make_odt(path: Path) -> Path:
    content = (
        '<?xml version="1.0" encoding="UTF-8"?>'
        '<office:document-content %s office:version="1.2"><office:body><office:text>'
        '<text:tracked-changes><text:changed-region text:id="c1"><text:deletion>'
        "<office:change-info><dc:creator>%s</dc:creator><dc:date>2026-09-22T08:00:00</dc:date></office:change-info>"
        "<text:p>%s</text:p></text:deletion></text:changed-region></text:tracked-changes>"
        '<text:h text:outline-level="1">angebot</text:h>'
        "<text:p>vertrag mit <text:span>%s</text:span>%s<text:s/>vom 1. oktober</text:p>"
        "<text:p>kommentar<office:annotation><dc:creator>%s</dc:creator><text:p>prüfen</text:p></office:annotation></text:p>"
        "<table:table table:name=\"liste\"><table:table-column/><table:table-column/>"
        "<table:table-row><table:table-cell><text:p>kunde</text:p></table:table-cell>"
        "<table:table-cell><text:p>%s</text:p></table:table-cell></table:table-row></table:table>"
        "</office:text></office:body></office:document-content>"
        % (ODF_NS, PERSON_SHORT, ORG, CUSTOMER[:3], CUSTOMER[3:], PERSON, CUSTOMER_SHORT)
    )
    styles = (
        '<?xml version="1.0" encoding="UTF-8"?>'
        '<office:document-styles %s office:version="1.2"><office:master-styles>'
        '<style:master-page style:name="Standard"><style:header><text:p>vertraulich für %s</text:p></style:header>'
        "<style:footer><text:p>%s, seite 1</text:p></style:footer></style:master-page>"
        "</office:master-styles></office:document-styles>" % (ODF_NS, LAWFIRM_SHORT, PLACE)
    )
    meta = (
        '<?xml version="1.0" encoding="UTF-8"?>'
        '<office:document-meta %s office:version="1.2"><office:meta>'
        "<meta:initial-creator>%s</meta:initial-creator><dc:creator>%s</dc:creator>"
        "<dc:title>angebot %s</dc:title><meta:document-statistic meta:page-count=\"2\"/>"
        "</office:meta></office:document-meta>" % (ODF_NS, PERSON, PERSON_SHORT, fx.TENDER_ID)
    )
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as zf:
        zf.writestr(zipfile.ZipInfo("mimetype"), "application/vnd.oasis.opendocument.text", compress_type=zipfile.ZIP_STORED)
        zf.writestr("content.xml", content)
        zf.writestr("styles.xml", styles)
        zf.writestr("meta.xml", meta)
    return path


def test_odt_body_text_and_meta_author(tmp_path):
    path = make_odt(tmp_path / "angebot.odt")
    assert sniff(path) == "odt"
    ex = extract(path)
    assert ex.state == "ok"
    assert ex.kind == "odt"
    assert "vertrag mit %s vom 1. oktober" % CUSTOMER in ex.text
    assert "# angebot" in lines_of(ex.text)
    assert "| kunde | %s |" % CUSTOMER_SHORT in lines_of(ex.text)
    assert ex.meta["author"] == PERSON
    assert ex.meta["last_modified_by"] == PERSON_SHORT
    assert ex.meta["title"] == "angebot %s" % fx.TENDER_ID
    assert ex.meta["page_count"] == 2
    assert "author: %s" % PERSON in ex.detect_text
    assert_clean_notes(ex)


def test_odt_header_footer_annotation_and_deletion_in_detect_text_only(tmp_path):
    ex = extract(make_odt(tmp_path / "angebot.odt"))
    assert "vertraulich für %s" % LAWFIRM_SHORT in ex.detect_text
    assert "%s, seite 1" % PLACE in ex.detect_text
    assert "comment by %s: prüfen" % PERSON in ex.detect_text
    assert "deleted: %s" % ORG in ex.detect_text
    assert "revision by %s" % PERSON_SHORT in ex.detect_text
    for hidden in (LAWFIRM_SHORT, PLACE, ORG, "prüfen"):
        assert hidden not in ex.text


# --------------------------------------------------------------------------- failure states


def test_office_reader_rejects_a_broken_container(tmp_path):
    path = tmp_path / "kaputt.docx"
    path.write_bytes(b"PK\x03\x04" + b"\x00" * 64)
    ex = extract(path)
    assert ex.state == "failed"
    assert ex.text == ""
    assert_clean_notes(ex)


def test_office_reader_notes_carry_no_values(tmp_path):
    path = make_docx(tmp_path / "angebot.docx")

    def edit(name, data):
        return b"<w:hdr><broken" if name == "word/header1.xml" else data

    with zipfile.ZipFile(path) as zf:
        has_header = "word/header1.xml" in zf.namelist()
    assert has_header
    rewrite_zip(path, edit)
    ex = office.extract_office(path, "docx")
    assert ex.state == "ok"
    assert "1 part(s) not well-formed, skipped" in ex.notes
    assert_clean_notes(ex)


# --------------------------------------------------------------------------- review findings (2026-09-22): the
# second pass over every part


def _plain_docx(path: Path) -> Path:
    d = docx.Document()
    d.add_paragraph("angebot ohne namen")
    d.save(str(path))
    return path


def _add_parts(path: Path, parts: dict[str, bytes | str], edit=None) -> Path:
    with zipfile.ZipFile(path) as zf:
        members = [(i.filename, zf.read(i.filename)) for i in zf.infolist()]
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as zf:
        for name, data in members:
            zf.writestr(name, edit(name, data) if edit else data)
        for name, data in parts.items():
            zf.writestr(name, data)
    return path


def _in_body(xml: str):
    def edit(name, data):
        if name == "word/document.xml":
            return data.replace(b"</w:body>", xml.encode("utf-8") + b"</w:body>")
        return data
    return edit


W = 'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"'

HIDDEN_DOCX = [
    ("tooltip", {}, _in_body('<w:p><w:hyperlink w:tooltip="%s"><w:r><w:t>x</w:t></w:r></w:hyperlink></w:p>' % CUSTOMER)),
    ("fallback", {}, _in_body('<mc:AlternateContent xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006">'
                              '<mc:Fallback><w:p><w:r><w:t>%s</w:t></w:r></w:p></mc:Fallback></mc:AlternateContent>'
                              % CUSTOMER)),
    ("document variable", {"word/docvars.xml": '<w:settings %s><w:docVars><w:docVar w:name="k" w:val="%s"/>'
                           '</w:docVars></w:settings>' % (W, CUSTOMER)}, None),
    ("renamed header part", {"word/kopf1.xml": '<w:hdr %s><w:p><w:r><w:t>%s</w:t></w:r></w:p></w:hdr>' % (W, CUSTOMER)},
     None),
    ("custom xml", {"customXml/item9.xml": "<root><kunde>%s</kunde></root>" % CUSTOMER}, None),
    ("chart", {"word/charts/chart1.xml": '<c:chartSpace xmlns:c="urn:c"><c:v>%s</c:v></c:chartSpace>' % CUSTOMER}, None),
    ("alt chunk", {"word/afchunk.htm": "<html><body><p>%s</p></body></html>" % CUSTOMER}, None),
]


@pytest.mark.parametrize("label,parts,edit", HIDDEN_DOCX, ids=[h[0] for h in HIDDEN_DOCX])
def test_hidden_places_of_a_docx_reach_scan_text(tmp_path, label, parts, edit):
    path = _add_parts(_plain_docx(tmp_path / "a.docx"), parts, edit)
    ex = extract(path)
    assert ex.state == "ok"
    assert CUSTOMER in ex.scan_text, label
    assert CUSTOMER not in ex.text


def test_embedded_workbook_is_read_as_a_child(tmp_path):
    inner = make_xlsx(tmp_path / "inner.xlsx").read_bytes()
    path = _add_parts(_plain_docx(tmp_path / "a.docx"), {"word/embeddings/Microsoft_Excel_Worksheet.xlsx": inner})
    ex = extract(path)
    assert len(ex.children) == 1 and ex.children[0].kind == "xlsx"
    assert ex.children[0].meta["member"] == "word/embeddings/Microsoft_Excel_Worksheet.xlsx"
    assert "1 embedded package(s) read as parts of this file" in ex.notes
    assert not ex.meta.get("incomplete")


def test_ole_object_is_incomplete_and_its_strings_are_scanned(tmp_path):
    ole = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\x00" * 64 + CUSTOMER.encode("utf-16-le") + b"\x00" * 64
    path = _add_parts(_plain_docx(tmp_path / "a.docx"), {"word/embeddings/oleObject1.bin": ole})
    ex = extract(path)
    assert ex.meta["incomplete"] is True
    assert CUSTOMER in ex.scan_text
    assert "1 embedded object(s) not read, review the original" in ex.notes


def test_workbook_validation_and_number_format_reach_scan_text(tmp_path):
    from openpyxl.worksheet.datavalidation import DataValidation

    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = 5
    ws["A1"].number_format = '"%s "0' % PERSON
    rule = DataValidation(type="list", formula1='"%s,andere"' % CUSTOMER_SHORT, prompt=ORG)
    ws.add_data_validation(rule)
    rule.add("B2")
    path = tmp_path / "a.xlsx"
    wb.save(str(path))
    ex = extract(path)
    for form in (PERSON, CUSTOMER_SHORT, ORG):
        assert form in ex.scan_text


def test_odt_user_field_reaches_scan_text(tmp_path):
    path = tmp_path / "a.odt"
    with zipfile.ZipFile(path, "w") as zf:
        zf.writestr("mimetype", "application/vnd.oasis.opendocument.text")
        zf.writestr("content.xml",
                    '<office:document-content xmlns:office="urn:oasis:names:tc:opendocument:xmlns:office:1.0" '
                    'xmlns:text="urn:oasis:names:tc:opendocument:xmlns:text:1.0"><office:body><office:text>'
                    '<text:user-field-decls><text:user-field-decl text:name="k" office:string-value="%s"/>'
                    '</text:user-field-decls><text:p>angebot ohne namen</text:p></office:text></office:body>'
                    '</office:document-content>' % CUSTOMER)
    ex = extract(path)
    assert ex.kind == "odt" and CUSTOMER in ex.scan_text


def test_xml_part_above_the_element_cap_is_refused(tmp_path, monkeypatch):
    monkeypatch.setattr(office, "MAX_ELEMENTS", 50)
    d = docx.Document()
    for i in range(40):
        d.add_paragraph("zeile %d" % i)
    path = tmp_path / "gross.docx"
    d.save(str(path))
    ex = extract(path)
    assert ex.meta["incomplete"] is True
    assert any("not well-formed" in n for n in ex.notes)


# --------------------------------------------------------------------------- red team findings (2026-09-27): what
# the reader sees reaches text, parts are found through their relationships, side channels reach detect_text

R_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
REL_T = R_NS + "/"
MC_NS = "http://schemas.openxmlformats.org/markup-compatibility/2006"
V_NS = "urn:schemas-microsoft-com:vml"
A_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"
C_NS = "http://schemas.openxmlformats.org/drawingml/2006/chart"
DGM_NS = "http://schemas.openxmlformats.org/drawingml/2006/diagram"
WP_NS = "http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing"
M_NS = "http://schemas.openxmlformats.org/officeDocument/2006/math"
P_NS = "http://schemas.openxmlformats.org/presentationml/2006/main"
X_NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"


def _rewrite(path: Path, parts=None, edit=None, drop=()) -> Path:
    """Rewrite a zip: `edit(name, data)` per member, the names in `drop` removed, `parts` added."""
    with zipfile.ZipFile(path) as zf:
        members = [(i.filename, zf.read(i.filename)) for i in zf.infolist()]
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as zf:
        for name, data in members:
            if name in drop:
                continue
            zf.writestr(name, edit(name, data) if edit else data)
        for name, data in (parts or {}).items():
            zf.writestr(name, data)
    return path


def _add_rel(data: bytes, rid: str, rtype: str, target: str) -> bytes:
    rel = "<Relationship " + 'Id="%s" Type="%s" Target="%s"/>' % (rid, rtype, target)
    return data.replace(b"</Relationships>", rel.encode("utf-8") + b"</Relationships>")


def _mc(requires: str, choice: str, fallback: str) -> str:
    """An mc:AlternateContent element with one Choice needing `requires` and one Fallback."""
    return ('<mc:AlternateContent xmlns:mc="%s"><mc:Choice %s="%s">%s</mc:Choice><mc:Fallback>%s</mc:Fallback>'
            '</mc:AlternateContent>' % (MC_NS, "Requires", requires, choice, fallback))


def _w_p(text: str) -> str:
    return '<w:p><w:r><w:t xml:space="preserve">%s</w:t></w:r></w:p>' % text


def _docx_with(tmp_path: Path, body_xml: str, parts=None, rels=(), drop=()) -> Path:
    """A python-docx file with `body_xml` before the section properties, extra parts and document rels."""
    path = _plain_docx(tmp_path / "a.docx")

    def edit(name, data):
        if name == "word/document.xml":
            i = data.rfind(b"<w:sectPr")
            return data[:i] + body_xml.encode("utf-8") + data[i:]
        if name == "word/_rels/document.xml.rels":
            for rid, kind, target in rels:
                data = _add_rel(data, rid, REL_T + kind, target)
        return data
    return _rewrite(path, parts, edit, drop)


def _chart_xml(title: str, series: str, cats: list[str], label: str) -> str:
    """A bar chart part: rich title, series name and categories in string caches, one rich data label."""
    return (
        '<c:chartSpace xmlns:c="%s" xmlns:a="%s"><c:chart><c:title><c:tx><c:rich><a:p><a:r><a:t>%s</a:t></a:r></a:p>'
        '</c:rich></c:tx></c:title><c:plotArea><c:barChart><c:ser><c:tx><c:strRef><c:f>s!$B$1</c:f><c:strCache>'
        '<c:pt idx="0"><c:v>%s</c:v></c:pt></c:strCache></c:strRef></c:tx><c:dLbls><c:dLbl><c:tx><c:rich><a:p><a:r>'
        '<a:t>%s</a:t></a:r></a:p></c:rich></c:tx></c:dLbl></c:dLbls><c:cat><c:strRef><c:f>s!$A$2:$A$3</c:f>'
        '<c:strCache><c:pt idx="0"><c:v>%s</c:v></c:pt><c:pt idx="1"><c:v>%s</c:v></c:pt></c:strCache></c:strRef>'
        '</c:cat><c:val><c:numRef><c:numCache><c:pt idx="0"><c:v>1</c:v></c:pt></c:numCache></c:numRef></c:val>'
        '</c:ser></c:barChart></c:plotArea></c:chart></c:chartSpace>' % (C_NS, A_NS, title, series, label, cats[0], cats[1])
    )


def _drawing(inner: str, uri: str) -> str:
    return ('<w:p><w:r><w:drawing><wp:inline xmlns:wp="%s"><wp:docPr id="7" name="grafik"/><a:graphic xmlns:a="%s">'
            '<a:graphicData uri="%s">%s</a:graphicData></a:graphic></wp:inline></w:drawing></w:r></w:p>' % (WP_NS, A_NS, uri, inner))


def test_docx_chart_text_stands_where_the_chart_is(tmp_path):
    drawing = _drawing('<c:chart xmlns:c="%s" xmlns:r="%s" r:id="rIdChart1"/>' % (C_NS, R_NS), C_NS)
    chart = _chart_xml("umsatz nach kunde " + CUSTOMER, ORG_SHORT, ["nord", PLACE], "label " + PERSON_SHORT)
    path = _docx_with(tmp_path, drawing + _w_p("nach dem diagramm"), {"word/charts/chart1.xml": chart},
                      rels=[("rIdChart1", "chart", "charts/chart1.xml")])
    ex = extract(path)
    lines = lines_of(ex.text)
    assert "chart: umsatz nach kunde %s" % CUSTOMER in lines
    assert "chart: %s" % ORG_SHORT in lines
    assert "chart: nord | %s" % PLACE in lines
    assert "chart: label %s" % PERSON_SHORT in lines
    assert lines.index("chart: %s" % ORG_SHORT) < lines.index("nach dem diagramm")
    assert ex.text.count(ORG_SHORT) == 1


def test_docx_smartart_text_reaches_text(tmp_path):
    drawing = _drawing('<dgm:relIds xmlns:dgm="%s" xmlns:r="%s" r:dm="rIdDm1"/>' % (DGM_NS, R_NS), DGM_NS)
    data = ('<dgm:dataModel xmlns:dgm="%s" xmlns:a="%s"><dgm:ptLst><dgm:pt modelId="{1}" type="doc"><dgm:t><a:bodyPr/>'
            '<a:p/></dgm:t></dgm:pt><dgm:pt modelId="{2}"><dgm:t><a:bodyPr/><a:p><a:r><a:t>schritt eins %s</a:t></a:r>'
            '</a:p></dgm:t></dgm:pt></dgm:ptLst></dgm:dataModel>' % (DGM_NS, A_NS, CUSTOMER_SHORT))
    path = _docx_with(tmp_path, drawing, {"word/diagrams/data1.xml": data}, rels=[("rIdDm1", "diagramData", "diagrams/data1.xml")])
    ex = extract(path)
    assert "diagram: schritt eins %s" % CUSTOMER_SHORT in lines_of(ex.text)
    assert ex.text.count(CUSTOMER_SHORT) == 1


def test_docx_legacy_wordart_reaches_text(tmp_path):
    art = ('<w:p><w:r><w:pict><v:shape xmlns:v="%s" type="#_x0000_t136"><v:textpath fitpath="t" string="willkommen bei %s"/>'
           '</v:shape></w:pict></w:r></w:p>' % (V_NS, CUSTOMER))
    ex = extract(_docx_with(tmp_path, art))
    assert "willkommen bei %s" % CUSTOMER in lines_of(ex.text)


def test_docx_block_level_equation_reaches_text(tmp_path):
    math = '<m:oMathPara xmlns:m="%s"><m:oMath><m:r><m:t>x = umsatz %s</m:t></m:r></m:oMath></m:oMathPara>' % (M_NS, ORG_SHORT)
    cell = ('<w:tbl><w:tr><w:tc>%s</w:tc><w:tc><m:oMathPara xmlns:m="%s"><m:oMath><m:r><m:t>y = %s</m:t></m:r></m:oMath>'
            '</m:oMathPara></w:tc></w:tr></w:tbl>' % (_w_p("links"), M_NS, PLACE))
    ex = extract(_docx_with(tmp_path, math + cell))
    lines = lines_of(ex.text)
    assert "x = umsatz %s" % ORG_SHORT in lines
    assert "| links | y = %s |" % PLACE in lines


def test_docx_alternate_content_shows_the_branch_the_reader_shows(tmp_path):
    run = "<w:r><w:t>%s</w:t></w:r>"
    known = "<w:p>%s</w:p>" % _mc("w14", run % ("neuer text " + CUSTOMER_SHORT), run % ("alter text " + ORG_SHORT))
    unknown = "<w:p>%s</w:p>" % _mc("zz", run % ("unbekannt " + PERSON_SHORT), run % ("ersatz " + PLACE))
    ex = extract(_docx_with(tmp_path, known + unknown))
    det = lines_of(ex.detect_text)
    assert "neuer text %s" % CUSTOMER_SHORT in ex.text
    assert ORG_SHORT not in ex.text
    assert "alternate content: alter text %s" % ORG_SHORT in det
    assert "ersatz %s" % PLACE in ex.text
    assert PERSON_SHORT not in ex.text
    assert "alternate content: unbekannt %s" % PERSON_SHORT in det


def test_docx_alt_chunks_are_read_as_children(tmp_path):
    html = "<html><body><p>importierter absatz %s</p></body></html>" % CUSTOMER
    rtf = "{\\rtf1\\ansi\\pard rtf absatz %s\\par}" % ORG_SHORT
    txt = "text absatz %s\n" % PLACE
    mht = ("MIME-Version: 1.0\r\nContent-Type: multipart/related; boundary=\"b1\"\r\n\r\n--b1\r\n"
           "Content-Transfer-Encoding: quoted-printable\r\nContent-Type: text/html; charset=\"utf-8\"\r\n\r\n"
           "<html><body><p>mht absatz %s=\r\n%s</p></body></html>\r\n\r\n--b1--\r\n" % (PERSON[:6], PERSON[6:]))
    parts = {"word/c1.htm": html, "word/c2.rtf": rtf, "word/c3.txt": txt, "word/c4.mht": mht}
    body = "".join('<w:altChunk xmlns:r="%s" r:id="rIdC%d"/>' % (R_NS, i) for i in range(1, 5))
    rels = [("rIdC%d" % i, "aFChunk", name.rpartition("/")[2]) for i, name in enumerate(parts, 1)]
    ex = extract(_docx_with(tmp_path, body, parts, rels=rels))
    assert ex.text.count("[alt chunk, read as a separate part of this file]") == 4
    assert "4 alt chunk(s) read as parts of this file" in ex.notes
    assert not ex.meta.get("incomplete")
    texts = {c.meta["member"]: c.text for c in ex.children}
    assert "importierter absatz %s" % CUSTOMER in texts["word/c1.htm"]
    assert "rtf absatz %s" % ORG_SHORT in texts["word/c2.rtf"]
    assert "text absatz %s" % PLACE in texts["word/c3.txt"]
    assert "mht absatz %s" % PERSON in texts["word/c4.mht"]
    for form in (CUSTOMER, ORG_SHORT, PLACE, PERSON):
        assert form not in ex.text


def test_docx_note_and_header_parts_are_found_through_the_relationships(tmp_path):
    fn = '<w:footnotes %s><w:footnote w:id="1">%s</w:footnote></w:footnotes>' % (W, _w_p("fussnote zu %s" % CUSTOMER))
    hdr = '<w:hdr %s>%s</w:hdr>' % (W, _w_p("kopfzeile %s" % ORG_SHORT))
    ref = '<w:p><w:r><w:t>siehe note</w:t></w:r><w:r><w:footnoteReference w:id="1"/></w:r></w:p>'
    path = _docx_with(tmp_path, ref, {"word/fn.xml": fn, "word/kopf.xml": hdr},
                      rels=[("rIdFn", "footnotes", "fn.xml"), ("rIdHdr", "header", "kopf.xml")])
    ex = extract(path)
    assert "[footnote 1] fussnote zu %s" % CUSTOMER in lines_of(ex.text)
    assert "kopfzeile %s" % ORG_SHORT in lines_of(ex.detect_text)
    assert ORG_SHORT not in ex.text
    assert ex.meta["header_footer_parts"] >= 1


def test_docx_repeated_part_name_is_noted_as_incomplete(tmp_path):
    import warnings

    path = _plain_docx(tmp_path / "a.docx")
    with zipfile.ZipFile(path) as zf:
        members = [(i.filename, zf.read(i.filename)) for i in zf.infolist()]
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as zf:
            for name, data in members:
                zf.writestr(name, data)
                if name == "word/document.xml":
                    zf.writestr(name, data.replace(b"angebot ohne namen", b"zweite kopie"))
    ex = extract(path)
    assert ex.meta["incomplete"] is True
    assert "1 part name(s) repeated in the container, one copy read" in ex.notes


def test_docx_macro_project_and_embedded_font_are_noted(tmp_path):
    vba = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\x00" * 120
    font = b"\x00" * 32 + b"name" + b"\x00" * 64
    path = _docx_with(tmp_path, "", {"word/vbaProject.bin": vba, "word/fonts/font1.odttf": font},
                      rels=[("rIdVba", "vbaProject", "vbaProject.bin")])
    ex = extract(path)
    assert ex.meta["incomplete"] is True
    assert "macro project not read, review the original" in ex.notes
    assert "1 embedded font(s) not read, the document brings its own glyphs" in ex.notes
    assert ex.meta["macro_count"] == 1 and ex.meta["font_count"] == 1


def test_docx_binary_custom_property_is_decoded_into_detect_text(tmp_path):
    import base64

    blob = base64.b64encode(("kontakt %s" % PERSON).encode("utf-16-le")).decode("ascii")
    custom = ('<Properties xmlns="http://schemas.openxmlformats.org/officeDocument/2006/custom-properties" '
              'xmlns:vt="http://schemas.openxmlformats.org/officeDocument/2006/docPropsVTypes">'
              '<property pid="2" name="kontakt"><vt:blob>%s</vt:blob></property>'
              '<property pid="3" name="stufe"><vt:lpwstr>zwei</vt:lpwstr></property></Properties>' % blob)
    path = _rewrite(_plain_docx(tmp_path / "a.docx"), {"docProps/custom.xml": custom}, drop=("docProps/custom.xml",))
    ex = extract(path)
    assert "property kontakt: kontakt %s" % PERSON in lines_of(ex.detect_text)
    assert blob not in ex.detect_text
    assert ex.meta["custom"]["stufe"] == "zwei"
    assert PERSON not in ex.text


def test_docx_control_items_table_caption_and_numbering_reach_detect_text(tmp_path):
    sdt = ('<w:sdt><w:sdtPr><w:alias w:val="kunde"/><w:dropDownList><w:listItem w:displayText="%s" w:value="1"/>'
           '<w:listItem w:displayText="andere" w:value="2"/></w:dropDownList></w:sdtPr><w:sdtContent>%s</w:sdtContent></w:sdt>'
           % (CUSTOMER_SHORT, _w_p("bitte wählen")))
    tbl = ('<w:tbl><w:tblPr><w:tblCaption w:val="tabelle %s"/><w:tblDescription w:val="beschreibung %s"/></w:tblPr>'
           '<w:tr><w:tc>%s</w:tc></w:tr></w:tbl>' % (ORG_SHORT, PLACE, _w_p("zelle")))
    numbering = ('<w:numbering %s><w:abstractNum w:abstractNumId="0"><w:lvl w:ilvl="0"><w:lvlText w:val="%s %%1."/></w:lvl>'
                 '</w:abstractNum><w:num w:numId="1"><w:abstractNumId w:val="0"/></w:num></w:numbering>' % (W, PERSON_SHORT))
    para = ('<w:p><w:pPr><w:numPr><w:ilvl w:val="0"/><w:numId w:val="1"/></w:numPr></w:pPr><w:r><w:t>erster punkt</w:t></w:r></w:p>')
    path = _docx_with(tmp_path, sdt + tbl + para, {"word/numbering.xml": numbering},
                      rels=[("rIdNum", "numbering", "numbering.xml")], drop=("word/numbering.xml",))
    ex = extract(path)
    det = lines_of(ex.detect_text)
    assert "list item: %s" % CUSTOMER_SHORT in det
    assert "content control: kunde" in det
    assert "alt text: tabelle %s" % ORG_SHORT in det
    assert "alt text: beschreibung %s" % PLACE in det
    assert "numbering: %s %%1." % PERSON_SHORT in det
    assert "- erster punkt" in lines_of(ex.text)
    for form in (CUSTOMER_SHORT, ORG_SHORT, PLACE, PERSON_SHORT):
        assert form not in ex.text


def _xlsx_with_cells(tmp_path: Path, cells_xml: str, parts=None, rels=(), drop=()) -> Path:
    """An openpyxl workbook whose first sheet holds the hand-written `cells_xml` in row 1, plus parts and
    workbook rels."""
    wb = openpyxl.Workbook()
    wb.active["A1"] = "platzhalter"
    path = tmp_path / "a.xlsx"
    wb.save(str(path))

    def edit(name, data):
        if name == "xl/worksheets/sheet1.xml":
            row = b'<sheetData><row r="1">' + cells_xml.encode("utf-8") + b"</row></sheetData>"
            return re.sub(rb"<sheetData>.*?</sheetData>|<sheetData/>", row, data, count=1, flags=re.DOTALL)
        if name == "xl/_rels/workbook.xml.rels":
            for rid, kind, target in rels:
                data = _add_rel(data, rid, REL_T + kind, target)
        return data
    return _rewrite(path, parts, edit, drop)


def test_xlsx_shared_strings_are_found_through_the_workbook_relationships(tmp_path):
    sst = '<sst xmlns="%s" count="2" uniqueCount="2"><si><t>kunde</t></si><si><t>%s</t></si></sst>' % (X_NS, CUSTOMER)
    path = _xlsx_with_cells(tmp_path, '<c r="A1" t="s"><v>0</v></c><c r="B1" t="s"><v>1</v></c>', {"xl/strings.xml": sst},
                            rels=[("rIdSst", "sharedStrings", "strings.xml")], drop=("xl/sharedStrings.xml",))
    ex = extract(path)
    assert "| kunde | %s |" % CUSTOMER in lines_of(ex.text)
    assert not ex.meta.get("incomplete")


def test_xlsx_string_cells_without_a_table_are_noted(tmp_path):
    path = _xlsx_with_cells(tmp_path, '<c r="A1" t="s"><v>0</v></c><c r="B1" t="s"><v>7</v></c>', drop=("xl/sharedStrings.xml",))
    ex = extract(path)
    assert ex.meta["incomplete"] is True
    assert "2 string cell(s) point past the shared string table, shown empty" in ex.notes


def test_xlsx_phonetic_runs_reach_detect_text(tmp_path):
    sst = ('<sst xmlns="%s" count="1" uniqueCount="1"><si><t>kunde</t><rPh sb="0" eb="5"><t>lesung %s</t></rPh>'
           '<phoneticPr fontId="1"/></si></sst>' % (X_NS, CUSTOMER_SHORT))
    path = _xlsx_with_cells(tmp_path, '<c r="A1" t="s"><v>0</v></c>', {"xl/sharedStrings.xml": sst},
                            rels=[("rIdSst", "sharedStrings", "sharedStrings.xml")], drop=("xl/sharedStrings.xml",))
    ex = extract(path)
    assert "| kunde |" in lines_of(ex.text)
    assert "phonetic: lesung %s" % CUSTOMER_SHORT in lines_of(ex.detect_text)
    assert CUSTOMER_SHORT not in ex.text


def test_xlsx_chart_caches_reach_detect_text(tmp_path):
    chart = _chart_xml("umsatz " + CUSTOMER_SHORT, "reihe " + ORG_SHORT, ["nord", PLACE], "label")
    ex = extract(_xlsx_with_cells(tmp_path, '<c r="A1" t="inlineStr"><is><t>zelle</t></is></c>', {"xl/charts/chart1.xml": chart}))
    det = lines_of(ex.detect_text)
    assert "chart: umsatz %s" % CUSTOMER_SHORT in det
    assert "chart: reihe %s" % ORG_SHORT in det
    assert "chart: nord | %s" % PLACE in det
    for form in (CUSTOMER_SHORT, ORG_SHORT, PLACE):
        assert form not in ex.text


def test_xlsx_formats_validation_rules_and_hyperlink_texts_reach_detect_text(tmp_path):
    from openpyxl.formatting.rule import FormulaRule
    from openpyxl.styles import PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation
    from openpyxl.worksheet.hyperlink import Hyperlink

    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = "sichtbar"
    ws["A2"] = 5
    ws["A2"].number_format = '"betrag %s "0' % CUSTOMER_SHORT
    rule = DataValidation(type="list", formula1='"%s,andere"' % ORG_SHORT, prompt="wahl für %s" % PLACE, error="falsch")
    ws.add_data_validation(rule)
    rule.add("B1")
    ws.conditional_formatting.add("A1:A9", FormulaRule(formula=['A1="%s"' % PERSON_SHORT], fill=PatternFill("solid", fgColor="FFFF00")))
    ws["A3"] = "portal"
    ws["A3"].hyperlink = Hyperlink(ref="A3", location="A1", tooltip="hinweis %s" % LAWFIRM_SHORT, display="anzeige")
    path = tmp_path / "a.xlsx"
    wb.save(str(path))
    ex = extract(path)
    det = lines_of(ex.detect_text)
    assert "number format: betrag %s" % CUSTOMER_SHORT in det
    assert any(line.startswith("validation: ") and ORG_SHORT in line and "wahl für %s" % PLACE in line for line in det)
    assert 'conditional format: =A1="%s"' % PERSON_SHORT in det
    assert any(line.startswith("hyperlink: ") and "hinweis %s" % LAWFIRM_SHORT in line and "anzeige" in line for line in det)
    for form in (CUSTOMER_SHORT, ORG_SHORT, PLACE, PERSON_SHORT, LAWFIRM_SHORT):
        assert form not in ex.text


def test_xlsx_linked_cells_text_boxes_pivot_query_and_slicer_reach_detect_text(tmp_path):
    ext = ('<externalLink xmlns="%s"><externalBook><sheetNames><sheetName val="quelle"/></sheetNames><sheetDataSet>'
           '<sheetData sheetId="0"><row r="1"><cell r="A1" t="str"><v>%s</v></cell></row></sheetData></sheetDataSet>'
           '</externalBook></externalLink>' % (X_NS, CUSTOMER))
    vml = '<xml xmlns:v="%s"><v:shape><v:textbox><div>kasten %s<br>zeile zwei</div></v:textbox></v:shape></xml>' % (V_NS, ORG_SHORT)
    pivot = ('<pivotCacheDefinition xmlns="%s" refreshedBy="%s"><cacheSource type="worksheet"><worksheetSource ref="A1:B3" '
             'sheet="quelle"/></cacheSource><cacheFields><cacheField name="feld"><sharedItems><s v="wert"/></sharedItems>'
             '</cacheField></cacheFields></pivotCacheDefinition>' % (X_NS, PERSON))
    qt = ('<queryTable xmlns="%s" name="abfrage %s"><queryTableRefresh><queryTableFields><queryTableField id="1" name="spalte"/>'
          '</queryTableFields></queryTableRefresh></queryTable>' % (X_NS, PLACE))
    slicer = ('<slicers xmlns="http://schemas.microsoft.com/office/spreadsheetml/2009/9/main"><slicer name="kunde" cache="c1" '
              'caption="filter %s"/></slicers>' % LAWFIRM_SHORT)
    parts = {"xl/externalLinks/externalLink1.xml": ext, "xl/drawings/vmlDrawing1.vml": vml,
             "xl/pivotCache/pivotCacheDefinition1.xml": pivot, "xl/queryTables/queryTable1.xml": qt, "xl/slicers/slicer1.xml": slicer}
    ex = extract(_xlsx_with_cells(tmp_path, '<c r="A1" t="inlineStr"><is><t>zelle</t></is></c>', parts))
    det = lines_of(ex.detect_text)
    assert "external link: quelle | %s" % CUSTOMER in det
    assert "text box: kasten %s zeile zwei" % ORG_SHORT in det
    assert "pivot: %s" % PERSON in det
    assert "query table: abfrage %s | spalte" % PLACE in det
    assert "slicer: kunde | filter %s" % LAWFIRM_SHORT in det
    for form in (CUSTOMER, ORG_SHORT, PERSON, PLACE, LAWFIRM_SHORT):
        assert form not in ex.text


def test_pptx_slides_come_from_the_presentation_relationships(tmp_path):
    path = make_pptx(tmp_path / "v.pptx")
    with zipfile.ZipFile(path) as zf:
        slide = zf.read("ppt/slides/slide1.xml")
        rels = zf.read("ppt/slides/_rels/slide1.xml.rels")
    orphan = slide.replace(CUSTOMER.encode("utf-8"), ("verwaist %s" % LAWFIRM_SHORT).encode("utf-8"))

    def edit(name, data):
        if name == "ppt/_rels/presentation.xml.rels":
            return data.replace(b"slides/slide1.xml", b"slides/deck-a.xml")
        if name == "[Content_Types].xml":
            return data.replace(b"/ppt/slides/slide1.xml", b"/ppt/slides/deck-a.xml")
        return data
    parts = {"ppt/slides/deck-a.xml": slide, "ppt/slides/_rels/deck-a.xml.rels": rels,
             "ppt/slides/slide9.xml": orphan, "ppt/slides/_rels/slide9.xml.rels": rels}
    ex = extract(_rewrite(path, parts, edit, drop=("ppt/slides/slide1.xml", "ppt/slides/_rels/slide1.xml.rels")))
    assert "angebot für %s" % CUSTOMER in ex.text
    assert ex.text.index(CUSTOMER) < ex.text.index("zeitplan")
    assert ex.meta["slide_count"] == 2
    assert LAWFIRM_SHORT not in ex.text
    assert "unlisted slide: angebot für verwaist %s" % LAWFIRM_SHORT in lines_of(ex.detect_text)


def test_pptx_chart_and_smartart_text_reach_text(tmp_path):
    path = make_pptx(tmp_path / "v.pptx")
    frame = ('<p:graphicFrame><p:nvGraphicFramePr><p:cNvPr id="31" name="diagramm"/><p:cNvGraphicFramePr/><p:nvPr/>'
             '</p:nvGraphicFramePr><p:xfrm><a:off x="0" y="0"/><a:ext cx="1" cy="1"/></p:xfrm><a:graphic><a:graphicData uri="%s">'
             '<c:chart xmlns:c="%s" r:id="rIdChart1"/></a:graphicData></a:graphic></p:graphicFrame>' % (C_NS, C_NS))
    # the base deck already carries the customer, the person and the place, so other forms are planted here
    chart = _chart_xml("absatz " + ORG_SHORT, "reihe", ["nord", "süd"], "label")
    # a diagram part nothing references (no relationship at all) still reaches the text
    data = ('<dgm:dataModel xmlns:dgm="%s" xmlns:a="%s"><dgm:ptLst><dgm:pt modelId="{2}"><dgm:t><a:bodyPr/><a:p><a:r>'
            '<a:t>knoten %s</a:t></a:r></a:p></dgm:t></dgm:pt></dgm:ptLst></dgm:dataModel>' % (DGM_NS, A_NS, LAWFIRM_SHORT))

    def edit(name, data_):
        if name == "ppt/slides/slide1.xml":
            return data_.replace(b"</p:spTree>", frame.encode("utf-8") + b"</p:spTree>", 1)
        if name == "ppt/slides/_rels/slide1.xml.rels":
            return _add_rel(data_, "rIdChart1", REL_T + "chart", "../charts/chart1.xml")
        return data_
    ex = extract(_rewrite(path, {"ppt/charts/chart1.xml": chart, "ppt/diagrams/data1.xml": data}, edit))
    lines = lines_of(ex.text)
    assert "chart: absatz %s" % ORG_SHORT in lines
    assert lines.index("chart: absatz %s" % ORG_SHORT) < lines.index("## Slide 2")
    assert "diagram: knoten %s" % LAWFIRM_SHORT in lines
    assert ex.text.count(ORG_SHORT) == 1


def test_pptx_fallback_shape_tags_and_sections_reach_detect_text(tmp_path):
    path = make_pptx(tmp_path / "v.pptx")
    shape = ('<p:sp><p:nvSpPr><p:cNvPr id="%d" name="box"/><p:cNvSpPr txBox="1"/><p:nvPr/></p:nvSpPr><p:spPr/><p:txBody>'
             '<a:bodyPr/><a:p><a:r><a:t>%s</a:t></a:r></a:p></p:txBody></p:sp>')
    alt = _mc("p14", shape % (20, "neue form"), shape % (21, "alte form %s" % ORG_SHORT))
    tag = '<p:tagLst xmlns:p="%s"><p:tag name="kunde" val="%s"/></p:tagLst>' % (P_NS, LAWFIRM_SHORT)
    ext = ('<p:extLst><p:ext><p14:sectionLst xmlns:p14="http://schemas.microsoft.com/office/powerpoint/2010/main">'
           '<p14:section name="abschnitt %s"><p14:sldIdLst/></p14:section></p14:sectionLst></p:ext></p:extLst>' % CUSTOMER_EN)

    def edit(name, data):
        if name == "ppt/slides/slide1.xml":
            return data.replace(b"</p:spTree>", alt.encode("utf-8") + b"</p:spTree>", 1)
        if name == "ppt/presentation.xml":
            return data.replace(b"</p:presentation>", ext.encode("utf-8") + b"</p:presentation>")
        return data
    ex = extract(_rewrite(path, {"ppt/tags/tag1.xml": tag}, edit))
    assert "neue form" in ex.text
    det = lines_of(ex.detect_text)
    assert "alternate content: alte form %s" % ORG_SHORT in det
    assert "tag kunde: %s" % LAWFIRM_SHORT in det
    assert "section: abschnitt %s" % CUSTOMER_EN in det
    for form in (ORG_SHORT, LAWFIRM_SHORT, CUSTOMER_EN):
        assert form not in ex.text


def _odf_with(path: Path, kind: str, body: str, decls: str = "") -> Path:
    mime = {"odt": "application/vnd.oasis.opendocument.text", "ods": "application/vnd.oasis.opendocument.spreadsheet"}[kind]
    inner = {"odt": "text", "ods": "spreadsheet"}[kind]
    content = ('<?xml version="1.0" encoding="UTF-8"?><office:document-content %s office:version="1.2"><office:body>'
               '<office:%s>%s%s</office:%s></office:body></office:document-content>' % (ODF_NS, inner, decls, body, inner))
    with zipfile.ZipFile(path, "w") as zf:
        zf.writestr(zipfile.ZipInfo("mimetype"), mime, compress_type=zipfile.ZIP_STORED)
        zf.writestr("content.xml", content)
    return path


def test_odt_hidden_and_conditional_text_values_reach_text(tmp_path):
    body = ('<text:p>preis für <text:hidden-text text:condition="ooow:FALSE" text:string-value="%s" text:is-hidden="false"/>'
            ' und <text:hidden-text text:condition="ooow:TRUE" text:string-value="geheim %s" text:is-hidden="true"/></text:p>'
            '<text:p>partner: <text:conditional-text text:condition="ooow:1 == 1" text:string-value-if-true="%s" '
            'text:string-value-if-false="niemand %s" text:current-value="true"/></text:p>'
            % (CUSTOMER, PLACE, ORG_SHORT, PERSON_SHORT))
    ex = extract(_odf_with(tmp_path / "a.odt", "odt", body))
    assert "preis für %s und" % CUSTOMER in ex.text
    assert "partner: %s" % ORG_SHORT in ex.text
    assert PLACE not in ex.text and PERSON_SHORT not in ex.text
    det = lines_of(ex.detect_text)
    assert "hidden text: geheim %s" % PLACE in det
    assert "conditional text: niemand %s" % PERSON_SHORT in det


def test_odt_field_declarations_and_form_controls(tmp_path):
    decls = ('<text:user-field-decls><text:user-field-decl office:value-type="string" office:string-value="%s" text:name="kunde"/>'
             '</text:user-field-decls><office:forms xmlns:form="urn:oasis:names:tc:opendocument:xmlns:form:1.0"><form:form '
             'form:name="f1"><form:text form:name="ort" form:current-value="sitz %s"/></form:form></office:forms>'
             % (CUSTOMER_SHORT, PLACE))
    body = '<text:p>kunde: <text:user-field-get text:name="kunde"/></text:p>'
    ex = extract(_odf_with(tmp_path / "a.odt", "odt", body, decls))
    assert "kunde: %s" % CUSTOMER_SHORT in ex.text
    det = lines_of(ex.detect_text)
    assert "field kunde: %s" % CUSTOMER_SHORT in det
    assert "form control: sitz %s" % PLACE in det
    assert PLACE not in ex.text


def test_ods_text_boxes_on_a_sheet_reach_text(tmp_path):
    body = ('<table:table table:name="plan"><table:shapes><draw:frame xmlns:draw="urn:oasis:names:tc:opendocument:xmlns:drawing:1.0">'
            '<draw:text-box><text:p>kasten %s</text:p></draw:text-box></draw:frame></table:shapes><table:table-column/>'
            '<table:table-row><table:table-cell office:value-type="string"><text:p>zelle</text:p></table:table-cell>'
            '</table:table-row></table:table>' % CUSTOMER_SHORT)
    ex = extract(_odf_with(tmp_path / "a.ods", "ods", body))
    lines = lines_of(ex.text)
    assert "## plan" in lines
    assert "kasten %s" % CUSTOMER_SHORT in lines
    assert "| zelle |" in lines
