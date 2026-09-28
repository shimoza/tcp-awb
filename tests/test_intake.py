"""The intake end to end, the candidate rules and the guarantees around the outputs.

Every name here comes from tests/fixtures.py. The inbox is built in a temporary Workbench with python-docx,
reportlab and the standard library email package.
"""
from __future__ import annotations

import base64
import io
import re
import stat
import zipfile
from collections import Counter
from email.message import EmailMessage
from pathlib import Path

import pytest

from awb import check, codes, intake, normalize, patterns, register
from awb.matcher import Matcher, Span
from tests import fixtures as fx

FULL, SHORT, ACRONYM, ENGLISH = fx.CUSTOMER_FORMS
PERSON, SURNAME = fx.PERSON_FORMS
ORG = fx.ORG_FORMS[0]
LAWFIRM = fx.LAWFIRM_FORMS[0]
PLACE = fx.PLACE_FORMS[0]
DOMAIN = fx.CUSTOMER_DOMAIN
MAIL = PERSON.lower().replace(" ", ".") + "@" + DOMAIN
IBAN = "DE89 3704 0044 0532 0130 00"
PHONE = "+49 40 123456-78"
PNG = b"\x89PNG\r\n\x1a\n" + b"\x00\x00\x00\x0dIHDR" + b"\x00" * 17

FILE_ID = re.compile(r"F-[A-Z2-7]{4}")
ROW = re.compile(r"^\| (?P<a>[^|]*?) \| (?P<b>(?:[^|\\]|\\.)*?) \| (?P<c>[^|]*?) \| (?P<d>[^|]*?) \|", re.M)


# --------------------------------------------------------------------------- building the inbox


def _docx_bytes(paragraphs: list[list[str]], header: str | None = None) -> bytes:
    import docx

    d = docx.Document()
    if header:
        d.sections[0].header.paragraphs[0].text = header
    for runs in paragraphs:
        p = d.add_paragraph()
        for r in runs:
            p.add_run(r)
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def _pdf(path: Path) -> None:
    from reportlab.pdfgen import canvas

    c = canvas.Canvas(str(path))
    c.setAuthor(LAWFIRM)
    c.setTitle("Angebot")
    c.drawString(72, 760, "Angebot fuer %s am Standort %s, Aktenzeichen %s." % (SHORT, PLACE, fx.FILE_NUMBER))
    c.drawString(72, 740, "Umsetzung mit Terraform und PostgreSQL hinter dem Load Balancer.")
    c.drawString(72, 720, "Zahlung auf %s, Server 203.0.113.7 im Netz." % IBAN)
    c.save()


def _eml(path: Path, attachment: bytes) -> None:
    msg = EmailMessage()
    msg["From"] = "%s <%s>" % (PERSON, MAIL)
    msg["To"] = "einkauf@example.org"
    msg["Date"] = "Tue, 22 Sep 2026 09:00:00 +0200"
    msg.set_content("guten morgen,\n\nanbei das Angebot zur Netzwerkbrücke.\n\nviele Grüße\n")
    msg.add_attachment(attachment, maintype="application",
                       subtype="vnd.openxmlformats-officedocument.wordprocessingml.document", filename="anlage.docx")
    subject = "=?utf-8?b?%s?=" % base64.b64encode(("Angebot " + ENGLISH).encode("utf-8")).decode("ascii")
    received = "Received: from %s ([192.0.2.10])\n\tby mx.example.net; Tue, 22 Sep 2026 09:00:01 +0200\n" % DOMAIN
    path.write_bytes(received.encode("ascii") + ("Subject: %s\n" % subject).encode("ascii") + msg.as_bytes())


def _zip(path: Path) -> None:
    with zipfile.ZipFile(path, "w") as zf:
        zf.writestr("%s_Netzplan.txt" % SHORT, "netzplan %s: 198.51.100.0/24 hinter dem Load Balancer.\n" % ACRONYM)


def build_inbox(home) -> dict[str, Path]:
    """The inbox of the end to end test. Keys name the role of each file, never its content."""
    inbox = home.inbox
    files = {}
    files["docx"] = inbox / "angebot-entwurf.docx"
    files["docx"].write_bytes(_docx_bytes(
        [
            ["Ansprechpartner ist ", "Tobi", "as Beispiel", "mann", " vom Standort ", PLACE, "."],
            ["Unterauftrag an die %s fuer den Netzaufbau." % fx.PLANTED_CANDIDATE],
            ["Kontakt per Mail an %s oder telefonisch unter %s." % (MAIL, PHONE)],
            ["Vergabe %s mit Terraform." % fx.TENDER_ID],
        ],
        header=FULL,
    ))
    files["pdf"] = inbox / "angebot.pdf"
    _pdf(files["pdf"])
    files["eml"] = inbox / "nachricht.eml"
    _eml(files["eml"], _docx_bytes([["Lieferant ist die %s aus dem Umland." % ORG]]))
    files["zip"] = inbox / "plaene.zip"
    _zip(files["zip"])
    files["png"] = inbox / "scan.png"
    files["png"].write_bytes(PNG)
    files["named"] = inbox / ("%s-Notizen.txt" % SHORT)
    files["named"].write_text(
        "notizen zum termin in %s mit %s, danach PostgreSQL und Terraform.\n" % (fx.CONTROL_UNREGISTERED, SURNAME),
        encoding="utf-8",
    )
    return files


# --------------------------------------------------------------------------- reading the reports


def _section(text: str, title: str) -> str:
    m = re.search(r"^## %s\n(.*?)(?=^## |\Z)" % re.escape(title), text, re.M | re.S)
    assert m, "section missing: %s" % title
    return m.group(1)


def _unescape(cell: str) -> str:
    return re.sub(r"\\(.)", r"\1", cell)


def files_table(private: str) -> dict[str, dict]:
    """file id -> {name, kind, state} from the private report."""
    out = {}
    for m in ROW.finditer(_section(private, "Files")):
        if FILE_ID.fullmatch(m.group("a")):
            out[m.group("a")] = {"name": _unescape(m.group("b")), "kind": m.group("c"), "state": m.group("d")}
    return out


def replacements(private: str) -> list[tuple[str, str, str]]:
    """(token, class, value) rows of the private report."""
    rows = []
    for m in ROW.finditer(_section(private, "Replacements")):
        if codes.is_code(m.group("a")):
            rows.append((m.group("a"), m.group("b"), _unescape(m.group("c"))))
    return rows


def candidates(private: str) -> list[str]:
    body = _section(private, "Candidates")
    return [_unescape(m.group(2)) for m in re.finditer(r"^\| (F-[A-Z2-7]{4}[.\d]*[^|]*?) \| (.*) \|$", body, re.M)]


# --------------------------------------------------------------------------- end to end


def test_blocked_without_force_writes_only_the_private_report(home):
    files = build_inbox(home)
    res = intake.run(list(files.values()), "new", home)

    assert res.blocked is True
    assert res.outputs == []
    assert res.candidates >= 1
    assert codes.is_code(res.customer) and codes.kind_of(res.customer) == "CUST"
    assert res.customer not in register.codes(register.load(home.register))
    assert not res.public_report.exists()
    assert not (home.outbox / res.customer).exists()
    assert not (home.originals / res.customer).exists()
    assert res.private_report.is_file()
    assert stat.S_IMODE(res.private_report.stat().st_mode) == 0o600
    assert list(home.private_reports.joinpath(res.customer).iterdir()) == [res.private_report]
    for f in files.values():
        assert f.is_file(), "a blocked run moves no original"
    private = res.private_report.read_text(encoding="utf-8")
    assert fx.PLANTED_CANDIDATE in candidates(private)
    assert "blocked: yes" in private
    assert set(res.states) == set(files_table(private))


def test_forced_run_end_to_end(home):
    files = build_inbox(home)
    first = intake.run(list(files.values()), "new", home)
    assert first.blocked

    res = intake.run(list(files.values()), first.customer, home, force=True)
    assert res.blocked is False
    assert res.customer == first.customer
    out_dir = home.outbox / res.customer
    private = res.private_report.read_text(encoding="utf-8")
    public = res.public_report.read_text(encoding="utf-8")
    table = files_table(private)
    by_name = {v["name"]: k for k, v in table.items()}

    # every output and the public report are free of every registered fixture form
    assert res.outputs
    for out in res.outputs:
        fx.assert_no_fixture_name(out.read_text(encoding="utf-8"), "an output")
    fx.assert_no_fixture_name(public, "the public report")

    # output file names are file ids only. The outbox holds nothing else
    for out in res.outputs:
        assert re.fullmatch(r"F-[A-Z2-7]{4}\.md", out.name)
        assert out.parent == out_dir
    assert sorted(p.name for p in out_dir.iterdir()) == sorted([o.name for o in res.outputs] + ["intake-report.md"])

    # at least 10 distinct planted forms were found and replaced, counted from the private report
    planted = {f.casefold() for f in fx.ALL_REGISTERED + [DOMAIN, fx.FILE_NUMBER, fx.TENDER_ID]}
    found = {" ".join(v.split()).casefold() for _, cls, v in replacements(private) if cls == "name"}
    assert len(found & planted) >= 10, "only %d distinct planted forms replaced" % len(found & planted)
    assert fx.PLANTED_CANDIDATE in candidates(private)

    # the image is unreadable in the public report
    png_id = by_name[files["png"].name]
    assert re.search(r"^\| %s \| image \| unreadable \|" % png_id, public, re.M)
    assert res.states[png_id] == "unreadable"
    assert "review the original" in public

    # originals left the inbox and sit in the vault under their file id
    assert list(home.inbox.iterdir()) == []
    for fid, row in table.items():
        suffix = Path(row["name"]).suffix.lower()
        moved = home.originals / res.customer / (fid + suffix)
        assert moved.is_file(), "original of %s missing in the vault" % fid
        assert stat.S_IMODE(moved.stat().st_mode) == 0o600

    # the private report is mode 600 and holds the original names, the public report none of them
    assert stat.S_IMODE(res.private_report.stat().st_mode) == 0o600
    for f in files.values():
        assert f.name in private
        assert f.name not in public
        assert f.stem not in public
    assert "%d candidates were reviewed." % res.candidates in public
    for c in candidates(private):
        assert c not in public

    # the same mail address got the same token in two files
    tokens = {t for t, cls, v in replacements(private) if cls == "mail" and v.casefold() == MAIL}
    assert len(tokens) == 1
    token = tokens.pop()
    for key in ("docx", "eml"):
        text = (out_dir / ("%s.md" % by_name[files[key].name])).read_text(encoding="utf-8")
        assert token in text
        assert MAIL not in text.casefold()

    # the name check finds nothing in any output or in the public report
    for out in res.outputs + [res.public_report]:
        assert check.check_file(out, home.register) == [], "check hit in %s" % out.name

    # what must stay, stays: the control name and the keep words
    named = (out_dir / ("%s.md" % by_name[files["named"].name])).read_text(encoding="utf-8")
    assert fx.CONTROL_UNREGISTERED in named
    assert "PostgreSQL" in named and "Terraform" in named
    pdf_out = (out_dir / ("%s.md" % by_name[files["pdf"].name])).read_text(encoding="utf-8")
    assert "Load Balancer" in pdf_out
    assert fx.CUSTOMER_CODE in pdf_out and fx.PLACE_CODE in pdf_out

    # the mail attachment and the zip member are parts of their file, sanitised in the same output
    eml_out = (out_dir / ("%s.md" % by_name[files["eml"].name])).read_text(encoding="utf-8")
    assert "## %s.1" % by_name[files["eml"].name] in eml_out
    assert fx.ORG_CODE in eml_out and fx.PERSON_CODE in eml_out
    zip_out = (out_dir / ("%s.md" % by_name[files["zip"].name])).read_text(encoding="utf-8")
    assert "## %s.1" % by_name[files["zip"].name] in zip_out
    assert "_Netzplan.txt" in zip_out and fx.CUSTOMER_CODE in zip_out


def test_states_and_ids(home):
    files = build_inbox(home)
    res = intake.run(list(files.values()), fx.CUSTOMER_CODE, home, force=True)
    assert res.customer == fx.CUSTOMER_CODE
    assert len(res.states) == len(files)
    assert all(FILE_ID.fullmatch(fid) for fid in res.states)
    assert res.states.values() and set(res.states.values()) <= {"ok", "unreadable"}
    assert sum(1 for s in res.states.values() if s == "unreadable") == 1


# --------------------------------------------------------------------------- customer codes


def test_customer_must_be_a_known_cust_code(home):
    f = home.inbox / "a.txt"
    f.write_text("nichts besonderes hier\n", encoding="utf-8")
    for bad in (fx.ORG_CODE, fx.PERSON_CODE, "CUST-AAAA", "zyx", "", "CUSTOMER"):
        with pytest.raises(intake.IntakeError) as err:
            intake.run([f], bad, home)
        fx.assert_no_fixture_name(str(err.value), "an error")
    register.add(home.register, "CUST-RTRD", "CUST", fx.PLANTED_PERSON)
    register.retire(home.register, "CUST-RTRD")
    with pytest.raises(intake.IntakeError, match="retired"):
        intake.run([f], "CUST-RTRD", home)
    assert f.is_file()


def test_missing_file_error_names_position_not_path(home):
    missing = home.inbox / ("%s.txt" % SHORT)
    with pytest.raises(intake.IntakeError) as err:
        intake.run([missing], fx.CUSTOMER_CODE, home)
    assert "file 1 of 1" in str(err.value)
    fx.assert_no_fixture_name(str(err.value), "an error")


def test_clean_run_with_known_customer_needs_no_force(home):
    f = home.inbox / "notiz.txt"
    f.write_text("termin mit %s am Standort %s, Server 198.51.100.4\n" % (SURNAME, PLACE), encoding="utf-8")
    res = intake.run([f], fx.CUSTOMER_CODE, home)
    assert not res.blocked and res.candidates == 0
    (out,) = res.outputs
    text = out.read_text(encoding="utf-8")
    assert fx.PERSON_CODE in text and fx.PLACE_CODE in text and re.search(r"\bIP-[A-Z2-7]{4}\b", text)
    fx.assert_no_fixture_name(text)
    assert "0 candidates were reviewed." in res.public_report.read_text(encoding="utf-8")


def test_same_value_same_token_and_different_values_differ(home):
    a = home.inbox / "a.txt"
    b = home.inbox / "b.txt"
    a.write_text("server 198.51.100.4 und 198.51.100.5\n", encoding="utf-8")
    b.write_text("wieder 198.51.100.4\n", encoding="utf-8")
    res = intake.run([a, b], fx.CUSTOMER_CODE, home)
    ta = re.findall(r"\bIP-[A-Z2-7]{4}\b", res.outputs[0].read_text(encoding="utf-8"))
    tb = re.findall(r"\bIP-[A-Z2-7]{4}\b", res.outputs[1].read_text(encoding="utf-8"))
    assert len(ta) == 2 and ta[0] != ta[1]
    assert tb == [ta[0]]


def test_final_check_failure_drops_output_and_keeps_original(home, monkeypatch):
    f = home.inbox / "leak.txt"
    f.write_text("vertrag mit %s\n" % FULL, encoding="utf-8")
    monkeypatch.setattr(intake._Engine, "sanitize", lambda self, text, where: (text, Counter()))
    # the self-test refuses this broken engine before anything is read
    # (test_the_self_test_refuses_a_broken_sanitiser); here the final check is under test
    monkeypatch.setattr(intake, "selftest", lambda: [])
    res = intake.run([f], fx.CUSTOMER_CODE, home)
    (fid,) = res.states
    assert res.states[fid] == "failed"
    assert res.outputs == []
    assert not (home.outbox / fx.CUSTOMER_CODE / ("%s.md" % fid)).exists()
    assert f.is_file(), "the original stays in the inbox when the final check fails"
    assert not (home.originals / fx.CUSTOMER_CODE).exists() or not any((home.originals / fx.CUSTOMER_CODE).iterdir())
    public = res.public_report.read_text(encoding="utf-8")
    assert "final check failed" in public and fid in public
    fx.assert_no_fixture_name(public)
    assert "final check found name" in res.private_report.read_text(encoding="utf-8")


def test_second_run_gets_fresh_ids_and_a_second_private_report(home):
    a = home.inbox / "a.txt"
    a.write_text("erste notiz zu Terraform\n", encoding="utf-8")
    r1 = intake.run([a], fx.CUSTOMER_CODE, home)
    b = home.inbox / "a.txt"
    b.write_text("zweite notiz zu Terraform\n", encoding="utf-8")
    r2 = intake.run([b], fx.CUSTOMER_CODE, home)
    assert set(r1.states).isdisjoint(r2.states)
    assert r1.private_report != r2.private_report
    assert len(list((home.originals / fx.CUSTOMER_CODE).iterdir())) == 2


def test_codes_and_file_ids_never_carry_a_register_form(home, monkeypatch):
    entries = register.load(home.register)
    ids = iter(["ZLG2", "Q2Q2"])
    monkeypatch.setattr(codes, "_id", lambda: next(ids))
    assert intake.safe_new_code(entries, "CUST") == "CUST-Q2Q2"


def test_base64_encoded_name_is_decoded_and_replaced(home):
    block = base64.b64encode(("Vertragspartner ist die %s in %s." % (FULL, PLACE)).encode("utf-8")).decode("ascii")
    f = home.inbox / "blob.txt"
    f.write_text("anhang:\n%s\nende\n" % block, encoding="utf-8")
    res = intake.run([f], fx.CUSTOMER_CODE, home)
    (out,) = res.outputs
    text = out.read_text(encoding="utf-8")
    assert block not in text
    assert fx.CUSTOMER_CODE in text
    assert check.check_file(out, home.register) == []


def test_a_mailbox_is_sanitised_message_by_message(home):
    first = ("From office@example.org 2026-09-22 10:00:00\r\nFrom: office@example.org\r\nTo: sender@example.org\r\n"
             "Subject: first\r\nContent-Type: text/plain; charset=utf-8\r\n\r\nthe first message.\r\n")
    second = ("From t.b@%s 2026-09-22 11:00:00\r\nFrom: %s <%s>\r\nTo: sender@example.org\r\n"
              "Subject: offer for %s\r\nContent-Type: text/plain; charset=utf-8\r\n\r\npay to %s\r\n"
              % (DOMAIN, PERSON, MAIL, FULL, IBAN))
    f = home.inbox / "box.mbox"
    f.write_bytes((first + "\r\n" + second).encode("utf-8"))
    res = intake.run([f], fx.CUSTOMER_CODE, home)
    assert not res.blocked
    (out,) = res.outputs
    (fid,) = res.states
    text = out.read_text(encoding="utf-8")
    assert "## %s.1" % fid in text and "## %s.2" % fid in text
    assert "the first message." in text
    assert fx.CUSTOMER_CODE in text and fx.PERSON_CODE in text and re.search(r"\bIBAN-[A-Z2-7]{4}\b", text)
    fx.assert_no_fixture_name(text)
    assert check.check_file(out, home.register) == []


# --------------------------------------------------------------------------- candidates


def _known(text: str, register_path: Path):
    m = Matcher(register.forms_for_matching(register.load(register_path)))
    n = normalize.normalize(text).text
    return n, m.find(n) + patterns.find_structured(n)


def test_candidate_company_form(register_path):
    text, known = _known("Wir beauftragen die %s mit dem Aufbau." % fx.PLANTED_CANDIDATE, register_path)
    assert intake.unknown_candidates(text, known) == [fx.PLANTED_CANDIDATE]


def test_registered_forms_are_not_candidates(register_path):
    text, known = _known("Vertrag der %s mit der %s, %s schreibt an %s." % (FULL, ORG, PERSON, MAIL), register_path)
    assert intake.unknown_candidates(text, known) == []
    # without the known spans the same text yields candidates: the rule, not the text, keeps it clean
    assert intake.unknown_candidates(text, []) != []


def test_candidate_capitalised_run_and_sentence_start(register_path):
    first, last = fx.PLANTED_PERSON.split()
    assert intake.unknown_candidates("ein termin mit %s morgen" % fx.PLANTED_PERSON, []) == [fx.PLANTED_PERSON]
    # a line start is not a sentence start: a name alone on a line is found
    assert intake.unknown_candidates("viele grüße\n%s\nvertrieb" % fx.PLANTED_PERSON, []) == [fx.PLANTED_PERSON]
    # after a full stop the pair is a candidate too (his decision of 2026-09-22: treat them as names; this
    # assertion said [] before)
    assert intake.unknown_candidates("es regnet. %s kommt" % fx.PLANTED_PERSON, []) == [fx.PLANTED_PERSON]
    # a title or article in front is dropped, the name stays
    assert intake.unknown_candidates("mit Herrn %s gesprochen" % fx.PLANTED_PERSON, []) == [fx.PLANTED_PERSON]


def test_candidate_stop_words_and_single_words(register_path):
    assert intake.unknown_candidates("wir nutzen Load Balancer und Microsoft Azure mit Terraform", []) == []
    assert intake.unknown_candidates("die Netzwerkbrücke in %s" % fx.CONTROL_UNREGISTERED, []) == []


def test_candidate_tool_names_in_metadata_are_not_names(register_path):
    meta = "application: Microsoft Macintosh Word\nproducer: LibreOffice Writer\n"
    assert intake.unknown_candidates(meta, []) == []
    assert intake.unknown_candidates(meta + "author: %s\n" % fx.PLANTED_PERSON, []) == [fx.PLANTED_PERSON]


def test_candidate_after_label(register_path):
    assert intake.unknown_candidates("Kunde: %s\n" % fx.CONTROL_UNREGISTERED, []) == [fx.CONTROL_UNREGISTERED]
    assert intake.unknown_candidates("Customer: %s" % fx.CUSTOMER_CODE, []) == []


def test_candidate_person_like_mail_local_part(register_path):
    local = fx.PLANTED_PERSON.lower().replace(" ", ".")
    text = "bitte an %s@example.org und info@example.org schreiben" % local
    assert intake.unknown_candidates(text, []) == [local]
    text, known = _known("bitte an %s schreiben" % MAIL, register_path)
    assert intake.unknown_candidates(text, known) == []


def test_candidate_kanzlei_prefix(register_path):
    last = fx.PLANTED_PERSON.split()[1]
    got = intake.unknown_candidates("vertreten durch die Kanzlei %s und Partner" % last, [])
    assert got == ["Kanzlei %s und Partner" % last]
    text, known = _known("vertreten durch die %s" % LAWFIRM, register_path)
    assert intake.unknown_candidates(text, known) == []


# --------------------------------------------------------------------------- review findings (2026-09-22)

import json as _json
import os as _os
import unicodedata as _ud

_LOOKALIKE = str.maketrans("аеорсухѕіјкΑΒΕΖΗΙΚΜΝΟΡΤΥΧο", "aeopcyxsijkABEZHIKMNOPTYXo")
_SKELETONS = sorted({re.sub(r"[\W_]+", "", f.casefold()) for f in fx.ALL_REGISTERED + [fx.FILE_NUMBER, fx.TENDER_ID]
                     + ["Qvoertz", "Qvortz"] if len(re.sub(r"[\W_]+", "", f)) >= 5})


def _skeleton(text: str) -> str:
    """The letters and digits of `text` as a reader would recognise them: compatibility forms and look-alike
    letters folded, accents and every separator dropped. Independent of the matcher on purpose."""
    t = _ud.normalize("NFKD", text).translate(_LOOKALIKE).replace("œ", "oe")
    t = "".join(c for c in t if not _ud.combining(c))
    return re.sub(r"[\W_]+", "", t.casefold())


def assert_nothing_readable(text: str, where: str) -> None:
    """No registered fixture form of five or more letters, however it is spread out or disguised, also with
    html tags taken out the way a Markdown viewer renders them."""
    for view in (text, re.sub(r"<[^<>]*>", "", text)):
        s = _skeleton(view)
        found = [k for k in _SKELETONS if k in s]
        assert not found, "%d fixture forms readable in %s" % (len(found), where)


_B64_FULL = base64.b64encode(FULL.encode()).decode()
_B64_ORG_LATIN = base64.b64encode(ORG.encode("latin-1")).decode()
_B64_FULL_UTF16 = base64.b64encode(FULL.encode("utf-16-le")).decode()
_HEX_FULL = FULL.encode().hex()
_JSON_ESCAPED = "\\u%04X" % ord(SHORT[0]) + SHORT[1:]
_DOUBLE_PCT = "%25" + "%25".join("%02X" % ord(c) for c in SHORT)
_RTF_ORG = ORG.replace("ö", "\\'f6").replace("ä", "\\'e4")


def _spread(word: str, sep: str) -> str:
    return sep.join(word)


def _docx_paragraphs(paragraphs: list[str]) -> bytes:
    return _docx_bytes([[p] for p in paragraphs])


# every way the red team got a registered form into an output; each file goes through a run without --force.
# The fourth field is the encoded payload that must not survive in the output.
DISGUISED = [
    ("spaced letters", "note.txt", "kunde ist %s, ansprechpartner %s." % (_spread(SHORT, " "), _spread(SURNAME, " ")), None),
    ("dotted letters", "note.txt", "kunde %s." % _spread(SHORT, "."), None),
    ("invisible separator", "note.txt", "kunde %s ok" % _spread(SHORT, "\u2063"), None),
    ("grapheme joiner", "note.txt", "kunde %s ok" % _spread(SHORT, "\u034f"), None),
    ("tag space", "note.txt", "kunde %s ok" % _spread(SHORT, "\U000e0020"), None),
    ("full width", "note.txt", "kunde %s ok" % "".join(chr(ord(c) + 0xFEE0) for c in SHORT), None),
    ("cyrillic letters", "note.txt", "kunde %s, %s" % (SHORT[:-1] + "\u043e", SURNAME.replace("a", "\u0430")), None),
    ("oe ligature", "note.txt", "lieferant %s" % fx.ORG_FORMS[1].replace("ö", "\u0153"), None),
    ("non-breaking hyphens", "note.txt", "vergabe %s" % fx.TENDER_ID.replace("-", "\u2011"), None),
    ("glued numbers", "note.txt", "akte%s und ref%s" % (fx.FILE_NUMBER, fx.TENDER_ID), None),
    ("line break inside", "note.txt", "kunde %s\n%s fertig" % (SHORT[:3], SHORT[3:]), None),
    ("table border inside", "note.md", "| %s|%s |\n" % (SHORT[:3], SHORT[3:]), None),
    ("markdown emphasis", "note.md", "%s**%s** und %s*%s*" % (SHORT[:3], SHORT[3:], SURNAME[:8], SURNAME[8:]), None),
    ("empty tags", "note.md", "%s<b></b>%s" % (SHORT[:3], SHORT[3:]), None),
    ("capitals hyphenated", "note.txt", "%s-\n%s" % (SHORT[:3].upper(), SHORT[3:].upper()), None),
    ("html comment inside", "p.html", "<html><body><p>%s<!-- x -->%s</p></body></html>" % (SHORT[:3], SHORT[3:]), None),
    ("short base64", "note.txt", "tag: %s" % _B64_FULL, _B64_FULL),
    ("base64 of latin-1", "note.txt", "x %s" % _B64_ORG_LATIN, _B64_ORG_LATIN),
    ("base64 of utf-16", "note.txt", "y %s" % _B64_FULL_UTF16, _B64_FULL_UTF16),
    ("hex of utf-8", "note.txt", "h %s" % _HEX_FULL, _HEX_FULL),
    ("encoded word", "note.txt", "s =?utf-8?b?%s?=" % _B64_FULL, _B64_FULL),
    ("latin-1 url encoding", "note.txt", "?q=" + ORG.replace("ö", "%F6").replace("ä", "%E4").replace(" ", "%20"), None),
    ("double url encoding", "note.txt", "id " + _DOUBLE_PCT, _DOUBLE_PCT),
    ("json escapes", "data.json", _json.dumps({"kunde": SHORT, "partner": ORG}).replace(SHORT, _JSON_ESCAPED), _JSON_ESCAPED),
    ("quoted-printable", "note.txt", ORG.replace("ö", "=C3=B6").replace("ä", "=C3=A4"), None),
    ("rtf escapes", "brief.rtf", "{\\rtf1 %s}" % _RTF_ORG, _RTF_ORG),
    ("docx body", "a.docx", _docx_paragraphs(["kunde %s" % _spread(SHORT, " ")]), None),
]


@pytest.mark.parametrize("label,name,content,payload", DISGUISED, ids=[d[0] for d in DISGUISED])
def test_a_disguised_form_never_reaches_an_output(home, label, name, content, payload):
    f = home.inbox / name
    f.write_bytes(content if isinstance(content, bytes) else content.encode("utf-8"))
    res = intake.run([f], fx.CUSTOMER_CODE, home)
    assert not res.blocked, label
    (out,) = res.outputs
    text = out.read_text(encoding="utf-8")
    assert_nothing_readable(text, label)
    if payload:
        assert payload not in text, "the encoded form stayed in the output"
    assert_nothing_readable(res.public_report.read_text(encoding="utf-8"), "the public report")


def test_a_payload_that_hides_a_form_behind_control_bytes_drops_the_output(home):
    blob = base64.b64encode(b"\x01\x02\x03\x04\x05" + FULL.encode()).decode()
    f = home.inbox / "note.txt"
    f.write_text("z %s\n" % blob, encoding="utf-8")
    res = intake.run([f], fx.CUSTOMER_CODE, home)
    (fid,) = res.states
    assert res.states[fid] == "failed" and res.outputs == []
    assert f.is_file()


def test_resolve_merges_a_chain_of_overlaps():
    # a name overlapped on both edges by two structured hits becomes one span
    names = [Span(10, 30, "name", fx.CUSTOMER_CODE)]
    structured = [Span(0, 15, "url"), Span(25, 40, "url")]
    got = intake._resolve(names, structured)
    assert [(s.start, s.end) for s in got] == [(0, 40)]
    assert got[0].cls == "url"
    # a structured hit inside a name is dropped, one that holds a name wins
    assert [(s.start, s.end, s.cls) for s in intake._resolve(names, [Span(12, 20, "url")])] == [(10, 30, "name")]
    assert [(s.start, s.end, s.cls) for s in intake._resolve(names, [Span(5, 35, "mail")])] == [(5, 35, "mail")]


def test_a_form_between_two_structured_hits_is_replaced_whole():
    form = "%s %s" % (PERSON, SHORT)
    entries = [register.Entry(fx.CUSTOMER_CODE, "CUST", form, fx.TODAY, "active")]
    engine = intake._Engine(entries)
    first, last = PERSON.split()
    text = "https://portal.test/%s %s %s.example" % (first.lower(), last, SHORT.lower())
    out, _ = engine.sanitize(text, "F-TEST output")
    assert last not in out and first.lower() not in out
    assert engine.final_hits(out) == []


def test_a_symbolic_link_is_refused(home):
    target = home.shared / "elsewhere.txt"
    target.write_text("inhalt\n", encoding="utf-8")
    _os.chmod(target, 0o644)
    link = home.inbox / "link.txt"
    link.symlink_to(target)
    with pytest.raises(intake.IntakeError) as err:
        intake.run([link], fx.CUSTOMER_CODE, home)
    assert "symbolic link" in str(err.value)
    assert target.is_file() and stat.S_IMODE(target.stat().st_mode) == 0o644
    assert link.is_symlink()


def test_a_failed_run_leaves_no_output_behind(home, monkeypatch):
    a = home.inbox / "a.txt"
    b = home.inbox / "b.txt"
    a.write_text("erste notiz zu Terraform\n", encoding="utf-8")
    b.write_text("zweite notiz zu Terraform\n", encoding="utf-8")
    real = intake._report.write_new
    calls = []

    def flaky(path, text, mode, dir_mode):
        calls.append(path)
        if len(calls) == 2:
            raise OSError("disk full")
        return real(path, text, mode, dir_mode)

    monkeypatch.setattr(intake._report, "write_new", flaky)
    with pytest.raises(OSError):
        intake.run([a, b], fx.CUSTOMER_CODE, home)
    assert len(calls) == 2
    assert not list((home.outbox / fx.CUSTOMER_CODE).glob("F-*.md"))
    assert a.is_file() and b.is_file()


def test_the_public_report_is_checked_again_after_notes_are_withheld(home, monkeypatch):
    f = home.inbox / "a.txt"
    f.write_text("notiz zu Terraform\n", encoding="utf-8")
    real = intake._report.render_public
    monkeypatch.setattr(intake._report, "render_public", lambda run: real(run) + "\n" + FULL + "\n")
    res = intake.run([f], fx.CUSTOMER_CODE, home)
    public = res.public_report.read_text(encoding="utf-8")
    fx.assert_no_fixture_name(public, "the public report")
    assert "withheld" in public and "1 files" not in public and "- files: 1" in public


def test_a_failed_reader_does_not_claim_the_original_stayed(home):
    f = home.inbox / "kaputt.zip"
    f.write_bytes(b"PK\x03\x04" + b"\x00" * 40)
    res = intake.run([f], fx.CUSTOMER_CODE, home)
    (fid,) = res.states
    assert res.states[fid] == "failed"
    assert not f.exists(), "the original of a failed reader moves to the vault"
    public = res.public_report.read_text(encoding="utf-8")
    assert "Originals still in the inbox" not in public
    assert "Originals moved to the vault although no output was written: %s." % fid in public


def test_positions_are_explained_and_temporary_files_stay_in_the_vault(home):
    f = home.inbox / "a.zip"
    with zipfile.ZipFile(f, "w") as zf:
        zf.writestr("teil.txt", "kunde %s\n" % SHORT)
    res = intake.run([f], fx.CUSTOMER_CODE, home)
    private = res.private_report.read_text(encoding="utf-8")
    assert "They are not offsets into the original file." in private
    assert (home.vault / "tmp").is_dir() and list((home.vault / "tmp").iterdir()) == []
    assert stat.S_IMODE((home.vault / "tmp").stat().st_mode) == 0o700


def test_detect_positions_point_into_the_text_as_given(home):
    entries = register.load(home.register)
    engine = intake._Engine(entries)
    text = "x&amp;y %s" % SHORT
    engine.detect("F-TEST detect", text)
    (hit,) = [h for h in engine.hits.values() if h.cls == "name"]
    assert hit.where["F-TEST detect"] == [text.index(SHORT)]


# --------------------------------------------------------------------------- candidate rules added after the review

FIRST, LAST = fx.PLANTED_PERSON.split()
BRAND = fx.PLANTED_CANDIDATE.split()[0]

NEW_CANDIDATES = [
    ("company in capitals", "lieferant %s liefert" % fx.PLANTED_CANDIDATE.upper(), fx.PLANTED_CANDIDATE.upper()),
    ("person in capitals", "kontakt %s schreibt" % fx.PLANTED_PERSON.upper(), fx.PLANTED_PERSON.upper()),
    ("lower-case brand", "partner %s GmbH" % BRAND.lower(), "%s GmbH" % BRAND.lower()),
    ("mbH", "partner %s mbH" % BRAND, "%s mbH" % BRAND),
    ("S.p.A.", "partner %s S.p.A." % BRAND, "%s S.p.A." % BRAND),
    ("s.r.o.", "partner %s s.r.o." % BRAND, "%s s.r.o." % BRAND),
    ("eG", "partner %s eG" % BRAND, "%s eG" % BRAND),
    ("e.K.", "partner %s e.K." % BRAND, "%s e.K." % BRAND),
    ("PLC", "partner %s PLC" % BRAND, "%s PLC" % BRAND),
    ("salutation", "sehr geehrter " + "Herr " + LAST + ",", LAST),
    ("title chain", "mit " + "Frau" + " Dr. " + LAST, LAST),
    ("initial", "kontakt %s. %s schreibt" % (FIRST[0], LAST), "%s. %s" % (FIRST[0], LAST)),
    ("accent", "kontakt %s %s kam" % (FIRST, LAST.replace("e", "è")), "%s %s" % (FIRST, LAST.replace("e", "è"))),
    ("apostrophe", "kontakt %s O'%s kam" % (FIRST, LAST), "%s O'%s" % (FIRST, LAST)),
    ("spaced company", "  ".join(_spread(w, " ") for w in fx.PLANTED_CANDIDATE.split()), None),
    ("spaced person, pdf gaps", "   ".join(fx.PLANTED_PERSON.replace(" ", "")), None),
]


@pytest.mark.parametrize("label,text,want", NEW_CANDIDATES, ids=[c[0] for c in NEW_CANDIDATES])
def test_new_candidate_rules(label, text, want):
    got = intake.unknown_candidates(text, [])
    assert len(got) == 1, label
    if want is not None:
        assert got == [want], label


ORDINARY = [
    "die neue GmbH", "der gegründeten GmbH", "zur GmbH", "die umwandlung als AG", "sie ist AG", "z. B. Speicher und u. a. Netze", "I N H A L T",
    "sehr geehrte Damen und Herren", "Herr der Lage", "wir nutzen AWS VPN mit Terraform",
    "DIESE NACHRICHT IST VERTRAULICH", "T e r r a f o r m und mehr",
]


@pytest.mark.parametrize("text", ORDINARY)
def test_new_candidate_rules_leave_ordinary_text(text):
    assert intake.unknown_candidates(text, []) == []


def test_new_candidate_rules_block_an_intake(home):
    f = home.inbox / "note.txt"
    f.write_text("sehr geehrter " + "Herr " + LAST + ",\nwir liefern.\n", encoding="utf-8")
    res = intake.run([f], fx.CUSTOMER_CODE, home)
    assert res.blocked and res.outputs == []


# --------------------------------------------------------------------------- candidate rules of release 2
# his decision of 2026-09-22: treat them as names (a pair at a sentence start, First von Last, Last, First)


def _j(*words: str) -> str:
    return " ".join(words)


SENTENCE_START = [
    ("after a full stop", "es regnet. %s kommt" % fx.PLANTED_PERSON, [fx.PLANTED_PERSON]),
    ("after a question mark", "wer kommt? %s kommt" % fx.PLANTED_PERSON, [fx.PLANTED_PERSON]),
    ("after an exclamation mark", "gut! %s schreibt" % fx.PLANTED_PERSON, [fx.PLANTED_PERSON]),
    ("an opener in front is dropped", "es regnet. %s kommt" % _j("Heute", FIRST, LAST), [fx.PLANTED_PERSON]),
    ("an article in front is dropped", "es regnet. %s kommt" % _j("Die", FIRST, LAST), [fx.PLANTED_PERSON]),
    ("three words", "ende. %s schreibt" % _j(FIRST, LAST, "Beratung"), [_j(FIRST, LAST, "Beratung")]),
]


@pytest.mark.parametrize("label,text,want", SENTENCE_START, ids=[c[0] for c in SENTENCE_START])
def test_a_pair_at_a_sentence_start_is_a_candidate(label, text, want):
    assert intake.unknown_candidates(text, []) == want, label


OPENING_PAIRS = [
    "es regnet. %s steigt." % _j("The", "Load"),
    "es regnet. %s kommt er." % _j("Heute", "Montag"),
    "danke. %s" % _j("Viele", "Grüße"),
    "ende. %s nutzen." % _j("Bitte", "Terraform"),
    "gut. %s ist fertig." % _j("Das", "Angebot"),
    _j("Viele", "Grüße") + "\n" + FIRST + "\n",
]


@pytest.mark.parametrize("text", OPENING_PAIRS)
def test_a_pair_of_openers_and_stop_words_at_a_sentence_start_is_not(text):
    assert intake.unknown_candidates(text, []) == []


def test_the_openers_come_from_the_rules_file():
    lines = (patterns.RULES_DIR / "sentence-openers.txt").read_text(encoding="utf-8").splitlines()
    words = [w.strip() for w in lines if w.strip() and not w.startswith("#")]
    assert "Heute" in words and "Viele" in words
    for w in words:
        # an opener in front of a person-like pair is dropped, the pair stays
        assert intake.unknown_candidates("ende. %s kommt" % _j(w, FIRST, LAST), []) == [fx.PLANTED_PERSON], w


def test_the_strong_rules_catch_a_pair_at_a_sentence_start_only():
    strong = intake.STRONG_RULES
    assert intake.unknown_candidates("plan it. %s signs off." % fx.PLANTED_PERSON, [], rules=strong) \
        == [fx.PLANTED_PERSON]
    goal = _j("Landing", "Zone") + " for the first workload"
    assert intake.unknown_candidates(goal, [], rules=strong) == []
    assert intake.unknown_candidates(goal, []) == [_j("Landing", "Zone")]
    assert intake.unknown_candidates("mit %s gesprochen" % fx.PLANTED_PERSON, [], rules=strong) == []


PARTICLES = [
    ("von", "termin mit %s von %s morgen" % (FIRST, LAST), "%s von %s" % (FIRST, LAST)),
    ("van", "termin mit %s van %s morgen" % (FIRST, LAST), "%s van %s" % (FIRST, LAST)),
    ("de", "termin mit %s de %s morgen" % (FIRST, LAST), "%s de %s" % (FIRST, LAST)),
    ("zu", "termin mit %s zu %s morgen" % (FIRST, LAST), "%s zu %s" % (FIRST, LAST)),
    ("vom", "termin mit %s vom %s morgen" % (FIRST, LAST), "%s vom %s" % (FIRST, LAST)),
    ("zur", "termin mit %s zur %s morgen" % (FIRST, LAST), "%s zur %s" % (FIRST, LAST)),
    ("von der", "termin mit %s von der %s morgen" % (FIRST, LAST), "%s von der %s" % (FIRST, LAST)),
    ("at a sentence start", "gut. %s von %s kommt" % (FIRST, LAST), "%s von %s" % (FIRST, LAST)),
    ("after a title", "mit Herrn %s von %s" % (FIRST, LAST), "%s von %s" % (FIRST, LAST)),
]


@pytest.mark.parametrize("label,text,want", PARTICLES, ids=[c[0] for c in PARTICLES])
def test_first_von_last_is_a_candidate(label, text, want):
    assert intake.unknown_candidates(text, []) == [want], label
    assert intake.unknown_candidates(text, [], rules=intake.STRONG_RULES) == [want], label


PARTICLE_ORDINARY = [
    "die Migration von Servern",
    "die Anforderungen zur Umsetzung",
    "die Sicherheit von Daten",
    "die Architektur von Netzen",
    "das Management von Clustern",
    "die Technik zur Anbindung",
    "die Informationen zur Lage",
    "das Angebot zur Netzwerkbrücke",
    "die Fragen zu Terraform",
    _j("Von", "Montag") + " bis Freitag",
    "der Server von Kunden",
]


@pytest.mark.parametrize("text", PARTICLE_ORDINARY)
def test_a_german_noun_before_von_is_not(text):
    assert intake.unknown_candidates(text, []) == []


INVERTED = [
    ("plain", "anwesend: %s, %s und andere" % (LAST, FIRST), "%s, %s" % (LAST, FIRST)),
    ("mail header", "From: %s, %s <x@example.org>" % (LAST, FIRST), "%s, %s" % (LAST, FIRST)),
    ("surname in capitals", "an %s, %s gesendet" % (LAST.upper(), FIRST), "%s, %s" % (LAST.upper(), FIRST)),
    ("line of its own", "\n%s, %s\n" % (LAST, FIRST), "%s, %s" % (LAST, FIRST)),
]


@pytest.mark.parametrize("label,text,want", INVERTED, ids=[c[0] for c in INVERTED])
def test_last_comma_first_is_a_candidate(label, text, want):
    assert intake.unknown_candidates(text, []) == [want], label
    assert intake.unknown_candidates(text, [], rules=intake.STRONG_RULES) == [want], label


INVERTED_ORDINARY = [
    "wir nutzen Linux, Windows, Terraform",
    "wir nutzen Linux, Windows und Terraform",
    "wir nutzen Ansible, Packer und Vault",
    "wir nutzen Ansible, Packer, Vault und Consul",
    "am Montag, Dienstag",
    "wir nutzen Linux, Windows heute",
    "danke, %s" % FIRST,
    "%s, 22. September" % fx.CONTROL_UNREGISTERED,
    "%s, Qx" % LAST,
]


@pytest.mark.parametrize("text", INVERTED_ORDINARY)
def test_a_comma_list_or_ordinary_words_around_a_comma_are_not(text):
    assert intake.unknown_candidates(text, []) == []


def test_a_registered_word_around_a_comma_is_known(register_path):
    text, known = _known("anwesend: %s, %s" % (SURNAME, FIRST), register_path)
    assert intake.unknown_candidates(text, known) == []
    assert intake.unknown_candidates(text, []) != []


# --------------------------------------------------------------------------- the keep list

KEPT = _j("Technisches", "Konzept")


def test_unknown_candidates_drops_kept_phrases():
    text = "anbei das %s und ein termin mit %s" % (KEPT, fx.PLANTED_PERSON)
    assert intake.unknown_candidates(text, []) == [KEPT, fx.PLANTED_PERSON]
    # case and whitespace do not matter; a phrase that is only part of a candidate keeps nothing
    assert intake.unknown_candidates(text, [], keep=[KEPT.upper().replace(" ", "  ")]) == [fx.PLANTED_PERSON]
    assert intake.unknown_candidates(text, [], keep=[FIRST]) == [KEPT, fx.PLANTED_PERSON]


def test_keep_phrase_writes_one_phrase_per_line_mode_600_in_the_vault(home):
    assert intake.load_keep(home) == []
    assert intake.keep_phrase(home, KEPT) == (True, 1)
    assert intake.keep_phrase(home, "  %s  " % KEPT.lower()) == (False, 1)
    assert intake.keep_phrase(home, "%s\t%s" % tuple(fx.PLANTED_PERSON.split())) == (True, 2)
    assert home.keep_list.parent == home.vault
    assert stat.S_IMODE(home.keep_list.stat().st_mode) == 0o600
    assert home.keep_list.read_text(encoding="utf-8") == "%s\n%s\n" % (KEPT, fx.PLANTED_PERSON)
    assert intake.load_keep(home) == [KEPT, fx.PLANTED_PERSON]
    # nothing else was left in the vault folder
    assert not [f for f in home.vault.iterdir() if f.name.startswith(".keep-")]


def test_keep_phrase_refuses_only_an_empty_phrase_and_never_echoes(home):
    for phrase in (FULL, fx.CUSTOMER_CODE, "Kunde: " + fx.PLANTED_CANDIDATE, "a | b # c (d)"):
        added, _ = intake.keep_phrase(home, phrase)
        assert added
    for empty in ("", "   ", "\t\n"):
        with pytest.raises(intake.IntakeError) as err:
            intake.keep_phrase(home, empty)
        assert "empty" in str(err.value)
    assert len(intake.load_keep(home)) == 4


def test_a_kept_phrase_is_no_candidate_in_the_next_intake(home):
    f = home.inbox / "konzept.txt"
    f.write_text("anbei das %s, bitte pruefen.\n" % KEPT, encoding="utf-8")
    first = intake.run([f], fx.CUSTOMER_CODE, home)
    assert first.blocked and first.candidates == 1
    assert KEPT in candidates(first.private_report.read_text(encoding="utf-8"))
    intake.keep_phrase(home, KEPT)
    second = intake.run([f], fx.CUSTOMER_CODE, home)
    assert not second.blocked and second.candidates == 0
    (out,) = second.outputs
    assert KEPT in out.read_text(encoding="utf-8"), "a kept phrase is not a name: it stays in the output"


def _fake_daemon(home, monkeypatch, fail_seal: bool = False):
    """The vault daemon's open_file and seal_file for the keep list, without gpg: sealed content is held in
    memory and the .gpg file on disk is a marker."""
    from awb import vault

    sealed: dict[str, bytes] = {}
    calls: list[str] = []

    def admin_call(op, sock=None, **fields):
        calls.append(op)
        path = Path(fields["path"])
        assert path.parent == home.vault
        if op == "seal_file":
            if fail_seal:
                raise vault.VaultLocked("locked")
            sealed[path.name + ".gpg"] = path.read_bytes()
            path.with_name(path.name + ".gpg").write_bytes(b"sealed")
            path.unlink()
            return {"ok": True}
        if op == "open_file":
            return {"ok": True, "data": base64.b64encode(sealed[path.name]).decode("ascii")}
        raise AssertionError("unexpected op")

    monkeypatch.setattr(vault, "admin_call", admin_call)
    return sealed, calls


def test_keep_list_in_an_encrypted_vault_goes_through_the_daemon(home, monkeypatch):
    home.register_encrypted.write_bytes(b"sealed")
    sealed, calls = _fake_daemon(home, monkeypatch)
    assert intake.keep_phrase(home, KEPT) == (True, 1)
    assert not home.keep_list.exists(), "no plaintext keep list stays in an encrypted vault"
    assert intake.keep_phrase(home, fx.PLANTED_PERSON) == (True, 2)
    assert not home.keep_list.exists()
    assert sealed["keep.tsv.gpg"].decode("utf-8") == "%s\n%s\n" % (KEPT, fx.PLANTED_PERSON)
    assert intake.load_keep(home) == [KEPT, fx.PLANTED_PERSON]
    assert calls.count("seal_file") == 2 and "open_file" in calls


def test_keep_list_seal_failure_leaves_no_plaintext_and_no_value(home, monkeypatch):
    home.register_encrypted.write_bytes(b"sealed")
    _fake_daemon(home, monkeypatch, fail_seal=True)
    with pytest.raises(intake.IntakeError) as err:
        intake.keep_phrase(home, KEPT)
    assert "locked" in str(err.value) and KEPT.split()[1] not in str(err.value)
    assert not home.keep_list.exists()


# --------------------------------------------------------------------------- an encrypted vault: real gpg and a daemon

VAULT_PASS = "fixture passphrase of the intake 9"


@pytest.fixture
def sealed_vault(home, monkeypatch):
    """The fixture vault encrypted with gpg and a vault daemon in threads of this process, unlocked. The sockets
    lie in a short folder under /tmp that is removed afterwards; the daemon is stopped at the end."""
    import shutil
    import tempfile
    import threading

    from awb import config, vault

    monkeypatch.setattr(vault, "S2K_COUNT", 65536)
    short = Path(tempfile.mkdtemp(prefix="awb", dir="/tmp"))
    monkeypatch.setenv("AWB_CHECK_SOCKET", str(short / "run" / "check.sock"))
    monkeypatch.setenv("AWB_ADMIN_SOCKET", str(short / "admin.sock"))
    p = config.paths()
    daemon = vault.Daemon(p)
    try:
        daemon.start()
        vault.encrypt_vault(p, VAULT_PASS)
        vault.admin_call("unlock", p.admin_sock, passphrase=VAULT_PASS)
        yield p
    finally:
        daemon.stop()
        shutil.rmtree(short, ignore_errors=True)
    assert not [t for t in threading.enumerate() if t.name.startswith("awb-vault")]


def _opened(p, path: Path) -> bytes:
    from awb import vault

    return base64.b64decode(vault.admin_call("open_file", p.admin_sock, path=str(path))["data"])


def _names_under(folder: Path) -> list[str]:
    return sorted(f.name for f in folder.rglob("*") if f.is_file())


def test_intake_in_an_encrypted_vault_seals_originals_and_the_private_report(sealed_vault):
    p = sealed_vault
    assert not p.register.exists() and intake.vault_encrypted(p)
    data = _docx_bytes([["das angebot fuer ", FULL, " geht an ", PERSON, " in ", PLACE, "."]])
    f = p.inbox / "angebot.docx"
    f.write_bytes(data)
    res = intake.run([f], fx.CUSTOMER_CODE, p)
    assert not res.blocked and res.unsealed == 0 and not f.exists()
    (out,) = res.outputs
    text = out.read_text(encoding="utf-8")
    fx.assert_no_fixture_name(text, "output")
    assert fx.CUSTOMER_CODE in text
    # the original: sealed under its file id, no plaintext copy anywhere in the vault
    (sealed,) = (p.originals / fx.CUSTOMER_CODE).iterdir()
    assert sealed.name == "%s.docx.gpg" % out.stem
    assert not sealed.read_bytes().startswith(b"PK")
    assert _opened(p, sealed) == data
    # the private report: sealed and only the daemon opens it
    assert res.private_report.name.endswith(".md.gpg") and res.private_report.is_file()
    assert _names_under(p.private_reports) == [res.private_report.name]
    assert out.stem in _opened(p, res.private_report).decode("utf-8")
    assert all(n.endswith(".gpg") for n in _names_under(p.originals) + _names_under(p.private_reports))
    # the public side names the report path by codes and dates only
    fx.assert_no_fixture_name(res.public_report.read_text(encoding="utf-8"), "public report")
    fx.assert_no_fixture_name(str(res.private_report), "private report path")


def test_a_blocked_intake_in_an_encrypted_vault_seals_its_private_report(sealed_vault):
    p = sealed_vault
    f = p.inbox / "notiz.txt"
    f.write_text("das angebot geht an die %s morgen.\n" % fx.PLANTED_CANDIDATE, encoding="utf-8")
    res = intake.run([f], fx.CUSTOMER_CODE, p)
    assert res.blocked and res.candidates >= 1 and res.unsealed == 0
    assert f.exists(), "a blocked run leaves the file in the inbox"
    assert res.private_report.name.endswith(".md.gpg")
    assert _names_under(p.private_reports) == [res.private_report.name]
    assert fx.PLANTED_CANDIDATE in _opened(p, res.private_report).decode("utf-8")
    assert not (p.outbox / fx.CUSTOMER_CODE).exists() or not list((p.outbox / fx.CUSTOMER_CODE).iterdir())


def test_a_second_report_in_the_same_second_never_replaces_a_sealed_one(home):
    first = home.private_reports / fx.CUSTOMER_CODE / "2026-09-22-120000.md"
    first.parent.mkdir(parents=True, exist_ok=True)
    first.with_name(first.name + ".gpg").write_bytes(b"an older sealed report")
    assert intake._free_private_path(first) == first.with_name("2026-09-22-120000-2.md")


def test_intake_with_a_locked_vault_writes_nothing(sealed_vault):
    from awb import vault

    p = sealed_vault
    vault.admin_call("lock", p.admin_sock)
    f = p.inbox / "angebot.txt"
    f.write_text("das angebot fuer %s.\n" % FULL, encoding="utf-8")
    with pytest.raises(register.RegisterError, match="locked"):
        intake.run([f], fx.CUSTOMER_CODE, p)
    assert f.exists()
    assert _names_under(p.originals) == [] and _names_under(p.private_reports) == []
    assert not (p.outbox / fx.CUSTOMER_CODE).exists()


def test_a_failed_seal_is_counted_and_reported_without_a_value(sealed_vault, monkeypatch, capsys):
    from awb import cli, vault

    p = sealed_vault
    real = vault.admin_call

    def no_seal(op, sock=None, **fields):
        if op == "seal_file":
            raise vault.VaultUnavailable("no vault daemon")
        return real(op, sock, **fields)

    monkeypatch.setattr(vault, "admin_call", no_seal)
    f = p.inbox / ("%s.txt" % SHORT.lower())
    f.write_text("das angebot fuer %s.\n" % FULL, encoding="utf-8")
    assert cli.main(["intake", "--customer", fx.CUSTOMER_CODE, str(f)]) == 2
    cap = capsys.readouterr()
    assert "2 files of this run are not sealed" in cap.err
    fx.assert_no_fixture_name(cap.out + cap.err, "intake output")
    # nothing is lost: the plaintext stays inside the vault with mode 600 and the report says why
    (original,) = (p.originals / fx.CUSTOMER_CODE).iterdir()
    assert original.suffix == ".txt" and stat.S_IMODE(original.stat().st_mode) == 0o600
    (report,) = (p.private_reports / fx.CUSTOMER_CODE).iterdir()
    assert report.suffix == ".md"
    assert "original not sealed (no vault daemon)" in report.read_text(encoding="utf-8")


def test_keep_list_with_a_real_daemon_in_an_encrypted_vault(sealed_vault):
    p = sealed_vault
    f = p.inbox / "konzept.txt"
    f.write_text("anbei das %s, bitte pruefen.\n" % KEPT, encoding="utf-8")
    assert intake.run([f], fx.CUSTOMER_CODE, p).blocked
    assert intake.keep_phrase(p, KEPT) == (True, 1)
    assert not p.keep_list.exists() and (p.vault / "keep.tsv.gpg").is_file()
    assert KEPT.encode("utf-8") not in (p.vault / "keep.tsv.gpg").read_bytes()
    assert intake.load_keep(p) == [KEPT]
    second = intake.run([f], fx.CUSTOMER_CODE, p)
    assert not second.blocked and second.unsealed == 0
    assert all(n.endswith(".gpg") for n in _names_under(p.originals) + _names_under(p.private_reports))


# --------------------------------------------------------------------------- the self-test and the second check


def test_the_planted_set_uses_the_invented_form_of_the_fixtures():
    from awb import planted
    assert planted.FORM == fx.SELFTEST_FORM
    labels = [c[0] for c in planted.CASES]
    assert len(labels) == len(set(labels)) >= 20


def test_the_self_test_passes_on_the_real_engine():
    assert intake.selftest() == []


def test_the_self_test_refuses_a_broken_sanitiser(home, monkeypatch):
    f = home.inbox / "note.txt"
    f.write_text("vertrag mit %s\n" % FULL, encoding="utf-8")
    monkeypatch.setattr(intake._Engine, "sanitize", lambda self, text, where: (text, Counter()))
    with pytest.raises(intake.IntakeError, match="self-test"):
        intake.run([f], fx.CUSTOMER_CODE, home)
    assert f.is_file()
    assert not (home.outbox / fx.CUSTOMER_CODE).exists() or not any((home.outbox / fx.CUSTOMER_CODE).iterdir())
    assert not (home.originals / fx.CUSTOMER_CODE).exists()


def test_the_self_test_refuses_a_blind_matcher(monkeypatch):
    monkeypatch.setattr(intake._Engine, "final_hits", lambda self, text: [])
    failed = intake.selftest()
    assert failed and all("does not report" in f or "clean" not in f for f in failed)


def test_the_second_check_catches_what_a_blind_matcher_lets_through(home, monkeypatch):
    from awb.matcher import Matcher
    f = home.inbox / "leak.txt"
    f.write_text("vertrag mit %s\n" % FULL, encoding="utf-8")
    monkeypatch.setattr(intake, "selftest", lambda: [])          # a blind matcher is what is simulated here
    monkeypatch.setattr(Matcher, "find", lambda self, text: [])   # the matcher finds no registered form at all
    monkeypatch.setattr(intake, "unknown_candidates", lambda *a, **k: [])
    res = intake.run([f], fx.CUSTOMER_CODE, home)
    (fid,) = res.states
    assert res.states[fid] == "failed" and res.outputs == []
    assert f.is_file(), "the original stays in the inbox"
    assert intake.SECOND_CLASS in res.private_report.read_text(encoding="utf-8")


def test_the_second_check_leaves_ordinary_lower_case_short_words_alone():
    from awb import register
    engine = intake._Engine([register.Entry("ORG-SELF", "ORG", "Apex", "2026-09-24", "active")])
    assert engine.second_hits("the apex of the curve") == []
    assert [h.cls for h in engine.second_hits("an offer from Apex today")] == [intake.SECOND_CLASS]
    assert [h.cls for h in engine.second_hits("an offer from A p e x today")] == [intake.SECOND_CLASS]


def test_the_second_check_follows_the_keep_list():
    from awb import planted
    engine = intake._Engine(planted.entries())
    assert engine.second_hits("offer for %s" % planted.FORM)
    engine.keep = (planted.FORM,)
    assert engine.second_hits("offer for %s" % planted.FORM) == []


def test_the_public_report_never_writes_through_a_link(tmp_path):
    """The review of 2026-09-27 (SEC-SEAL-2): the temp file of a public report has a new name every time, is
    opened exclusively and never through a link; a link on the way to its folder is refused."""
    import os

    from awb import report

    outbox = tmp_path / "outbox"
    outbox.mkdir()
    owners = tmp_path / "owners-file.txt"
    owners.write_text("keep\n", encoding="utf-8")
    os.symlink(owners, outbox / ".report.md.tmp")            # the fixed temp name of the old code, planted
    report.write_replace(outbox / "report.md", "public\n", 0o640, 0o750)
    assert owners.read_text(encoding="utf-8") == "keep\n"
    assert (outbox / "report.md").read_text(encoding="utf-8") == "public\n"
    assert [f.name for f in outbox.iterdir() if f.name.endswith(".tmp")] == [".report.md.tmp"]
    os.symlink(owners, outbox / "linked.md")                 # a link under the report's own name is replaced
    report.write_replace(outbox / "linked.md", "public\n", 0o640, 0o750)
    assert owners.read_text(encoding="utf-8") == "keep\n" and not (outbox / "linked.md").is_symlink()
    elsewhere = tmp_path / "elsewhere"
    elsewhere.mkdir()
    os.symlink(elsewhere, outbox / "cust")                    # a link on the way to the folder
    with pytest.raises(FileExistsError, match="a link stands"):
        report.write_replace(outbox / "cust" / "report.md", "public\n", 0o640, 0o750)
    assert list(elsewhere.iterdir()) == []
